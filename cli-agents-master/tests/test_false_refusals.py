"""The formula checks must not refuse valid Excel (2026-09-19).

Before: a text cell near any bracket refused every lookup by label and every
IF on a text selector (and looked on the wrong sheet for Sheet!A1); the
hand-kept function list refused YEARFRAC, STDEV.S, MROUND, ...; the circular
check ignored sheet names; and post-2007 functions were stored without their
"_xlfn." prefix, so XLOOKUP/IFS/LET were "#NAME?" in Excel and LibreOffice.
Rule now: when in doubt, allow - every write reports the recalculated value.
"""
from openpyxl import Workbook

from excel_mcp_server import formula_validator as fv
from excel_mcp_server.helpers.cell_validation import _validate_cell_references

VALID = [  # (cell, formula) - all valid Excel, all refused before the fix
    ("D10", "=YEARFRAC(DATE(2026,1,1),DATE(2027,1,1),1)"), ("D11", "=STDEV.S(B1:B4)"), ("D11", "=STDEV.P(B1:B4)"),
    ("D11", "=VAR.P(B1:B4)"), ("D17", "=MROUND(B2,5)"), ("D18", "=NORM.S.DIST(0,TRUE)"), ("D18", "=NORM.S.INV(0.95)"),
    ("D22", "=AGGREGATE(9,6,B1:B4)"), ("D22", "=CEILING.MATH(B1)"), ("D22", "=FLOOR.MATH(B1)"),
    ("E2", "=SUM(Ratios!E1:E3)"), ("B2", "=SUM('Other Sheet'!B1:B3)"), ("D20", "=ROWS($D$20:D20)"),
    ("D19", '="Note (see "&A1'), ("D7", '=IF(ABS(B1-B1)<0.01,"OK","ERROR!")'),
    ("D2", "=_xlfn.XLOOKUP(A2,A1:A4,B1:B4)"), ("D12", "=_xlfn.LET(_xlpm.x,B1*2,_xlpm.x+1)"),
]
INVALID = [  # must still be refused
    ("D28", "=SUM(D20:D30)", "CIRCULAR"), ("D35", "=SUM($D$30:$D$40)", "CIRCULAR"), ("B1", "=B1*2", "CIRCULAR"),
    ("B1", "=Calc!B1+1", "CIRCULAR"), ("D29", "=SUMPMT(B1)", "Invalid function"), ("D1", "=SUM((B1:B4)", "parentheses"),
]


def test_valid_formulas_are_accepted():
    for cell, formula in VALID:
        result = fv.validate_formula(formula, cell=cell, worksheet="Calc")
        assert result["valid"], (formula, result["errors"])


def test_real_errors_are_still_refused():
    for cell, formula, needle in INVALID:
        result = fv.validate_formula(formula, cell=cell, worksheet="Calc")
        assert not result["valid"] and needle in " ".join(result["errors"]), (formula, result["errors"])


def _workbook():
    wb = Workbook()
    calc = wb.active; calc.title = "Calc"
    for row, (label, number) in enumerate([("a", 10), ("b", 20), ("c", 30), ("d", 40)], start=1):
        calc[f"A{row}"], calc[f"B{row}"] = label, number
    calc["C1"], calc["C2"], calc["C3"] = "1,200", "12%", "2026-01-31"     # text Excel reads as numbers
    wb.create_sheet("Assumptions")["B1"] = "Base"
    questions = wb.create_sheet("Questions"); questions["C20"] = "USD"
    ratios = wb.create_sheet("Ratios"); ratios["C20"], ratios["A1"] = 12.3456, "label"
    return wb


def test_text_cells_may_be_looked_up_compared_and_joined():
    wb = _workbook()
    for sheet, formula in [
        ("Questions", "=ROUND(Ratios!C20,2)"),                 # Questions!C20 is a unit label - not the cell referenced
        ("Calc", '=IF(Assumptions!B1="Base",1,2)'), ("Calc", '=SWITCH(Assumptions!B1,"Base",1,"Upside",2,0)'),
        ("Calc", "=XLOOKUP(A2,A1:A4,B1:B4)"), ("Calc", "=INDEX(B1:B4,MATCH(A3,A1:A4,0))"), ("Calc", "=SUMIF(A1:A4,A2,B1:B4)"),
        ("Calc", "=COUNTIF(A1:A4,A1)"), ("Calc", "=MAXIFS(B1:B4,A1:A4,A2)"), ("Calc", '=XMATCH("c",A1:A4)'),
        ("Calc", "=LEN(A1)"), ("Calc", "=CONCAT(A1,A2)"), ("Calc", '=TEXTJOIN(", ",TRUE,A1:A4)'), ("Calc", '=A1&" total"'),
        ("Calc", "=SUM(A1:B4)"), ("Calc", "=SUM(B1:B4)/COUNTA(A1:A4)"), ("Calc", '=B1*(A1="a")'),
        ("Calc", "=N(A1)+B1"), ("Calc", "=IFERROR(A1*2,0)"), ("Calc", "=C1*2+C2*B1+C3+1"),
    ]:
        assert _validate_cell_references(formula, wb, sheet) == [], formula


def test_arithmetic_on_a_label_is_still_refused_on_the_right_sheet():
    wb = _workbook()
    for sheet, formula, where in [("Calc", "=A1*5", "A1"), ("Calc", "=B1-A2", "A2"), ("Calc", "=-A1", "A1"),
                                  ("Calc", "=Ratios!A1*2", "Ratios!A1"), ("Questions", "=Calc!A3/2", "Calc!A3")]:
        errors = _validate_cell_references(formula, wb, sheet)
        assert len(errors) == 1 and f"Cell {where} contains text" in errors[0], (formula, errors)
    # same address, other sheet: Calc!B1 is a number although Assumptions!B1 is text
    assert _validate_cell_references("=Calc!B1*2", wb, "Assumptions") == []


def test_post_2007_functions_are_stored_with_their_prefixes():
    cases = {
        "=XLOOKUP(A2,A1:A4,B1:B4)": "=_xlfn.XLOOKUP(A2,A1:A4,B1:B4)",
        "=xmatch(A2,A1:A4)": "=_xlfn.XMATCH(A2,A1:A4)",
        "=IFS(B1>5,1,TRUE,0)": "=_xlfn.IFS(B1>5,1,TRUE,0)",
        "=LET(rev,B1,cost,B2,rev-cost)": "=_xlfn.LET(_xlpm.rev,B1,_xlpm.cost,B2,_xlpm.rev-_xlpm.cost)",
        "=LET(x, B1*2, y, LET(z, x+1, z*2), x+y)":
            "=_xlfn.LET(_xlpm.x, B1*2, _xlpm.y, _xlfn.LET(_xlpm.z, _xlpm.x+1, _xlpm.z*2), _xlpm.x+_xlpm.y)",
        "=FILTER(A1:A4,B1:B4>0)": "=_xlfn._xlws.FILTER(A1:A4,B1:B4>0)",
        '=IF(A1="XLOOKUP(",STDEV.S(B1:B4),0)': '=IF(A1="XLOOKUP(",_xlfn.STDEV.S(B1:B4),0)',
    }
    for plain, stored in cases.items():
        assert fv.add_future_function_prefixes(plain) == stored, plain
        assert fv.add_future_function_prefixes(stored) == stored, "already prefixed text must not change"
    for untouched in ("=SUM(B1:B4)", "=VLOOKUP(A1,A1:B4,2,FALSE)", "=IFERROR(B1/B2,0)", "=SUMIF(A:A, '>100', B:B)"):
        assert fv.add_future_function_prefixes(untouched) == untouched


def test_text_literals_are_not_read_as_sheets_or_brackets():
    assert fv.strip_string_literals('=IF(A1>0,"ERROR!","say ""(hi""")') == '=IF(A1>0,"","")'
    for name in ("XLOOKUP", "XMATCH", "IFS", "SWITCH", "LET", "YEARFRAC", "STDEV.S", "MROUND", "AGGREGATE", "EOMONTH"):
        assert name in fv.VALID_EXCEL_FUNCTIONS, name
    assert "SUMPMT" not in fv.VALID_EXCEL_FUNCTIONS
