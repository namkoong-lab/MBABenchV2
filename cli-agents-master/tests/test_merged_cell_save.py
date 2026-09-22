"""Saving a workbook whose column starts inside a merged range must not crash
(2026-09-19: copy_file died with "'MergedCell' object has no attribute
'column_letter'" on CoastalAggregates, so solution.xlsx was never created)."""
from openpyxl import Workbook, load_workbook

from excel_mcp_server.core.workbook_io import _auto_fit_columns, _save_workbook_sync


def test_a_merged_title_row_does_not_break_the_save(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "A long merged title across the top of the sheet"
    ws.merge_cells("A1:D2")                      # B1, C1, D1 become MergedCell - the first cell of their columns
    ws["B3"], ws["C3"] = "label", 12345.678
    _auto_fit_columns(wb)                        # raised AttributeError before the fix
    assert ws.column_dimensions["B"].width >= 10 and ws.column_dimensions["D"].width >= 10
    out = tmp_path / "merged.xlsx"
    _save_workbook_sync(wb, out)
    assert "A1:D2" in [str(r) for r in load_workbook(out).active.merged_cells.ranges]


def test_a_save_killed_mid_write_leaves_the_previous_workbook_intact(tmp_path, monkeypatch):
    """2026-09-22: a tool call that overran its timeout was killed while openpyxl was
    rewriting solution.xlsx in place, leaving a truncated zip - the attempt's whole
    deliverable (gpt-6-astra task 89) could not be opened. The save is atomic now."""
    import os
    import zipfile

    import openpyxl
    from excel_mcp_server.core import workbook_io

    out = tmp_path / "solution.xlsx"
    wb = Workbook(); wb.active["A1"] = "first good version"
    _save_workbook_sync(wb, out)
    good = out.read_bytes()

    real_save = openpyxl.workbook.workbook.Workbook.save

    def die_midway(self, path):                       # what a SIGKILL during the write looks like
        with open(path, "wb") as f:
            f.write(b"PK\x03\x04" + b"\x00" * 4096)   # a zip that stops before its directory
        raise KeyboardInterrupt("killed mid-save")

    monkeypatch.setattr(openpyxl.workbook.workbook.Workbook, "save", die_midway)
    wb2 = Workbook(); wb2.active["A1"] = "second version, never finished"
    try:
        _save_workbook_sync(wb2, out)
    except KeyboardInterrupt:
        pass
    monkeypatch.setattr(openpyxl.workbook.workbook.Workbook, "save", real_save)

    assert out.read_bytes() == good                   # untouched by the killed save
    assert load_workbook(out).active["A1"].value == "first good version"
    assert zipfile.ZipFile(out).namelist()            # still a valid xlsx
    assert not [p for p in os.listdir(tmp_path) if p.startswith(".")]   # no temp left behind

    wb3 = Workbook(); wb3.active["A1"] = "third version"
    _save_workbook_sync(wb3, out)                     # a normal save still replaces it
    assert load_workbook(out).active["A1"].value == "third version"
    assert os.listdir(tmp_path) == ["solution.xlsx"]
