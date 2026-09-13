"""Offline tests for infra.run.check_output_quality — the post-engine gate that
turns a recorded 'success' into 'failed' when the workbook is not a deliverable
(2026-09-11: starting file / formula-free checkpoint picked up as the artifact)."""
import shutil
import sys
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from infra.run import check_output_quality, QUALITY_MIN_BYTES  # noqa: E402


def _book(path: Path, formulas: bool, pad: bool = True) -> Path:
    wb = openpyxl.Workbook()
    for name in ("Instructions", "Assumptions", "Questions"):
        ws = wb.create_sheet(name)
        for r in range(1, 400 if pad else 5):
            ws.cell(r, 1, f"row {r} " * 8)
        if formulas and name == "Questions":
            ws["B2"] = "=SUM(1,2)"
    del wb["Sheet"]
    wb.save(path)
    assert path.stat().st_size >= QUALITY_MIN_BYTES
    return path


def test_real_model_passes(tmp_path):
    sol = _book(tmp_path / "sol.xlsx", formulas=True)
    inp = _book(tmp_path / "start.xlsx", formulas=False)
    ok, reason = check_output_quality(sol, [inp])
    assert ok, reason


def test_identical_to_input_fails(tmp_path):
    inp = _book(tmp_path / "start.xlsx", formulas=True)
    sol = tmp_path / "sol.xlsx"
    shutil.copyfile(inp, sol)
    ok, reason = check_output_quality(sol, [inp])
    assert not ok and "byte-identical" in reason


def test_no_formulas_fails(tmp_path):
    sol = _book(tmp_path / "sol.xlsx", formulas=False)
    ok, reason = check_output_quality(sol, [])
    assert not ok and "no formulas" in reason


def test_namespaced_xml_formulas_count(tmp_path):
    # Writers such as the ChatGPT sandbox emit <x:f>; the count must see them.
    import zipfile
    src = _book(tmp_path / "plain.xlsx", formulas=True)
    dst = tmp_path / "ns.xlsx"
    with zipfile.ZipFile(src) as zi, zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zo:
        for item in zi.infolist():
            data = zi.read(item.filename)
            if "worksheets/" in item.filename:
                data = data.replace(b"<f>", b"<x:f>").replace(b"</f>", b"</x:f>")
            zo.writestr(item, data)
    from infra.run import _count_formulas
    assert _count_formulas(dst) == 1


# --- usage-cap lane stop -------------------------------------------------
from infra.run import usage_cap_hit  # noqa: E402
import json  # noqa: E402


def test_usage_cap_from_completion_json(tmp_path):
    (tmp_path / "json_logs").mkdir()
    (tmp_path / "json_logs" / "completion_x.json").write_text(
        json.dumps({"tasks": [{"task_status": "rate_limited"}]})
    )
    assert "rate_limited" in (usage_cap_hit(tmp_path) or "")


def test_usage_cap_from_engine_log(tmp_path):
    (tmp_path / "logs").mkdir()
    (tmp_path / "logs" / "claude_web_x.log").write_text(
        "... - ERROR - Rate limit persisted — aborting wait\n"
    )
    assert "Rate limit persisted" in (usage_cap_hit(tmp_path) or "")


def test_no_cap_on_ordinary_failure(tmp_path):
    (tmp_path / "json_logs").mkdir(); (tmp_path / "logs").mkdir()
    (tmp_path / "json_logs" / "completion_x.json").write_text(
        json.dumps({"tasks": [{"task_status": "download_failed"}]})
    )
    (tmp_path / "logs" / "claude_web_x.log").write_text("Download failed after 2 attempt(s)\n")
    assert usage_cap_hit(tmp_path) is None
