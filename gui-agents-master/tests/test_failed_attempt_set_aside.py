"""Offline check: a failed attempt's archived workbooks can never be picked as
the task's solution. No browser.

Live 2026-09-21 (DailyCash): attempt 1 died with "Thinking failed"; the
archival download then dropped its sandbox files — including a 1.2 MB
DailyCash_Completed_Model.xlsx the model never delivered — into the same
solutions/ folder attempt 2 delivers into, where the runner's finder ranks
name-matching workbooks by quality, then SIZE.

Run from gui-agents-master:  python -m pytest tests/test_failed_attempt_set_aside.py
"""
import sys
from datetime import datetime, timedelta
from pathlib import Path

import openpyxl

REPO = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO))

from claude_web_agent.claude_web_engine import (  # noqa: E402
    rename_solution_file,
    set_aside_failed_attempt_files,
)
from infra.run import collect_extra_workbooks, find_solution_file  # noqa: E402


def workbook(path: Path, sheets: int, pad: int) -> Path:
    wb = openpyxl.Workbook()
    wb.active.title = "S1"
    for i in range(2, sheets + 1):
        wb.create_sheet(f"S{i}")
    for r in range(1, pad + 1):
        wb["S1"].cell(row=r, column=1, value=f"row {r} " + "x" * 40)
    wb.save(path)
    return path


def test_leftovers_lose_to_the_delivered_workbook(tmp_path):
    started = datetime.now() - timedelta(minutes=5)
    sol = tmp_path / "solutions"
    sol.mkdir()
    # attempt 1 (failed): archival download — a BIG never-delivered model, the input copy, a scratch file
    big = workbook(sol / "DailyCash_Completed_Model.xlsx", sheets=6, pad=4000)
    inp = workbook(sol / "DailyCash(5).xlsx", sheets=3, pad=50)
    scr = workbook(sol / "build_checkpoint.xlsx", sheets=2, pad=50)

    # without the fix the finder would take the big leftover
    assert find_solution_file(tmp_path, "DailyCash", None, started) in (big, inp)

    moved = set_aside_failed_attempt_files([big, inp, scr], 1, "DailyCash")
    assert [p.name for p in moved] == [
        "failed_attempt1__TASK_Completed_Model.xlsx",
        "failed_attempt1__TASK(5).xlsx",
        "failed_attempt1__build_checkpoint.xlsx",
    ]
    assert find_solution_file(tmp_path, "DailyCash", None, started) is None

    # attempt 2 (success): a smaller, delivered workbook
    delivered = rename_solution_file(
        workbook(sol / "DailyCash_Model.xlsx", sheets=5, pad=300),
        "DailyCash", "chatgpt_gpt_6_pro",
    )
    chosen = find_solution_file(tmp_path, "DailyCash", None, started)
    assert chosen == delivered
    assert delivered.stat().st_size < moved[0].stat().st_size   # size alone would have lost

    # the leftovers still travel to S3 as extras
    extras = {p.name for p in collect_extra_workbooks(tmp_path, chosen)}
    assert extras == {p.name for p in moved}


def test_it_is_idempotent_and_never_raises(tmp_path):
    sol = tmp_path / "solutions"
    sol.mkdir()
    f = workbook(sol / "P&L Budget - Meridian v2.xlsx", sheets=2, pad=5)
    first = set_aside_failed_attempt_files([f], 2, "Meridian")
    assert first[0].name == "failed_attempt2__P&L Budget - TASK v2.xlsx"
    assert set_aside_failed_attempt_files(first, 2, "Meridian") == []      # already set aside
    assert set_aside_failed_attempt_files(["/nope/missing.xlsx"], 1, "Meridian") == []
    assert set_aside_failed_attempt_files(None, 1, "Meridian") == []


# ---- a copy of the input must never be chosen as the deliverable ------------
# Live 2026-09-23 (task 99 TAM): the downloads held the 4.4 MB starting
# workbook alongside the model's 2.6 MB answer (87,198 formulas). Ranking is
# quality-then-SIZE, so the input copy won and the quality gate failed the
# attempt — with the real deliverable sitting next to it in solutions/.


def test_an_input_copy_loses_to_the_real_deliverable(tmp_path):
    started = datetime.now() - timedelta(minutes=5)
    sol = tmp_path / "solutions"
    sol.mkdir()
    src = tmp_path / "TAM.xlsx"
    workbook(src, sheets=4, pad=3000)                      # the big starting file
    copy = sol / "20260923_TAM_Solution_x_Model.xlsx"
    copy.write_bytes(src.read_bytes())                     # downloaded straight back
    deliverable = workbook(sol / "20260923_TAM_Solution_x_Model_2.xlsx", sheets=6, pad=900)
    assert copy.stat().st_size > deliverable.stat().st_size, "size alone would pick the copy"

    assert find_solution_file(tmp_path, "TAM", None, started) == copy      # old behaviour
    chosen = find_solution_file(tmp_path, "TAM", None, started, inputs=[str(src)])
    assert chosen == deliverable
    # the copy still travels to S3 as an extra
    assert copy in collect_extra_workbooks(tmp_path, chosen)


def test_an_input_copy_is_still_returned_when_it_is_all_there_is(tmp_path):
    """Never leave the caller with nothing: a lone copy is a bad attempt to
    report, not a missing file to crash on."""
    started = datetime.now() - timedelta(minutes=5)
    sol = tmp_path / "solutions"
    sol.mkdir()
    src = tmp_path / "TAM.xlsx"
    workbook(src, sheets=3, pad=200)
    copy = sol / "20260923_TAM_Solution_x_Model.xlsx"
    copy.write_bytes(src.read_bytes())
    assert find_solution_file(tmp_path, "TAM", None, started, inputs=[str(src)]) == copy


def test_unreadable_inputs_do_not_break_the_pick(tmp_path):
    started = datetime.now() - timedelta(minutes=5)
    sol = tmp_path / "solutions"
    sol.mkdir()
    a = workbook(sol / "20260923_TAM_Solution_x_Model.xlsx", sheets=5, pad=400)
    assert find_solution_file(
        tmp_path, "TAM", None, started, inputs=["/nope/missing.xlsx"]
    ) == a
