"""Reader / engine tests on synthetic workbooks (openpyxl-built and raw XML zips).

    cd /Users/patrick/MBABench-deterministic-checks
    /Users/patrick/MBABenchV2/.venv/bin/python -m detchecks.tests.test_reader

Plain asserts; also collectable by pytest.  Temporary files go to detchecks/scratch/.
"""
from __future__ import annotations

import json
import os
import shutil
import sys
import tempfile
import traceback
import zipfile

import openpyxl
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from detchecks.api import Engine, grade
from detchecks.checks import REGISTRY
from detchecks.checks.base import Check, Mistakes, cells_to_ranges, make_verdict, validate_verdict
from detchecks.core import refs
from detchecks.core.package import Package
from detchecks.core.sheet import ExcelError, SheetStream, load_sheet
from detchecks.core.styles import BLACK, WHITE, Color, Styles, excel_tint
from detchecks.core.values import detect_provenance, detect_writer, make_context
from detchecks.errors import GradingError

HERE = os.path.dirname(os.path.abspath(__file__))
SCRATCH = os.path.join(os.path.dirname(HERE), "scratch")
RUBRIC = os.path.join(os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))), "prompts", "rubrics", "rubric_9.json")

MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PKG_REL = "http://schemas.openxmlformats.org/package/2006/relationships"
S_MAIN = "http://purl.oclc.org/ooxml/spreadsheetml/main"
S_REL = "http://purl.oclc.org/ooxml/officeDocument/relationships"
CT_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"
CT_XLSM = "application/vnd.ms-excel.sheet.macroEnabled.main+xml"
CT_WS = "application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"

# Excel's own XML declaration (every part Excel writes starts with it)
XL_DECL = '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\r\n'
PY_DECL = "<?xml version='1.0' encoding='UTF-8'?>"      # Python ElementTree style

_TMP = None


def tmp(name: str) -> str:
    global _TMP
    if _TMP is None:
        os.makedirs(SCRATCH, exist_ok=True)
        _TMP = tempfile.mkdtemp(prefix="test_reader_", dir=SCRATCH)
    return os.path.join(_TMP, name)


# ============================================================================ raw builder
def build(path, sheets, *, styles=None, sst=None, app=None, core=None, wb_extra="", wb_pre="",
          main_ns=MAIN, rel_ns=REL, rel_type_base="http://schemas.openxmlformats.org/officeDocument/2006/relationships",
          main_ct=CT_XLSX, extra_parts=None, extra_wb_rels="", sheet_rels=None, raw_sheets=False,
          workbook_name="xl/workbook.xml", decl=XL_DECL):
    """sheets: [(name, state|None, worksheet_xml_or_inner)].  When raw_sheets is False the
    inner XML is wrapped in <worksheet xmlns=main_ns xmlns:r=rel_ns>.  sheet_rels: {index: rels xml body}.
    decl: XML declaration written before workbook.xml, docProps and wrapped sheets (Excel's by default)."""
    extra_parts = extra_parts or {}
    sheet_rels = sheet_rels or {}
    z = zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED)
    ov = [f'<Override PartName="/{workbook_name}" ContentType="{main_ct}"/>']
    for i in range(len(sheets)):
        ov.append(f'<Override PartName="/xl/worksheets/sheet{i + 1}.xml" ContentType="{CT_WS}"/>')
    z.writestr("[Content_Types].xml",
               '<?xml version="1.0" encoding="UTF-8"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
               '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
               '<Default Extension="xml" ContentType="application/xml"/>' + "".join(ov) + "</Types>")
    root_rels = [f'<Relationship Id="rId1" Type="{rel_type_base}/officeDocument" Target="{workbook_name}"/>']
    if app is not None:
        root_rels.append(f'<Relationship Id="rId2" Type="{rel_type_base}/extended-properties" Target="docProps/app.xml"/>')
        z.writestr("docProps/app.xml", decl + '<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/extended-properties">'
                   + app + "</Properties>")
    if core is not None:
        root_rels.append('<Relationship Id="rId3" Type="http://schemas.openxmlformats.org/package/2006/relationships/metadata/core-properties" Target="docProps/core.xml"/>')
        z.writestr("docProps/core.xml", decl + '<cp:coreProperties xmlns:cp="http://schemas.openxmlformats.org/package/2006/metadata/core-properties" '
                   'xmlns:dc="http://purl.org/dc/elements/1.1/">' + core + "</cp:coreProperties>")
    z.writestr("_rels/.rels", f'<Relationships xmlns="{PKG_REL}">' + "".join(root_rels) + "</Relationships>")
    wrels, sh = [], []
    for i, (name, state, xml) in enumerate(sheets):
        st = f' state="{state}"' if state else ""
        sh.append(f'<sheet name="{name}" sheetId="{i + 1}"{st} r:id="rId{i + 1}"/>')
        wrels.append(f'<Relationship Id="rId{i + 1}" Type="{rel_type_base}/worksheet" Target="worksheets/sheet{i + 1}.xml"/>')
        body = xml if raw_sheets else f'{decl}<worksheet xmlns="{main_ns}" xmlns:r="{rel_ns}">{xml}</worksheet>'
        z.writestr(f"xl/worksheets/sheet{i + 1}.xml", body)
        if i in sheet_rels:
            z.writestr(f"xl/worksheets/_rels/sheet{i + 1}.xml.rels",
                       f'<Relationships xmlns="{PKG_REL}">{sheet_rels[i]}</Relationships>')
    if styles is not None:
        wrels.append(f'<Relationship Id="rIdS" Type="{rel_type_base}/styles" Target="styles.xml"/>')
        z.writestr("xl/styles.xml", f'<styleSheet xmlns="{main_ns}">{styles}</styleSheet>')
    if sst is not None:
        wrels.append(f'<Relationship Id="rIdT" Type="{rel_type_base}/sharedStrings" Target="sharedStrings.xml"/>')
        z.writestr("xl/sharedStrings.xml", f'<sst xmlns="{main_ns}">' + "".join(sst) + "</sst>")
    wrels.append(extra_wb_rels)
    z.writestr(workbook_name, f'{decl}<workbook xmlns="{main_ns}" xmlns:r="{rel_ns}">{wb_pre}<sheets>' + "".join(sh)
               + f"</sheets>{wb_extra}</workbook>")
    d = os.path.dirname(workbook_name)
    z.writestr(f"{d}/_rels/{os.path.basename(workbook_name)}.rels",
               f'<Relationships xmlns="{PKG_REL}">' + "".join(wrels) + "</Relationships>")
    for k, v in extra_parts.items():
        z.writestr(k, v)
    z.close()
    return path


def cells_by_ref(cells):
    return {c.ref: c for c in cells}


def _raises(fn, *needles):
    """fn() must raise GradingError whose message contains every needle; returns the error."""
    try:
        fn()
    except GradingError as e:
        msg = str(e) + " " + " ".join(str(v) for v in e.failures.values())
        for n in needles:
            assert n in msg, (n, msg[:500])
        return e
    raise AssertionError(f"expected GradingError containing {needles}")


def _rewrite(path, **parts):
    """Replace zip members of a built workbook: _rewrite(p, **{"xl/workbook.xml": b"..."})."""
    with zipfile.ZipFile(path) as z:
        old = {n: z.read(n) for n in z.namelist()}
    old.update({k: (v.encode() if isinstance(v, str) else v) for k, v in parts.items()})
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for k, v in old.items():
            z.writestr(k, v)
    return path


# ============================================================================ tests
def test_refs():
    assert refs.col_to_index("A") == 1 and refs.col_to_index("xfd") == 16384
    assert refs.index_to_col(28) == "AB"
    assert refs.split_ref("$B$7") == (7, 2) and refs.split_ref("Name") is None
    assert refs.parse_range("C5:A1") == (1, 1, 5, 3)
    assert refs.parse_range("A:B") == (1, 1, refs.MAX_ROW, 2)
    assert refs.parse_range("$3:$4") == (3, 1, 4, refs.MAX_COL)
    assert refs.parse_range("A1:5") is None
    assert refs.quote_sheet("Solution Model") == "'Solution Model'" and refs.quote_sheet("Sheet1") == "Sheet1"
    assert refs.quote_sheet("A1") == "'A1'"
    cells = [(1, 1), (1, 2), (2, 1), (2, 2), (5, 5), (3, 1)]
    assert cells_to_ranges(cells) == ["A1:B2", "A3", "E5"]
    assert cells_to_ranges([(1, 1), (1, 3)]) == ["A1", "C1"]
    assert cells_to_ranges([(r, 2) for r in range(1, 11)]) == ["B1:B10"]


KNOWN_TINTS = [
    ("000000", 0.0499893185216834, "0D0D0D"), ("000000", 0.249977111117893, "404040"),
    ("000000", 0.499984740745262, "808080"), ("000000", 0.49998, "808080"),
    ("FFFFFF", -0.0499893185216834, "F2F2F2"), ("FFFFFF", -0.149998474074526, "D9D9D9"),
    ("FFFFFF", -0.499984740745262, "808080"), ("44546A", 0.799981688894314, "D6DCE4"),
    ("4472C4", 0.799981688894314, "D9E1F2"), ("4472C4", -0.249977111117893, "305496"),
    ("ED7D31", 0.599993896298105, "F8CBAD"), ("FFC000", 0.799981688894314, "FFF2CC"),
    ("FFC000", -0.499984740745262, "806000"), ("5B9BD5", 0.399975585192419, "9BC2E6"),
    ("70AD47", -0.249977111117893, "548235"), ("A5A5A5", -0.499984740745262, "525252"),
]


def test_excel_tint():
    for base, t, exp in KNOWN_TINTS:
        assert excel_tint(base, t) == exp, (base, t, excel_tint(base, t), exp)


STYLES_XML = (
    '<numFmts count="1"><numFmt numFmtId="164" formatCode="#,##0;(#,##0);&quot;-&quot;"/></numFmts>'
    '<fonts count="5"><font><sz val="11"/><name val="Calibri"/></font>'
    '<font><b/><color rgb="00FF0000"/><sz val="10"/><name val="Arial"/></font>'
    '<font><color theme="1" tint="0.49998"/></font>'
    '<font><color indexed="13"/></font>'
    '<font><color theme="0"/></font></fonts>'
    '<fills count="5"><fill><patternFill patternType="none"/></fill><fill><patternFill patternType="gray125"/></fill>'
    '<fill><patternFill patternType="solid"><fgColor indexed="13"/><bgColor indexed="64"/></patternFill></fill>'
    '<fill><patternFill patternType="solid"><fgColor theme="7" tint="0.79998168889431442"/></patternFill></fill>'
    '<fill><patternFill/></fill></fills>'
    '<borders count="1"><border><left/><right/><top/><bottom/><diagonal/></border></borders>'
    '<cellStyleXfs count="1"><xf numFmtId="0" fontId="0" fillId="0" borderId="0"/></cellStyleXfs>'
    '<cellXfs count="6"><xf numFmtId="0" fontId="0" fillId="0" borderId="0" xfId="0"/>'
    '<xf numFmtId="164" fontId="1" fillId="2" borderId="0" xfId="0" applyFont="1"/>'
    '<xf numFmtId="44" fontId="2" fillId="3" borderId="0" xfId="0"><alignment horizontal="centerContinuous" wrapText="1"/></xf>'
    '<xf numFmtId="0" fontId="3" fillId="4" borderId="0" xfId="0"/>'
    '<xf numFmtId="10" fontId="4" fillId="0" borderId="0" xfId="0"><protection locked="0" hidden="1"/></xf>'
    '<xf numFmtId="0" fontId="0" fillId="2" borderId="0" xfId="0"/></cellXfs>'
    '<cellStyles count="1"><cellStyle name="Normal" xfId="0" builtinId="0"/></cellStyles>'
    '<dxfs count="3"><dxf><font><color rgb="FF9C0006"/></font><fill><patternFill><bgColor rgb="FFFFC7CE"/></patternFill></fill></dxf>'
    '<dxf><fill><patternFill patternType="solid"><fgColor rgb="FFFFFF00"/></patternFill></fill></dxf>'
    '<dxf><numFmt numFmtId="165" formatCode=";;;"/></dxf></dxfs>'
)


def test_styles_and_colours():
    st = Styles.parse(STYLES_XML.join(['<styleSheet xmlns="%s">' % MAIN, "</styleSheet>"]).encode(), None)
    assert st.indexed_palette[13] == "FFFF00"
    assert st.font_color(0) == BLACK                       # no <color> = automatic
    assert st.font_color(1) == "FFFF0000"                  # alpha 00 ignored
    assert st.font_color(2) == "FF808080"                  # theme 1 + tint 0.49998 (Excel: grey)
    assert st.font_color(3) == "FFFFFF00"                  # indexed 13 = bright yellow
    assert st.font_color(4) == WHITE                       # theme 0 = lt1 = white
    assert st.fill_color(0) is None and st.fill_color(1) == "FFFFFF00"
    assert st.fill_color(2) == "FFFFF2CC"                  # accent4 lighter 80% (pale yellow)
    assert st.fill_color(3) is None                        # patternFill without patternType = none
    assert st.cell_fill(1).pattern == "solid"
    assert st.num_fmt_code(1) == "#,##0;(#,##0);\"-\"" and st.num_fmt_id(2) == 44
    assert st.num_fmt_code(2).startswith('_("$"* #,##0.00_);') and st.num_fmt_code(4) == "0.00%"
    assert st.xf(2).alignment.horizontal == "centerContinuous" and st.xf(2).alignment.wrap_text
    assert st.xf(4).hidden and not st.xf(4).locked
    assert st.xf(999).font_id == 0                         # out of range -> default xf
    assert st.font(1).b and st.font(1).name == "Arial" and st.font(1).sz == 10.0
    assert st.style_name(1) == "Normal"
    assert st.dxf_fill_color(0) == "FFFFC7CE"              # dxf: solid colour lives in bgColor
    assert st.dxf_font_color(0) == "FF9C0006"
    assert st.dxf_fill_color(1) is None                    # dxf with fgColor only: no fill (Excel 2026-10-03)
    assert st.dxf(2).num_fmt_code == ";;;" and st.dxf_font_color(2) is None
    assert st.resolve(Color(indexed=64), "font") == BLACK and st.resolve(Color(indexed=65), "fill") == WHITE
    assert st.resolve(Color(auto=True), "fill") == WHITE and st.resolve(Color(auto=True), "font") == BLACK
    assert st.resolve(Color(theme=1)) == BLACK and st.resolve(Color(theme=99)) is None
    assert st.resolve(Color(indexed=200)) is None
    # custom indexed palette overrides the default
    custom = STYLES_XML + '<colors><indexedColors>' + '<rgbColor rgb="FF000000"/>' * 13 + '<rgbColor rgb="FF00B050"/></indexedColors></colors>'
    st2 = Styles.parse(f'<styleSheet xmlns="{MAIN}">{custom}</styleSheet>'.encode(), None)
    assert st2.custom_indexed and st2.font_color(3) == "FF00B050" and st2.fill_color(1) == "FF00B050"
    # theme part read from the workbook's own theme
    theme = (b'<a:theme xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><a:themeElements>'
             b'<a:clrScheme name="x"><a:dk1><a:sysClr val="windowText" lastClr="111111"/></a:dk1>'
             b'<a:lt1><a:srgbClr val="FEFEFE"/></a:lt1></a:clrScheme></a:themeElements></a:theme>')
    st3 = Styles.parse(None, theme)
    assert st3.resolve(Color(theme=1)) == "FF111111" and st3.resolve(Color(theme=0)) == "FFFEFEFE"
    assert st3.resolve(Color(theme=4)) == "FF4472C4"      # missing slot -> Office default
    assert Styles.effective_style(None, 7, 3) == 7 and Styles.effective_style(None, None, 3) == 3
    assert Styles.effective_style(5, 7, 3) == 5 and Styles.effective_style(None, None, None) == 0


def _openpyxl_book(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Model"
    ws["A1"] = "Title"
    ws["A1"].font = Font(bold=True, color="FF0000FF")
    ws["A2"] = 10
    ws["A3"] = 0
    ws["A4"] = "=A2*2"
    ws["B2"] = "=SUM(A2:A3)"
    ws["C2"] = True
    ws.merge_cells("D1:F1")
    ws["D1"].alignment = Alignment(horizontal="center")
    ws.freeze_panes = "B3"
    ws.sheet_view.zoomScale = 85
    ws["G5"].fill = PatternFill("solid", fgColor="FFFFFF00")
    ws.conditional_formatting.add("A2:A10", CellIsRule(operator="equal", formula=["0"],
                                                       fill=PatternFill(bgColor="FFFFC7CE")))
    ws.conditional_formatting.add("B2:B10", FormulaRule(formula=["TODAY()>1"], font=Font(color="FFFF0000")))
    dv = DataValidation(type="list", formula1="=Lists!$A$1:$A$3", allow_blank=True)
    ws.add_data_validation(dv)
    dv.add("C5:C7")
    ws["H1"] = "link"
    ws["H1"].hyperlink = "https://example.com/x"
    ws.column_dimensions["J"].width = 0.5
    ws.column_dimensions["K"].hidden = True
    ws.row_dimensions[8].hidden = True
    ws.row_dimensions[9].height = 40
    lists = wb.create_sheet("Lists")
    for i, v in enumerate(["a", "b", "c"], 1):
        lists.cell(i, 1, v)
    lists.sheet_state = "hidden"
    vh = wb.create_sheet("Secret")
    vh.sheet_state = "veryHidden"
    vh["A1"] = "x"
    wb.create_sheet("Instructions").merge_cells("B2:K2")
    from openpyxl.workbook.defined_name import DefinedName
    wb.defined_names["Rate"] = DefinedName("Rate", attr_text="Model!$A$2")
    lists.defined_names["LocalList"] = DefinedName("LocalList", attr_text="Lists!$A$1:$A$3")
    wb.defined_names["HiddenOne"] = DefinedName("HiddenOne", attr_text="Model!$A$3", hidden=True)
    wb.save(path)
    return path


def test_openpyxl_workbook():
    p = _openpyxl_book(tmp("op.xlsx"))
    pkg = Package.open(p)
    assert pkg.is_spreadsheetml and pkg.format == "xlsx" and pkg.conformance == "transitional"
    assert [(s.name, s.state) for s in pkg.sheets] == [("Model", "visible"), ("Lists", "hidden"),
                                                      ("Secret", "veryHidden"), ("Instructions", "visible")]
    names = {d.name: d for d in pkg.defined_names}
    assert names["Rate"].scope is None and names["Rate"].text == "Model!$A$2"
    assert names["LocalList"].scope == "Lists" and names["LocalList"].local_sheet_id == 1
    assert names["HiddenOne"].hidden and not names["Rate"].hidden
    prov = detect_provenance(pkg)
    assert prov.writer == "openpyxl" and not prov.formula_caches_trusted
    head, rows, cells, tail = load_sheet(pkg, "Model")
    c = cells_by_ref(cells)
    assert c["A1"].value == "Title" and c["A1"].t == "inlineStr" and c["A1"].value_trusted   # openpyxl 3.1 writes inline strings
    assert c["A2"].value == 10.0 and c["A3"].value == 0.0
    assert c["A4"].has_formula and c["A4"].formula_text == "A2*2"
    assert c["A4"].cached_value is None and c["A4"].unverified_value is None
    assert not c["A4"].value_trusted and c["A4"].value_source == "cached"
    _raises(lambda: c["A4"].value, "untrusted")              # no fallback: untrusted formula value
    assert c["C2"].value is True
    assert pkg.styles.font_color(c["A1"].s) == "FF0000FF"
    assert pkg.styles.fill_color(c["G5"].s) == "FFFFFF00" and c["G5"].is_blank
    v = head.view
    assert v.zoom == 85 and v.pane.state == "frozen" and v.pane.top_left_cell == "B3"
    assert v.active_pane == "bottomRight"
    assert head.col_info(10).width == 0.5 and head.col_info(11).hidden and head.col_info(1) is None
    rws = {r.r: r for r in rows}
    assert rws[8].hidden and rws[9].ht == 40 and rws[9].custom_height
    assert [m.ref for m in tail.merges] == ["D1:F1"]
    cf = {f.sqref: f for f in tail.conditional_formats}
    rule = cf["A2:A10"].rules[0]
    assert rule.type == "cellIs" and rule.operator == "equal" and rule.formulas == ["0"]
    assert pkg.styles.dxf_fill_color(rule.dxf_id) == "FFFFC7CE"
    assert cf["B2:B10"].rules[0].formulas == ["TODAY()>1"]
    dv = tail.data_validations[0]
    # formula text is kept exactly as stored (this openpyxl file stores a leading '=')
    assert dv.type == "list" and dv.formula1 == "=Lists!$A$1:$A$3" and dv.ranges == [(5, 3, 7, 3)]
    hl = tail.hyperlinks[0]
    assert hl.ref == "H1" and hl.target == "https://example.com/x" and hl.external
    # skeleton mode reads the same head / tail
    with SheetStream(pkg, pkg.sheet("Model")) as ss:
        h2 = ss.read_head()
        t2 = ss.skip_body()
    assert [m.ref for m in t2.merges] == ["D1:F1"] and len(t2.conditional_formats) == 2
    assert h2.view.pane.top_left_cell == "B3" and len(t2.hyperlinks) == 1
    pkg.close()


PREFIXED_SHEET = (
    XL_DECL + '<x:worksheet xmlns:x="%s" xmlns:r="%s">'
    '<x:sheetPr><x:tabColor rgb="FFFF0000"/><x:outlinePr summaryBelow="0"/></x:sheetPr>'
    '<x:dimension ref="A1:E6"/>'
    '<x:sheetViews><x:sheetView tabSelected="1" zoomScale="90" workbookViewId="0">'
    '<x:pane xSplit="1" ySplit="1" topLeftCell="B2" activePane="bottomRight" state="frozen"/>'
    '<x:selection pane="topRight" activeCell="B1" sqref="B1"/>'
    '<x:selection pane="bottomRight" activeCell="C4" sqref="C4"/></x:sheetView></x:sheetViews>'
    '<x:sheetFormatPr defaultRowHeight="15" zeroHeight="1" baseColWidth="10"/>'
    '<x:cols><x:col min="2" max="4" width="12" customWidth="1" style="3"/></x:cols>'
    '<x:sheetData>'
    '<x:row><x:c t="inlineStr"><x:is><x:r><x:t>Hel</x:t></x:r><x:r><x:rPr><x:b/><x:color rgb="FFFFFFFF"/></x:rPr><x:t>lo</x:t></x:r>'
    '<x:rPh><x:t>IGNORED</x:t></x:rPh></x:is></x:c><x:c><x:v>1</x:v></x:c><x:c><x:v>2</x:v></x:c></x:row>'
    '<x:row r="3" s="5" customFormat="1" ht="30" customHeight="1">'
    '<x:c r="A3"><x:f t="shared" ref="A3:C3" si="0">A1+B$1*2</x:f><x:v>5</x:v></x:c>'
    '<x:c r="B3"><x:f t="shared" si="0"/><x:v>6</x:v></x:c><x:c><x:f t="shared" si="0"/><x:v>7</x:v></x:c>'
    '<x:c r="E3" t="b"><x:v>1</x:v></x:c></x:row>'
    '<x:row r="4"><x:c r="A4" t="e"><x:v>#DIV/0!</x:v></x:c><x:c r="B4" t="str"><x:f>"a"&amp;"b"</x:f><x:v>ab</x:v></x:c>'
    '<x:c r="C4" cm="1"><x:f t="array" ref="C4:D5">SEQUENCE(2,2)</x:f><x:v>1</x:v></x:c><x:c r="D4"><x:v>2</x:v></x:c>'
    '<x:c r="E4" t="d"><x:v>2024-01-31T00:00:00</x:v></x:c></x:row>'
    '<x:row r="5"><x:c r="C5"><x:v>3</x:v></x:c><x:c r="D5"><x:f ca="1"/><x:v>4</x:v></x:c><x:c r="E5" t="s"><x:v>0</x:v></x:c></x:row>'
    '</x:sheetData>'
    '<x:sheetProtection sheet="1" objects="1"/>'
    '<x:mergeCells count="2"><x:mergeCell ref="$A$6:$C$6"/><x:mergeCell ref="E6:E6"/></x:mergeCells>'
    '<x:conditionalFormatting sqref="A1:E5"><x:cfRule type="dataBar" priority="2"><x:dataBar showValue="0">'
    '<x:cfvo type="min"/><x:cfvo type="max"/><x:color rgb="FF638EC6"/></x:dataBar>'
    '<x:extLst><x:ext uri="{B025F937-C7B1-47D3-B67F-A62EFF666E3E}"><x14:id xmlns:x14="http://schemas.microsoft.com/office/spreadsheetml/2009/9/main">{ID-1}</x14:id></x:ext></x:extLst>'
    '</x:cfRule></x:conditionalFormatting>'
    '<x:dataValidations count="1"><x:dataValidation type="whole" operator="between" sqref="A1 B2:B3"><x:formula1>1</x:formula1><x:formula2>9</x:formula2></x:dataValidation></x:dataValidations>'
    '<x:hyperlinks><x:hyperlink ref="A1" location="\'Other\'!A1" display="go"/></x:hyperlinks>'
    '<x:ignoredErrors><x:ignoredError sqref="A1:A3" numberStoredAsText="1"/></x:ignoredErrors>'
    '<x:drawing r:id="rId1"/><x:legacyDrawing r:id="rId2"/>'
    '<x:extLst><x:ext uri="{78C0D931-6437-407d-A8EE-F0AAD7539E65}" xmlns:x14="http://schemas.microsoft.com/office/spreadsheetml/2009/9/main">'
    '<x14:conditionalFormattings><x14:conditionalFormatting xmlns:xm="http://schemas.microsoft.com/office/excel/2006/main">'
    '<x14:cfRule type="expression" priority="1" id="{ID-2}"><xm:f>$A$1=1</xm:f><x14:dxf><fill xmlns="%s"><patternFill><bgColor rgb="FFFFFF00"/></patternFill></fill></x14:dxf></x14:cfRule>'
    '<xm:sqref>A2:A9</xm:sqref></x14:conditionalFormatting></x14:conditionalFormattings></x:ext>'
    '<x:ext uri="{CCE6A557-97BC-4b89-ADB6-D9C93CAAB3DF}" xmlns:x14="http://schemas.microsoft.com/office/spreadsheetml/2009/9/main">'
    '<x14:dataValidations count="1" xmlns:xm="http://schemas.microsoft.com/office/excel/2006/main"><x14:dataValidation type="list" allowBlank="1">'
    '<x14:formula1><xm:f>Other!$A$1:$A$3</xm:f></x14:formula1><xm:sqref>D2:D4</xm:sqref></x14:dataValidation></x14:dataValidations></x:ext>'
    '</x:extLst></x:worksheet>') % (MAIN, REL, MAIN)

DRAWING = ('<xdr:wsDr xmlns:xdr="http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing" '
           'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" xmlns:r="%s">'
           '<xdr:twoCellAnchor><xdr:sp textlink="$A$1"/><xdr:graphicFrame><a:graphic><a:graphicData>'
           '<c:chart xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart" r:id="rId1"/>'
           '</a:graphicData></a:graphic></xdr:graphicFrame></xdr:twoCellAnchor></xdr:wsDr>') % REL
CHART = ('<c:chartSpace xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart"><c:chart><c:plotArea>'
         '<c:lineChart><c:ser><c:val><c:numRef><c:f>Data!$B$2:$B$9</c:f></c:numRef></c:val></c:ser></c:lineChart>'
         '</c:plotArea></c:chart></c:chartSpace>')
EXTLINK = ('<externalLink xmlns="%s" xmlns:r="%s"><externalBook r:id="rId1"><sheetNames><sheetName val="Prices"/>'
           '</sheetNames><sheetDataSet><sheetData sheetId="0"/></sheetDataSet></externalBook></externalLink>') % (MAIN, REL)
COMMENTS = ('<comments xmlns="%s"><authors><author>Ann</author></authors><commentList>'
            '<comment ref="B2" authorId="0"><text><r><t>Note </t></r><r><t>text</t></r></text></comment>'
            '</commentList></comments>') % MAIN


def test_prefixed_package():
    p = build(tmp("prefixed.xlsx"), [("Data", None, PREFIXED_SHEET),
                                     ("Other", None, f'<?xml version="1.0"?><x:worksheet xmlns:x="{MAIN}"><x:sheetData/></x:worksheet>')],
              raw_sheets=True, styles=STYLES_XML, sst=['<si><t>shared</t></si>'],
              app="<Application>Microsoft Excel</Application><AppVersion>16.0300</AppVersion>",
              wb_pre='<fileVersion appName="xl" lastEdited="7" rupBuild="27328"/>',
              wb_extra='<definedNames><definedName name="_xlnm.Print_Area" localSheetId="0">Data!$A$1:$E$6</definedName></definedNames>'
                       '<calcPr calcId="191029"/><externalReferences><externalReference r:id="rIdX"/></externalReferences>',
              extra_wb_rels=f'<Relationship Id="rIdX" Type="{REL}/externalLink" Target="externalLinks/externalLink1.xml"/>',
              sheet_rels={0: f'<Relationship Id="rId1" Type="{REL}/drawing" Target="../drawings/drawing1.xml"/>'
                             f'<Relationship Id="rId2" Type="{REL}/vmlDrawing" Target="../drawings/vmlDrawing1.vml"/>'
                             f'<Relationship Id="rId3" Type="{REL}/comments" Target="../comments1.xml"/>'},
              extra_parts={"xl/drawings/drawing1.xml": DRAWING, "xl/comments1.xml": COMMENTS,
                           "xl/drawings/_rels/drawing1.xml.rels":
                               f'<Relationships xmlns="{PKG_REL}"><Relationship Id="rId1" Type="{REL}/chart" Target="../charts/chart1.xml"/></Relationships>',
                           "xl/charts/chart1.xml": CHART, "xl/externalLinks/externalLink1.xml": EXTLINK,
                           "xl/externalLinks/_rels/externalLink1.xml.rels":
                               f'<Relationships xmlns="{PKG_REL}"><Relationship Id="rId1" Type="{REL}/externalLinkPath" '
                               f'Target="file:///C:/data/prices.xlsx" TargetMode="External"/></Relationships>'})
    pkg = Package.open(p)
    assert pkg.is_spreadsheetml, pkg.unsupported_reason
    assert detect_writer(pkg)[0] == "excel"
    d = pkg.defined_names[0]
    assert d.builtin and d.scope == "Data"
    ext = pkg.external_links[0]
    assert ext.index == 1 and ext.kind == "externalBook" and ext.target == "file:///C:/data/prices.xlsx"
    assert ext.target_external and ext.sheet_names == ["Prices"] and ext.has_cached_data
    assert pkg.charts[0].formulas == ["Data!$B$2:$B$9"] and pkg.charts[0].sheet == "Data"
    assert pkg.drawings[0].textlinks == ["$A$1"]
    com = pkg.comments(pkg.sheet("Data"))
    assert com[0].ref == "B2" and com[0].text == "Note text" and com[0].author == "Ann"
    head, rows, cells, tail = load_sheet(pkg, "Data")
    c = cells_by_ref(cells)
    # rows / cells without r
    assert [r.r for r in rows] == [1, 3, 4, 5] and "A1" in c and "B1" in c and "C3" in c
    assert c["A1"].value == "Hello" and c["A1"].t == "inlineStr"
    runs = c["A1"].rich_runs
    assert [r.text for r in runs] == ["Hel", "lo"] and runs[1].font.b and runs[1].font.color.rgb == "FFFFFFFF"
    assert c["C1"].value == 2.0
    # shared formulas expanded
    assert c["A3"].formula.is_shared_master and c["A3"].formula_text == "A1+B$1*2"
    assert c["B3"].formula.is_shared_child and c["B3"].formula_text == "B1+C$1*2"
    assert c["C3"].formula_text == "C1+D$1*2" and c["C3"].formula.master_ref == "A3"
    assert c["E3"].value is True and rows[1].style == 5 and rows[1].ht == 30
    assert isinstance(c["A4"].value, ExcelError) and c["A4"].value == "#DIV/0!"
    assert c["B4"].value == "ab" and c["B4"].formula_text == '"a"&"b"'
    # array / spill: anchor + members
    assert c["C4"].formula.kind == "array" and c["C4"].cm == 1 and c["C4"].array is None
    assert c["D4"].array is not None and c["D4"].array.anchor_ref == "C4" and not c["D4"].has_formula
    assert c["C5"].array.anchor_ref == "C4" and c["C5"].is_formula_result
    assert c["D5"].formula.is_empty_marker and c["D5"].array.anchor_ref == "C4" and not c["D5"].has_formula
    assert c["E4"].value == "2024-01-31T00:00:00" and c["E5"].value == "shared"
    # head
    assert head.sheet_pr.tab_color.rgb == "FFFF0000" and not head.sheet_pr.summary_below
    assert head.dimension == "A1:E6" and head.format.zero_height and head.format.base_col_width == 10
    v = head.view
    assert v.zoom == 90 and v.tab_selected and v.active_cell == "C4"     # active pane, not the first selection
    assert v.selections[0].active_cell == "B1"
    assert head.col_style(3) == 3 and head.col_style(5) is None and head.col_info(4).width == 12
    # tail
    assert [m.ref for m in tail.merges] == ["A6:C6", "E6:E6"] and tail.merges[1].is_single_cell
    assert tail.is_protected and tail.drawing_rid == "rId1" and tail.legacy_drawing_rid == "rId2"
    main_cf = [f for f in tail.conditional_formats if f.source == "main"][0]
    r0 = main_cf.rules[0]
    assert r0.type == "dataBar" and r0.show_value is False and r0.id == "{ID-1}" and r0.colors[0].rgb == "FF638EC6"
    x14 = [f for f in tail.conditional_formats if f.source == "x14"][0]
    assert x14.sqref == "A2:A9" and x14.rules[0].formulas == ["$A$1=1"]
    assert pkg.styles.fill_paint(x14.rules[0].dxf.fill, dxf=True).effective == "FFFFFF00"
    dvs = {d.source: d for d in tail.data_validations}
    assert dvs["main"].formula1 == "1" and dvs["main"].formula2 == "9" and len(dvs["main"].ranges) == 2
    assert dvs["x14"].type == "list" and dvs["x14"].formula1 == "Other!$A$1:$A$3" and dvs["x14"].sqref == "D2:D4"
    assert tail.hyperlinks[0].location == "'Other'!A1" and tail.hyperlinks[0].target is None
    assert tail.ignored_errors[0].flags == {"numberStoredAsText": True}
    # skeleton mode gives the same tail
    with SheetStream(pkg, pkg.sheet("Data")) as ss:
        ss.read_head()
        t2 = ss.skip_body()
    assert [m.ref for m in t2.merges] == ["A6:C6", "E6:E6"] and len(t2.conditional_formats) == 2
    assert len(t2.data_validations) == 2 and t2.is_protected
    # self-closing <sheetData/>
    h3, rows3, cells3, t3 = load_sheet(pkg, "Other")
    assert cells3 == [] and h3.has_sheet_data
    pkg.close()


STRICT_SHEET = ('<sheetViews><sheetView workbookViewId="0"/></sheetViews><sheetData><row r="1">'
                '<c r="A1" t="s"><v>0</v></c><c r="B1"><f>A2*2</f><v>4</v></c></row>'
                '<row r="2"><c r="A2"><v>2</v></c></row></sheetData><mergeCells count="1"><mergeCell ref="A5:B5"/></mergeCells>')


def test_strict_ooxml():
    p = build(tmp("strict.xlsx"), [("S1", "hidden", STRICT_SHEET), ("S2", None, "<sheetData/>")],
              main_ns=S_MAIN, rel_ns=S_REL, rel_type_base="http://purl.oclc.org/ooxml/officeDocument/relationships",
              sst=["<si><t>strict text</t></si>"], styles='<cellXfs count="1"><xf/></cellXfs>',
              app="<Application>Microsoft Excel</Application>", wb_pre='<fileVersion appName="xl"/>')
    pkg = Package.open(p)
    assert pkg.is_spreadsheetml and pkg.conformance == "strict"
    assert [(s.name, s.state, s.kind) for s in pkg.sheets] == [("S1", "hidden", "worksheet"), ("S2", "visible", "worksheet")]
    head, rows, cells, tail = load_sheet(pkg, "S1")
    c = cells_by_ref(cells)
    assert c["A1"].value == "strict text" and c["B1"].formula_text == "A2*2" and c["B1"].value == 4.0
    assert c["B1"].value_trusted
    assert tail.merges[0].ref == "A5:B5"
    v = grade(p, checks=[92, 74])
    assert v["Potential Dangers/No hidden sheets"]["decision"] == "fail"
    assert v["Formatting/No merged cells"]["decision"] == "fail"
    pkg.close()


def test_formats():
    # legacy .xls: OLE2 magic, not a zip
    xls = tmp("legacy.xls")
    with open(xls, "wb") as fh:
        fh.write(bytes.fromhex("d0cf11e0a1b11ae1") + b"\0" * 1024)
    pkg = Package.open(xls)
    assert pkg.container == "ole2" and pkg.format == "xls" and not pkg.is_spreadsheetml
    # .xlsb: zip with xl/workbook.bin
    xlsb = tmp("bin.xlsb")
    with zipfile.ZipFile(xlsb, "w") as z:
        z.writestr("[Content_Types].xml", '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
                   '<Override PartName="/xl/workbook.bin" ContentType="application/vnd.ms-excel.sheet.binary.macroEnabled.main"/></Types>')
        z.writestr("_rels/.rels", f'<Relationships xmlns="{PKG_REL}"><Relationship Id="rId1" Type="{REL}/officeDocument" Target="xl/workbook.bin"/></Relationships>')
        z.writestr("xl/workbook.bin", b"\x83\x01\x00")
    pkg = Package.open(xlsb)
    assert pkg.format == "xlsb" and not pkg.is_spreadsheetml
    # .xlsm with and without a VBA project
    vba_rel = f'<Relationship Id="rIdV" Type="http://schemas.microsoft.com/office/2006/relationships/vbaProject" Target="vbaProject.bin"/>'
    with_vba = build(tmp("m1.xlsm"), [("A", None, "<sheetData/>")], main_ct=CT_XLSM, extra_wb_rels=vba_rel,
                     extra_parts={"xl/vbaProject.bin": b"\xd0\xcf\x11\xe0"})
    no_vba = build(tmp("m2.xlsm"), [("A", None, "<sheetData/>")], main_ct=CT_XLSM)
    p1, p2 = Package.open(with_vba), Package.open(no_vba)
    assert p1.format == "xlsm" and p1.has_vba and p1.is_spreadsheetml and p1.vba_parts == ["xl/vbaProject.bin"]
    assert p2.format == "xlsm" and not p2.has_vba and not p2.vba_file_present
    # upper-case extension
    up = build(tmp("UPPER.XLSX"), [("A", None, "<sheetData/>")])
    pu = Package.open(up)
    assert pu.extension == ".xlsx" and pu.format == "xlsx" and pu.is_spreadsheetml
    # corrupt zip / csv / empty
    bad = tmp("bad.xlsx")
    with open(bad, "wb") as fh:
        fh.write(b"PK\x03\x04garbage")
    assert Package.open(bad).format == "corrupt"
    csv = tmp("x.csv")
    with open(csv, "w") as fh:
        fh.write("a,b\n1,2\n")
    assert Package.open(csv).format == "csv"
    # a malformed workbook.xml makes the package unparseable, not a crash
    broken = build(tmp("broken.xlsx"), [("A", None, "<sheetData/>")])
    with zipfile.ZipFile(broken, "a") as z:
        z.writestr("xl/workbook2.xml", "x")
    with zipfile.ZipFile(broken) as z:
        parts = {n: z.read(n) for n in z.namelist()}
    parts["xl/workbook.xml"] = b"<workbook><sheets>"
    with zipfile.ZipFile(broken, "w") as z:
        for k, v in parts.items():
            z.writestr(k, v)
    pb = Package.open(broken)
    assert not pb.is_spreadsheetml and "parse error" in pb.unsupported_reason
    # missing file -> GradingError
    try:
        Package.open(tmp("nope.xlsx"))
        raise AssertionError("expected GradingError")
    except GradingError:
        pass
    # non-xlsx files: 92 / 74 raise GradingError (only 77 may grade them)
    for path in (xls, xlsb, bad, csv):
        try:
            grade(path, checks=[92, 74])
            raise AssertionError("expected GradingError for " + path)
        except GradingError as e:
            assert set(e.failures) == {REGISTRY[92].key, REGISTRY[74].key}, e.failures


def test_writer_detection():
    def wb(app=None, pre="", extra="", core=None):
        return Package.open(build(tmp("w.xlsx"), [("A", None, "<sheetData/>")], app=app, wb_pre=pre, wb_extra=extra, core=core))
    assert detect_writer(wb("<Application>Microsoft Excel</Application><AppVersion>16.0300</AppVersion>",
                            '<fileVersion appName="xl" lastEdited="7" rupBuild="27328"/>'))[0] == "excel"
    assert detect_writer(wb("<Application>Microsoft Excel Online</Application><AppVersion>16.0300</AppVersion>",
                            '<fileVersion appName="xl"/>'))[0] == "excel"
    assert detect_writer(wb("<Application>Microsoft Macintosh Excel</Application><AppVersion>16.0300</AppVersion>",
                            '<fileVersion appName="xl"/>'))[0] == "excel"
    assert detect_writer(wb("<Application>Microsoft Excel</Application><AppVersion>12.0000</AppVersion>",
                            '<fileVersion appName="xl" lastEdited="4" lowestEdited="4" rupBuild="4505"/>',
                            '<calcPr calcId="124519" fullCalcOnLoad="1"/>'))[0] == "xlsxwriter"
    assert detect_writer(wb("<Application>Microsoft Excel Compatible / Openpyxl 3.1.5</Application>"))[0] == "openpyxl"
    assert detect_writer(wb("<Application>LibreOffice/7.4.7.2$Linux_AARCH64</Application>", '<fileVersion appName="Calc"/>'))[0] == "libreoffice"
    assert detect_writer(wb(None, "", '<calcPr calcMode="auto" fullCalcOnLoad="1" forceFullCalc="1"/>'))[0] == "unknown"
    assert detect_writer(wb("<Application>Microsoft Excel</Application>", "", '<calcPr calcId="124519" fullCalcOnLoad="1"/>',
                            "<dc:creator>openpyxl</dc:creator>"))[0] == "openpyxl"


VAL_SHEET = ('<sheetData><row r="1"><c r="A1"><v>2</v></c><c r="B1"><f>A1*3</f><v>0</v></c>'
             '<c r="C1"><f>NOSUCH(A1)</f><v>1</v></c><c r="D1"><f>A1</f></c></row></sheetData>')
RECALC_SHEET = ('<sheetData><row r="1"><c r="A1"><v>2</v></c><c r="B1"><f>A1*3</f><v>6</v></c>'
                '<c r="C1" t="e"><f>NOSUCH(A1)</f><v>#NAME?</v></c></row></sheetData>')


class _ValueCheck(Check):
    number = 9001
    key = "Test/values"
    needs_cells = True
    needs_values = True

    def start(self, wb):
        super().start(wb)
        self.seen = {}

    def cell(self, cell):
        v = cell.value if cell.value_trusted else ("untrusted", cell.unverified_value)
        self.seen[cell.ref] = (v, cell.value_trusted, cell.value_source)

    def finish(self):
        return self.verdict("ok", "bad")


class _StrictValueCheck(_ValueCheck):
    def cell(self, cell):
        self.require_value(cell)


def _run(path, chk, value_path=None):
    return Engine(path, [chk], value_path=value_path).run()


def test_value_provenance_and_recalc():
    xw = build(tmp("xw.xlsx"), [("S", None, VAL_SHEET)],
               app="<Application>Microsoft Excel</Application><AppVersion>12.0000</AppVersion>",
               wb_pre='<fileVersion appName="xl" lastEdited="4" lowestEdited="4" rupBuild="4505"/>',
               wb_extra='<calcPr calcId="124519" fullCalcOnLoad="1"/>')
    lo = build(tmp("lo.xlsx"), [("S", None, RECALC_SHEET)],
               app="<Application>LibreOffice/7.4.7.2$Linux</Application>", wb_pre='<fileVersion appName="Calc"/>')
    chk = _ValueCheck()
    _run(xw, chk)
    assert chk.seen["A1"] == (2.0, True, "constant")
    assert chk.seen["B1"] == (("untrusted", 0.0), False, "cached")    # XlsxWriter all-zero cache: untrusted
    assert chk.seen["D1"] == (("untrusted", None), False, "cached")   # no cache at all
    chk = _ValueCheck()
    _run(xw, chk, value_path=lo)
    assert chk.seen["B1"] == (6.0, True, "recalc")
    assert chk.seen["C1"] == (("untrusted", "#NAME?"), False, "recalc")   # LibreOffice #NAME? = unmeasured
    assert chk.seen["D1"] == (("untrusted", None), False, "recalc")       # missing in the copy
    assert chk.seen["A1"] == (2.0, True, "constant")
    # require_value raises -> GradingError naming the check (no fallback)
    try:
        _run(xw, _StrictValueCheck())
        raise AssertionError("expected GradingError")
    except GradingError as e:
        assert "Test/values" in e.failures and "untrusted" in e.failures["Test/values"]
    # LibreOffice-written file: #NAME? cached is unmeasured, numbers trusted
    chk = _ValueCheck()
    _run(lo, chk)
    assert chk.seen["B1"] == (6.0, True, "cached") and chk.seen["C1"] == (("untrusted", "#NAME?"), False, "cached")



def test_values_policy_ruling_2026_10_03():
    """Ruling 2026-10-03 (A5): agent-tool caches are never trusted (an openpyxl-labelled file's
    numeric cache stays untrusted even when it looks right); production supplies a LibreOffice
    recalculation copy for every file not saved by Excel, and formula values then come from it."""
    opx = build(tmp("opx_val.xlsx"), [("S", None, RECALC_SHEET.replace('t="e"', "").replace("#NAME?", "1"))],
                app="<Application>Microsoft Excel Compatible / Openpyxl 3.1.5</Application>")
    lo = build(tmp("lo_val.xlsx"), [("S", None, RECALC_SHEET)],
               app="<Application>LibreOffice/7.4.7.2$Linux</Application>", wb_pre='<fileVersion appName="Calc"/>')
    chk = _ValueCheck()
    _run(opx, chk)
    assert chk.seen["B1"] == (("untrusted", 6.0), False, "cached"), chk.seen           # a right-looking agent cache: untrusted
    chk = _ValueCheck()
    _run(opx, chk, value_path=lo)
    assert chk.seen["B1"] == (6.0, True, "recalc"), chk.seen                            # the LibreOffice copy decides
    assert chk.seen["C1"][1] is False                                                   # its #NAME? stays a loud error

class _Boom(Check):
    number = 9002
    key = "Test/boom"
    needs_cells = True

    def cell(self, cell):
        raise ZeroDivisionError("kaboom")

    def finish(self):
        return self.verdict("ok", "bad")


class _BadVerdict(Check):
    number = 9003
    key = "Test/badverdict"

    def finish(self):
        return {"engine": "harness", "decision": "pass", "summary": "x", "stats": {},
                "mistakes": [{"location": "A1", "description": "d", "severity": "major"}]}


class _SecondPass(Check):
    number = 9004
    key = "Test/second"
    needs_cells = True

    def start(self, wb):
        super().start(wb)
        self.got = []
        self.first_cells = 0

    def cell(self, cell):
        self.first_cells += 1

    def sheet_end(self, head, tail):
        if head.name == "B":
            self.request_second_pass("A", [(1, 2)])

    def second_pass_cell(self, cell):
        self.got.append(cell.ref)

    def finish(self):
        return self.verdict("ok", "bad")


class _OptOut(Check):
    number = 9005
    key = "Test/optout"
    needs_cells = True

    def start(self, wb):
        super().start(wb)
        self.sheets = []

    def sheet_start(self, head):
        return head.name != "B"

    def cell(self, cell):
        self.sheets.append(cell.sheet)

    def finish(self):
        return self.verdict("ok", "bad")


def test_engine_contract():
    p = build(tmp("eng.xlsx"), [("A", None, VAL_SHEET), ("B", None, VAL_SHEET)])
    try:
        Engine(p, [_Boom(), REGISTRY[92]()]).run()
        raise AssertionError("expected GradingError")
    except GradingError as e:
        assert list(e.failures) == ["Test/boom"] and "ZeroDivisionError" in e.failures["Test/boom"]
        assert "Potential Dangers/No hidden sheets" in e.verdicts       # others finished, call still fails
    try:
        Engine(p, [_BadVerdict()]).run()
        raise AssertionError("expected GradingError")
    except GradingError as e:
        assert "invariant" in e.failures["Test/badverdict"]
    sp = _SecondPass()
    Engine(p, [sp]).run()
    assert sp.got == ["B1"] and sp.first_cells == 8
    oo = _OptOut()
    Engine(p, [oo]).run()
    assert set(oo.sheets) == {"A"}
    # unknown check number
    try:
        grade(p, checks=[1])
        raise AssertionError("expected GradingError")
    except GradingError:
        pass
    # verdict helpers
    m = Mistakes(cap=2)
    for i in range(5):
        m.add(f"A{i}", "d")
    v = make_verdict(m, "5 bad")
    assert v["decision"] == "fail" and len(v["mistakes"]) == 2 and v["stats"]["n_mistakes"] == 5
    assert v["stats"]["mistakes_truncated"]
    validate_verdict(make_verdict(Mistakes(), "fine"))
    chk = _OptOut()
    chk.add_cell_mistakes("My Sheet", [(1, 1), (1, 2), (3, 3)], "bad {range} ({n})")
    assert [x["location"] for x in chk.mistakes.items] == ["'My Sheet'!A1:B1", "'My Sheet'!C3"]
    assert chk.mistakes.items[0]["description"] == "bad A1:B1 (2)"


def test_utf16_sheet():
    xml = ('<?xml version="1.0" encoding="UTF-16"?><worksheet xmlns="%s"><sheetData><row r="1"><c r="A1" t="inlineStr">'
           '<is><t>ü16</t></is></c></row></sheetData><mergeCells><mergeCell ref="A1:B1"/></mergeCells></worksheet>') % MAIN
    p = tmp("u16.xlsx")
    build(p, [("U", None, "")], raw_sheets=True)
    with zipfile.ZipFile(p) as z:
        parts = {n: z.read(n) for n in z.namelist()}
    parts["xl/worksheets/sheet1.xml"] = xml.encode("utf-16")
    with zipfile.ZipFile(p, "w") as z:
        for k, v in parts.items():
            z.writestr(k, v)
    pkg = Package.open(p)
    head, rows, cells, tail = load_sheet(pkg, "U")
    assert cells[0].value == "ü16" and tail.merges[0].ref == "A1:B1"


def test_reference_checks():
    p = _openpyxl_book(tmp("op2.xlsx"))
    v = grade(p, checks=[92, 74])
    v92 = v["Potential Dangers/No hidden sheets"]
    assert v92["decision"] == "fail" and v92["stats"]["n_hidden"] == 1 and v92["stats"]["n_very_hidden"] == 1
    assert {m["location"] for m in v92["mistakes"]} == {"Lists", "Secret"}
    v74 = v["Formatting/No merged cells"]
    assert v74["decision"] == "fail" and [m["location"] for m in v74["mistakes"]] == ["Model!D1:F1"]
    assert v74["stats"]["instructions_exempt"][0]["sheet"] == "Instructions"
    # 1x1 merges ignored, Instructions exempt case-insensitively, hidden sheets count
    q = build(tmp("m.xlsx"), [("instructions ", None, '<sheetData/><mergeCells><mergeCell ref="A1:C1"/></mergeCells>'),
                              ("Calc", "hidden", '<sheetData/><mergeCells><mergeCell ref="B2:B2"/></mergeCells>'),
                              ("H", "veryHidden", '<sheetData/><mergeCells><mergeCell ref="A1:A2"/></mergeCells>')])
    v = grade(q, checks=[74])["Formatting/No merged cells"]
    assert v["decision"] == "fail" and [m["location"] for m in v["mistakes"]] == ["H!A1:A2"]
    assert v["stats"]["n_single_cell_merges_ignored"] == 1
    clean = build(tmp("clean.xlsx"), [("A", None, "<sheetData/>")])
    v = grade(clean, checks=[92, 74])
    assert all(x["decision"] == "pass" and x["mistakes"] == [] for x in v.values())
    # a sheet whose relationship is broken is not silently skipped: 74 raises, 92 (workbook-level) still grades
    broken = build(tmp("brokenrel.xlsx"), [("A", None, "<sheetData/>"), ("B", None, "<sheetData/>")])
    with zipfile.ZipFile(broken) as z:
        parts = {n: z.read(n) for n in z.namelist()}
    parts["xl/_rels/workbook.xml.rels"] = parts["xl/_rels/workbook.xml.rels"].replace(b'Id="rId2"', b'Id="rIdZZ"')
    with zipfile.ZipFile(broken, "w") as z:
        for k, val in parts.items():
            z.writestr(k, val)
    pk = Package.open(broken)
    assert pk.sheets[1].kind == "unknown" and pk.sheets[1].part is None
    try:
        grade(broken, checks=[92, 74])
        raise AssertionError("expected GradingError")
    except GradingError as e:
        assert list(e.failures) == ["Formatting/No merged cells"] and "missing" in e.failures["Formatting/No merged cells"]
        assert "Potential Dangers/No hidden sheets" in e.verdicts


def test_rubric_keys():
    if not os.path.exists(RUBRIC):
        return
    with open(RUBRIC) as fh:
        rub = json.load(fh)
    keys = [f"{cat}/{it['name']}" for cat, items in rub.items() for it in items]
    for n, cls in REGISTRY.items():
        assert keys[n - 1] == cls.key, (n, cls.key, keys[n - 1])


# ============================================================================ review fixes (R1-R11)
K92, K74 = "Potential Dangers/No hidden sheets", "Formatting/No merged cells"
WORD_CT = "application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"


def test_r1_not_a_workbook():
    # (a) a Word package renamed .xlsx is not a workbook: every check raises
    p = tmp("word_renamed.xlsx")
    with zipfile.ZipFile(p, "w") as z:
        z.writestr("[Content_Types].xml", '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
                   '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
                   f'<Default Extension="xml" ContentType="application/xml"/><Override PartName="/word/document.xml" ContentType="{WORD_CT}"/></Types>')
        z.writestr("_rels/.rels", f'<Relationships xmlns="{PKG_REL}"><Relationship Id="rId1" Type="{REL}/officeDocument" Target="word/document.xml"/></Relationships>')
        z.writestr("word/document.xml", '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body/></w:document>')
    pk = Package.open(p)
    assert pk.format == "unknown" and not pk.is_spreadsheetml and "wordprocessingml" in pk.unsupported_reason
    e = _raises(lambda: grade(p, checks=[92, 74]), "not a SpreadsheetML workbook")
    assert set(e.failures) == {K92, K74}
    # (b) no <sheet> at all: Excel cannot open it, so nothing is graded as an empty workbook
    q = build(tmp("nosheets.xlsx"), [])
    assert not Package.open(q).is_spreadsheetml
    e = _raises(lambda: grade(q, checks=[92, 74]), "lists no sheets")
    assert set(e.failures) == {K92, K74}
    # (c) generic content type: the main part's root element decides
    r = build(tmp("genericct.xlsx"), [("A", None, "<sheetData/>")], main_ct="application/xml")
    assert Package.open(r).is_spreadsheetml and grade(r, checks=[92])[K92]["decision"] == "pass"
    _rewrite(r, **{"xl/workbook.xml": f'<document xmlns="{MAIN}"><sheets/></document>'})
    pr = Package.open(r)
    assert not pr.is_spreadsheetml and "not a SpreadsheetML <workbook>" in pr.unsupported_reason
    _raises(lambda: grade(r, checks=[92]), "not a SpreadsheetML <workbook>")


class _CaseSecond(Check):
    number = 9006
    key = "Test/second pass case"
    needs_cells = True

    def __init__(self, names):
        super().__init__()
        self.names = names
        self.got = []

    def sheet_end(self, head, tail):
        for n, cells in self.names:
            self.request_second_pass(n, cells)

    def second_pass_cell(self, cell):
        self.got.append((cell.sheet, cell.ref))

    def finish(self):
        return self.verdict("ok", "bad")


def test_r2_second_pass_sheet_names():
    p = build(tmp("second.xlsx"), [("Model", None, '<sheetData><row r="1"><c r="A1"><v>1</v></c><c r="B1"><v>2</v></c></row></sheetData>')])
    chk = _CaseSecond([("model", None)])          # differs from 'Model' only in case
    Engine(p, [chk]).run()
    assert chk.got == [("Model", "A1"), ("Model", "B1")]
    chk = _CaseSecond([("model", [(1, 1)]), ("MODEL", [(1, 2)])])   # two spellings, one sheet: merged
    Engine(p, [chk]).run()
    assert chk.got == [("Model", "A1"), ("Model", "B1")]
    _raises(lambda: Engine(p, [_CaseSecond([("Nope", None)])]).run(), "unknown sheet 'Nope'")


SHARED_BATTERY = ["SUM(B5:INDEX(B5:B20,C2))", "XFD1", "A1048576", "SUM(A:A)", "SUM($1:$1)", "'My Sheet'!A1:B2",
                  "A1:INDEX(A:A,5)", "Table1[@Col]", '"A1"&A1', "_xlfn.LET(_xlpm.x,A1,_xlpm.x+1)", "$A$1:A1",
                  "[1]Sheet1!A1", "SUM(Sheet1:Sheet3!A1)", "_xlfn._TRO_TRAILING(A:A)", "'O''Brien'!A1", "Rate*A1",
                  "a1+b$1", "A1#", "SUM(XFC1:XFD1)"]


def test_r3_shared_formula_expansion():
    from detchecks.core import formula as fm
    L = refs.index_to_col
    cells = {4: [], 5: [], 6: [], 8: []}
    for i, f in enumerate(SHARED_BATTERY):
        c0 = 4 + 3 * i
        esc = f.replace("&", "&amp;").replace("<", "&lt;").replace('"', "&quot;")
        cells[5].append(f'<c r="{L(c0)}5"><f t="shared" ref="{L(c0)}5:{L(c0 + 2)}8" si="{i}">{esc}</f><v>0</v></c>')
        cells[5].append(f'<c r="{L(c0 + 1)}5"><f t="shared" si="{i}"/><v>0</v></c>')
        cells[6].append(f'<c r="{L(c0)}6"><f t="shared" si="{i}"/><v>0</v></c>')
        cells[8].append(f'<c r="{L(c0 + 2)}8"><f t="shared" si="{i}"/><v>0</v></c>')
    xml = "<sheetData>" + "".join(f'<row r="{r}">{"".join(c)}</row>' for r, c in cells.items() if c) + "</sheetData>"
    pkg = Package.open(build(tmp("shared.xlsx"), [("S", None, xml)]))
    by = cells_by_ref(load_sheet(pkg, "S")[2])
    for i, f in enumerate(SHARED_BATTERY):
        c0 = 4 + 3 * i
        master = f"{L(c0)}5"
        for ref in (f"{L(c0 + 1)}5", f"{L(c0)}6", f"{L(c0 + 2)}8"):
            assert by[ref].formula_text == fm.translate(f, master, ref), (f, ref, by[ref].formula_text)
    # the two Translator errors of the review, at their exact positions
    assert by["D6"].formula_text == "SUM(B6:INDEX(B6:B21,C3))"          # left end of 'B5:INDEX(' shifted
    assert by["H5"].formula_text == "#REF!"                              # XFD1 moved right
    assert by["J6"].formula_text == "#REF!"                              # A1048576 moved down
    pkg.close()


class _Lazy(Check):
    """needs_values=True, but reads cell.value without testing value_trusted."""
    number = 9007
    key = "Test/lazy values"
    needs_cells = True
    needs_values = True

    def cell(self, cell):
        cell.value

    def finish(self):
        return self.verdict("ok", "bad")


class _NoFlag(_Lazy):
    """Forgets needs_values but reads formula values."""
    key = "Test/no needs_values"
    needs_values = False


class _Constants(_Lazy):
    """No needs_values: constants only, which is allowed."""
    key = "Test/constants"
    needs_values = False

    def cell(self, cell):
        if not cell.is_formula_result:
            cell.value


def test_r4_values_no_fallback():
    xw = build(tmp("xw4.xlsx"), [("S", None, VAL_SHEET)],
               app="<Application>Microsoft Excel</Application><AppVersion>12.0000</AppVersion>",
               wb_pre='<fileVersion appName="xl" lastEdited="4" lowestEdited="4" rupBuild="4505"/>',
               wb_extra='<calcPr calcId="124519" fullCalcOnLoad="1"/>')
    lo = build(tmp("lo4.xlsx"), [("S", None, RECALC_SHEET)],
               app="<Application>LibreOffice/7.4.7.2$Linux</Application>", wb_pre='<fileVersion appName="Calc"/>')
    # a check that reads .value directly cannot use XlsxWriter's placeholder zero
    _raises(lambda: Engine(xw, [_Lazy()]).run(), "S!B1", "untrusted", "writer=xlsxwriter")
    # with a recalc copy, the trusted recalc value is used; the copy's #NAME? still raises
    _raises(lambda: Engine(xw, [_Lazy()], value_path=lo).run(), "S!C1", "untrusted", "source=recalc")
    # a check that forgot needs_values cannot read formula values, trusted or not, copy or not
    _raises(lambda: Engine(lo, [_NoFlag()]).run(), "S!B1", "needs_values")
    _raises(lambda: Engine(xw, [_NoFlag()], value_path=lo).run(), "S!B1", "needs_values")
    # constants never need the flag
    assert Engine(xw, [_Constants()]).run()["Test/constants"]["decision"] == "pass"


def test_r5_writer_evidence():
    excel_app = "<Application>Microsoft Excel</Application><AppVersion>16.0300</AppVersion>"
    xl_fv = '<fileVersion appName="xl" lastEdited="7" lowestEdited="7" rupBuild="29822"/>'

    def writer(**kw):
        sheets = kw.pop("sheets", [("A", None, "<sheetData/>")])
        pk = Package.open(build(tmp("w5.xlsx"), sheets, **kw))
        try:
            return detect_writer(pk)
        finally:
            pk.close()
    assert writer(app=excel_app, wb_pre=xl_fv)[0] == "excel"
    # 2290-style: Excel's app.xml and fileVersion copied, parts re-serialised by Python
    w, ev = writer(app="<Application>Microsoft Excel</Application>", wb_pre=xl_fv, decl=PY_DECL,
                   core="<dc:creator>GPT-6 Astra Pro</dc:creator>")
    assert w == "unknown" and "no AppVersion" in ev and "GPT-6" in ev, ev
    w, ev = writer(app=excel_app, wb_pre=xl_fv, decl=PY_DECL)
    assert w == "unknown" and "single-quoted" in ev, ev
    w, ev = writer(app=excel_app, wb_pre=xl_fv, decl="")                     # openpyxl writes no declaration
    assert w == "unknown" and "no XML declaration" in ev, ev
    w, ev = writer(app=excel_app, wb_pre="")                                 # no fileVersion
    assert w == "unknown" and "fileVersion" in ev, ev
    # one sheet part re-serialised by Python is enough
    raw = PY_DECL + f'<worksheet xmlns="{MAIN}"><sheetData/></worksheet>'
    w, ev = writer(app=excel_app, wb_pre=xl_fv, sheets=[("A", None, raw)], raw_sheets=True)
    assert w == "unknown" and "sheet 'A'" in ev, ev
    # untrusted caches stay untrusted (XlsxWriter signature still wins over the declarations)
    assert writer(app="<Application>Microsoft Excel</Application><AppVersion>12.0000</AppVersion>",
                  wb_pre='<fileVersion appName="xl" lastEdited="4" lowestEdited="4" rupBuild="4505"/>')[0] == "xlsxwriter"


def test_r6_failing_sheets_gate():
    from detchecks.tools.run_toys import failing_sheets_problem, sheet_of_location
    names = ["Overview", "Solution Model", "O'Brien", "A!B"]
    assert sheet_of_location("'Solution Model'!A4", names) == "Solution Model"
    assert sheet_of_location("Overview!A4:B5", names) == "Overview"
    assert sheet_of_location("'O''Brien'!C3", names) == "O'Brien"
    assert sheet_of_location("Solution Model", names) == "Solution Model"      # sheet-level finding
    assert sheet_of_location("A!B", names) == "A!B" and sheet_of_location("'A!B'!A1", names) == "A!B"
    try:
        sheet_of_location("'Nope'!A1", names)
        raise AssertionError("expected ValueError")
    except ValueError:
        pass

    def verdict(*locs, truncated=False):
        m = Mistakes()
        for loc in locs:
            m.add(loc, "d")
        v = make_verdict(m, "x")
        v["stats"]["mistakes_truncated"] = truncated
        return v
    assert failing_sheets_problem(verdict("Overview!A4", "'Solution Model'!F4"), ["Overview", "Solution Model"], names) is None
    # a check that fails every sheet is told apart from the right one
    prob = failing_sheets_problem(verdict("Overview!A4", "'Solution Model'!F4", "'O''Brien'!A1"), ["Overview", "Solution Model"], names)
    assert prob and "unexpected [\"O'Brien\"]" in prob, prob
    assert "missing ['Solution Model']" in failing_sheets_problem(verdict("Overview!A4"), ["Overview", "Solution Model"], names)
    assert "truncated" in failing_sheets_problem(verdict("Overview!A4", truncated=True), ["Overview"], names)
    # every 62 toy carries an expected sheet set, so 'always fail' cannot pass that gate
    with open(os.path.join(os.path.dirname(HERE), "tools", "expected.json")) as fh:
        exp62 = json.load(fh)["62"]
    assert len(exp62) == 8 and all(e.get("failing_sheets") for e in exp62.values())


X14 = "http://schemas.microsoft.com/office/spreadsheetml/2009/9/main"
XM = "http://schemas.microsoft.com/office/excel/2006/main"
EXTLST = (f'<extLst><ext uri="{{05C60535-1F16-4fd2-B633-F4F36F0B64E0}}" xmlns:x14="{X14}">'
          f'<x14:sparklineGroups xmlns:xm="{XM}"><x14:sparklineGroup type="column" dateAxis="1"><x14:colorSeries rgb="FF376092"/>'
          '<xm:f>Data!A1:E1</xm:f><x14:sparklines><x14:sparkline><xm:f>[2]Data!A2:E2</xm:f><xm:sqref>F2</xm:sqref></x14:sparkline>'
          '<x14:sparkline><xm:f>Data!$A$3:$E$3</xm:f><xm:sqref>F3</xm:sqref></x14:sparkline></x14:sparklines></x14:sparklineGroup>'
          '</x14:sparklineGroups></ext>'
          f'<ext uri="{{CCE6A557-97BC-4b89-ADB6-D9C93CAAB3DF}}" xmlns:x14="{X14}"><x14:dataValidations count="1" xmlns:xm="{XM}">'
          '<x14:dataValidation type="list"><x14:formula1><xm:f>Data!$A$1:$A$3</xm:f></x14:formula1><xm:sqref>B2</xm:sqref>'
          '</x14:dataValidation></x14:dataValidations></ext>'
          f'<ext uri="{{3A4CF648-6AED-40f4-86FF-DC5316D8AED3}}" xmlns:x14="{X14}"><x14:slicerList><x14:slicer r:id="rId9"/></x14:slicerList></ext>'
          '</extLst>')


def test_r7_sparklines():
    p = build(tmp("spark.xlsx"), [("S", None, '<sheetData><row r="1"><c r="A1"><v>1</v></c></row></sheetData>' + EXTLST)])
    pkg = Package.open(p)
    tails = [load_sheet(pkg, "S")[3]]
    with SheetStream(pkg, pkg.sheet("S")) as ss:
        ss.read_head()
        tails.append(ss.skip_body())
    for tail in tails:                                   # streamed and skipped tails agree
        g = tail.sparkline_groups[0]
        assert g.type == "column" and g.date_range == "Data!A1:E1"
        assert [(s.sqref, s.formula) for s in g.sparklines] == [("F2", "[2]Data!A2:E2"), ("F3", "Data!$A$3:$E$3")]
        assert g.formulas == ["Data!A1:E1", "[2]Data!A2:E2", "Data!$A$3:$E$3"]
        assert tail.data_validations[0].formula1 == "Data!$A$1:$A$3"
        assert tail.other == ["extLst/slicerList"]       # uninterpreted extension content is listed
    pkg.close()


def test_r8_cursor_memory():
    import tracemalloc
    from detchecks.core.values import ValueSource
    n = 30000
    body = "".join(f'<row r="{r}"><c r="A{r}"><v>{r}</v></c><c r="B{r}"><f>A{r}*2</f><v>{2 * r}</v></c></row>'
                   for r in range(1, n + 1))
    copy = build(tmp("bigcopy.xlsx"), [("S", None, f"<sheetData>{body}</sheetData>")],
                 app="<Application>LibreOffice/7.6</Application>", wb_pre='<fileVersion appName="Calc"/>')
    vs = ValueSource(copy)
    try:
        cur = vs.cursor("S")
        tracemalloc.start()
        base = tracemalloc.get_traced_memory()[0]
        for r in range(1, n + 1):
            assert cur.get(r, 2) == ("ok", 2.0 * r)
        peak = tracemalloc.get_traced_memory()[1] - base
        tracemalloc.stop()
    finally:
        vs.close()
    # finished rows are dropped from <sheetData>: no per-row residue (was ~90 bytes x 30k rows)
    assert peak < 1_000_000, peak


def test_r9_silent_skips():
    # (a) relationship type in another letter case is still a worksheet; an unknown type raises
    p = build(tmp("relcase.xlsx"), [("S", None, '<sheetData/><mergeCells><mergeCell ref="A1:B1"/></mergeCells>')])
    with zipfile.ZipFile(p) as z:
        rels = z.read("xl/_rels/workbook.xml.rels")
    _rewrite(p, **{"xl/_rels/workbook.xml.rels": rels.replace(b"relationships/worksheet", b"relationships/Worksheet")})
    assert Package.open(p).sheets[0].kind == "worksheet"
    v = grade(p, checks=[74])[K74]
    assert v["decision"] == "fail" and v["stats"]["n_sheets_checked"] == 1
    _rewrite(p, **{"xl/_rels/workbook.xml.rels": rels.replace(b"relationships/worksheet", b"relationships/fooSheet")})
    pk = Package.open(p)
    assert pk.sheets[0].kind == "unknown" and pk.sheets[0].rel_type.endswith("/fooSheet")
    e = _raises(lambda: grade(p, checks=[92, 74]), "fooSheet")
    assert list(e.failures) == [K74] and K92 in e.verdicts
    # (b) unreadable or off-grid merge refs raise instead of being dropped
    for bad in ("A1:B2:C3", "A1;B2", "A0:B1", "XFE1:XFF1", "Name"):
        q = build(tmp("badmerge.xlsx"), [("S", None, f'<sheetData/><mergeCells><mergeCell ref="{bad}"/></mergeCells>')])
        _raises(lambda: grade(q, checks=[74]), "unreadable mergeCell ref", bad)


def test_r10_prolog_before_root():
    for prolog in ("<!-- written by <tool> v1 -->", "<?mso-application progid='Excel.Sheet'?><!-- <x> -->",
                   "<!DOCTYPE worksheet>"):
        xml = (f'<?xml version="1.0"?>{prolog}<worksheet xmlns="{MAIN}" xmlns:r="{REL}"><sheetData>'
               '<row r="1"><c r="A1"><v>1</v></c></row></sheetData><mergeCells><mergeCell ref="A1:B1"/></mergeCells></worksheet>')
        p = build(tmp("prolog.xlsx"), [("S", None, xml)], raw_sheets=True)
        v = grade(p, checks=[74])[K74]                  # skip_body path
        assert v["decision"] == "fail" and v["mistakes"][0]["location"] == "S!A1:B1", (prolog, v)
        pkg = Package.open(p)
        head, rows, cells, tail = load_sheet(pkg, "S")  # streaming path
        assert [m.ref for m in tail.merges] == ["A1:B1"] and cells[0].value == 1.0
        pkg.close()
    # a '>' inside a quoted attribute of the root start tag does not end the tag
    xml = (f'<worksheet xmlns="{MAIN}" xmlns:mc="urn:x" mc:Note="a>b"><sheetData/>'
           '<mergeCells><mergeCell ref="C1:D1"/></mergeCells></worksheet>')
    p = build(tmp("gtattr.xlsx"), [("S", None, xml)], raw_sheets=True)
    assert grade(p, checks=[74])[K74]["mistakes"][0]["location"] == "S!C1:D1"


def test_r11_edge_cases():
    # a defined name whose localSheetId names no sheet is neither global nor local: .scope raises
    p = build(tmp("names.xlsx"), [("S", None, "<sheetData/>")],
              wb_extra='<definedNames><definedName name="L" localSheetId="5">S!$A$1</definedName>'
                       '<definedName name="M" localSheetId="x">S!$A$2</definedName>'
                       '<definedName name="G">S!$A$3</definedName><definedName name="K" localSheetId="0">S!$A$4</definedName></definedNames>')
    pk = Package.open(p)
    names = {d.name: d for d in pk.defined_names}
    assert names["G"].scope is None and names["K"].scope == "S"
    _raises(lambda: names["L"].scope, "localSheetId 5 does not name a sheet")
    _raises(lambda: names["M"].scope, "not an integer")
    assert grade(p, checks=[92])[K92]["decision"] == "pass"      # checks that do not use names still grade
    # grade() argument errors
    _raises(lambda: grade(p, checks=[]), "no checks selected")
    _raises(lambda: grade(p, checks=[74, 74]), "more than once")
    # a missing file names every selected check
    e = _raises(lambda: grade(tmp("missing.xlsx"), checks=[92, 74]), "workbook not found")
    assert set(e.failures) == {K92, K74} and e.check in (K92, K74) and e.path.endswith("missing.xlsx")


TESTS = [test_refs, test_excel_tint, test_styles_and_colours, test_openpyxl_workbook, test_prefixed_package,
         test_strict_ooxml, test_formats, test_writer_detection, test_value_provenance_and_recalc,
         test_engine_contract, test_utf16_sheet, test_reference_checks, test_rubric_keys,
         test_r1_not_a_workbook, test_r2_second_pass_sheet_names, test_r3_shared_formula_expansion,
         test_r4_values_no_fallback, test_r5_writer_evidence, test_r6_failing_sheets_gate, test_r7_sparklines,
         test_r8_cursor_memory, test_r9_silent_skips, test_r10_prolog_before_root, test_r11_edge_cases, test_values_policy_ruling_2026_10_03]


def main() -> int:
    failed = 0
    try:
        for t in TESTS:
            try:
                t()
                print(f"PASS {t.__name__}")
            except Exception:  # noqa: BLE001
                failed += 1
                print(f"FAIL {t.__name__}")
                traceback.print_exc()
    finally:
        if _TMP and os.path.isdir(_TMP):
            shutil.rmtree(_TMP, ignore_errors=True)
    print(f"\n{len(TESTS) - failed}/{len(TESTS)} tests passed")
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main())
