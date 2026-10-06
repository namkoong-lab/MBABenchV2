"""Unit tests for checks 47 (No bright-yellow highlighting) and 94 (No white-on-white hiding)
on synthetic micro-workbooks (raw SpreadsheetML zips, plus openpyxl where its labelling matters).

    cd judge && python -m detchecks.tests.test_checks_fills

Plain asserts; also collectable by pytest.  Temporary files go to detchecks/scratch/fills/.
"""
from __future__ import annotations

import os
import shutil
import sys
import tempfile
import traceback
import zipfile
from xml.sax.saxutils import escape, quoteattr

from detchecks.api import Engine
from detchecks.checks.c47 import (C47, classify_label, classify_legend, is_bright_yellow, is_key_label,
                                  is_wip_content, mentions_yellow)
from detchecks.checks._cfeval import is_unevaluable
from detchecks.checks.c94 import C94, CONCEAL_LC, apca_lc, conceals, eval_rule, parse_operand, UNKNOWN
from detchecks.core.package import Package
from detchecks.core.sheet import CfRule, ExcelError
from detchecks.errors import GradingError

HERE = os.path.dirname(os.path.abspath(__file__))
SCRATCH = os.path.join(os.path.dirname(HERE), "scratch", "fills")
MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PKG_REL = "http://schemas.openxmlformats.org/package/2006/relationships"
CT_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"
CT_WS = "application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"
DECL = '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\r\n'

_TMP = None


def tmp(name: str) -> str:
    global _TMP
    if _TMP is None:
        os.makedirs(SCRATCH, exist_ok=True)
        _TMP = tempfile.mkdtemp(prefix="test_fills_", dir=SCRATCH)
    return os.path.join(_TMP, name)


# ============================================================================ raw workbook builder
class Styles:
    """styles.xml builder: font / fill / numfmt / xf / dxf return their index."""

    def __init__(self):
        self.fonts = ['<font><sz val="11"/><name val="Calibri"/></font>']
        self.fills = ['<fill><patternFill patternType="none"/></fill>', '<fill><patternFill patternType="gray125"/></fill>']
        self.numfmts = {}
        self.xfs = ['<xf numFmtId="0" fontId="0" fillId="0" borderId="0" xfId="0"/>']
        self.dxfs = []

    def font(self, color: str | None = None) -> int:
        """color: an ARGB/RGB hex, or raw attributes like 'theme="0"'."""
        if color is None:
            c = ""
        elif "=" in color:
            c = f"<color {color}/>"
        else:
            c = f'<color rgb="{color if len(color) == 8 else "FF" + color}"/>'
        self.fonts.append(f'<font><sz val="11"/>{c}<name val="Calibri"/></font>')
        return len(self.fonts) - 1

    def fill(self, fg: str | None = None, pattern: str = "solid", fg_attr: str | None = None, bg: str | None = None) -> int:
        fgx = f"<fgColor {fg_attr}/>" if fg_attr else (f'<fgColor rgb="FF{fg[-6:]}"/>' if fg else "")
        bgx = f'<bgColor rgb="FF{bg[-6:]}"/>' if bg else ""
        self.fills.append(f'<fill><patternFill patternType="{pattern}">{fgx}{bgx}</patternFill></fill>')
        return len(self.fills) - 1

    def gradient(self, *stops: str) -> int:
        n = len(stops)
        st = "".join(f'<stop position="{i / max(n - 1, 1):g}"><color rgb="FF{c}"/></stop>' for i, c in enumerate(stops))
        self.fills.append(f'<fill><gradientFill degree="90">{st}</gradientFill></fill>')
        return len(self.fills) - 1

    def numfmt(self, code: str) -> int:
        for k, v in self.numfmts.items():
            if v == code:
                return k
        k = 164 + len(self.numfmts)
        self.numfmts[k] = code
        return k

    def xf(self, font: int = 0, fill: int = 0, numfmt: int | str = 0) -> int:
        nf = self.numfmt(numfmt) if isinstance(numfmt, str) else numfmt
        self.xfs.append(f'<xf numFmtId="{nf}" fontId="{font}" fillId="{fill}" borderId="0" xfId="0" '
                        f'applyFont="1" applyFill="1" applyNumberFormat="1"/>')
        return len(self.xfs) - 1

    def dxf(self, font: str | None = None, fill: str | None = None, numfmt: str | None = None,
            fill_xml: str | None = None) -> int:
        parts = []
        if font:
            parts.append(f'<font><color rgb="FF{font[-6:]}"/></font>')
        if numfmt is not None:
            parts.append(f'<numFmt numFmtId="{200 + len(self.dxfs)}" formatCode={quoteattr(numfmt)}/>')
        if fill:
            parts.append(f'<fill><patternFill><bgColor rgb="FF{fill[-6:]}"/></patternFill></fill>')
        if fill_xml:
            parts.append(f"<fill>{fill_xml}</fill>")
        self.dxfs.append("<dxf>" + "".join(parts) + "</dxf>")
        return len(self.dxfs) - 1

    def xml(self) -> str:
        nf = "".join(f'<numFmt numFmtId="{k}" formatCode={quoteattr(v)}/>' for k, v in self.numfmts.items())
        return (f'<styleSheet xmlns="{MAIN}">'
                + (f'<numFmts count="{len(self.numfmts)}">{nf}</numFmts>' if nf else "")
                + f'<fonts count="{len(self.fonts)}">{"".join(self.fonts)}</fonts>'
                + f'<fills count="{len(self.fills)}">{"".join(self.fills)}</fills>'
                + '<borders count="1"><border/></borders><cellStyleXfs count="1"><xf/></cellStyleXfs>'
                + f'<cellXfs count="{len(self.xfs)}">{"".join(self.xfs)}</cellXfs>'
                + '<cellStyles count="1"><cellStyle name="Normal" xfId="0" builtinId="0"/></cellStyles>'
                + f'<dxfs count="{len(self.dxfs)}">{"".join(self.dxfs)}</dxfs></styleSheet>')


def c(ref, s=0, v=None, t=None, f=None, extra=""):
    """One <c> element.  v: number / str (inline string unless t given) / bool; f: formula text."""
    attrs = f'r="{ref}"' + (f' s="{s}"' if s else "")
    if f is not None:
        tt = t or ("str" if isinstance(v, str) else "b" if isinstance(v, bool) else None)
        vv = "" if v is None else ("1" if v is True else "0" if v is False else escape(str(v)))
        return f'<c {attrs}{f" t=\"{tt}\"" if tt else ""}{extra}><f>{escape(f)}</f><v>{vv}</v></c>'
    if v is None:
        return f"<c {attrs}{extra}/>"
    if t == "e":
        return f'<c {attrs} t="e"{extra}><v>{escape(v)}</v></c>'
    if t == "s":
        return f'<c {attrs} t="s"{extra}><v>{v}</v></c>'
    if isinstance(v, bool):
        return f'<c {attrs} t="b"{extra}><v>{int(v)}</v></c>'
    if isinstance(v, str):
        return f'<c {attrs} t="inlineStr"{extra}><is><t xml:space="preserve">{escape(v)}</t></is></c>'
    return f"<c {attrs}{extra}><v>{v!r}</v></c>"


def sheet(cells, *, head="", tail="", rows_attr=None):
    """cells: list of c() strings (any order).  rows_attr: {row: ' hidden="1"'}."""
    import re as _re
    rows = {}
    for x in cells:
        ref = _re.search(r'r="([A-Z]+)(\d+)"', x)
        rows.setdefault(int(ref.group(2)), []).append((ref.group(1), x))
    rows_attr = rows_attr or {}
    for r in rows_attr:
        rows.setdefault(r, [])

    def colkey(col):
        n = 0
        for ch in col:
            n = n * 26 + ord(ch) - 64
        return n
    body = "".join(f'<row r="{r}"{rows_attr.get(r, "")}>' + "".join(x for _k, x in sorted(rows[r], key=lambda y: colkey(y[0])))
                   + "</row>" for r in sorted(rows))
    return f"{head}<sheetData>{body}</sheetData>{tail}"


def book(path, sheets, styles: Styles, *, sst=None, theme=None, sheet_parts=None, states=None):
    """sheets: [(name, inner_xml)]; sheet_parts: {sheet index: [(rel_type, target_name, xml)]}
    (target written to xl/<target_name>, related from the sheet); no docProps (writer 'unknown':
    formula caches written here are trusted)."""
    sheet_parts = sheet_parts or {}
    states = states or {}
    z = zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED)
    ov = [f'<Override PartName="/xl/workbook.xml" ContentType="{CT_XLSX}"/>']
    ov += [f'<Override PartName="/xl/worksheets/sheet{i + 1}.xml" ContentType="{CT_WS}"/>' for i in range(len(sheets))]
    z.writestr("[Content_Types].xml", '<?xml version="1.0" encoding="UTF-8"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
               '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
               '<Default Extension="xml" ContentType="application/xml"/>' + "".join(ov) + "</Types>")
    z.writestr("_rels/.rels", f'<Relationships xmlns="{PKG_REL}"><Relationship Id="rId1" Type="{REL}/officeDocument" Target="xl/workbook.xml"/></Relationships>')
    wrels, sh = [], []
    for i, (name, xml) in enumerate(sheets):
        st = f' state="{states[i]}"' if i in states else ""
        sh.append(f'<sheet name={quoteattr(name)} sheetId="{i + 1}"{st} r:id="rId{i + 1}"/>')
        wrels.append(f'<Relationship Id="rId{i + 1}" Type="{REL}/worksheet" Target="worksheets/sheet{i + 1}.xml"/>')
        extra = ""
        rels = []
        for k, (rtype, target, pxml) in enumerate(sheet_parts.get(i, [])):
            rid = f"rIdP{k + 1}"
            rels.append(f'<Relationship Id="{rid}" Type="{REL}/{rtype}" Target="../{target}"/>')
            z.writestr(f"xl/{target}", pxml)
            if rtype == "drawing":
                extra += f'<drawing r:id="{rid}"/>'
            elif rtype == "table":
                extra += f'<tableParts count="1"><tablePart r:id="{rid}"/></tableParts>'
        z.writestr(f"xl/worksheets/sheet{i + 1}.xml", f'{DECL}<worksheet xmlns="{MAIN}" xmlns:r="{REL}">{xml}{extra}</worksheet>')
        if rels:
            z.writestr(f"xl/worksheets/_rels/sheet{i + 1}.xml.rels", f'<Relationships xmlns="{PKG_REL}">{"".join(rels)}</Relationships>')
    wrels.append(f'<Relationship Id="rIdS" Type="{REL}/styles" Target="styles.xml"/>')
    z.writestr("xl/styles.xml", styles.xml())
    if sst is not None:
        wrels.append(f'<Relationship Id="rIdT" Type="{REL}/sharedStrings" Target="sharedStrings.xml"/>')
        z.writestr("xl/sharedStrings.xml", f'<sst xmlns="{MAIN}">' + "".join(sst) + "</sst>")
    if theme is not None:
        wrels.append(f'<Relationship Id="rIdTh" Type="{REL}/theme" Target="theme/theme1.xml"/>')
        z.writestr("xl/theme/theme1.xml", theme)
    z.writestr("xl/workbook.xml", f'{DECL}<workbook xmlns="{MAIN}" xmlns:r="{REL}"><sheets>' + "".join(sh) + "</sheets></workbook>")
    z.writestr("xl/_rels/workbook.xml.rels", f'<Relationships xmlns="{PKG_REL}">' + "".join(wrels) + "</Relationships>")
    z.close()
    return path


def run(path, check_cls):
    return Engine(path, [check_cls()]).run()[check_cls.key]


def graded(kind, fn, *needles):
    """Patrick 2026-10-05 (every attempt graded): where the check used to raise, it decides by a default
    recorded in stats.defaults[kind]; needles that name a location ('S!A1') must appear in that record."""
    import json
    v = fn()
    d = (v["stats"].get("defaults") or {}).get(kind)
    assert d and d["count"] >= 1, (kind, v["stats"].get("defaults"))
    blob = json.dumps(d)
    for n in needles:
        if "!" in n:
            assert n.split("!")[-1] in blob, (n, d)
    return v


def raises(fn, *needles):
    try:
        fn()
    except GradingError as e:
        msg = str(e) + " " + " ".join(str(v) for v in e.failures.values())
        for n in needles:
            assert n in msg, (n, msg[:600])
        return e
    raise AssertionError(f"expected GradingError containing {needles}")


def locs(v):
    return [m["location"] for m in v["mistakes"]]


def comments_part(ref_text):
    items = "".join(f'<comment ref="{r}" authorId="0"><text><t>{escape(t)}</t></text></comment>' for r, t in ref_text)
    return ("comments", "comments1.xml",
            f'<comments xmlns="{MAIN}"><authors><author>me</author></authors><commentList>{items}</commentList></comments>')


def textbox_part(*paras):
    ps = "".join(f"<a:p><a:r><a:t>{escape(p)}</a:t></a:r></a:p>" for p in paras)
    return ("drawing", "drawings/drawing1.xml",
            '<xdr:wsDr xmlns:xdr="http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing" '
            'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><xdr:twoCellAnchor><xdr:sp>'
            f"<xdr:txBody><a:bodyPr/>{ps}</xdr:txBody></xdr:sp></xdr:twoCellAnchor></xdr:wsDr>")


def theme_xml(accent4="FFC000"):
    slots = [("dk1", "000000"), ("lt1", "FFFFFF"), ("dk2", "44546A"), ("lt2", "E7E6E6"), ("accent1", "4472C4"),
             ("accent2", "ED7D31"), ("accent3", "A5A5A5"), ("accent4", accent4), ("accent5", "5B9BD5"),
             ("accent6", "70AD47"), ("hlink", "0563C1"), ("folHlink", "954F72")]
    cs = "".join(f'<a:{n}><a:srgbClr val="{v}"/></a:{n}>' for n, v in slots)
    return (f'<a:theme xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" name="T"><a:themeElements>'
            f'<a:clrScheme name="T">{cs}</a:clrScheme></a:themeElements></a:theme>')


# ============================================================================ 47: pure functions
def test_47_colour_band():
    bright = ["FFFFFF00", "FFFF00", "FFFFFF66", "FFFFD700", "FFFFCC00", "FFFFFF7F"]   # FFFF7F: s = 0.502
    not_bright = ["FFFFFF80", "FFFFFFCC", "FFFFF2CC", "FFFFFF99", "FFFFEB9C", "FFFFFFE0", "FFFFC000", "FFCCCC00", "FF808000",
                  "FFF5E6C0", "FFFEF3E2", "FFFFEB84", "FF92D050", "FFFFFFFF", None, ""]
    for x in bright:
        assert is_bright_yellow(x), x
    for x in not_bright:
        assert not is_bright_yellow(x), x


def test_47_legend_classifier():
    table = [
        # the toys' wording
        ("Yellow fill = TBD / to be confirmed", "WIP"),                  # 47 T4 Fail legend
        ("Yellow fill = key output cells", "DEFINES"),                   # 47 T3 Pass-style legend line
        # explains
        ("Yellow cells are inputs", "DEFINES"),
        ("Inputs are highlighted in yellow", "DEFINES"),
        ("Enter assumptions only in the yellow input cells", "DEFINES"),
        ("No hard-coded numbers outside the yellow input cells", "DEFINES"),   # exception phrasing, not a negation
        ("Yellow = assumptions to be reviewed quarterly", "DEFINES"),          # rubric: periodic review
        ("Yellow fill on Checks = a check is failing", "DEFINES"),              # rubric: error-check CF
        ("Bright yellow fill (#FFFF00) marks error-check rows", "DEFINES"),
        ("Assumptions (yellow) are updated each quarter", "DEFINES"),
        ("Input cells = yellow", "DEFINES"),
        ("Yellow cells = key outputs; nothing in this file is TBD", "DEFINES"),
        # real wording (GPT-6 2237 Summary!B61)
        ("Yellow on Questions is the preserved source answer-cell fill, not unfinished work.", "TEMPLATE"),
        # negates
        ("No yellow highlighting is used in this workbook", "NEGATED"),
        ("Yellow: none", "NEGATED"),
        ("Yellow – not used", "NEGATED"),
        ("All yellow highlighting has been removed", "NEGATED"),
        ("Blue, not yellow, marks inputs", "NEGATED"),
        ("Yellow = Unfinished - none in this file", "NEGATED"),
        # unfinished
        ("Yellow = placeholder", "WIP"),
        ("Yellow cells: to do", "WIP"),
        ("Yellow = review me", "WIP"),
        ("Cells in yellow need review", "WIP"),
        ("Yellow cells are placeholders - update before sending", "WIP"),
        ("Yellow = work in progress", "WIP"),
        ("Yellow = to be confirmed with the client", "WIP"),
        ("Revenue growth (yellow = TBD)", "WIP"),
        # mentions without a definition
        ("The logo is yellow", "MENTION"),
        ("Yellow and green are the brand colours", "MENTION"),
        # pale yellow does not explain bright yellow
        ("Light yellow cells are inputs", None),
        ("Pale yellow = inputs", None),
        ("Inputs in light-yellow fill", None),
        ("Total revenue", None),
    ]
    bad = [(t, e, classify_legend(t)) for t, e in table if classify_legend(t) != e]
    assert not bad, bad
    assert mentions_yellow("FFFF00 fill") and not mentions_yellow("FFFF00AA") and not mentions_yellow("pale yellow")
    assert classify_label("Key outputs") == "DEFINES" and classify_label("TBD") == "WIP"
    assert classify_label("Discount rate - TBD") == "WIP" and classify_label("%") == "MENTION"
    assert classify_label("Not used") == "NEGATED"
    assert is_key_label("Key output cells") and is_key_label("Hard-coded inputs") and not is_key_label("Working notes")
    assert is_wip_content("TBD") and is_wip_content("to be confirmed") and is_wip_content("???")
    assert not is_wip_content("Total") and not is_wip_content("Debt")


# ============================================================================ 47: workbooks
def _y47(cells, st, **kw):
    return book(tmp(kw.pop("name", "y.xlsx")), [("S", sheet(cells, **kw.pop("sheet_kw", {})))], st, **kw)


def test_47_own_fill_pale_and_bright():
    st = Styles()
    pale = [st.xf(fill=st.fill(x)) for x in ("FFFFCC", "FFF2CC", "FFFF99", "FFC000", "FFEB9C")]
    p = _y47([c(f"A{i + 1}", s, 1.0) for i, s in enumerate(pale)], st, name="pale.xlsx")
    v = run(p, C47)
    assert v["decision"] == "pass" and v["stats"]["bright_cells"] == 0, v
    yb = st.xf(fill=st.fill("FFFF00"))
    p = _y47([c(f"A{i + 1}", s, 1.0) for i, s in enumerate(pale)] + [c("C3", yb, 5.0), c("D3", yb), c("C4", yb, f="1+1", v=2)],
             st, name="bright.xlsx")
    v = run(p, C47)
    assert v["decision"] == "fail" and locs(v) == ["S!C3:D3", "S!C4"], locs(v)


def test_47_theme_indexed_gradient():
    st = Styles()
    s_idx13 = st.xf(fill=st.fill(fg_attr='indexed="13"'))           # legacy palette 13 = FFFF00
    s_th4 = st.xf(fill=st.fill(fg_attr='theme="7"'))                # accent4 (custom theme: FFFF00)
    s_grad = st.xf(fill=st.gradient("FFFFFF", "FFFF00"))
    p = book(tmp("theme.xlsx"), [("S", sheet([c("A1", s_idx13, 1.0), c("A3", s_th4, 1.0), c("A5", s_grad, 1.0)]))],
             st, theme=theme_xml(accent4="FFFF00"))
    v = run(p, C47)
    assert v["decision"] == "fail" and locs(v) == ["S!A1", "S!A3", "S!A5"], locs(v)
    # Office theme accent4 = FFC000 (gold): not bright yellow
    p = book(tmp("theme2.xlsx"), [("S", sheet([c("A3", s_th4, 1.0)]))], st, theme=theme_xml())
    assert run(p, C47)["decision"] == "pass"
    # yellow font, yellow tab: not highlighting
    st2 = Styles()
    sf = st2.xf(font=st2.font("FFFF00"))
    p = book(tmp("font.xlsx"), [("S", sheet([c("A1", sf, "x")], head='<sheetPr><tabColor rgb="FFFFFF00"/></sheetPr>'))], st2)
    assert run(p, C47)["decision"] == "pass"


def test_47_row_col_default_styles():
    st = Styles()
    yb = st.xf(fill=st.fill("FFFF00"))
    xml = (f'<cols><col min="3" max="4" width="9" style="{yb}" customWidth="1"/></cols>'
           + sheet([c("A1", 0, 1.0)], rows_attr={5: f' s="{yb}" customFormat="1"', 6: f' s="{yb}" customFormat="1"'}))
    p = book(tmp("rowcol.xlsx"), [("S", xml)], st)
    v = run(p, C47)
    assert v["decision"] == "fail" and sorted(locs(v)) == ["S!5:6", "S!C:D"], locs(v)
    # default style 0 painted bright yellow: every sheet's empty grid
    st2 = Styles()
    st2.xfs[0] = f'<xf numFmtId="0" fontId="0" fillId="{st2.fill("FFFF00")}" borderId="0" xfId="0"/>'
    p = book(tmp("default.xlsx"), [("S", sheet([])), ("T", sheet([]))], st2)
    v = run(p, C47)
    assert v["decision"] == "fail" and locs(v) == ["S", "T"], locs(v)


def test_47_conditional_formats():
    st = Styles()
    dy, dr = st.dxf(fill="FFFF00"), st.dxf(fill="FF0000")
    cf = (f'<conditionalFormatting sqref="B2:B9 D2"><cfRule type="cellIs" dxfId="{dy}" priority="1" operator="equal">'
          f'<formula>"never"</formula></cfRule></conditionalFormatting>'
          f'<conditionalFormatting sqref="C2:C9"><cfRule type="cellIs" dxfId="{dr}" priority="2" operator="lessThan">'
          f'<formula>0</formula></cfRule></conditionalFormatting>')
    p = book(tmp("cf.xlsx"), [("S", sheet([c("B2", 0, 1.0)], tail=cf))], st)
    v = run(p, C47)        # counted although it can never fire today (CF_COUNTS_ALWAYS)
    assert v["decision"] == "fail" and locs(v) == ["S!B2:B9,D2"], locs(v)
    # dxf solid fills: Excel paints bgColor ONLY (measured 2026-10-03).  ChatGPT tool: fg = colour, bg indexed 64
    # -> BLACK, not yellow; openpyxl start_color only: bg 00000000 -> BLACK; fgColor only -> no fill
    amb = st.dxf(fill_xml='<patternFill patternType="solid"><fgColor rgb="FFFFFF00"/><bgColor indexed="64"/></patternFill>')
    same = st.dxf(fill_xml='<patternFill patternType="solid"><fgColor rgb="FFFFFF00"/><bgColor rgb="FFFFFF00"/></patternFill>')
    rule = '<conditionalFormatting sqref="E1"><cfRule type="cellIs" dxfId="%d" priority="3" operator="equal"><formula>1</formula></cfRule></conditionalFormatting>'
    v = run(book(tmp("cf_amb.xlsx"), [("S", sheet([c("B2", 0, 1.0)], tail=rule % amb))], st), C47)
    assert v["decision"] == "pass", locs(v)
    fgonly = st.dxf(fill_xml='<patternFill patternType="solid"><fgColor rgb="FFFFFF00"/></patternFill>')
    v = run(book(tmp("cf_fgonly.xlsx"), [("S", sheet([c("B2", 0, 1.0)], tail=rule % fgonly))], st), C47)
    assert v["decision"] == "pass", locs(v)
    blue = st.dxf(fill_xml='<patternFill patternType="solid"><fgColor rgb="FFFFFF00"/><bgColor rgb="FF0000FF"/></patternFill>')
    v = run(book(tmp("cf_blue.xlsx"), [("S", sheet([c("B2", 0, 1.0)], tail=rule % blue))], st), C47)
    assert v["decision"] == "pass", locs(v)
    v = run(book(tmp("cf_same.xlsx"), [("S", sheet([c("B2", 0, 1.0)], tail=rule % same))], st), C47)
    assert v["decision"] == "fail" and locs(v) == ["S!E1"], locs(v)
    # Excel's own layout (Excel Online, corpus 1375): fgColor indexed 64 = automatic placeholder, colour in bgColor
    xl = st.dxf(fill_xml='<patternFill patternType="solid"><fgColor indexed="64"/><bgColor rgb="FFFFFF00"/></patternFill>')
    v = run(book(tmp("cf_xl.xlsx"), [("S", sheet([c("B2", 0, 1.0)], tail=rule % xl))], st), C47)
    assert v["decision"] == "fail" and locs(v) == ["S!E1"], locs(v)
    # openpyxl PatternFill(start_color=X, fill_type="solid"): bgColor 00000000 paints black (alpha ignored)
    opx = st.dxf(fill_xml='<patternFill patternType="solid"><fgColor rgb="00FFFF00"/><bgColor rgb="00000000"/></patternFill>')
    v = run(book(tmp("cf_opx.xlsx"), [("S", sheet([c("B2", 0, 1.0)], tail=rule % opx))], st), C47)
    assert v["decision"] == "pass", locs(v)
    scale = ('<conditionalFormatting sqref="A1:A9"><cfRule type="colorScale" priority="1"><colorScale>'
             '<cfvo type="min"/><cfvo type="percentile" val="50"/><cfvo type="max"/>'
             '<color rgb="FF63BE7B"/><color rgb="FFFFFF00"/><color rgb="FFF8696B"/></colorScale></cfRule></conditionalFormatting>')
    p = book(tmp("scale.xlsx"), [("S", sheet([c("A1", 0, 1.0)], tail=scale))], st)
    assert run(p, C47)["decision"] == "pass"


def test_47_spill_member_needle():
    st = Styles()
    yb = st.xf(fill=st.fill("FFFF00"))
    cells = [c("A1", 0, 1.0, f="SEQUENCE(5)", extra=' cm="1"').replace("<f>", '<f t="array" ref="A1:A5">')]
    cells += [c(f"A{r}", 0, float(r)) for r in (2, 3, 5)] + [c("A4", yb, 4.0)]
    p = book(tmp("spill.xlsx"), [("S", sheet(cells))], st)
    v = run(p, C47)
    assert v["decision"] == "fail" and locs(v) == ["S!A4"], locs(v)


def test_47_legend_excuses_and_wip_override():
    st = Styles()
    yb = st.xf(fill=st.fill("FFFF00"))
    base = [c("A1", 0, "Legend"), c("C5", yb, 10.0), c("C6", yb, 11.0)]
    v = run(_y47(base + [c("A2", 0, "Yellow fill = key output cells")], st, name="pos.xlsx"), C47)
    assert v["decision"] == "pass" and v["stats"]["excused_by_legend"], v
    # positive legend does not excuse a yellow cell that says TBD
    v = run(_y47(base + [c("A2", 0, "Yellow fill = key output cells"), c("C7", yb, "TBD")], st, name="pos_tbd.xlsx"), C47)
    assert v["decision"] == "fail" and locs(v) == ["S!C7"], locs(v)
    for txt in ("Yellow fill = TBD / to be confirmed", "No yellow is used", "Light yellow = inputs", "Yellow = review me"):
        v = run(_y47(base + [c("A2", 0, txt)], st, name="neg.xlsx"), C47)
        assert v["decision"] == "fail" and locs(v) == ["S!C5:C6"], (txt, locs(v))
    # bare colour label + its meaning in the next cell
    v = run(_y47(base + [c("A2", 0, "Yellow"), c("B2", 0, "Key outputs")], st, name="bare.xlsx"), C47)
    assert v["decision"] == "pass", v["stats"]["legend_lines"]
    # legend in a comment / in a text box (whole workbook text)
    p = book(tmp("cm.xlsx"), [("S", sheet(base))], st, sheet_parts={0: [comments_part([("C5", "Yellow cells = inputs")])]})
    assert run(p, C47)["decision"] == "pass"
    p = book(tmp("tb.xlsx"), [("S", sheet(base))], st, sheet_parts={0: [textbox_part("Formatting key", "Yellow fill: key outputs")]})
    assert run(p, C47)["decision"] == "pass"
    # a WIP comment on a yellow cell is not excused by a positive legend
    p = book(tmp("cmwip.xlsx"), [("S", sheet(base + [c("A2", 0, "Yellow = inputs")]))], st,
             sheet_parts={0: [comments_part([("C6", "TBD - confirm with client")])]})
    v = run(p, C47)
    assert v["decision"] == "fail" and locs(v) == ["S!C6"], locs(v)
    # legend on another (hidden) sheet still counts; yellow on a hidden sheet still counts
    p = book(tmp("hid.xlsx"), [("S", sheet([c("A1", 0, 1.0)])), ("H", sheet([c("B2", yb, 1.0)]))], st, states={1: "hidden"})
    v = run(p, C47)
    assert v["decision"] == "fail" and locs(v) == ["H!B2"], locs(v)


def test_47_swatches():
    st = Styles()
    yb = st.xf(fill=st.fill("FFFF00"))
    navy = st.xf(fill=st.fill("1F4E78"))
    data = [c("H10", yb, 1.0), c("H11", yb, 2.0)]
    # swatch beside a colour-key label without the colour word; the row also holds data further right
    v = run(_y47(data + [c("A4", yb), c("B4", 0, "Key output cells"), c("H4", 0, 5.0, f="1+4")], st, name="sw1.xlsx"), C47)
    assert v["decision"] == "pass" and v["stats"]["swatches"] == 1, (locs(v), v["stats"]["legend_lines"])
    # plain note beside an empty yellow cell, no legend context: the cell is highlighting
    v = run(_y47(data + [c("A4", yb), c("B4", 0, "Working notes")], st, name="sw2.xlsx"), C47)
    assert v["decision"] == "fail" and locs(v) == ["S!A4", "S!H10:H11"], locs(v)
    # ... but under a 'Legend' header it is a sample
    v = run(_y47(data + [c("A2", 0, "Legend"), c("A4", yb), c("B4", 0, "Working notes")], st, name="sw3.xlsx"), C47)
    assert v["decision"] == "pass", locs(v)
    # a section header band (navy A7 + title B7 in the same fill) is not a colour key
    v = run(_y47(data + [c("A7", navy), c("B7", navy, "REVENUE"), c("A4", yb), c("B4", 0, "Working notes")], st,
                 name="sw4.xlsx"), C47)
    assert v["decision"] == "fail" and "S!A4" in locs(v), locs(v)
    # another colour's sample with a label two rows below makes a key block
    v = run(_y47(data + [c("A6", navy), c("B6", 0, "Formulas"), c("A4", yb), c("B4", 0, "Working notes")], st,
                 name="sw5.xlsx"), C47)
    assert v["decision"] == "pass", locs(v)
    # swatch labelled 'Yellow = TBD' is documentation, but does not excuse the other yellow
    v = run(_y47(data + [c("A4", yb), c("B4", 0, "Yellow = TBD")], st, name="sw6.xlsx"), C47)
    assert v["decision"] == "fail" and locs(v) == ["S!H10:H11"], locs(v)
    # attempt-4060 layout: 'Yellow | Unfinished - none in this file | [sample]', no other yellow: pass
    v = run(_y47([c("A4", 0, "Yellow"), c("B4", 0, "Unfinished - none in this file"), c("C4", yb)], st, name="sw7.xlsx"), C47)
    assert v["decision"] == "pass" and v["stats"]["swatches"] == 1, locs(v)
    # self-labelled sample (attempt 3439): the yellow cell's own text is the legend line
    v = run(_y47([c("A19", yb, "Yellow fill = Unfinished / review (none in delivered file)")], st, name="sw11.xlsx"), C47)
    assert v["decision"] == "pass" and v["stats"]["swatches"] == 1, locs(v)
    v = run(_y47(data + [c("A19", yb, "Yellow fill = Unfinished / review (none in delivered file)")], st, name="sw12.xlsx"), C47)
    assert v["decision"] == "fail" and locs(v) == ["S!H10:H11"], locs(v)
    # ... but a yellow data label that merely contains the word is highlighting
    v = run(_y47([c("A5", yb, "Revenue growth (yellow = TBD)")], st, name="sw13.xlsx"), C47)
    assert v["decision"] == "fail" and locs(v) == ["S!A5"], locs(v)
    # an empty yellow input placeholder: label on its left, 'TBD' on its right
    v = run(_y47([c("B4", 0, "Discount rate"), c("C4", yb), c("D4", 0, "TBD"), c("A1", 0, "Legend")], st, name="sw8.xlsx"), C47)
    assert v["decision"] == "fail" and locs(v) == ["S!C4"], locs(v)
    # 'Discount rate - TBD' as the label of an empty yellow cell under a legend header: not a sample
    v = run(_y47([c("A1", 0, "Conventions"), c("A4", yb), c("B4", 0, "Discount rate - TBD")], st, name="sw9.xlsx"), C47)
    assert v["decision"] == "fail" and locs(v) == ["S!A4"], locs(v)
    # two adjacent yellow cells are a band, not a sample
    v = run(_y47([c("A4", yb), c("B4", yb), c("C4", 0, "Key outputs")], st, name="sw10.xlsx"), C47)
    assert v["decision"] == "fail" and locs(v) == ["S!A4:B4"], locs(v)


def test_47_legend_swatch_rulings_2026_10_06():
    """Patrick 2026-10-06 on the 23 check-47 failures of the v13 run (all legend artefacts): (A) a legend
    sample that holds a short text beside a label mentioning yellow (anywhere in it) is documentation
    (1605, 1800, 1924, 1886, 1622 and 9 more: PASS); (B) yellow on the row above / below the yellow legend
    line stays a mistake (1647 and 7 more: FAIL); (C) a conditional-format review flag whose legend quotes
    the flag's name ('"review me" flag', 'nothing unfinished') is a documented convention (4026: PASS)."""
    st = Styles()
    yb = st.xf(fill=st.fill("FFFF00"))
    # A: text sample, label to the right starting with the colour word (1605)
    v = run(_y47([c("B55", yb, "Unfinished / review me"),
                  c("C55", 0, "Yellow fill — reserved for unfinished items; none present in the delivered file")],
                 st, name="rl1.xlsx"), C47)
    assert v["decision"] == "pass" and v["stats"]["swatches"] == 1, (locs(v), v["stats"]["legend_lines"])
    # A: label two columns to the right, sample text 'Sample' (1800)
    v = run(_y47([c("B52", yb, "Sample"), c("D52", 0, "Yellow = unfinished / review me - none in the delivered file")],
                 st, name="rl2.xlsx"), C47)
    assert v["decision"] == "pass" and v["stats"]["swatches"] == 1, locs(v)
    # A: label two columns to the LEFT with the colour word mid-text, sample 'none' (1924)
    v = run(_y47([c("B22", 0, "Unfinished / review me (yellow fill) - none in the delivered file"), c("D22", yb, "none")],
                 st, name="rl3.xlsx"), C47)
    assert v["decision"] == "pass" and v["stats"]["swatches"] == 1, locs(v)
    # A: self-labelled sample with the colour word in parentheses (1886)
    v = run(_y47([c("C28", yb, "Unfinished / review me (yellow) — none in this file")], st, name="rl4.xlsx"), C47)
    assert v["decision"] == "pass" and v["stats"]["swatches"] == 1, locs(v)
    # A: 'Yellow = ...' | 'review' sample | explanation (1622)
    v = run(_y47([c("B73", 0, "Yellow = unfinished / review me"), c("C73", yb, "review"), c("D73", 0, "None left in this file")],
                 st, name="rl5.xlsx"), C47)
    assert v["decision"] == "pass" and v["stats"]["swatches"] == 1, locs(v)
    # a LONG yellow text beside a legend line is a note, not a sample
    v = run(_y47([c("B5", yb, "Discount rate assumption still to be confirmed with the client before the final board pack"),
                  c("C5", 0, "Yellow fill = inputs")], st, name="rl6.xlsx"), C47)
    assert v["decision"] == "fail" and locs(v) == ["S!B5"], locs(v)
    # the label line still governs the rest of the yellow: a WIP line excuses nothing elsewhere
    v = run(_y47([c("B55", yb, "Unfinished / review me"), c("C55", 0, "Yellow fill = unfinished"), c("H10", yb, 1.0)],
                 st, name="rl7.xlsx"), C47)
    assert v["decision"] == "fail" and locs(v) == ["S!H10"] and v["stats"]["swatches"] == 1, locs(v)
    # B: yellow on the row below the yellow legend line (1647): a misplaced swatch is still a mistake
    v = run(_y47([c("B29", 0, "Yellow fill"), c("C29", 0, "Unfinished / review - none in the delivered file"),
                  c("B30", yb, "Abbreviations"), c("C30", 0, "A = actual, E = estimate")], st, name="rl8.xlsx"), C47)
    assert v["decision"] == "fail" and locs(v) == ["S!B30"], locs(v)
    # C: a conditional-format review flag documented as a flag (4026): DEFINES, excused
    dy = st.dxf(fill="FFFF00")
    cf = (f'<conditionalFormatting sqref="D5:D26"><cfRule type="expression" dxfId="{dy}" priority="1">'
          f'<formula>$D5="REVIEW"</formula></cfRule></conditionalFormatting>')
    note = ('REVIEW flags (yellow) mark sense-band breaches or a changed assumption — "review me", nothing unfinished; '
            'FAIL (red) feeds the master count in C2')
    v = run(book(tmp("rl9.xlsx"), [("S", sheet([c("A28", 0, note)], tail=cf))], st), C47)
    assert v["decision"] == "pass" and v["stats"]["legend_lines"][0]["class"] == "DEFINES", (locs(v), v["stats"]["legend_lines"])
    assert classify_legend('yellow fill = "review me" flag on the Checks tab only (REVIEW = sense-band breach)') == "DEFINES"
    # ... a soft marker that names no flag convention is still unfinished work
    v = run(book(tmp("rl10.xlsx"), [("S", sheet([c("A28", 0, "Yellow = needs review")], tail=cf))], st), C47)
    assert v["decision"] == "fail" and v["stats"]["legend_lines"][0]["class"] == "WIP", v["stats"]["legend_lines"]
    assert classify_legend('Yellow = "TBD"') == "WIP"


# ============================================================================ 94: pure functions
def test_94_contrast_anchors():
    assert apca_lc("FFFFFF", "FFFFFF") == 0.0
    assert abs(apca_lc("BFBFBF", "FFFFFF") - 34.5) < 0.2       # report anchor
    visible = [("D9D9D9", "FFFFFF"), ("CCCCCC", "FFFFFF"), ("BFBFBF", "FFFFFF"), ("808080", "FFFFFF"),
               ("E7E6E6", "FFFFFF"),                            # palest theme grey (lt2), Lc 12.2
               ("000000", "1F4E78"),                            # Lc 14.5: dark-blue header, black text
               ("FFFFFF", "1F4E78"), ("FFFFFF", "D9E2F3"), ("FFFFFF", "013E12"), ("00FF00", "FFFFFF")]
    hidden = [("FFFFFF", "FFFFFF"), ("FFFFFF", "F2F2F2"), ("F2F2F2", "FFFFFF"), ("FFFFFF", "FFFFCC"),
              ("FFFFFF", "E9EBEE"),                             # Lc 11.1
              ("FFFF00", "FFFFFF"), ("000000", "002060"), ("000000", "1B365D")]
    for t, b in visible:
        assert not conceals(t, (b,)), (t, b, apca_lc(t, b))
    for t, b in hidden:
        assert conceals(t, (b,)), (t, b, apca_lc(t, b))
    assert CONCEAL_LC == 12.0
    # gradients: concealed only if invisible against every candidate
    assert not conceals("FFFFFF", ("FFFFFF", "1F4E78"))
    assert conceals("FFFFFF", ("FFFFFF", "FAFAFA"))


def test_94_cf_rule_evaluation():
    def rule(t, op=None, formulas=(), text=None):
        return CfRule(type=t, operator=op, formulas=list(formulas), text=text)
    assert parse_operand("0") == 0.0 and parse_operand('"a""b"') == 'a"b' and parse_operand("TRUE") is True
    assert parse_operand("$A$1") is UNKNOWN and parse_operand("=1.5E3") == 1500.0
    assert eval_rule(rule("cellIs", "greaterThan", ["1"]), 5.0) is True
    assert eval_rule(rule("cellIs", "greaterThan", ["1"]), 0.0) is False
    assert eval_rule(rule("cellIs", "greaterThan", ["1"]), "text") is True          # text > numbers in Excel
    assert eval_rule(rule("cellIs", "equal", ['"abc"']), "ABC") is True
    assert eval_rule(rule("cellIs", "between", ["10", "1"]), 5.0) is True
    assert eval_rule(rule("cellIs", "notBetween", ["1", "10"]), 5.0) is False
    assert eval_rule(rule("cellIs", "equal", ["A1"]), 5.0) is UNKNOWN
    assert eval_rule(rule("cellIs", "equal", ["0"]), ExcelError("#N/A")) is False
    assert eval_rule(rule("containsText", text="tb"), "TBD item") is True
    assert eval_rule(rule("beginsWith", text="x"), "abc") is False
    assert eval_rule(rule("containsErrors"), ExcelError("#DIV/0!")) is True
    assert eval_rule(rule("notContainsBlanks"), 3.0) is True
    assert eval_rule(rule("expression", formulas=["TRUE"]), 1.0) is True
    assert eval_rule(rule("expression", formulas=["$A1>0"]), 1.0) is UNKNOWN
    # self references (relative to the first range's top-left cell) and banding
    assert eval_rule(rule("expression", formulas=['C4="PASS"']), "pass", (4, 3, 4, 3)) is True
    assert eval_rule(rule("expression", formulas=['I31="FAIL"']), "FAIL", (33, 9, 31, 9)) is True
    assert eval_rule(rule("expression", formulas=["0<A1"]), 5.0, (7, 1, 1, 1)) is True
    # another cell without a value source (no env): cannot be evaluated - with one it is read (test_cfeval.py)
    assert is_unevaluable(eval_rule(rule("expression", formulas=["$A$1>0"]), 5.0, (2, 1, 1, 1)))
    assert is_unevaluable(eval_rule(rule("expression", formulas=["$B1>0"]), 5.0, (1, 1, 1, 1)))
    assert eval_rule(rule("expression", formulas=["MOD(ROW(),2)=0"]), 5.0, (32, 1, 31, 1)) is True
    assert eval_rule(rule("expression", formulas=["ISODD(COLUMN())"]), 5.0, (1, 2, 1, 1)) is False
    assert is_unevaluable(eval_rule(rule("top10"), 1.0))      # Patrick 2026-10-05: assumed visible where it matters


# ============================================================================ 94: workbooks
def _w94(cells, st, name, **kw):
    sk = kw.pop("sheet_kw", {})
    return book(tmp(name), [("S", sheet(cells, **sk))], st, **kw)


def test_94_font_against_own_background():
    st = Styles()
    white = st.font("FFFFFF")
    th0 = st.font('theme="0"')
    s_w = st.xf(font=white)
    s_th0 = st.xf(font=th0)
    s_w_navy = st.xf(font=white, fill=st.fill("1F4E78"))
    s_w_wfill = st.xf(font=white, fill=st.fill("FFFFFF"))
    s_w_grad = st.xf(font=white, fill=st.gradient("007B00", "020024"))
    s_grey = st.xf(font=st.font("D9D9D9"))
    s_f2 = st.xf(font=st.font("F2F2F2"))
    cells = [c("A1", s_w_navy, "Year"), c("B1", s_w_navy, 2025.0), c("A3", s_w_grad, "Title"),
             c("A5", s_grey, "note"), c("C5", s_w), c("C6", s_w), c("D7", s_w, "", f='""')]
    v = run(_w94(cells, st, "vis.xlsx"), C94)
    assert v["decision"] == "pass", locs(v)        # white on dark fill / gradient, pale grey, empty cells, "" result
    cells += [c("E2", s_w, 101.0), c("E3", s_th0, 5.0, f="2+3"), c("F2", s_w_wfill, "Year"), c("G2", s_f2, "x")]
    v = run(_w94(cells, st, "hid.xlsx"), C94)
    assert v["decision"] == "fail" and locs(v) == ["S!E2:E3", "S!F2", "S!G2"], locs(v)   # rgb white + theme 0 merge


def test_94_number_formats():
    st = Styles()
    hide = st.xf(numfmt=";;;")
    zero_blank = st.xf(numfmt="0;-0;;@")
    hashes = st.xf(numfmt="#,###")
    white_tag = st.xf(numfmt="[White]0")
    red_tag = st.xf(numfmt="[Red]0")
    acct = st.xf(numfmt=44)
    cells = [c("A1", zero_blank, 5.0), c("A2", red_tag, 5.0), c("A3", hide, "#N/A", t="e"), c("A4", acct, True),
             c("A5", hashes, 12.0), c("A6", hide), c("A7", acct, 0.0)]
    v = run(_w94(cells, st, "nf_ok.xlsx"), C94)
    assert v["decision"] == "pass", locs(v)
    cells2 = cells + [c("B1", hide, 1568455.5), c("B2", hide, "text"), c("B3", zero_blank, 0.0), c("B4", hashes, 0.0),
                      c("B5", white_tag, 5.0), c("B6", hide, 7.0, f="3+4")]
    v = run(_w94(cells2, st, "nf_bad.xlsx"), C94)
    assert v["decision"] == "fail" and sorted(locs(v)) == ["S!B1:B2", "S!B3", "S!B4", "S!B5", "S!B6"], locs(v)
    # logicals go through the TEXT section (Excel 2026-10-03): ';;;' hides TRUE; '[White]0' (a numeric
    # section's colour) does not colour FALSE, so it stays visible
    v = run(_w94([c("A1", hide, True)], st, "nf_bool.xlsx"), C94)
    assert v["decision"] == "fail" and locs(v) == ["S!A1"], locs(v)
    assert run(_w94([c("A1", white_tag, False)], st, "nf_bool2.xlsx"), C94)["decision"] == "pass"
    assert run(_w94([c("A1", white_tag, "text")], st, "nf_text_tag.xlsx"), C94)["decision"] == "pass"


def test_94_rich_text_runs():
    st = Styles()
    s_w = st.xf(font=st.font("FFFFFF"))
    sst = ['<si><r><t>Visible </t></r><r><rPr><color rgb="FFFFFFFF"/></rPr><t>hidden</t></r></si>',
           '<si><r><rPr><color rgb="FF000000"/></rPr><t>all black runs</t></r></si>']
    v = run(book(tmp("rich.xlsx"), [("S", sheet([c("A1", 0, 0, t="s"), c("A2", s_w, 1, t="s")]))], st, sst=sst), C94)
    assert v["decision"] == "fail" and locs(v) == ["S!A1"], locs(v)


def test_94_conditional_formatting():
    st = Styles()
    s_w = st.xf(font=st.font("FFFFFF"))
    d_white, d_black, d_navy, d_red = st.dxf(font="FFFFFF"), st.dxf(font="000000"), st.dxf(fill="1F4E78"), st.dxf(fill="FF0000")

    def cf(sqref, *rules):
        return f'<conditionalFormatting sqref="{sqref}">' + "".join(rules) + "</conditionalFormatting>"

    def cellis(dxf, prio, op, f, stop=False):
        return (f'<cfRule type="cellIs" dxfId="{dxf}" priority="{prio}" operator="{op}"'
                f'{" stopIfTrue=\"1\"" if stop else ""}><formula>{escape(f)}</formula></cfRule>')
    # (a) CF turns the font white where the value > 1: fires on 5, not on 0
    tail = cf("A1:A2", cellis(d_white, 1, "greaterThan", "1"))
    v = run(_w94([c("A1", 0, 5.0), c("A2", 0, 0.0)], st, "cf_a.xlsx", sheet_kw={"tail": tail}), C94)
    assert v["decision"] == "fail" and locs(v) == ["S!A1"], locs(v)
    # (b) base white font revealed by a CF black font where > 0
    tail = cf("A1:A2", cellis(d_black, 1, "greaterThan", "0"))
    v = run(_w94([c("A1", s_w, 5.0), c("A2", s_w, -1.0, f="0-1")], st, "cf_b.xlsx", sheet_kw={"tail": tail}), C94)
    assert v["decision"] == "fail" and locs(v) == ["S!A2"], locs(v)
    # (c) navy CF fill under white text reveals it
    tail = cf("A1", '<cfRule type="notContainsBlanks" dxfId="%d" priority="1"><formula>LEN(TRIM(A1))&gt;0</formula></cfRule>' % d_navy)
    assert run(_w94([c("A1", s_w, 5.0)], st, "cf_c.xlsx", sheet_kw={"tail": tail}), C94)["decision"] == "pass"
    # (d) priority + stopIfTrue: black (prio 1, stop) beats white (prio 2); reversed priorities conceal
    tail = cf("A1", cellis(d_black, 1, "greaterThan", "0", stop=True), cellis(d_white, 2, "greaterThan", "0"))
    assert run(_w94([c("A1", 0, 5.0)], st, "cf_d1.xlsx", sheet_kw={"tail": tail}), C94)["decision"] == "pass"
    tail = cf("A1", cellis(d_white, 1, "greaterThan", "0"), cellis(d_black, 2, "greaterThan", "0"))
    assert run(_w94([c("A1", 0, 5.0)], st, "cf_d2.xlsx", sheet_kw={"tail": tail}), C94)["decision"] == "fail"
    # (e) an expression rule reading another cell is evaluated through the value source (Patrick 2026-10-05;
    #     until then GradingError): B1 empty -> the rule is off, B1 = "x" -> white text, concealed
    tail = cf("A1", f'<cfRule type="expression" dxfId="{d_white}" priority="1"><formula>$B1="x"</formula></cfRule>')
    v = run(_w94([c("A1", 0, 5.0)], st, "cf_e.xlsx", sheet_kw={"tail": tail}), C94)
    assert v["decision"] == "pass" and v["stats"]["cf_assumptions"]["cells"] == 0, v["stats"]
    v = run(_w94([c("A1", 0, 5.0), c("B1", 0, "x")], st, "cf_e2.xlsx", sheet_kw={"tail": tail}), C94)
    assert v["decision"] == "fail" and locs(v) == ["S!A1"], locs(v)
    # (f) a rule that cannot change visibility (red fill under black text) is never evaluated
    tail = cf("A1:A9", f'<cfRule type="expression" dxfId="{d_red}" priority="1"><formula>$B1="x"</formula></cfRule>')
    v = run(_w94([c("A1", 0, 5.0), c("A2", 0, 1.0, f="1")], st, "cf_f.xlsx", sheet_kw={"tail": tail}), C94)
    assert v["decision"] == "pass" and v["stats"]["cf_second_pass_sheets"] == [], v["stats"]
    # (h) self-referencing expressions and row banding are evaluated (GPT-6 2582 pattern: green 'PASS' text,
    #     banding F2F2F2 on even rows, white-on-red when the cell says FAIL)
    s_g = st.xf(font=st.font("008000"))
    d_band = st.dxf(fill="F2F2F2")
    d_fail = st.dxf(font="FFFFFF", fill="FF0000")
    tail = (cf("A31:A34", f'<cfRule type="expression" dxfId="{d_band}" priority="1"><formula>MOD(ROW(),2)=0</formula></cfRule>')
            + cf("A31:A34", f'<cfRule type="expression" dxfId="{d_fail}" priority="2"><formula>A31="FAIL"</formula></cfRule>'))
    cells = [c("A31", s_g, "PASS", f='IF(1,"PASS")'), c("A32", s_g, "PASS", f='IF(1,"PASS")'),
             c("A33", s_g, "FAIL", f='IF(1,"FAIL")'), c("A34", s_g, "FAIL", f='IF(1,"FAIL")')]
    v = run(_w94(cells, st, "cf_h.xlsx", sheet_kw={"tail": tail}), C94)
    # A34 (even row): banding F2F2F2 (priority 1) wins the fill, FAIL rule gives white text -> invisible
    assert v["decision"] == "fail" and locs(v) == ["S!A34"], locs(v)
    # (i) a dxf fill with fgColor X and bgColor indexed 64 paints BLACK (Excel 2026-10-03): black text hidden
    d_amb = st.dxf(fill_xml='<patternFill patternType="solid"><fgColor rgb="FFEEEAF2"/><bgColor indexed="64"/></patternFill>')
    tail = cf("C4", f'<cfRule type="expression" dxfId="{d_amb}" priority="1"><formula>C4="PASS"</formula></cfRule>')
    v = run(_w94([c("C4", 0, "PASS")], st, "cf_i.xlsx", sheet_kw={"tail": tail}), C94)
    assert v["decision"] == "fail" and locs(v) == ["S!C4"], locs(v)
    assert run(_w94([c("C4", 0, "FAIL")], st, "cf_i2.xlsx", sheet_kw={"tail": tail}), C94)["decision"] == "pass"
    # (g) a colour scale ending in white under white text: a built-in conditional format never counts as hiding
    #     text (Patrick 2026-10-05; until then GradingError: the fill depends on the value's rank)
    tail = cf("A1:A3", '<cfRule type="colorScale" priority="1"><colorScale><cfvo type="min"/><cfvo type="max"/>'
                       '<color rgb="FFFFFFFF"/><color rgb="FF1F4E78"/></colorScale></cfRule>')
    v = run(_w94([c("A1", s_w, 1.0), c("A2", s_w, 5.0)], st, "cf_g.xlsx", sheet_kw={"tail": tail}), C94)
    assert v["decision"] == "pass" and v["stats"]["cf_builtin_visible"]["cells"] == 2, v["stats"]


def test_94_values_needed_only_where_formatting_could_conceal():
    import openpyxl
    from openpyxl.styles import Font
    p = tmp("opx_ok.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "=1+1"                               # openpyxl: no cached value (untrusted)
    ws["A2"] = "=2+2"
    ws["A2"].font = Font(color="FF0000")
    wb.save(p)
    assert run(p, C94)["decision"] == "pass"        # visible styles: no value needed, no error
    ws["A3"] = "=3+3"
    ws["A3"].font = Font(color="FFFFFF")            # white formula: its value is needed and unknown
    wb.save(p)
    graded("untrusted_value", lambda: run(p, C94), "needs the value", "A3")


def test_94_table_style_and_hidden_row():
    st = Styles()
    s_w = st.xf(font=st.font("FFFFFF"))
    table = ("table", "tables/table1.xml",
             f'<table xmlns="{MAIN}" id="1" name="T1" displayName="T1" ref="A1:B3"><tableColumns count="2">'
             '<tableColumn id="1" name="H1"/><tableColumn id="2" name="H2"/></tableColumns>'
             '<tableStyleInfo name="TableStyleMedium2" showRowStripes="1"/></table>')
    p = book(tmp("table.xlsx"), [("S", sheet([c("A1", s_w, "H1"), c("B1", s_w, "H2"), c("A2", 0, 1.0)]))], st,
             sheet_parts={0: [table]})
    graded("table_style_visible", lambda: run(p, C94), "table", "TableStyleMedium2")
    # white-on-white in a hidden row still counts (whole workbook; hiding is check 93's matter)
    p = _w94([c("A2", s_w, 3.0)], st, "hidrow.xlsx", sheet_kw={"rows_attr": {2: ' hidden="1"'}})
    assert locs(run(p, C94)) == ["S!A2"]


# ============================================================================ review regressions (2026-10-03)
def _o32(st, yb):
    """T6-like offenders: eight FFFF00 check-total formulas O32:O39."""
    return [c(f"O{r}", yb, 100.0 * r, f=f"SUM(B{r}:N{r})") for r in range(32, 40)]


def test_47_review_wip_vocab():
    """47-wip-vocab: legends that define yellow as unfinished in other words do not excuse."""
    wip = ["Yellow = to confirm", "Yellow = assumptions to confirm with client", "Yellow = to update", "Yellow = to complete",
           "Yellow = still to finalise", "Yellow = not yet confirmed", "Yellow = unconfirmed assumptions",
           "Yellow = unverified figures", "Yellow = dummy values", "Yellow = missing data",
           "Yellow = estimates, replace with actuals", "Yellow = figures to revisit", "Yellow = items requiring follow-up",
           "Yellow highlight = needs input", "Yellow cells: fill in later", "Yellow = preliminary numbers",
           "Yellow fill = figures not yet confirmed"]
    bad = [(t, classify_legend(t)) for t in wip if classify_legend(t) != "WIP"]
    assert not bad, bad
    # periodic or conditional wording stays an ongoing convention
    for t in ("Yellow = assumptions to update each quarter", "Error checks turn yellow when data is missing",
              "Yellow = needs review if the balance check fails", "Yellow = key assumptions, reviewed each quarter"):
        assert classify_legend(t) == "DEFINES", (t, classify_legend(t))
    st = Styles()
    yb = st.xf(fill=st.fill("FFFF00"))
    for txt in ("Yellow = to confirm", "Yellow highlight = needs input", "Yellow = not yet confirmed"):
        v = run(_y47(_o32(st, yb) + [c("A2", 0, txt)], st, name="wipv.xlsx"), C47)
        assert v["decision"] == "fail" and locs(v) == ["S!O32:O39"], (txt, locs(v))


def test_47_review_hex_colour_label():
    """47-hex-colour-label / 47-legend-bare-hex-label: a colour cell with its hex code is a bare label."""
    for t in ("Yellow (FFFF00)", "Yellow Fill (#FFFF00)", "Highlight colour: FFFF00", "Bright yellow (#FFFF00)"):
        assert classify_legend(t) == "MENTION", (t, classify_legend(t))
    assert classify_legend("Key outputs (yellow fill)") == "DEFINES"
    st = Styles()
    yb = st.xf(fill=st.fill("FFFF00"))
    cases = [
        [c("A3", 0, "Yellow (FFFF00)"), c("B3", 0, "TBD - to be confirmed")],
        [c("A3", 0, "Yellow fill (#FFFF00)"), c("B3", 0, "Placeholder values")],
        [c("A3", yb), c("B3", 0, "Yellow (FFFF00)"), c("C3", 0, "TBD")],
        # attempt 2226 Cover!A29:C29 layout: meaning | colour | note
        [c("A29", 0, "Review / Unfinished"), c("B29", 0, "Yellow Fill (#FFFF00)"),
         c("C29", 0, "None remaining in final deliverable per House Standards")],
    ]
    for i, extra in enumerate(cases):
        v = run(_y47(_o32(st, yb) + extra, st, name=f"hex{i}.xlsx"), C47)
        assert v["decision"] == "fail" and "S!O32:O39" in locs(v) and not v["stats"]["excused_by_legend"], (i, locs(v))
    # a hex-coded label followed by a real meaning still explains yellow
    v = run(_y47(_o32(st, yb) + [c("A3", 0, "Yellow (FFFF00)"), c("B3", 0, "Key outputs")], st, name="hexok.xlsx"), C47)
    assert v["decision"] == "pass", v["stats"]["legend_lines"]


def test_47_review_swatch_label_promotion():
    """47-swatch-promotes-any-label: an ordinary row label beside a yellow flag cell is not a colour key."""
    st = Styles()
    yb = st.xf(fill=st.fill("FFFF00"))
    gr, rd = st.xf(fill=st.fill("00B050")), st.xf(fill=st.fill("FF0000"))
    off = _o32(st, yb)
    v = run(_y47(off + [c("A5", yb), c("B5", 0, "Revenue assumptions"), c("D5", 0, 0.05), c("E5", 0, 0.06)], st,
                 name="prom1.xlsx"), C47)
    assert v["decision"] == "fail" and locs(v) == ["S!A5", "S!O32:O39"], locs(v)
    v = run(_y47(off + [c("A9", yb), c("B9", 0, "Total costs")], st, name="prom2.xlsx"), C47)
    assert v["decision"] == "fail" and locs(v) == ["S!A9", "S!O32:O39"], locs(v)
    # a red/amber/green status column is not a colour key (its labels are section names)
    rag = [c("A4", 0, "Status"), c("A5", gr), c("B5", 0, "Revenue schedule"), c("A6", yb), c("B6", 0, "Cost schedule"),
           c("A7", rd), c("B7", 0, "Debt schedule")]
    v = run(_y47(rag + [c(f"F{r}", yb, 1.0 * r) for r in range(20, 24)], st, name="rag.xlsx"), C47)
    assert v["decision"] == "fail" and locs(v) == ["S!A6", "S!F20:F23"], locs(v)
    # ... a real colour key block still is one ([green] Formulas / [yellow] Key outputs)
    key = [c("A5", gr), c("B5", 0, "Formulas"), c("A6", yb), c("B6", 0, "Key outputs")]
    v = run(_y47(key + [c(f"F{r}", yb, 1.0 * r) for r in range(20, 24)], st, name="keyblock.xlsx"), C47)
    assert v["decision"] == "pass" and v["stats"]["swatches"] == 1, (locs(v), v["stats"]["legend_lines"])


def test_47_review_selflabel_wip_note():
    """47-selflabel-skips-wip: a yellow note that starts with 'Yellow' but marks a TBD item is highlighting."""
    st = Styles()
    yb = st.xf(fill=st.fill("FFFF00"))
    v = run(_y47([c("B3", yb, "Yellow flag: discount rate TBD - confirm with client")], st, name="selfwip.xlsx"), C47)
    assert v["decision"] == "fail" and locs(v) == ["S!B3"], locs(v)
    v = run(_y47([c("B3", yb, "Yellow Pages advertising revenue")], st, name="selfmention.xlsx"), C47)
    assert v["decision"] == "fail" and locs(v) == ["S!B3"], locs(v)
    # self-labelled legend samples still pass (attempt 3439, bare colour label + meaning beside it)
    v = run(_y47([c("A19", yb, "Yellow fill"), c("B19", 0, "Unfinished / review (none in delivered file)")], st,
                 name="self3439.xlsx"), C47)
    assert v["decision"] == "pass" and v["stats"]["swatches"] == 1, locs(v)


def test_47_review_positive_legends():
    """47-positive-legends-missed: plain conventions explain yellow; scoped negations are not negations."""
    defines = ["Inputs are in yellow", "Input cells in yellow", "Assumptions are yellow", "Inputs use a yellow fill",
               "Enter your assumptions in the yellow cells", "Change only the yellow cells",
               "Error-check rows turn yellow when a check fails",
               "Conditional formatting turns the cell yellow if the balance sheet does not balance",
               "Hard-coded cells (without formulas) are shaded yellow", "Cells with no formula are shaded yellow",
               "Yellow fill = inputs (none on output sheets)", "No yellow highlighting is used except for the input cells",
               "Yellow is used only for inputs", "Assumptions to be updated annually are in yellow"]
    bad = [(t, classify_legend(t)) for t in defines if classify_legend(t) != "DEFINES"]
    assert not bad, bad
    # GPT-6 sentences about the case template's retained (pale) answer fill: classified alike (1322 vs 2237),
    # and they do not excuse bright yellow (TEMPLATE_LINES_EXCUSE = False)
    for t in ("Its yellow answer fill is a preserved template style, not unfinished work.",
              "Yellow on Questions is the preserved source answer-cell fill, not unfinished work.",
              "Questions retains its original layout, font and yellow answer fill",
              "Yellow = Used only on the reserved answer cells of the Questions tab, which were shaded in the file as received",
              "All yellow answer cells are completed formulas.", "Yellow means reserved answers here."):
        assert classify_legend(t) == "TEMPLATE", (t, classify_legend(t))
    for t in ("No cells are yellow", "Without yellow highlighting", "Yellow fill = Unfinished - none in this file",
              "No external workbook links, macros, hidden items, circular references or unfinished yellow cells",
              "Yellow: none", "All yellow flags were cleaned up", "Yellow highlighting has been cleared"):
        assert classify_legend(t) == "NEGATED", (t, classify_legend(t))
    assert classify_legend("The Yellow River project indicates strong growth") == "MENTION"
    st = Styles()
    yb = st.xf(fill=st.fill("FFFF00"))
    v = run(_y47(_o32(st, yb) + [c("A2", 0, "Inputs are in yellow")], st, name="posin.xlsx"), C47)
    assert v["decision"] == "pass" and v["stats"]["excused_by_legend"], v["stats"]["legend_lines"]
    v = run(_y47(_o32(st, yb) + [c("A2", 0, "Questions retains its original layout and yellow answer fill.")], st,
                 name="tmpl.xlsx"), C47)
    assert v["decision"] == "fail" and locs(v) == ["S!O32:O39"], locs(v)
    assert "template's retained yellow" in v["mistakes"][0]["description"]
    d = st.dxf(fill="FFFF00")
    cf = (f'<conditionalFormatting sqref="B2:B9"><cfRule type="cellIs" dxfId="{d}" priority="1" operator="notEqual">'
          f'<formula>0</formula></cfRule></conditionalFormatting>')
    v = run(_y47([c("A1", 0, "Error-check rows turn yellow when a check fails"), c("B2", 0, 0.0)], st, name="poscf.xlsx",
                 sheet_kw={"tail": cf}), C47)
    assert v["decision"] == "pass", locs(v)


def test_47_review_left_and_self_keys():
    """47-left-label-and-selflabel-keys: label left of the sample, and samples holding their own meaning."""
    st = Styles()
    yb = st.xf(fill=st.fill("FFFF00"))
    blue = st.xf(fill=st.fill("DDEBF7"))
    outs = [c(f"F{r}", yb, 100.0 * r, f=f"{r}*100") for r in range(10, 14)]
    layouts = {
        "left_under_header": [c("A2", 0, "Legend"), c("A3", 0, "Inputs"), c("B3", blue), c("A4", 0, "Key outputs"), c("B4", yb)],
        "left_colon": [c("A3", 0, "Key outputs:"), c("B3", yb)],
        "self_under_header": [c("A2", 0, "Legend"), c("A3", blue, "Input"), c("A4", yb, "Key output")],
        "self_key_block": [c("A3", blue, "Inputs"), c("A4", yb, "Key outputs")],
    }
    for name, cells in layouts.items():
        v = run(_y47(outs + cells, st, name=f"lk_{name}.xlsx"), C47)
        assert v["decision"] == "pass" and v["stats"]["swatches"] == 1, (name, locs(v), v["stats"]["legend_lines"])
    # an input placeholder under a legend header: the label on its left is not a cell category -> highlighting
    v = run(_y47([c("A1", 0, "Legend"), c("B4", 0, "Discount rate"), c("C4", yb)], st, name="lk_input.xlsx"), C47)
    assert v["decision"] == "fail" and locs(v) == ["S!C4"], locs(v)
    # a key label on the left followed by a value right of the sample is a data row, not a key line
    v = run(_y47([c("A3", 0, "Key outputs"), c("B3", yb), c("C3", 0, 5.0)], st, name="lk_row.xlsx"), C47)
    assert v["decision"] == "fail" and locs(v) == ["S!B3"], locs(v)
    # a yellow text cell naming a category, with no legend context, is highlighting
    v = run(_y47([c("A4", yb, "Key outputs")], st, name="lk_self_alone.xlsx"), C47)
    assert v["decision"] == "fail" and locs(v) == ["S!A4"], locs(v)
    # ... and so is one followed by numbers, even under a legend header (a highlighted row label)
    v = run(_y47([c("A2", 0, "Legend"), c("A4", yb, "Key outputs"), c("B4", 0, 5.0)], st, name="lk_self_row.xlsx"), C47)
    assert v["decision"] == "fail" and locs(v) == ["S!A4"], locs(v)
    # a self-labelled sample followed by a description under a legend header is a key line
    v = run(_y47(outs + [c("A2", 0, "Legend"), c("A4", yb, "Key outputs"), c("B4", 0, "cells feeding the summary")], st,
                 name="lk_self_desc.xlsx"), C47)
    assert v["decision"] == "pass", locs(v)


def test_47_review_ambiguous_cf_only_when_decisive():
    """Formerly 47-raise-when-decided; since Excel 2026-10-03 a dxf fill is bgColor only, so the
    ChatGPT-tool layout (fg yellow, bg indexed 64) paints black: never yellow, never a raise."""
    st = Styles()
    yb = st.xf(fill=st.fill("FFFF00"))
    d = st.dxf(fill_xml='<patternFill patternType="solid"><fgColor rgb="FFFFFF00"/><bgColor indexed="64"/></patternFill>')
    tail = (f'<conditionalFormatting sqref="B2:B9"><cfRule type="cellIs" dxfId="{d}" priority="1" operator="notEqual">'
            f'<formula>0</formula></cfRule></conditionalFormatting>')
    v = run(_y47([c("A1", 0, "Yellow cells are inputs"), c("B2", 0, 0.0), c("D5", yb, 3.0)], st, name="amb1.xlsx",
                 sheet_kw={"tail": tail}), C47)
    assert v["decision"] == "pass" and "cf_ambiguous" not in v["stats"], v["stats"]
    v = run(_y47([c("B2", 0, 0.0), c("D5", yb, 3.0)], st, name="amb2.xlsx", sheet_kw={"tail": tail}), C47)
    assert v["decision"] == "fail" and locs(v) == ["S!D5"], locs(v)
    v = run(_y47([c("B2", 0, 0.0)], st, name="amb3.xlsx", sheet_kw={"tail": tail}), C47)
    assert v["decision"] == "pass", locs(v)


def test_47_review_replaced_and_wip_blocks():
    """47-legend-yellow-replaced: 'yellow fill replaced' is a negation; 47-most-permissive-line-wins: a legend
    that defines yellow as unfinished is not overridden by another line (WIP_LEGEND_BLOCKS_EXCUSE)."""
    for t in ("Completed answers replace the template's yellow input fill.",
              "Yellow source input fills are replaced with pale blue."):
        assert classify_legend(t) == "NEGATED", (t, classify_legend(t))
    assert classify_legend("Yellow = scenario inputs; replace with your own values") == "DEFINES"
    st = Styles()
    yb = st.xf(fill=st.fill("FFFF00"))
    v = run(_y47(_o32(st, yb) + [c("A2", 0, "Completed answers replace the template's yellow input fill.")], st,
                 name="repl.xlsx"), C47)
    assert v["decision"] == "fail" and locs(v) == ["S!O32:O39"], locs(v)
    v = run(_y47(_o32(st, yb) + [c("A2", 0, "Yellow fill = TBD / to be confirmed"), c("A3", 0, "Yellow cells are inputs")],
                 st, name="mixed.xlsx"), C47)
    assert v["decision"] == "fail" and locs(v) == ["S!O32:O39"] and v["stats"]["legend_blocked_by_wip"], v["stats"]
    assert "defines yellow as unfinished" in v["mistakes"][0]["description"]


def test_94_review_undecided_only_when_decisive():
    """94-raise-when-decided: an undecidable cell raises only when nothing is certainly concealed."""
    st = Styles()
    w = st.xf(font=st.font("FFFFFF"))
    # a rule that cannot be evaluated decides whether A2 turns white: until 2026-10-05 undecided (raising when
    # nothing else was concealed); now assumed visible and recorded (Patrick 2026-10-05) - never undecided
    d_w = st.dxf(font="FFFFFF")
    tail = ('<conditionalFormatting sqref="A2"><cfRule type="expression" dxfId="%d" priority="1">'
            '<formula>COUNTIF($B:$B,"x")&gt;0</formula></cfRule></conditionalFormatting>' % d_w)
    v = run(_w94([c("A1", w, 5.0), c("A2", 0, 7.0)], st, "und1.xlsx", sheet_kw={"tail": tail}), C94)
    assert v["decision"] == "fail" and locs(v) == ["S!A1"] and v["stats"]["undecided_cells"] == 0, v["stats"]
    assert v["stats"]["cf_assumptions"]["cells"] == 1, v["stats"]["cf_assumptions"]
    v = run(_w94([c("A2", 0, 7.0)], st, "und2.xlsx", sheet_kw={"tail": tail}), C94)
    assert v["decision"] == "pass" and v["stats"]["cf_assumptions"]["cells"] == 1, v["stats"]
    import openpyxl
    from openpyxl.styles import Font
    p = tmp("und_opx.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "=1+1"
    ws["A1"].font = Font(color="FFFFFF")         # white formula, untrusted cache: undecided
    ws["B1"] = 7
    ws["B1"].font = Font(color="FFFFFF")         # white constant: certainly concealed
    wb.save(p)
    v = run(p, C94)
    assert v["decision"] == "fail" and locs(v) == ["Sheet!B1"] and v["stats"]["undecided_cells"] == 0, (locs(v), v["stats"])
    assert v["stats"]["defaults"]["untrusted_value"]["count"] == 1, v["stats"]          # skipped (Patrick 2026-10-05)


def test_94_review_hue_pairs():
    """94-hue-pairs: clearly different hues of similar lightness are visible (prototype v5b exemption)."""
    for t, b in (("00B050", "FF0000"), ("FF0000", "0070C0"), ("FF0000", "00B050"), ("00FFFF", "FFFFFF")):
        assert abs(apca_lc(t, b)) < CONCEAL_LC and not conceals(t, (b,)), (t, b)
    for t, b in (("FFFF00", "FFFFFF"), ("FFFFFF", "FFFFCC"), ("FFFFFF", "F2F2F2"), ("000000", "002060"),
                 ("000000", "404040"), ("FFFFFF", "E9EBEE")):
        assert conceals(t, (b,)), (t, b)
    st = Styles()
    g_on_r = st.xf(font=st.font("00B050"), fill=st.fill("FF0000"))
    v = run(_w94([c("A1", g_on_r, "PASS"), c("A2", g_on_r, 42.0)], st, "hue.xlsx"), C94)
    assert v["decision"] == "pass", locs(v)


def test_94_review_conditional_sections():
    """94-conditional-section-screen: a conditional section that prints nothing for large values."""
    st = Styles()
    a, b = st.xf(numfmt='[>2]"";0'), st.xf(numfmt='[>=1000]" ";#,##0')
    v = run(_w94([c("A1", a, 5.0), c("A2", b, 5000.0, f="5000"), c("A3", a, 1.0), c("A4", b, 12.0)], st, "cond.xlsx"), C94)
    assert v["decision"] == "fail" and locs(v) == ["S!A1", "S!A2"], locs(v)     # two formats: two reasons


def test_94_review_rich_run_cf_fill():
    """94-rich-run-cf-fill: a CF fill is tested against rich-text run colours, not only the cell font."""
    st = Styles()
    d = st.dxf(fill="C0C0C0")
    sst = ['<si><r><t>Visible </t></r><r><rPr><color rgb="FFC0C0C0"/></rPr><t>secret</t></r></si>']
    tail = ('<conditionalFormatting sqref="A1"><cfRule type="notContainsBlanks" dxfId="%d" priority="1">'
            '<formula>LEN(TRIM(A1))&gt;0</formula></cfRule></conditionalFormatting>' % d)
    v = run(book(tmp("richcf.xlsx"), [("S", sheet([c("A1", 0, 0, t="s")], tail=tail))], st, sst=sst), C94)
    assert v["decision"] == "fail" and locs(v) == ["S!A1"], locs(v)


def test_94_review_containstext_wildcards():
    """containsText uses SEARCH() wildcards (Excel 2026-10-03: "P*S" and "P?SS" match "PASS")."""
    def rule(t, text):
        return CfRule(type=t, operator=t, formulas=[], text=text)
    assert eval_rule(rule("containsText", "P*S"), "PASS") is True
    assert eval_rule(rule("containsText", "P?SS"), "pass") is True
    assert eval_rule(rule("containsText", "P~*S"), "PASS") is False and eval_rule(rule("containsText", "P~*S"), "P*S") is True
    assert eval_rule(rule("notContainsText", "P*S"), "PASS") is False
    assert eval_rule(rule("containsText", "P*S"), "FAIL") is False
    assert eval_rule(rule("notContainsText", "F?IL"), "OK") is True
    assert eval_rule(rule("containsText", "P*S"), "P*S here") is True      # both readings fire
    assert eval_rule(rule("beginsWith", "P*"), "PASS") is False             # LEFT()=: no wildcards
    st = Styles()
    d = st.dxf(font="FFFFFF")
    tail = (f'<conditionalFormatting sqref="A1"><cfRule type="containsText" dxfId="{d}" priority="1" operator="containsText" '
            f'text="P*S"><formula>NOT(ISERROR(SEARCH("P*S",A1)))</formula></cfRule></conditionalFormatting>')
    v = run(_w94([c("A1", 0, "PASS")], st, "wild1.xlsx", sheet_kw={"tail": tail}), C94)
    assert v["decision"] == "fail" and locs(v) == ["S!A1"], locs(v)
    assert run(_w94([c("A1", 0, "FAIL")], st, "wild2.xlsx", sheet_kw={"tail": tail}), C94)["decision"] == "pass"


def test_94_review_dark_table_style():
    """94-table-style-one-way: dark directly formatted text in a TableStyleDark* table is undecided."""
    st = Styles()
    blk = st.xf(font=st.font("000000"))

    def tbl(style):
        return ("table", "tables/table1.xml",
                f'<table xmlns="{MAIN}" id="1" name="T1" displayName="T1" ref="A1:B3"><tableColumns count="2">'
                '<tableColumn id="1" name="H1"/><tableColumn id="2" name="H2"/></tableColumns>'
                f'<tableStyleInfo name="{style}" showRowStripes="1"/></table>')
    cells = [c("A1", blk, "H1"), c("B1", blk, "H2"), c("A2", blk, 1.0), c("B2", blk, 2.0)]
    graded("table_style_visible", lambda: run(book(tmp("dark1.xlsx"), [("S", sheet(cells))], st, sheet_parts={0: [tbl("TableStyleDark1")]}), C94),
           "TableStyleDark1", "not resolved")
    # default font (the table style's own text colour shows): visible
    dflt = [c("A1", 0, "H1"), c("B1", 0, "H2"), c("A2", 0, 1.0)]
    assert run(book(tmp("dark2.xlsx"), [("S", sheet(dflt))], st, sheet_parts={0: [tbl("TableStyleDark1")]}),
               C94)["decision"] == "pass"
    # a light / medium table style over the same black text: visible, as before
    assert run(book(tmp("dark3.xlsx"), [("S", sheet(cells))], st, sheet_parts={0: [tbl("TableStyleMedium2")]}),
               C94)["decision"] == "pass"


# ============================================================================ Excel session 2026-10-03
def test_excel_2026_10_03_dxf_fill_is_bgcolor_only():
    """Q2: a dxf solid fill paints bgColor only: fgColor ignored, bgColor indexed 64 / 00000000 =
    black, no bgColor = no fill.  Core (styles) plus 47 and 94."""
    st = Styles()
    cases = {   # fill xml -> what Excel paints
        '<patternFill patternType="solid"><fgColor rgb="FFFFFF00"/></patternFill>': None,
        '<patternFill patternType="solid"><fgColor rgb="FFFFFF00"/><bgColor indexed="64"/></patternFill>': "FF000000",
        '<patternFill patternType="solid"><fgColor rgb="FFFFFF00"/><bgColor rgb="FF0000FF"/></patternFill>': "FF0000FF",
        '<patternFill patternType="solid"><fgColor indexed="64"/><bgColor rgb="FFFFFF00"/></patternFill>': "FFFFFF00",
        '<patternFill patternType="solid"><fgColor rgb="FFFFFF00"/><bgColor rgb="00000000"/></patternFill>': "FF000000",
        '<patternFill><bgColor rgb="FFFFFF00"/></patternFill>': "FFFFFF00",
    }
    ids = {x: st.dxf(fill_xml=x) for x in cases}
    p = _w94([c("A1", 0, 1.0)], st, "dxf_q2.xlsx")
    pkg = Package.open(p)
    try:
        for x, want in cases.items():
            assert pkg.styles.dxf_fill_color(ids[x]) == want, (x, pkg.styles.dxf_fill_color(ids[x]))
    finally:
        pkg.close()
    rule = ('<conditionalFormatting sqref="B2"><cfRule type="cellIs" dxfId="%d" priority="1" operator="equal">'
            '<formula>1</formula></cfRule></conditionalFormatting>')
    for x, want in cases.items():                                   # 47: yellow only where bgColor is yellow
        v = run(book(tmp("q2_47.xlsx"), [("S", sheet([c("B2", 0, 1.0)], tail=rule % ids[x]))], st), C47)
        assert v["decision"] == ("fail" if want == "FFFFFF00" else "pass"), (x, locs(v))
    for x, want in cases.items():                                   # 94: black text hidden only on the black paint
        v = run(_w94([c("B2", 0, 1.0)], st, "q2_94.xlsx", sheet_kw={"tail": rule % ids[x]}), C94)
        assert v["decision"] == ("fail" if want == "FF000000" else "pass"), (x, locs(v))


def test_excel_2026_10_03_cf_font_beats_format_colour():
    """Q8: a conditional-format font colour beats a number-format colour tag (5 under [Red]0 with
    a CF font blue shows blue)."""
    st = Styles()
    white_tag = st.xf(numfmt="[White]General")
    blue_tag = st.xf(numfmt="[Blue]0")
    d_blue, d_white = st.dxf(font="0000FF"), st.dxf(font="FFFFFF")
    rule = ('<conditionalFormatting sqref="A1"><cfRule type="cellIs" dxfId="%d" priority="1" operator="greaterThan">'
            '<formula>0</formula></cfRule></conditionalFormatting>')
    v = run(_w94([c("A1", white_tag, 5.0)], st, "cfw1.xlsx", sheet_kw={"tail": rule % d_blue}), C94)
    assert v["decision"] == "pass", locs(v)                         # white tag overridden by blue CF font
    v = run(_w94([c("A1", blue_tag, 5.0)], st, "cfw2.xlsx", sheet_kw={"tail": rule % d_white}), C94)
    assert v["decision"] == "fail" and locs(v) == ["S!A1"], locs(v)  # blue tag overridden by white CF font
    v = run(_w94([c("A1", white_tag, -5.0)], st, "cfw3.xlsx", sheet_kw={"tail": rule % d_blue}), C94)
    assert v["decision"] == "fail", locs(v)                          # rule does not fire: the white tag counts


def test_excel_2026_10_03_show_zeros_off_94():
    """Q5: showZeros=0 hides a zero only under formats without a zero section (not counted); a
    zero section is printed and judged as usual."""
    st = Styles()
    w3 = st.xf(font=st.font("FFFFFF"), numfmt='0.00;-0.00;0.00')    # white digits, explicit zero section
    w1 = st.xf(font=st.font("FFFFFF"), numfmt='0.00')               # white, single section
    empty_zero = st.xf(numfmt='0;-0;;@')                             # explicit EMPTY zero section
    head = '<sheetViews><sheetView workbookViewId="0" showZeros="0"/></sheetViews>'
    v = run(_w94([c("A1", w1, 0.0), c("A2", 0, 0.0)], st, "sz1.xlsx", sheet_kw={"head": head}), C94)
    assert v["decision"] == "pass" and v["stats"]["zeros_hidden_by_show_zeros_off"] >= 1, v["stats"]
    v = run(_w94([c("A1", w3, 0.0)], st, "sz2.xlsx", sheet_kw={"head": head}), C94)
    assert v["decision"] == "fail" and locs(v) == ["S!A1"], locs(v)   # printed '0.00' in white on white
    v = run(_w94([c("A1", empty_zero, 0.0)], st, "sz3.xlsx", sheet_kw={"head": head}), C94)
    assert v["decision"] == "fail" and locs(v) == ["S!A1"], locs(v)   # the format itself prints nothing
    v = run(_w94([c("A1", w1, 0.0)], st, "sz4.xlsx"), C94)            # showZeros on: white zero is hidden text
    assert v["decision"] == "fail", locs(v)



def test_rulings_2026_10_03_banners_and_near_empty():
    """Rulings 2026-10-03: (A4) a case banner damaged by the delivering tool fails 94 (gradient
    dropped -> white title on the white sheet; font colour dropped -> black on dark green);
    (A3) no completeness gate: a near-empty attempt gets its plain Python verdicts (pass here)."""
    st = Styles()
    white = st.xf(font=st.font("FFFFFF"))
    dark = st.xf(fill=st.fill("244020"))           # dark green banner (Lc 7.5 for black text)
    v = run(_w94([c("B2", white, "Case: EasyDCF"), c("B4", dark, "Instructions")], st, "banner.xlsx"), C94)
    assert v["decision"] == "fail" and sorted(locs(v)) == ["S!B2", "S!B4"], locs(v)
    v = run(_w94([c("A1", 0, "Answer:")], st, "near_empty94.xlsx"), C94)
    assert v["decision"] == "pass", locs(v)
    v = run(_y47([c("A1", 0, "Answer:")], st, name="near_empty47.xlsx"), C47)
    assert v["decision"] == "pass", locs(v)


# ============================================================================ second review (2026-10-04)
def test_47_rereview_swatch_between_colour_word_and_meaning():
    """47-R2-01: 'Yellow fill | [sample] | meaning' - the colour-word label left of the sample labels it; the
    plain text right of it is the meaning, not a key label.  The line decides whether yellow is excused."""
    st = Styles()
    yb = st.xf(fill=st.fill("FFFF00"))
    # corpus 3439 layout with the sample between colour word and meaning, no other yellow: documentation
    v = run(_y47([c("A19", 0, "Yellow fill"), c("B19", yb), c("C19", 0, "Unfinished / review (none in delivered file)")],
                 st, name="r201a.xlsx"), C47)
    assert v["decision"] == "pass" and v["stats"]["swatches"] == 1, (locs(v), v["stats"]["legend_lines"])
    assert [l["class"] for l in v["stats"]["legend_lines"]] == ["NEGATED"], v["stats"]["legend_lines"]
    # the same row with a definition excuses the other yellow
    v = run(_y47([c("A2", 0, "Legend"), c("A4", 0, "Yellow fill"), c("B4", yb), c("C4", 0, "Key outputs"), c("H9", yb, 3.0)],
                 st, name="r201b.xlsx"), C47)
    assert v["decision"] == "pass" and v["stats"]["excused_by_legend"] and v["stats"]["swatches"] == 1, locs(v)
    # ... with a WIP meaning the sample is still documentation but excuses nothing (as 'Yellow = TBD' beside a sample)
    v = run(_y47([c("A2", 0, "Legend"), c("A4", 0, "Yellow fill"), c("B4", yb), c("C4", 0, "TBD / to be confirmed"),
                  c("H9", yb, 3.0)], st, name="r201c.xlsx"), C47)
    assert v["decision"] == "fail" and locs(v) == ["S!H9"], locs(v)
    v = run(_y47([c("A4", 0, "Yellow fill"), c("B4", yb), c("C4", 0, "TBD / to be confirmed")], st, name="r201d.xlsx"), C47)
    assert v["decision"] == "pass" and v["stats"]["swatches"] == 1, locs(v)
    # a yellow input placeholder between a row label and a note is still highlighting
    v = run(_y47([c("B4", 0, "Discount rate"), c("C4", yb), c("D4", 0, "TBD")], st, name="r201e.xlsx"), C47)
    assert v["decision"] == "fail" and locs(v) == ["S!C4"], locs(v)
    # a right label that starts with the colour word still wins over a plain left text
    v = run(_y47([c("A4", 0, "Outputs"), c("B4", yb), c("C4", 0, "Yellow = TBD"), c("H9", yb, 3.0)], st, name="r201f.xlsx"), C47)
    assert v["decision"] == "fail" and locs(v) == ["S!H9"], locs(v)


def test_47_rereview_status_word_is_not_a_legend():
    """47-R2-02: '<x> = yellow' defines yellow only when x names cells; a RAG status word beside a row label
    ('Cost schedule | Yellow') is not a colour key."""
    for t in ("Cost schedule = Yellow", "Discount rate = Yellow", "Status: Yellow", "Highlight colour: FFFF00"):
        assert classify_legend(t) == "MENTION", (t, classify_legend(t))
    for t in ("Input cells = yellow", "Inputs: yellow fill", "Scenario selector = yellow", "Key outputs = yellow",
              "Hard-coded = yellow", "Total rows = yellow"):
        assert classify_legend(t) == "DEFINES", (t, classify_legend(t))
    st = Styles()
    yb = st.xf(fill=st.fill("FFFF00"))
    rag = [c("A1", 0, "Workstream"), c("B1", 0, "Status"), c("A2", 0, "Revenue schedule"), c("B2", 0, "Green"),
           c("A3", 0, "Cost schedule"), c("B3", 0, "Yellow"), c("A4", 0, "Debt schedule"), c("B4", 0, "Red")]
    v = run(_y47(rag + [c("F20", yb, 100.0), c("F21", yb, 200.0)], st, name="r202a.xlsx"), C47)
    assert v["decision"] == "fail" and locs(v) == ["S!F20:F21"] and not v["stats"]["excused_by_legend"], v["stats"]
    assert [l["class"] for l in v["stats"]["legend_lines"]] == ["MENTION"], v["stats"]["legend_lines"]
    # 'Inputs | Yellow' (meaning left of the colour word) is still a legend
    v = run(_y47([c("A3", 0, "Inputs"), c("B3", 0, "Yellow"), c("F20", yb, 100.0)], st, name="r202b.xlsx"), C47)
    assert v["decision"] == "pass" and v["stats"]["excused_by_legend"], v["stats"]["legend_lines"]


def test_47_rereview_back_reference_defines_nothing():
    """47-R2-03: 'its yellow fill is explained above' (GPT-6 2356) describes the text, not a meaning."""
    for t in ("Original Questions styling is preserved; its yellow fill is explained above.",
              "Yellow cells are not inputs", "Yellow fill is intentional", "Yellow cells are documented in the glossary",
              "Yellow fill is not unfinished work"):
        assert classify_legend(t) == "MENTION", (t, classify_legend(t))
    for t in ("Yellow cells are inputs", "Yellow fill is used for checks", "Yellow cells are hard-coded assumptions"):
        assert classify_legend(t) == "DEFINES", (t, classify_legend(t))
    st = Styles()
    yb = st.xf(fill=st.fill("FFFF00"))
    sentence = "House_Standards_v1(20260922-064408).md. Original Questions styling is preserved; its yellow fill is explained above."
    v = run(book(tmp("r203.xlsx"), [("Summary", sheet([c("B145", 0, sentence), c("F20", yb, 100.0)]))], st), C47)
    assert v["decision"] == "fail" and locs(v) == ["Summary!F20"], locs(v)
    assert "does not define what yellow means" in v["mistakes"][0]["description"], v["mistakes"][0]


def test_47_rereview_section_header_is_not_a_colour_key():
    """47-R2-04: a yellow marker beside a label that heads label / value rows ('Inputs' over 'Growth | 5%') is
    highlighting, not a colour-key sample, and the label is no legend line."""
    st = Styles()
    yb = st.xf(fill=st.fill("FFFF00"))
    table = [c("B11", 0, "Growth"), c("C11", 0, 0.05), c("B12", 0, "Margin"), c("C12", 0, 0.3)]
    v = run(_y47([c("A10", yb), c("B10", 0, "Inputs")] + table + [c("F3", yb, 9.0)], st, name="r204a.xlsx"), C47)
    assert v["decision"] == "fail" and sorted(locs(v)) == ["S!A10", "S!F3"] and v["stats"]["swatches"] == 0, locs(v)
    assert not v["stats"]["legend_lines"], v["stats"]["legend_lines"]
    # left label over a table: the same
    v = run(_y47([c("A10", 0, "Inputs"), c("B10", yb), c("A11", 0, "Growth"), c("B11", 0, 0.05)], st, name="r204b.xlsx"), C47)
    assert v["decision"] == "fail" and locs(v) == ["S!B10"], locs(v)
    # the same label with nothing beneath it stays a colour-key line (no-header path, doc Question 6)
    v = run(_y47([c("A10", yb), c("B10", 0, "Inputs"), c("F3", yb, 9.0)], st, name="r204c.xlsx"), C47)
    assert v["decision"] == "pass" and v["stats"]["swatches"] == 1, locs(v)
    # a legend header above keeps the sample even over a table
    v = run(_y47([c("A8", 0, "Legend"), c("A10", yb), c("B10", 0, "Inputs")] + table + [c("F3", yb, 9.0)], st,
                 name="r204d.xlsx"), C47)
    assert v["decision"] == "pass" and v["stats"]["swatches"] == 1, locs(v)
    # 47 T3-like layout: the legend row's neighbours below hold texts and numbers far to the right only
    v = run(_y47([c("A4", yb), c("B4", 0, "Key output cells"), c("F5", 0, "Year"), c("G5", 0, 2026.0), c("E6", 0, "Units"),
                  c("F6", 0, "Mode"), c("H10", yb, 1.0)], st, name="r204e.xlsx"), C47)
    assert v["decision"] == "pass" and v["stats"]["swatches"] == 1, locs(v)


def test_47_rereview_wip_vocabulary():
    """47-R2-07 / 47-R2-11: 'unaudited', a cross-reference to open items, 'WIP inventory' and 'were replaced
    with actuals' are not unfinished work; 'draft', 'open items', 'replace with actuals' still are."""
    for t in ("Yellow = unaudited FY2025 figures", "Yellow = management estimates (unaudited)",
              "Yellow fill = inputs (see Checks tab for open items)", "Yellow = inputs - see the Checks tab for open items",
              "Yellow marks inputs that were replaced with actuals in FY25", "Yellow = WIP inventory lines"):
        assert classify_legend(t) == "DEFINES", (t, classify_legend(t))
    for t in ("Yellow = draft budget columns", "Yellow = open items", "Yellow = inputs (open items)",
              "Yellow = estimates, replace with actuals", "Yellow = estimates to be replaced with actuals", "Yellow = WIP",
              "Yellow = unverified figures"):
        assert classify_legend(t) == "WIP", (t, classify_legend(t))
    assert not is_wip_content("WIP inventory") and not is_wip_content("WIP balance") and is_wip_content("WIP")
    assert is_wip_content("WIP - confirm with client")
    st = Styles()
    yb = st.xf(fill=st.fill("FFFF00"))
    v = run(_y47([c("A1", 0, "Yellow fill = inputs"), c("A5", yb, "WIP inventory"), c("B5", yb, 120.0)], st, name="r207a.xlsx"), C47)
    assert v["decision"] == "pass" and v["stats"]["excused_by_legend"], locs(v)
    v = run(_y47([c("A1", 0, "Yellow = unaudited FY2025 figures"), c("B5", yb, 120.0)], st, name="r207b.xlsx"), C47)
    assert v["decision"] == "pass", locs(v)
    v = run(_y47([c("A1", 0, "Yellow = draft budget columns"), c("B5", yb, 120.0)], st, name="r207c.xlsx"), C47)
    assert v["decision"] == "fail" and locs(v) == ["S!B5"], locs(v)


def test_47_rereview_classifier_phrasings():
    """47-R2-11: a pronoun clause continues the colour ('...; it marks the output cells'); the ARGB spelling
    FFFFFF00 is the colour too."""
    for t in ("Yellow is not used for inputs; it marks the output cells", "Yellow is not used for inputs. It marks the output cells.",
              "Fill FFFFFF00 = key outputs", "Yellow fill (#FFFFFF00) = inputs"):
        assert classify_legend(t) == "DEFINES", (t, classify_legend(t))
    for t in ("Yellow (FFFFFF00)", "Yellow fill (#FFFFFF00)", "Yellow = inputs; this is standard"):
        assert classify_legend(t) in ("MENTION", "DEFINES"), (t, classify_legend(t))
    assert classify_legend("Yellow (FFFFFF00)") == "MENTION"
    assert classify_legend("No yellow is used; it would mean unfinished work") == "WIP"
    assert mentions_yellow("FFFFFF00") and not mentions_yellow("FFFF00AA") and not mentions_yellow("00FFFF00")
    st = Styles()
    yb = st.xf(fill=st.fill("FFFF00"))
    v = run(_y47([c("A1", 0, "Fill FFFFFF00 = key outputs"), c("B5", yb, 1.0)], st, name="r211a.xlsx"), C47)
    assert v["decision"] == "pass" and v["stats"]["excused_by_legend"], locs(v)
    v = run(_y47([c("A1", 0, "Yellow is not used for inputs; it marks the output cells"), c("B5", yb, 1.0)], st, name="r211b.xlsx"), C47)
    assert v["decision"] == "pass" and v["stats"]["excused_by_legend"], locs(v)


def test_47_rereview_unreadable_text_decisive_only_with_yellow():
    """47-R2-09: a text cell that cannot be decoded (shared-string index out of range) raises only when the
    workbook has bright yellow (the text could be its legend); without yellow the verdict does not depend on it."""
    st = Styles()
    yb = st.xf(fill=st.fill("FFFF00"))
    sst = ['<si><t>x</t></si>']
    raises(lambda: run(book(tmp("r209a.xlsx"), [("S", sheet([c("A1", 0, 7, t="s"), c("B5", yb, 1.0)]))], st, sst=sst), C47),
           "could not be read", "S!A1", "bright yellow")
    v = run(book(tmp("r209b.xlsx"), [("S", sheet([c("A1", 0, 7, t="s"), c("B5", 0, 1.0)]))], st, sst=sst), C47)
    assert v["decision"] == "pass", locs(v)
    d = st.dxf(fill="FFFF00")
    cf = (f'<conditionalFormatting sqref="B2:B9"><cfRule type="cellIs" dxfId="{d}" priority="1" operator="equal">'
          f'<formula>1</formula></cfRule></conditionalFormatting>')
    raises(lambda: run(book(tmp("r209c.xlsx"), [("S", sheet([c("A1", 0, 7, t="s"), c("B2", 0, 1.0)], tail=cf))], st, sst=sst), C47),
           "could not be read")


def test_47_rereview_merged_swatch():
    """47-R2-10: a legend sample merged across 2-3 empty cells is one sample; unmerged adjacent yellow cells
    stay a band."""
    st = Styles()
    yb = st.xf(fill=st.fill("FFFF00"))
    merge = '<mergeCells count="1"><mergeCell ref="A4:B4"/></mergeCells>'
    cells = [c("A2", 0, "Legend"), c("A4", yb), c("B4", yb), c("C4", 0, "Key output cells"), c("H10", yb, 1.0)]
    v = run(_y47(cells, st, name="r210a.xlsx", sheet_kw={"tail": merge}), C47)
    assert v["decision"] == "pass" and v["stats"]["swatches"] == 1 and v["stats"]["bright_cells"] == 3, (locs(v), v["stats"])
    v = run(_y47(cells, st, name="r210b.xlsx"), C47)
    assert v["decision"] == "fail" and locs(v) == ["S!A4:B4", "S!H10"], locs(v)
    # colour-word label, three merged cells
    merge3 = '<mergeCells count="1"><mergeCell ref="A4:C4"/></mergeCells>'
    v = run(_y47([c("A4", yb), c("B4", yb), c("C4", yb), c("D4", 0, "Yellow fill = key outputs"), c("H10", yb, 1.0)], st,
                 name="r210c.xlsx", sheet_kw={"tail": merge3}), C47)
    assert v["decision"] == "pass" and v["stats"]["swatches"] == 1, locs(v)
    # a merged yellow title (text inside) and a four-cell merged band are highlighting
    v = run(_y47([c("B2", yb, "Total"), c("C2", yb), c("D2", yb)], st, name="r210d.xlsx",
                 sheet_kw={"tail": '<mergeCells count="1"><mergeCell ref="B2:D2"/></mergeCells>'}), C47)
    assert v["decision"] == "fail" and locs(v) == ["S!B2:D2"], locs(v)
    v = run(_y47([c("A2", 0, "Legend"), c("A4", yb), c("B4", yb), c("C4", yb), c("D4", yb), c("E4", 0, "Key output cells")], st,
                 name="r210e.xlsx", sheet_kw={"tail": '<mergeCells count="1"><mergeCell ref="A4:D4"/></mergeCells>'}), C47)
    assert v["decision"] == "fail" and locs(v) == ["S!A4:D4"], locs(v)
    # a merged sample that is not isolated (yellow data cell right of its label row) is still judged by its label
    v = run(_y47([c("A4", yb), c("B4", yb), c("C4", 0, "Working notes")], st, name="r210f.xlsx", sheet_kw={"tail": merge}), C47)
    assert v["decision"] == "fail" and locs(v) == ["S!A4:B4"], locs(v)


def _cfx(sqref, dxf, op="greaterThan", f="0", prio=1):
    return (f'<conditionalFormatting sqref="{sqref}"><cfRule type="cellIs" dxfId="{dxf}" priority="{prio}" '
            f'operator="{op}"><formula>{escape(f)}</formula></cfRule></conditionalFormatting>')


def test_47_overlapping_col_styles_later_entry_wins():
    """Column styles through the shared <col> reading (SheetHead.col_segments, later entry wins): a
    bright-yellow column style overridden by a later entry paints only the columns it still covers;
    fully overridden, it paints nothing.  Before 2026-10-04 the check read the raw entries (C:H yellow + E:E
    plain failed as C:H; a yellow E:E overridden by a later plain E:E failed at E)."""
    st = Styles()
    yb = st.xf(fill=st.fill("FFFF00"))
    plain = st.xf()
    xml = (f'<cols><col min="3" max="8" width="9" style="{yb}"/><col min="5" max="5" width="9" style="{plain}"/></cols>'
           + sheet([c("A1", 0, 1.0)]))
    v = run(book(tmp("colover.xlsx"), [("S", xml)], st), C47)
    assert v["decision"] == "fail" and locs(v) == ["S!C:D", "S!F:H"], locs(v)
    assert v["stats"]["col_style_hits"] == 2
    xml = (f'<cols><col min="5" max="5" width="9" style="{yb}"/><col min="5" max="5" width="9" style="{plain}"/></cols>'
           + sheet([c("A1", 0, 1.0)]))
    v = run(book(tmp("colover2.xlsx"), [("S", xml)], st), C47)
    assert v["decision"] == "pass" and v["stats"]["col_style_hits"] == 0, locs(v)
    # an unresolvable fill colour on an overridden column style is never examined (no fork, no raise)
    unk = st.xf(fill=st.fill(fg_attr='rgb="FFGGFF00"'))
    xml = (f'<cols><col min="3" max="3" width="9" style="{unk}"/><col min="3" max="3" width="9" style="{plain}"/></cols>'
           + sheet([c("A1", 0, 1.0)]))
    v = run(book(tmp("colover3.xlsx"), [("S", xml)], st), C47)
    assert v["decision"] == "pass" and "unresolved_fill_colours" not in v["stats"], v["stats"]


def test_47_unresolvable_fill_colours():
    """Finding 47-R2-08 (2026-10-04) and Patrick 2026-10-05 (every attempt graded): an unresolvable fill colour
    (rgb FFGGFF00, theme index 20 ...) is read as Excel's default for a fill, none - recorded per examined
    position in stats.defaults.unresolved_colour (until 2026-10-05 every reading was graded and the check
    raised when they disagreed)."""
    st = Styles()
    unk = st.xf(fill=st.fill(fg_attr='rgb="FFGGFF00"'))
    yb = st.xf(fill=st.fill("FFFF00"))
    # alone, with no legend: no fill -> pass (it raised "verdict depends on" until 2026-10-05)
    v = run(_y47([c("B2", unk, 5.0)], st, name="u1.xlsx"), C47)
    assert v["decision"] == "pass", v
    d = v["stats"]["defaults"]["unresolved_colour"]
    assert d["count"] == 1 and "S!B2" in d["examples"][0] and "rgb:FFGGFF00" in d["examples"][0], d
    assert "every attempt graded" in d["rule"] and "fill as none" in d["rule"], d
    assert v["stats"]["unresolved_fill_colours"] == [{"colour": "rgb:FFGGFF00", "why": "rgb is not 6 or 8 hex digits",
                                                      "first_used": "S!B2 (cell style 1)"}], v["stats"]
    # a certain bright-yellow offender elsewhere still fails
    v = run(_y47([c("B2", unk, 5.0), c("D9", yb, 1.0)], st, name="u2.xlsx"), C47)
    assert v["decision"] == "fail" and locs(v) == ["S!D9"], (locs(v), v["stats"])
    # a TBD cell in the unknown colour is not bright yellow (no fill)
    v = run(_y47([c("A1", 0, "Yellow cells are inputs"), c("B2", unk, "TBD")], st, name="u4.xlsx"), C47)
    assert v["decision"] == "pass", v
    # never recorded: a style no cell uses, a solid fill's unused bgColor, an unresolvable FONT colour
    st2 = Styles()
    st2.xf(fill=st2.fill(fg_attr='theme="40"'))                              # defined, never used
    st2.fills.append('<fill><patternFill patternType="solid"><fgColor rgb="FFDDEBF7"/><bgColor rgb="FFGGFF00"/>'
                     '</patternFill></fill>')
    used = st2.xf(fill=len(st2.fills) - 1)
    fnt = st2.xf(font=st2.font('indexed="81"'))
    v = run(_y47([c("A1", used, 1.0), c("A2", fnt, 2.0)], st2, name="u5.xlsx"), C47)
    assert v["decision"] == "pass" and "unresolved_fill_colours" not in v["stats"], v["stats"]
    assert "defaults" not in v["stats"], v["stats"]
    # a conditional format whose dxf bgColor cannot be resolved: the rule is off for this check, recorded
    # (Patrick 2026-10-05: conditional formats never stop a grading)
    st3 = Styles()
    d_unk = st3.dxf(fill_xml='<patternFill><bgColor theme="25"/></patternFill>')
    d_fg = st3.dxf(fill_xml='<patternFill patternType="solid"><fgColor rgb="FFGGFF00"/><bgColor rgb="FF0000FF"/></patternFill>')
    v = run(_y47([c("B2", 0, 1.0)], st3, name="u6.xlsx", sheet_kw={"tail": _cfx("B2:B5", d_unk)}), C47)
    assert v["decision"] == "pass" and "unresolved_fill_colours" not in v["stats"], v["stats"]
    assert v["stats"]["cf_assumptions"]["rules"] == 1 and "theme:25" in v["stats"]["cf_assumptions"]["examples"][0], \
        v["stats"]["cf_assumptions"]
    assert "S!B2:B5" in v["stats"]["cf_assumptions"]["examples"][0]
    v = run(_y47([c("B2", 0, 1.0)], st3, name="u7.xlsx", sheet_kw={"tail": _cfx("B2:B5", d_fg)}), C47)
    assert v["decision"] == "pass" and "unresolved_fill_colours" not in v["stats"]     # fgColor of a dxf: ignored
    # a column style with an unresolvable fill: no fill, recorded with where it was met
    v = run(_y47([c("A1", 0, 1.0)], st, name="u8.xlsx",
                 sheet_kw={"head": f'<cols><col min="3" max="3" width="9" style="{unk}"/></cols>'}), C47)
    assert v["decision"] == "pass" and "the column style of S!C:C" in \
        v["stats"]["defaults"]["unresolved_colour"]["examples"][0], v["stats"]
    # many distinct unresolvable colours (more than the 4 the readings allowed): all read as no fill
    many = [st.xf(fill=st.fill(fg_attr=f'theme="{20 + k}"')) for k in range(1, 6)]
    v = run(_y47([c("B2", unk, 5.0)] + [c(f"C{2 + k}", x, 1.0) for k, x in enumerate(many)] +
                 [c("D9", yb, 1.0)], st, name="u10.xlsx"), C47)
    assert v["decision"] == "fail" and locs(v) == ["S!D9"] and len(v["stats"]["unresolved_fill_colours"]) == 6, v
    assert v["stats"]["defaults"]["unresolved_colour"]["count"] == 6, v["stats"]["defaults"]


def test_94_unresolvable_colours():
    """Finding 47-R2-08 for No white-on-white hiding (94) and Patrick 2026-10-05 (every attempt graded): an
    unresolvable font / fill / conditional-format / rich-run / number-format-tag colour is read as Excel's
    default for its slot - a font colour as automatic (black), a fill as none - and recorded in
    stats.defaults.unresolved_colour (until 2026-10-05 such a cell was undecided and could raise)."""
    st = Styles()
    uf = st.xf(font=st.font('rgb="FFGGFF00"'))
    ufill = st.xf(fill=st.fill(fg_attr='theme="30"'))
    hid = st.xf(numfmt=";;;")
    uf_hid = st.xf(font=st.font('indexed="81"'), numfmt=";;;")
    dark = st.xf(font=st.font('rgb="FFGGFF00"'), fill=st.fill("000000"))
    v = run(_w94([c("A1", uf, 5.0)], st, "w1.xlsx"), C94)               # black on white: visible
    assert v["decision"] == "pass" and v["stats"]["undecided_cells"] == 0, v["stats"]
    d = v["stats"]["defaults"]["unresolved_colour"]
    assert d["count"] == 1 and "S!A1" in d["examples"][0] and "rgb:FFGGFF00" in d["examples"][0], d
    v = run(_w94([c("A1", ufill, "label")], st, "w2.xlsx"), C94)        # no fill: black on white
    assert v["decision"] == "pass" and "theme:30" in v["stats"]["defaults"]["unresolved_colour"]["examples"][0], v
    v = run(_w94([c("A1", dark, 5.0)], st, "w2b.xlsx"), C94)            # read as black text on a black fill
    assert v["decision"] == "fail" and locs(v) == ["S!A1"], v
    v = run(_w94([c("A1", uf, 5.0), c("B1", hid, 3.0)], st, "w3.xlsx"), C94)
    assert v["decision"] == "fail" and locs(v) == ["S!B1"] and v["stats"]["undecided_cells"] == 0, v["stats"]
    # empty cells and a format that prints nothing anyway; styles no cell uses are never recorded
    st.xf(font=st.font('theme="44"'), fill=st.fill(fg_attr='indexed="99"'))         # unused
    v = run(_w94([c("A1", uf), c("A2", uf, "", f='""'), c("A3", uf_hid, 7.0), c("A4", 0, 1.0)], st, "w4.xlsx"), C94)
    assert v["decision"] == "fail" and locs(v) == ["S!A3"] and v["stats"]["undecided_cells"] == 0, v["stats"]
    # a conditional-format font colour that cannot be resolved: automatic (black), recorded
    st2 = Styles()
    st2.dxfs.append('<dxf><font><color theme="77"/></font></dxf>')
    d = len(st2.dxfs) - 1
    v = run(_w94([c("A1", 0, 5.0)], st2, "w5.xlsx", sheet_kw={"tail": _cfx("A1:A3", d)}), C94)
    assert v["decision"] == "pass" and v["stats"]["undecided_cells"] == 0, v["stats"]
    assert "theme:77" in v["stats"]["defaults"]["unresolved_colour"]["examples"][0], v["stats"]["defaults"]
    v = run(_w94([c("A1", 0, 5.0)], st2, "w6.xlsx", sheet_kw={"tail": _cfx("A1:A3", d, "equal", "99")}), C94)
    assert v["decision"] == "pass" and v["stats"]["undecided_cells"] == 0, v["stats"]
    # a colour scale with an unresolvable stop: a built-in format never hides text
    scale = ('<conditionalFormatting sqref="A1:A3"><cfRule type="colorScale" priority="1"><colorScale>'
             '<cfvo type="min"/><cfvo type="max"/><color rgb="FFFFFFFF"/><color rgb="FFGGFF00"/></colorScale>'
             '</cfRule></conditionalFormatting>')
    v = run(_w94([c("A1", 0, 1.0), c("A2", 0, 2.0)], st2, "w7.xlsx", sheet_kw={"tail": scale}), C94)
    assert v["decision"] == "pass" and v["stats"]["undecided_cells"] == 0, v["stats"]
    # a rich-text run whose colour cannot be resolved: black, visible
    sst = ['<si><r><t>Visible </t></r><r><rPr><color theme="50"/></rPr><t>maybe</t></r></si>']
    v = run(book(tmp("w8.xlsx"), [("S", sheet([c("A1", 0, 0, t="s")]))], Styles(), sst=sst), C94)
    assert v["decision"] == "pass" and "rich-text run colour" in v["stats"]["defaults"]["unresolved_colour"]["examples"][0], v
    # a [Color10] tag naming an invalid custom palette entry (17): black, visible
    st3 = Styles()
    tag = st3.xf(numfmt="[Color10]0")
    pal = '<colors><indexedColors>' + '<rgbColor rgb="FF000000"/>' * 17 + '<rgbColor rgb="nope"/></indexedColors></colors>'
    p = book(tmp("w9.xlsx"), [("S", sheet([c("A1", tag, 5.0)]))], st3)
    with zipfile.ZipFile(p) as z:
        parts = {n: z.read(n) for n in z.namelist()}
    parts["xl/styles.xml"] = parts["xl/styles.xml"].replace(b"</styleSheet>", pal.encode() + b"</styleSheet>")
    with zipfile.ZipFile(p, "w") as z:
        for n, b in parts.items():
            z.writestr(n, b)
    v = run(p, C94)
    assert v["decision"] == "pass" and "[Color10]" in v["stats"]["defaults"]["unresolved_colour"]["examples"][0], v


def test_94_unverified_number_format():
    """Finding 66-S3 in the core: a code mixing digit placeholders with unquoted date letters ('0 days') renders
    certain=False; 94 used to ignore that flag.  Now both blank / shown and every colour the code could show are
    tried: undecided, raising only when nothing else is concealed."""
    st = Styles()
    days = st.xf(numfmt="0 days")
    hid = st.xf(numfmt=";;;")
    graded("unverified_number_format", lambda: run(_w94([c("A1", days, 25.0)], st, "nf1.xlsx"), C94), "cannot decide whether S!A1", "'0 days'")
    v = run(_w94([c("A1", days, 25.0), c("B1", hid, 1.0)], st, "nf2.xlsx"), C94)
    assert v["decision"] == "fail" and locs(v) == ["S!B1"] and v["stats"]["undecided_cells"] == 0, v["stats"]
    # text under a code without a text section is shown as typed whatever Excel makes of the code (certain)
    v = run(_w94([c("A1", days, "text"), c("A2", 0, 3.0)], st, "nf3.xlsx"), C94)
    assert v["decision"] == "pass" and v["stats"]["undecided_cells"] == 0, v
    v = run(_w94([c("A1", days)], st, "nf4.xlsx"), C94)                                  # empty: never judged
    assert v["decision"] == "pass", v


TESTS = [test_47_colour_band, test_47_legend_classifier, test_47_own_fill_pale_and_bright,
         test_47_theme_indexed_gradient, test_47_row_col_default_styles, test_47_conditional_formats,
         test_47_spill_member_needle, test_47_legend_excuses_and_wip_override, test_47_swatches,
         test_94_contrast_anchors, test_94_cf_rule_evaluation, test_94_font_against_own_background,
         test_94_number_formats, test_94_rich_text_runs, test_94_conditional_formatting,
         test_94_values_needed_only_where_formatting_could_conceal, test_94_table_style_and_hidden_row,
         test_47_legend_swatch_rulings_2026_10_06,
         test_47_review_wip_vocab, test_47_review_hex_colour_label, test_47_review_swatch_label_promotion,
         test_47_review_selflabel_wip_note, test_47_review_positive_legends, test_47_review_left_and_self_keys,
         test_47_review_ambiguous_cf_only_when_decisive, test_47_review_replaced_and_wip_blocks,
         test_94_review_undecided_only_when_decisive, test_94_review_hue_pairs, test_94_review_conditional_sections,
         test_94_review_rich_run_cf_fill, test_94_review_containstext_wildcards, test_94_review_dark_table_style,
         test_excel_2026_10_03_dxf_fill_is_bgcolor_only, test_excel_2026_10_03_cf_font_beats_format_colour,
         test_excel_2026_10_03_show_zeros_off_94,
         test_rulings_2026_10_03_banners_and_near_empty,
         test_47_rereview_swatch_between_colour_word_and_meaning, test_47_rereview_status_word_is_not_a_legend,
         test_47_rereview_back_reference_defines_nothing, test_47_rereview_section_header_is_not_a_colour_key,
         test_47_rereview_wip_vocabulary, test_47_rereview_classifier_phrasings,
         test_47_rereview_unreadable_text_decisive_only_with_yellow, test_47_rereview_merged_swatch,
         test_47_unresolvable_fill_colours, test_94_unresolvable_colours, test_94_unverified_number_format,
         test_47_overlapping_col_styles_later_entry_wins]


def main() -> int:
    failed = 0
    try:
        for t in TESTS:
            try:
                t()
                print(f"PASS {t.__name__}")
            except Exception:  # noqa: BLE001
                failed += 1
                print(f"FAIL {t.__name__}")
                traceback.print_exc()
    finally:
        if _TMP and os.path.isdir(_TMP):
            shutil.rmtree(_TMP, ignore_errors=True)
    print(f"\n{len(TESTS) - failed}/{len(TESTS)} tests passed")
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main())
