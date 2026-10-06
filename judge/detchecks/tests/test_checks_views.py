"""Unit tests for checks 61 (Consistent zoom level), 62 (Active cell reset to A1 on all sheets)
and 77 (File extension (.xlsx)) on synthetic micro-workbooks (raw SpreadsheetML zips and
hand-built OLE2 / CSV / zip bytes for 77).

    cd judge && python -m detchecks.tests.test_checks_views

Plain asserts; also collectable by pytest.  Temporary files go to detchecks/scratch/views/.
"""
from __future__ import annotations

import os
import re
import shutil
import struct
import sys
import tempfile
import traceback
import zipfile
from xml.sax.saxutils import quoteattr

from detchecks.api import grade
from detchecks.checks import c61, c62, c77
from detchecks.errors import GradingError

HERE = os.path.dirname(os.path.abspath(__file__))
SCRATCH = os.path.join(os.path.dirname(HERE), "scratch", "views")
MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
MSREL = "http://schemas.microsoft.com/office/2006/relationships"
PKG_REL = "http://schemas.openxmlformats.org/package/2006/relationships"
CT = {
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml",
    "xlsm": "application/vnd.ms-excel.sheet.macroEnabled.main+xml",
    "xltx": "application/vnd.openxmlformats-officedocument.spreadsheetml.template.main+xml",
    "generic": "application/xml",
}
CT_WS = "application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"
CT_CS = "application/vnd.openxmlformats-officedocument.spreadsheetml.chartsheet+xml"
CT_DS = "application/vnd.openxmlformats-officedocument.spreadsheetml.dialogsheet+xml"
CT_MS = "application/vnd.ms-excel.macrosheet+xml"
CT_VBA = "application/vnd.ms-office.vbaProject"
DECL = '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\r\n'
OLE2_MAGIC = bytes.fromhex("d0cf11e0a1b11ae1")

K61 = c61.C61.key
K62 = c62.C62.key
K77 = c77.C77.key

_TMP = None


def tmp(name: str) -> str:
    global _TMP
    if _TMP is None:
        os.makedirs(SCRATCH, exist_ok=True)
        _TMP = tempfile.mkdtemp(prefix="test_views_", dir=SCRATCH)
    return os.path.join(_TMP, name)


# ============================================================================ builders
def view(zoom=None, *, normal=None, page_layout=None, sheet_layout=None, mode=None, pane=None, sels=(),
         wv=0, top_left=None, extra="") -> str:
    """One <sheetView>.  pane: (xSplit, ySplit, topLeftCell, activePane, state) or None.
    sels: [(pane or None, activeCell or None, sqref or None)]."""
    a = [f'workbookViewId="{wv}"']
    for k, v in (("zoomScale", zoom), ("zoomScaleNormal", normal), ("zoomScalePageLayoutView", page_layout),
                 ("zoomScaleSheetLayoutView", sheet_layout), ("view", mode), ("topLeftCell", top_left)):
        if v is not None:
            a.append(f"{k}={quoteattr(str(v))}")
    body = ""
    if pane is not None:
        xs, ys, tl, ap, st = pane
        pa = []
        if xs:
            pa.append(f'xSplit="{xs}"')
        if ys:
            pa.append(f'ySplit="{ys}"')
        if tl:
            pa.append(f'topLeftCell="{tl}"')
        if ap:
            pa.append(f'activePane="{ap}"')
        if st:
            pa.append(f'state="{st}"')
        body += f"<pane {' '.join(pa)}/>"
    for p, ac, sq in sels:
        sa = []
        if p:
            sa.append(f'pane="{p}"')
        if ac:
            sa.append(f'activeCell="{ac}"')
        if sq:
            sa.append(f'sqref="{sq}"')
        body += f"<selection {' '.join(sa)}/>"
    return f"<sheetView {' '.join(a)}{extra}>{body}</sheetView>"


class S:
    """A sheet for book(): kind worksheet/chartsheet/dialogsheet/macrosheet, its sheetViews XML."""

    def __init__(self, name, views="", kind="worksheet", state=None):
        self.name, self.views, self.kind, self.state = name, views, kind, state


NOSTREAM = 0xFFFFFFFF


def cfb_bytes(entries, *, dir_start: int = 1) -> bytes:
    """A minimal Compound File Binary (v3, 512-byte sectors) with the given directory entries
    [(name, type, left, right, child)] (index 0 = Root Entry, type 5; 1 = storage, 2 = stream;
    all streams empty).  Sector 0 = FAT, sectors 1.. = directory.  dir_start points the
    directory elsewhere (e.g. beyond the end of the file)."""
    n_dir = max(1, (len(entries) + 3) // 4)
    hdr = bytearray(512)
    hdr[0:8] = OLE2_MAGIC
    struct.pack_into("<HHHHH", hdr, 0x18, 0x3E, 3, 0xFFFE, 9, 6)
    struct.pack_into("<I", hdr, 0x2C, 1)
    struct.pack_into("<I", hdr, 0x30, dir_start)
    struct.pack_into("<I", hdr, 0x38, 4096)
    struct.pack_into("<II", hdr, 0x3C, 0xFFFFFFFE, 0)
    struct.pack_into("<II", hdr, 0x44, 0xFFFFFFFE, 0)
    struct.pack_into("<109I", hdr, 0x4C, *([0] + [0xFFFFFFFF] * 108))
    fat = bytearray(b"\xff" * 512)
    chain = [0xFFFFFFFD] + [i + 2 for i in range(n_dir - 1)] + [0xFFFFFFFE]
    struct.pack_into(f"<{len(chain)}I", fat, 0, *chain)
    dirs = bytearray(512 * n_dir)
    for i, (name, typ, left, right, child) in enumerate(entries):
        e = bytearray(128)
        raw_name = name.encode("utf-16-le") + b"\0\0"
        e[0:len(raw_name)] = raw_name
        struct.pack_into("<H", e, 0x40, len(raw_name))
        e[0x42] = typ
        struct.pack_into("<III", e, 0x44, left, right, child)
        dirs[i * 128:(i + 1) * 128] = e
    return bytes(hdr) + bytes(fat) + bytes(dirs)


# a real VBA project's shape (MS-OVBA 2.2): Root -> {VBA (storage) -> {dir, _VBA_PROJECT}, PROJECT}
VBA_TREE = [("Root Entry", 5, NOSTREAM, NOSTREAM, 1), ("VBA", 1, NOSTREAM, 2, 3), ("PROJECT", 2, NOSTREAM, NOSTREAM, NOSTREAM),
            ("dir", 2, NOSTREAM, 4, NOSTREAM), ("_VBA_PROJECT", 2, NOSTREAM, NOSTREAM, NOSTREAM)]
VBA_BINS = {
    "ole2": cfb_bytes(VBA_TREE),                                            # a VBA project
    "junk": b"not a compound file",
    "magic": OLE2_MAGIC + b"\0" * 504,                                      # signature only, broken header
    "xls": cfb_bytes([("Root Entry", 5, NOSTREAM, NOSTREAM, 1), ("Workbook", 2, NOSTREAM, NOSTREAM, NOSTREAM)]),
    "nodir": cfb_bytes([("Root Entry", 5, NOSTREAM, NOSTREAM, 1), ("VBA", 1, NOSTREAM, 2, 3),
                        ("PROJECT", 2, NOSTREAM, NOSTREAM, NOSTREAM), ("_VBA_PROJECT", 2, NOSTREAM, NOSTREAM, NOSTREAM)]),
    "noproject": cfb_bytes([("Root Entry", 5, NOSTREAM, NOSTREAM, 1), ("VBA", 1, NOSTREAM, NOSTREAM, 2),
                            ("dir", 2, NOSTREAM, NOSTREAM, NOSTREAM)]),
    "badtree": cfb_bytes(VBA_TREE, dir_start=50),                           # directory beyond the end
}


def book(path: str, sheets, *, main_ct: str = "xlsx", vba: str | None = None, vba_related: bool = True,
         macrosheet: bool = False, app: str | None = None) -> str:
    """Write a minimal workbook package.  vba: None | a VBA_BINS key ('ole2' = a real VBA
    project structure; 'junk' = not OLE2; 'magic' = OLE2 signature + zeros; 'xls' = a
    compound file with a Workbook stream; 'nodir' / 'noproject' = VBA structure incomplete;
    'badtree' = valid header, unwalkable directory).  macrosheet: add an XLM macro sheet
    related from the workbook."""
    z = zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED)
    overrides = [f'<Override PartName="/xl/workbook.xml" ContentType="{CT.get(main_ct, main_ct)}"/>']
    wb_rels, sh = [], []
    sheets = list(sheets)
    if macrosheet:
        sheets.append(S("Macro1", "", kind="macrosheet"))
    for i, s in enumerate(sheets, 1):
        folder, root, ct, rtype = {
            "worksheet": ("worksheets", "worksheet", CT_WS, f"{REL}/worksheet"),
            "chartsheet": ("chartsheets", "chartsheet", CT_CS, f"{REL}/chartsheet"),
            "dialogsheet": ("dialogsheets", "dialogsheet", CT_DS, f"{REL}/dialogsheet"),
            "macrosheet": ("macrosheets", "worksheet", CT_MS, f"{MSREL}/xlMacrosheet"),
        }[s.kind]
        part = f"xl/{folder}/sheet{i}.xml"
        overrides.append(f'<Override PartName="/{part}" ContentType="{ct}"/>')
        wb_rels.append(f'<Relationship Id="rId{i}" Type="{rtype}" Target="{folder}/sheet{i}.xml"/>')
        st = f' state="{s.state}"' if s.state else ""
        sh.append(f'<sheet name={quoteattr(s.name)} sheetId="{i}"{st} r:id="rId{i}"/>')
        views = f"<sheetViews>{s.views}</sheetViews>" if s.views else ""
        if s.kind == "chartsheet":
            xml = f'{DECL}<chartsheet xmlns="{MAIN}" xmlns:r="{REL}">{views}</chartsheet>'
        elif s.kind == "dialogsheet":
            xml = f'{DECL}<dialogsheet xmlns="{MAIN}" xmlns:r="{REL}">{views}</dialogsheet>'
        else:
            xml = (f'{DECL}<{root} xmlns="{MAIN}" xmlns:r="{REL}">{views}'
                   f'<sheetData><row r="1"><c r="A1" t="inlineStr"><is><t>x</t></is></c></row></sheetData></{root}>')
        z.writestr(part, xml)
    if vba:
        z.writestr("xl/vbaProject.bin", VBA_BINS[vba])
        overrides.append(f'<Override PartName="/xl/vbaProject.bin" ContentType="{CT_VBA}"/>')
        if vba_related:
            wb_rels.append(f'<Relationship Id="rId900" Type="{MSREL}/vbaProject" Target="vbaProject.bin"/>')
    z.writestr("[Content_Types].xml",
               '<?xml version="1.0" encoding="UTF-8"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
               '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
               '<Default Extension="xml" ContentType="application/xml"/>' + "".join(overrides) + "</Types>")
    z.writestr("_rels/.rels", f'<Relationships xmlns="{PKG_REL}"><Relationship Id="rId1" Type="{REL}/officeDocument" '
                              f'Target="xl/workbook.xml"/></Relationships>')
    z.writestr("xl/_rels/workbook.xml.rels", f'<Relationships xmlns="{PKG_REL}">{"".join(wb_rels)}</Relationships>')
    z.writestr("xl/workbook.xml", f'{DECL}<workbook xmlns="{MAIN}" xmlns:r="{REL}"><sheets>{"".join(sh)}</sheets></workbook>')
    if app:
        z.writestr("docProps/app.xml", f'<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/extended-properties">'
                                       f'<Application>{app}</Application></Properties>')
    z.close()
    return path


def ole2(path: str, streams=("Workbook",)) -> str:
    """A minimal Compound File Binary (v3, 512-byte sectors): sector 0 = FAT, sector 1 =
    directory (Root Entry + the given stream names, all empty)."""
    hdr = bytearray(512)
    hdr[0:8] = OLE2_MAGIC
    struct.pack_into("<HHHHH", hdr, 0x18, 0x3E, 3, 0xFFFE, 9, 6)
    struct.pack_into("<I", hdr, 0x2C, 1)            # FAT sectors
    struct.pack_into("<I", hdr, 0x30, 1)            # first directory sector
    struct.pack_into("<I", hdr, 0x38, 4096)
    struct.pack_into("<II", hdr, 0x3C, 0xFFFFFFFE, 0)
    struct.pack_into("<II", hdr, 0x44, 0xFFFFFFFE, 0)
    difat = [0] + [0xFFFFFFFF] * 108
    struct.pack_into("<109I", hdr, 0x4C, *difat)
    fat = bytearray(b"\xff" * 512)
    struct.pack_into("<II", fat, 0, 0xFFFFFFFD, 0xFFFFFFFE)
    dirs = bytearray(512)
    for i, (name, typ) in enumerate([("Root Entry", 5)] + [(s, 2) for s in streams][:3]):
        e = bytearray(128)
        raw = name.encode("utf-16-le") + b"\0\0"
        e[0:len(raw)] = raw
        struct.pack_into("<H", e, 0x40, len(raw))
        e[0x42] = typ
        struct.pack_into("<III", e, 0x44, 0xFFFFFFFF, 0xFFFFFFFF, 0xFFFFFFFF)
        dirs[i * 128:(i + 1) * 128] = e
    with open(path, "wb") as fh:
        fh.write(bytes(hdr) + bytes(fat) + bytes(dirs))
    return path


def raw(path: str, data: bytes) -> str:
    with open(path, "wb") as fh:
        fh.write(data)
    return path


def late_views(path: str, n: int, views: str | None = None, *, keep_head: bool = False) -> str:
    """Rewrite worksheet n (1-based) of a book() workbook so that a <sheetViews> follows
    </sheetData> (schema-invalid element order).  By default the sheet's own <sheetViews> is
    moved there; `views` (one or more <sheetView> elements) puts that instead, and keep_head
    leaves the head's <sheetViews> in place as well."""
    part = f"xl/worksheets/sheet{n}.xml"
    out = path + ".tmp"
    with zipfile.ZipFile(path) as zi, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zo:
        for it in zi.infolist():
            data = zi.read(it.filename)
            if it.filename == part:
                xml = data.decode("utf-8")
                m = re.search(r"<sheetViews>.*?</sheetViews>", xml, re.S)
                head_views = m.group(0) if m else ""
                if head_views and not keep_head:
                    xml = xml.replace(head_views, "", 1)
                late = f"<sheetViews>{views}</sheetViews>" if views is not None else head_views
                assert late and "</sheetData>" in xml, (part, xml)
                xml = xml.replace("</sheetData>", "</sheetData>" + late, 1)
                data = xml.encode("utf-8")
            zo.writestr(it, data)
    os.replace(out, path)
    return path


def run(path, number, task_meta=None):
    return grade(path, checks=[number], task_meta=task_meta)[{61: K61, 62: K62, 77: K77}[number]]


def locs(v):
    return [m["location"] for m in v["mistakes"]]


def raises(path, number, needle="", task_meta=None):
    try:
        v = run(path, number, task_meta)
    except GradingError as e:
        msg = "\n".join(e.failures.values()) or str(e)
        assert needle in msg, f"GradingError without {needle!r}: {msg}"
        return msg
    raise AssertionError(f"expected GradingError, got {v['decision']}: {v['summary']}")


# ============================================================================ 61
def test_61_same_level_passes_and_mixed_fails():
    p = book(tmp("z_same.xlsx"), [S("A", view(80)), S("B", view(80, normal=80)), S("C", view(80))])
    v = run(p, 61)
    assert v["decision"] == "pass" and v["stats"]["reference_zoom"] == 80, v
    # T2 trap: 85 vs 80 on one sheet among many - no tolerance
    p = book(tmp("z_85.xlsx"), [S(f"S{i}", view(85 if i == 7 else 80)) for i in range(10)])
    v = run(p, 61)
    assert v["decision"] == "fail" and locs(v) == ["S7"], v
    assert "85%" in v["mistakes"][0]["description"] and "9 of 10" in v["mistakes"][0]["description"]
    # 79 vs 80 are two levels (ZOOM_TOLERANCE = 0)
    v = run(book(tmp("z_79.xlsx"), [S("A", view(79)), S("B", view(80)), S("C", view(80))]), 61)
    assert locs(v) == ["A"], v


def test_61_missing_zoomscale_is_100():
    # T3 trap: a sheet without zoomScale opens at 100%, so {90, 90, missing} fails ...
    p = book(tmp("z_missing.xlsx"), [S("A", view(90)), S("B", view(90)), S("EV", view())])
    v = run(p, 61)
    assert v["decision"] == "fail" and locs(v) == ["EV"], v
    assert "default 100%" in v["mistakes"][0]["description"]
    # ... while {explicit 100, zoomScaleNormal-only 100, nothing, no sheetView} passes
    p = book(tmp("z_100.xlsx"), [S("A", view(100)), S("B", view(normal=100)), S("C", view()), S("D", "")])
    v = run(p, 61)
    assert v["decision"] == "pass" and v["stats"]["reference_zoom"] == 100, v
    assert sorted(v["stats"]["sheets_without_zoomScale"]) == ["B", "C", "D"]
    # zoomScale="0" is treated as absent (100)
    v = run(book(tmp("z_zero.xlsx"), [S("A", view(0)), S("B", view(100))]), 61)
    assert v["decision"] == "pass", v


def test_61_zoomscalenormal_never_decides():
    # openpyxl shape: zoomScaleNormal=80 without zoomScale opens at 100 -> fail
    p = book(tmp("z_normal_only.xlsx"), [S("A", view(80)), S("B", view(normal=80)), S("C", view(80))])
    v = run(p, 61)
    assert locs(v) == ["B"], v
    assert "zoomScaleNormal=80 is not the opening zoom" in v["mistakes"][0]["description"]
    assert v["stats"]["other_reading"]["decision"] == "pass" and v["stats"]["reading_sensitive"] is True
    # mirror: zoomScale=80 + zoomScaleNormal=100 opens at 80 -> pass
    p = book(tmp("z_mirror.xlsx"), [S("A", view(80)), S("B", view(80, normal=100)), S("C", view(80))])
    v = run(p, 61)
    assert v["decision"] == "pass" and v["stats"]["zoomScaleNormal_differs"] == {"B": "100"}, v


def test_61_view_modes():
    # page break preview at 80 (its per-mode memories elsewhere) among normal 80s passes ...
    p = book(tmp("z_pbp80.xlsx"), [S("A", view(80)), S("B", view(80, normal=100, sheet_layout=80, mode="pageBreakPreview"))])
    v = run(p, 61)
    assert v["decision"] == "pass" and v["stats"]["non_normal_views"] == {"B": "pageBreakPreview"}, v
    # ... a page break preview opening at 60 fails, a page layout view at 80 passes
    p = book(tmp("z_pbp60.xlsx"), [S("A", view(80)), S("B", view(80)),
                                   S("C", view(60, normal=80, sheet_layout=60, mode="pageBreakPreview")),
                                   S("D", view(80, page_layout=80, mode="pageLayout"))])
    v = run(p, 61)
    assert locs(v) == ["C"] and "pageBreakPreview view" in v["mistakes"][0]["description"], v


def test_61_scope_hidden_chart_windows():
    # hidden and very hidden sheets count (whole workbook)
    p = book(tmp("z_hidden.xlsx"), [S("A", view(80)), S("B", view(80)), S("H", view(100), state="hidden"),
                                    S("VH", view(80), state="veryHidden")])
    v = run(p, 61)
    assert locs(v) == ["H"] and "(hidden sheet)" in v["mistakes"][0]["description"], v
    assert v["stats"]["hidden_sheets_counted"] == ["H", "VH"]
    # a chart sheet (no cell grid) at another zoom is not counted; it is listed in stats
    p = book(tmp("z_chart.xlsx"), [S("A", view(80)), S("Chart1", view(100), kind="chartsheet"), S("B", view(80))])
    v = run(p, 61)
    assert v["decision"] == "pass", v
    assert [e["sheet"] for e in v["stats"]["excluded_sheets"]] == ["Chart1"]
    # a second workbook window's view does not decide; window 0 does
    p = book(tmp("z_window.xlsx"), [S("A", view(80) + view(150, wv=1)), S("B", view(80))])
    v = run(p, 61)
    assert v["decision"] == "pass" and v["stats"]["sheets_with_extra_window_views"] == ["A"], v
    # dialog sheets have a grid and count
    p = book(tmp("z_dialog.xlsx"), [S("A", view(80)), S("B", view(80)), S("D1", view(100), kind="dialogsheet")])
    assert locs(run(p, 61)) == ["D1"]


def test_61_tie_and_summary():
    # 2 x 80 then 2 x 100: no majority -> the level of the earliest sheet is the reference
    p = book(tmp("z_tie.xlsx"), [S("A", view(80)), S("B", view(100)), S("C", view(100)), S("D", view(80))])
    v = run(p, 61)
    assert v["stats"]["reference_zoom"] == 80 and locs(v) == ["B", "C"], v
    assert "2 sheet(s) differ" in v["summary"]


def test_61_unreadable_zoom_raises_only_when_it_decides():
    raises(book(tmp("z_bad.xlsx"), [S("A", view("eighty")), S("B", view(80))]), 61, "not an integer zoom")
    raises(book(tmp("z_500.xlsx"), [S("A", view(500)), S("B", view(80))]), 61, "outside the valid zoom range")
    # a bad per-mode memory is stats only
    v = run(book(tmp("z_badnormal.xlsx"), [S("A", view(80, normal="abc")), S("B", view(80))]), 61)
    assert v["decision"] == "pass" and "A" in v["stats"]["unreadable_stats_only_attributes"], v
    # an excluded chart sheet never raises
    v = run(book(tmp("z_badchart.xlsx"), [S("A", view(80)), S("C", view("x"), kind="chartsheet")]), 61)
    assert v["decision"] == "pass", v


def test_61_unknown_zoom_with_certain_fail():
    # review 2026-10-03 (61-raises-although-already-failing): S1 80 vs S2 100 fail whatever S3 is
    for bad in ("abc", 500, 9):
        v = run(book(tmp(f"z_known_fail_{bad}.xlsx"), [S("S1", view(80)), S("S2", view(100)), S("S3", view(bad))]), 61)
        assert v["decision"] == "fail" and locs(v) == ["S2"], (bad, v)
        assert list(v["stats"]["undecidable_sheets"]) == ["S3"] and "S3" not in v["stats"]["zooms"], v
        assert "'S3'" in v["summary"] and "does not depend on it" in v["summary"], v
    # readable sheets agree: the unknown zoom decides -> GradingError
    raises(book(tmp("z_known_same.xlsx"), [S("S1", view(80)), S("S2", view(80)), S("S3", view(500))]), 61,
           "do not already differ")
    # only unknown zooms: GradingError
    raises(book(tmp("z_only_bad.xlsx"), [S("S1", view("x")), S("S2", view(500))]), 61, "unknown")
    # ... unless it is the only counted sheet (a chart sheet does not count): nothing to differ from
    v = run(book(tmp("z_single_bad.xlsx"), [S("S1", view("x")), S("C", view(100), kind="chartsheet")]), 61)
    assert v["decision"] == "pass" and list(v["stats"]["undecidable_sheets"]) == ["S1"], v
    assert "cannot differ from itself" in v["summary"], v
    # with a tolerance the readable spread must exceed 2 x tolerance
    with option(c61, "ZOOM_TOLERANCE", 2):
        raises(book(tmp("z_tol_close.xlsx"), [S("S1", view(79)), S("S2", view(82)), S("S3", view("x"))]), 61,
               "do not already differ")
        v = run(book(tmp("z_tol_far.xlsx"), [S("S1", view(75)), S("S2", view(80)), S("S3", view("x"))]), 61)
        assert v["decision"] == "fail" and list(v["stats"]["undecidable_sheets"]) == ["S3"], v
        # attempt 2624's levels (77/79/80) pass under a 2-point tolerance; the summary says "within"
        v = run(book(tmp("z_tol_2624.xlsx"), [S("Instructions", view(79)), S("Questions", view(77)),
                                              S("Assumptions", view(80))]), 61)
        assert v["decision"] == "pass" and "within 2 points of 79%" in v["summary"], v


def test_61_sheet_views_after_sheet_data_is_unknown():
    # review 2026-10-04 (61-sheetviews-after-sheetdata-silent): a <sheetViews> AFTER <sheetData> is
    # schema-invalid (Excel offers a repair), so that sheet's opening zoom is unknown - not 100.
    # Before the fix: 80/late 80/80 failed naming 'Late' (read as no view), 100/late 80/100 passed.
    # The readable sheets agree, so the unknown zoom decides -> GradingError
    p = late_views(book(tmp("z_late_same.xlsx"), [S("A", view(80)), S("Late", view(80)), S("B", view(80))]), 2)
    msg = raises(p, 61, "follows <sheetData>")
    assert "'Late'" in msg and "do not already differ" in msg, msg
    p = late_views(book(tmp("z_late_100.xlsx"), [S("A", view(100)), S("Late", view(80)), S("B", view(100))]), 2)
    raises(p, 61, "follows <sheetData>")
    # a head <sheetViews> plus a second one after <sheetData>: still unknown (neither 80 nor 100)
    p = late_views(book(tmp("z_late_dup.xlsx"), [S("A", view(80)), S("Late", view(80)), S("B", view(80))]), 2,
                   view(100), keep_head=True)
    raises(p, 61, "follows <sheetData>")
    # the readable sheets already differ -> fail; 'Late' gets no mistake and is listed as unknown
    p = late_views(book(tmp("z_late_diff.xlsx"), [S("A", view(80)), S("Late", view(80)), S("B", view(100))]), 2)
    v = run(p, 61)
    assert v["decision"] == "fail" and locs(v) == ["B"], v
    assert list(v["stats"]["undecidable_sheets"]) == ["Late"] and "Late" not in v["stats"]["zooms"], v
    assert v["stats"]["sheet_views_after_sheet_data"] == ["Late"] and "'Late'" in v["summary"], v
    # the only counted sheet -> pass (one sheet cannot differ from itself)
    v = run(late_views(book(tmp("z_late_single.xlsx"), [S("Late", view(80))]), 1), 61)
    assert v["decision"] == "pass" and list(v["stats"]["undecidable_sheets"]) == ["Late"], v
    # a sheet whose <sheetViews> is in its place is untouched by the detection
    v = run(book(tmp("z_not_late.xlsx"), [S("A", view(80)), S("B", view(80))]), 61)
    assert v["decision"] == "pass" and v["stats"]["sheet_views_after_sheet_data"] == [], v


def test_61_both_attributes_differ_reading_switch():
    # review 2026-10-04 (61-both-attrs-reading-unmeasured): zoomScale and zoomScaleNormal both present
    # and different.  The default reading follows ECMA-376 (zoomScale = zoom of the current view) and
    # the Excel-saved starting files; Excel itself has not been measured on this shape (question for
    # Patrick, docs/checks/61.md).  Pins the default and proves that ZOOM_READING = "per_view" flips
    # exactly this shape, so the switch alone covers the other outcome of the measurement.
    # attempt 1272's Instructions (zoomScale=100, zoomScaleNormal=80, zoomScaleSheetLayoutView=100) among 100s
    p = book(tmp("z_1272.xlsx"), [S("Summary", view(100)), S("Instructions", view(100, normal=80, sheet_layout=100)),
                                  S("DCF", view(100))])
    v = run(p, 61)
    assert v["decision"] == "pass" and v["stats"]["zoomScaleNormal_differs"] == {"Instructions": "80"}, v
    assert v["stats"]["reading_sensitive"] is True and v["stats"]["other_reading"]["deviating_sheets"] == ["Instructions"], v
    with option(c61, "ZOOM_READING", "per_view"):
        v2 = run(p, 61)
    assert v2["decision"] == "fail" and locs(v2) == ["Instructions"], v2
    assert "opens at 80% zoom" in v2["mistakes"][0]["description"] and "per-view reading" in v2["mistakes"][0]["description"], v2
    # an untouched Excel-saved Instructions sheet (zoomScale=79, zoomScaleNormal=80) next to agent sheets at 80
    p = book(tmp("z_79_80.xlsx"), [S("Instructions", view(79, normal=80)), S("Model", view(80)), S("Checks", view(80))])
    v = run(p, 61)
    assert v["decision"] == "fail" and locs(v) == ["Instructions"], v
    assert "zoomScaleNormal=80 is not the opening zoom" in v["mistakes"][0]["description"], v
    with option(c61, "ZOOM_READING", "per_view"):
        assert run(p, 61)["decision"] == "pass"


# ============================================================================ 62
FROZEN3 = (None, 3, "A4", "bottomLeft", "frozen")          # frozen below row 3


def test_62_a1_and_defaults_pass():
    p = book(tmp("c_ok.xlsx"), [S("NoView", ""), S("NoSel", view(80)), S("A1", view(sels=[(None, "A1", "A1")])),
                                S("Range", view(sels=[(None, "A1", "A1:D10")])),
                                S("Bare", view(pane=FROZEN3, sels=[("bottomLeft", None, None)])),
                                S("FrozenA1", view(pane=FROZEN3, sels=[("bottomLeft", "A1", "A1")]))])
    v = run(p, 62)
    assert v["decision"] == "pass", v
    assert v["stats"]["n_sheets_counted"] == 6


def test_62_strict_a1_frozen_corner_fails():
    # Patrick 2026-10-02: the frozen-pane corner (A4) is not A1 -> fail
    p = book(tmp("c_corner.xlsx"), [S("Inst", ""), S("Model", view(pane=FROZEN3, sels=[("bottomLeft", "A4", "A4")]))])
    v = run(p, 62)
    assert locs(v) == ["Model!A4"], v
    assert "frozen-pane corner" in v["mistakes"][0]["description"]
    # bottomRight corner F4 (xSplit 5, ySplit 3)
    pane = (5, 3, "F4", "bottomRight", "frozen")
    p = book(tmp("c_corner2.xlsx"), [S("M", view(pane=pane, sels=[("topRight", "F1", "F1"), ("bottomLeft", "A4", "A4"),
                                                                   ("bottomRight", "F4", "F4")]))])
    assert locs(run(p, 62)) == ["M!F4"]


def test_62_active_pane_not_first_selection():
    # T2 trap: a stale FIRST selection (topLeft I276) of a non-active pane must not decide
    p = book(tmp("c_stale.xlsx"), [S("U", view(pane=FROZEN3, sels=[("topLeft", "I276", "I276"), ("bottomLeft", "A1", "A1")]))])
    v = run(p, 62)
    assert v["decision"] == "pass" and v["stats"]["first_selection_not_active_pane"] == {"U": "topLeft I276"}, v
    # T3 trap: first selection topRight F1, active bottomRight F5 -> the cursor is F5
    pane = (5, 3, "F4", "bottomRight", "frozen")
    p = book(tmp("c_t3.xlsx"), [S("Solution Model", view(pane=pane, sels=[("topRight", "F1", "F1"),
                                                                          ("bottomLeft", "A4", "A4"),
                                                                          ("bottomRight", "F5", "F5")]))])
    v = run(p, 62)
    assert locs(v) == ["'Solution Model'!F5"], v
    # active pane without a selection (another pane has one at D10): the active pane's top-left cell D7
    # (Excel, measured 2026-10-03)
    p = book(tmp("c_nosel.xlsx"), [S("X", view(pane=(3, 6, "D7", "bottomRight", "frozen"), sels=[(None, "D10", "D10")]))])
    assert locs(run(p, 62)) == ["X!D7"]
    # no pane: the topLeft (unlabelled) selection decides
    p = book(tmp("c_plain.xlsx"), [S("Marketcap", view(sels=[(None, "B1", "B1")]))])
    assert locs(run(p, 62)) == ["Marketcap!B1"]


def test_62_unverified_shapes():
    # a selection without activeCell whose sqref is D10: A1 or D10 -> cannot decide
    raises(book(tmp("c_noac.xlsx"), [S("X", view(sels=[(None, None, "D10")]))]), 62, "cannot be decided")
    # activeCell A1 outside its selection D10:F12: A1 or D10 -> cannot decide
    raises(book(tmp("c_out.xlsx"), [S("X", view(sels=[(None, "A1", "D10:F12")]))]), 62, "outside its selection")
    # ... but activeCell B2 outside D10 is non-A1 under both readings -> fail
    v = run(book(tmp("c_out2.xlsx"), [S("X", view(sels=[(None, "B2", "D10")]))]), 62)
    assert locs(v) == ["X!B2"] and "or D10" in v["mistakes"][0]["description"], v
    # duplicate selections of the active pane: the LAST one decides (Patrick in Excel 2026-10-06,
    # attempts 2379 / 2402: A1 then B9 opens on B9, A1 then A4 on A4)
    dup = [("bottomRight", "A1", "A1"), ("bottomRight", "A1", "A1")]
    assert run(book(tmp("c_dup.xlsx"), [S("X", view(pane=(2, 9, "C10", "bottomRight", "frozen"), sels=dup))]), 62)["decision"] == "pass"
    dup = [("bottomRight", "A1", "A1"), ("bottomRight", "C5", "C5")]
    v = run(book(tmp("c_dup2.xlsx"), [S("X", view(pane=(2, 9, "C10", "bottomRight", "frozen"), sels=dup))]), 62)
    assert locs(v) == ["X!C5"] and "2 selections" in v["stats"]["cursors"]["X"]["notes"][0], v
    dup = [("bottomRight", "C5", "C5"), ("bottomRight", "A1", "A1")]
    assert run(book(tmp("c_dup3.xlsx"), [S("X", view(pane=(2, 9, "C10", "bottomRight", "frozen"), sels=dup))]), 62)["decision"] == "pass"
    # unreadable cursor data raises
    raises(book(tmp("c_badac.xlsx"), [S("X", view(sels=[(None, "Q", "Q")]))]), 62, "not a cell reference")
    raises(book(tmp("c_badpane.xlsx"), [S("X", view(pane=(0, 3, "A4", "middle", "frozen")))]), 62, "not a pane name")


def test_62_undecidable_sheet_with_certain_fail():
    # review 2026-10-03 (62-raises-although-already-failing): S1 fails for certain at D10, so an
    # undecidable S2 is recorded, not raised
    # (duplicate active-pane selections are decided since 2026-10-06: the last one; see test_62_unverified_shapes)
    shapes = {"outside": view(sels=[(None, "A1", "D10:F12")]),
              "badac": view(sels=[(None, "Q", "Q")]),
              "badpane": view(pane=(0, 3, "A4", "middle", "frozen")),
              "default_sqref": view(sels=[(None, "D10", None)])}
    for k, sv in shapes.items():
        v = run(book(tmp(f"c_und_{k}.xlsx"), [S("S1", view(sels=[(None, "D10", "D10")])), S("S2", sv)]), 62)
        assert v["decision"] == "fail" and locs(v) == ["'S1'!D10"], (k, v)        # S1 is quoted (cell-like name)
        assert list(v["stats"]["undecidable_sheets"]) == ["S2"] and v["stats"]["failing_sheets"] == ["S1"], (k, v)
        assert "does not depend on it" in v["summary"], (k, v)
        # alone (nothing fails for certain) the same sheet still raises
        raises(book(tmp(f"c_und_alone_{k}.xlsx"), [S("S1", ""), S("S2", sv)]), 62, "undecidable")


def test_62_absent_sqref_is_schema_default_a1():
    # review 2026-10-03 (62-default-sqref-inconsistent): <selection activeCell="D10"/> has
    # sqref "A1" by schema default, the same shape as an explicit sqref="A1"
    for name, sq in (("c_nosq.xlsx", None), ("c_sqA1.xlsx", "A1")):
        msg = raises(book(tmp(name), [S("X", view(sels=[(None, "D10", sq)]))]), 62, "outside its selection")
        assert "A1 or D10" in msg, msg
    # activeCell A1 without sqref is inside the default A1 -> pass; B2 inside its own sqref -> fail
    assert run(book(tmp("c_nosq_a1.xlsx"), [S("X", view(sels=[(None, "A1", None)]))]), 62)["decision"] == "pass"
    assert locs(run(book(tmp("c_b2.xlsx"), [S("X", view(sels=[(None, "B2", "B2")]))]), 62)) == ["X!B2"]


def test_62_active_pane_without_selection():
    # Excel, measured by Patrick 2026-10-03 (attempt 1408): a frozen sheet whose only selection (A1)
    # sits in a pane that is not the active pane opens on the ACTIVE pane's top-left cell (the
    # frozen corner / first scrolling cell), so strict A1 fails it
    corpus_shape = view(pane=(5, 9, "F10", "bottomRight", "frozen"), sels=[(None, "A1", "A1")])   # attempt 1408 'WACC'
    p = book(tmp("c_noactsel.xlsx"), [S("DCF", corpus_shape), S("BL", view(pane=FROZEN3, sels=[("topLeft", "D10", "D10")])),
                                      S("NoSel", view(pane=FROZEN3))])
    v = run(p, 62)
    assert v["decision"] == "fail" and locs(v) == ["DCF!F10", "BL!A4", "NoSel!A4"], v
    assert v["stats"]["active_pane_without_selection"] == {"DCF": "F10", "BL": "A4", "NoSel": "A4"}, v
    assert v["stats"]["active_pane_without_selection_reading"] == "pane_top_left"
    assert "pane's top-left cell" in v["mistakes"][0]["description"]
    # the pane's stored topLeftCell decides (scrolled pane: E20), not the frozen corner
    scrolled = view(pane=(3, 6, "E20", "bottomRight", "frozen"), sels=[(None, "A1", "A1")])
    assert locs(run(book(tmp("c_noactsel_scr.xlsx"), [S("X", scrolled)]), 62)) == ["X!E20"]
    with option(c62, "ACTIVE_PANE_WITHOUT_SELECTION", "A1"):           # the former reading, kept as a switch
        assert run(p, 62)["decision"] == "pass"
    # a frozen pane whose active pane is topLeft without a selection is plain A1 (not listed)
    v = run(book(tmp("c_noactsel_tl.xlsx"), [S("X", view(pane=(0, 3, "A4", "topLeft", "frozen")))]), 62)
    assert v["decision"] == "pass" and v["stats"]["active_pane_without_selection"] == {}, v
    with option(c62, "ACTIVE_PANE_WITHOUT_SELECTION", "unverified"):
        msg = raises(p, 62, "A1 or F10")
        assert "pane's top-left cell F10" in msg, msg
        # bottomLeft pane of a view scrolled right to column C: top-left is (pane row, view column)
        bl = view(pane=(None, 3, "C4", "bottomLeft", "frozen"), top_left="C1", sels=[("topLeft", "C1", "C1")])
        raises(book(tmp("c_noactsel_bl.xlsx"), [S("X", bl)]), 62, "A1 or C4")
        v = run(book(tmp("c_noactsel_fail.xlsx"), [S("F", view(sels=[(None, "B2", "B2")])), S("DCF", corpus_shape)]), 62)
        assert locs(v) == ["F!B2"] and list(v["stats"]["undecidable_sheets"]) == ["DCF"], v
    with option(c62, "ACTIVE_PANE_WITHOUT_SELECTION", "bogus"):
        raises(p, 62, "unknown ACTIVE_PANE_WITHOUT_SELECTION")


def test_62_stray_pane_labels_without_pane():
    # re-review 2026-10-04 (62-stray-pane-label-silent-a1): openpyxl keeps the pane labels after
    # ws.freeze_panes = None, so a view WITHOUT <pane> can carry <selection pane="bottomRight"
    # activeCell="D10"/>.  Excel's reading is not measured: A1 (no topLeft selection) or D10 ->
    # undecidable, not a silent pass.  Before the fix: pass, "no selection for the active pane topLeft"
    unfrozen = view(sels=[("topRight", "A1", "A1"), ("bottomLeft", "A1", "A1"), ("bottomRight", "D10", "D10")])
    msg = raises(book(tmp("c_stray.xlsx"), [S("Model", unfrozen)]), 62, "undecidable")
    assert "A1 or D10" in msg and "stray pane labels" in msg and "bottomRight D10" in msg, msg
    # every stray cell A1 (frozen then unfrozen with the cursor at A1): pass, listed in stats
    v = run(book(tmp("c_stray_a1.xlsx"), [S("Model", view(sels=[("topRight", "A1", "A1"), ("bottomRight", "A1", "A1")]))]), 62)
    assert v["decision"] == "pass", v
    assert v["stats"]["selections_labelled_without_pane"] == {"Model": "topRight A1, bottomRight A1"}, v
    # bare stray labels (an XlsxWriter / Excel Online frozen shape, unfrozen): default sqref A1 -> pass
    assert run(book(tmp("c_stray_bare.xlsx"), [S("M", view(sels=[("bottomLeft", None, None)]))]), 62)["decision"] == "pass"
    # a topLeft selection at C3 plus a stray at D10: no candidate is A1 -> fails for certain at C3
    v = run(book(tmp("c_stray_c3.xlsx"), [S("M", view(sels=[(None, "C3", "C3"), ("bottomRight", "D10", "D10")]))]), 62)
    assert locs(v) == ["M!C3"] and "(or D10)" in v["mistakes"][0]["description"], v
    # a topLeft selection at A1 plus a stray at D10: undecidable
    raises(book(tmp("c_stray_tl_a1.xlsx"), [S("M", view(sels=[(None, "A1", "A1"), ("bottomRight", "D10", "D10")]))]), 62, "A1 or D10")
    # a stray without activeCell whose sqref starts at D10: A1 or D10 as well
    raises(book(tmp("c_stray_sq.xlsx"), [S("M", view(sels=[("bottomLeft", None, "D10:E12")]))]), 62, "A1 or D10")
    # an unreadable stray activeCell makes the sheet undecidable (it may be the cursor)
    raises(book(tmp("c_stray_bad.xlsx"), [S("M", view(sels=[("bottomLeft", "Q", "Q")]))]), 62, "not a cell reference")
    # with another sheet failing for certain the file fails; the stray sheet is listed as undecidable
    v = run(book(tmp("c_stray_fail.xlsx"), [S("F", view(sels=[(None, "B2", "B2")])), S("Model", unfrozen)]), 62)
    assert locs(v) == ["F!B2"] and list(v["stats"]["undecidable_sheets"]) == ["Model"], v
    assert v["stats"]["cursors"]["Model"]["cursor"] == "A1/D10", v
    # the same labelled selection on a view WITH that pane is the ordinary active-pane reading
    v = run(book(tmp("c_not_stray.xlsx"), [S("M", view(pane=(1, 1, "B2", "bottomRight", "frozen"),
                                                         sels=[("bottomRight", "D10", "D10")]))]), 62)
    assert locs(v) == ["M!D10"] and v["stats"]["selections_labelled_without_pane"] == {}, v
    # a pane-less view whose only selections are topLeft (562 corpus sheets carry pane="topLeft") is untouched
    v = run(book(tmp("c_tl_label.xlsx"), [S("M", view(sels=[("topLeft", "A1", "A1")]))]), 62)
    assert v["decision"] == "pass" and v["stats"]["selections_labelled_without_pane"] == {}, v


def test_62_openpyxl_unfreeze_keeps_pane_labels():
    # the real writer shape behind 62-stray-pane-label-silent-a1 (openpyxl 3.1: freeze B2, move the
    # cursor to D10 in the bottomRight pane, unfreeze, save -> no <pane>, three labelled selections)
    import openpyxl                                     # test-only: building the synthetic workbook
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Model"
    ws["A1"] = 1
    ws.freeze_panes = "B2"
    ws.sheet_view.selection[-1].activeCell = "D10"
    ws.sheet_view.selection[-1].sqref = "D10"
    ws.freeze_panes = None
    p = tmp("c_opx_unfrozen.xlsx")
    wb.save(p)
    with zipfile.ZipFile(p) as z:
        xml = z.read("xl/worksheets/sheet1.xml").decode()
    assert "<pane" not in xml and 'pane="bottomRight" activeCell="D10"' in xml, xml
    msg = raises(p, 62, "A1 or D10")
    assert "stray pane labels" in msg, msg
    # the cursor left at A1 before unfreezing passes
    wb2 = openpyxl.Workbook()
    ws2 = wb2.active
    ws2["A1"] = 1
    ws2.freeze_panes = "A4"
    ws2.freeze_panes = None
    p2 = tmp("c_opx_unfrozen_a1.xlsx")
    wb2.save(p2)
    assert run(p2, 62)["decision"] == "pass"


def test_62_active_topleft_pane_of_scrolled_view_without_selection():
    # re-review 2026-10-04 (62-topleft-active-ignores-view-scroll): a frozen view scrolled to column C
    # (sheetView topLeftCell C1), active pane topLeft without a selection, the bottomLeft pane at D10.
    # A1 (Excel's default for a missing selection) or C1 (the active pane's top-left cell, by the 1408
    # reading): not measured -> undecidable.  Before the fix: pass (A1).  0 corpus sheets have this shape
    p18c = view(pane=(None, 3, "C4", "topLeft", "frozen"), top_left="C1", sels=[("bottomLeft", "D10", "D10")])
    msg = raises(book(tmp("c_tl_scrolled.xlsx"), [S("X", p18c)]), 62, "A1 or C1")
    assert "not measured" in msg, msg
    # the unscrolled sibling stays plain A1 (GPT-6 1750 / 1762, 15 sheets)
    v = run(book(tmp("c_tl_unscrolled.xlsx"), [S("X", view(pane=(None, 3, "A4", "topLeft", "frozen"),
                                                            sels=[("bottomLeft", "D10", "D10")]))]), 62)
    assert v["decision"] == "pass" and v["stats"]["active_pane_without_selection"] == {}, v
    # a topLeft selection decides as before; the scroll is stats only
    v = run(book(tmp("c_tl_sel.xlsx"), [S("X", view(pane=(None, 3, "C4", "topLeft", "frozen"), top_left="C1",
                                                     sels=[("topLeft", "A1", "A1")]))]), 62)
    assert v["decision"] == "pass" and v["stats"]["scrolled_views"] == {"X": {"view_top_left": "C1"}}, v
    # with another sheet failing for certain: fail, X undecidable, the C1 candidate in stats
    v = run(book(tmp("c_tl_scrolled_fail.xlsx"), [S("F", view(sels=[(None, "B2", "B2")])), S("X", p18c)]), 62)
    assert locs(v) == ["F!B2"] and list(v["stats"]["undecidable_sheets"]) == ["X"], v
    assert v["stats"]["active_pane_without_selection"] == {"X": "C1"} and v["stats"]["cursors"]["X"]["cursor"] == "A1/C1", v
    # the former "A1" reading keeps A1
    with option(c62, "ACTIVE_PANE_WITHOUT_SELECTION", "A1"):
        assert run(book(tmp("c_tl_scrolled_a1.xlsx"), [S("X", p18c)]), 62)["decision"] == "pass"
    # a pane-less scrolled view without a selection is A1: Excel itself omits the A1 selection even when
    # the view is scrolled (toy 94 T4 'NXTI IS', saved by Microsoft Excel, topLeftCell A63)
    v = run(book(tmp("c_plain_scrolled.xlsx"), [S("NXTI IS", view(top_left="A63"))]), 62)
    assert v["decision"] == "pass" and v["stats"]["cursors"]["NXTI IS"]["cursor"] == "A1", v


def test_62_scrolled_pane_without_selection_keeps_corner_candidate():
    # re-review 2026-10-04 (62-scrolled-pane-location): GPT-6 1353 'Statements' is frozen at F6
    # (xSplit 5, ySplit 5) with the scrolling pane parked at R6 and only an A1 selection for topLeft.
    # Excel was measured (1408) with the pane AT the corner only, so the cursor is R6 (the pane's
    # topLeftCell, primary, unchanged location) or F6 (the corner); both fail strict A1
    shape = view(pane=(5, 5, "R6", "bottomRight", "frozen"), sels=[(None, "A1", "A1")])
    v = run(book(tmp("c_1353.xlsx"), [S("Statements", shape)]), 62)
    assert locs(v) == ["Statements!R6"], v
    d = v["mistakes"][0]["description"]
    assert "cursor on R6 (or F6)" in d and "past the frozen corner F6" in d and "measured only" in d, d
    assert v["stats"]["cursors"]["Statements"]["cursor"] == "R6/F6", v
    assert v["stats"]["active_pane_without_selection"] == {"Statements": "R6"} and "Statements" in v["stats"]["scrolled_views"], v
    # the measured case (pane at the corner) has one candidate and no note
    v = run(book(tmp("c_1408.xlsx"), [S("DCF", view(pane=(3, 6, "D7", "bottomRight", "frozen"), sels=[(None, "A1", "A1")]))]), 62)
    assert locs(v) == ["DCF!D7"] and v["stats"]["cursors"]["DCF"]["cursor"] == "D7", v
    assert "notes" not in v["stats"]["cursors"]["DCF"] and "(or" not in v["mistakes"][0]["description"], v
    # bottomLeft pane scrolled down (pane topLeftCell A20, corner A4) and topRight pane scrolled right (H1, corner C1)
    v = run(book(tmp("c_bl_scrolled.xlsx"), [S("BL", view(pane=(None, 3, "A20", "bottomLeft", "frozen"), sels=[("topLeft", "A1", "A1")])),
                                              S("TR", view(pane=(2, None, "H1", "topRight", "frozen"), sels=[("topLeft", "A1", "A1")]))]), 62)
    assert locs(v) == ["BL!A20", "TR!H1"], v
    assert v["stats"]["cursors"]["BL"]["cursor"] == "A20/A4" and v["stats"]["cursors"]["TR"]["cursor"] == "H1/C1", v
    # a split (not frozen) pane has no corner: the pane's topLeftCell alone
    v = run(book(tmp("c_split_scrolled.xlsx"), [S("Sp", view(pane=(2000, 1200, "R6", "bottomRight", "split")))]), 62)
    assert locs(v) == ["Sp!R6"] and v["stats"]["cursors"]["Sp"]["cursor"] == "R6", v
    # under the "unverified" reading all three are candidates (A1, R6, F6) -> undecidable
    with option(c62, "ACTIVE_PANE_WITHOUT_SELECTION", "unverified"):
        raises(book(tmp("c_1353_unv.xlsx"), [S("Statements", shape)]), 62, "A1 or R6, F6")
    with option(c62, "ACTIVE_PANE_WITHOUT_SELECTION", "A1"):
        assert run(book(tmp("c_1353_a1.xlsx"), [S("Statements", shape)]), 62)["decision"] == "pass"


def test_62_scope_hidden_chart_scroll():
    # hidden sheets count (whole workbook); chart sheets have no cursor
    p = book(tmp("c_hidden.xlsx"), [S("A", ""), S("H", view(sels=[(None, "Z99", "Z99")]), state="hidden"),
                                    S("Ch", view(100), kind="chartsheet")])
    v = run(p, 62)
    assert locs(v) == ["H!Z99"] and "(hidden sheet)" in v["mistakes"][0]["description"], v
    assert [e["sheet"] for e in v["stats"]["excluded_sheets"]] == ["Ch"]
    # a view scrolled to row 200 with the cursor at A1: stats only (CHECK_SCROLL = False)
    p = book(tmp("c_scroll.xlsx"), [S("A", view(top_left="A200", sels=[(None, "A1", "A1")]))])
    v = run(p, 62)
    assert v["decision"] == "pass" and v["stats"]["scrolled_views"] == {"A": {"view_top_left": "A200"}}, v
    # only window 0 decides
    p = book(tmp("c_win.xlsx"), [S("A", view(sels=[(None, "A1", "A1")]) + view(wv=1, sels=[(None, "C3", "C3")]))])
    assert run(p, 62)["decision"] == "pass"


def test_62_mistake_per_sheet_and_summary():
    p = book(tmp("c_many.xlsx"), [S("I", ""), S("MPHG IS", view(sels=[(None, "D23", "D23")])),
                                  S("MPHG BS", view(sels=[(None, "B12", "B12")])),
                                  S("Model", view(pane=FROZEN3, sels=[("bottomLeft", "K37", "K37")]))])
    v = run(p, 62)
    assert locs(v) == ["'MPHG IS'!D23", "'MPHG BS'!B12", "Model!K37"], v
    assert v["summary"].startswith("3 of 4 sheet(s)") and v["stats"]["failing_sheets"] == ["MPHG IS", "MPHG BS", "Model"]
    assert all(m["severity"] == "minor" for m in v["mistakes"])


# ============================================================================ 77
def test_77_pass_cases():
    v = run(book(tmp("ok.xlsx"), [S("A", "")]), 77)
    assert v["decision"] == "pass" and v["stats"]["filename_source"] == "file_name", v
    v = run(book(tmp("UPPER.XLSX"), [S("A", "")]), 77)                      # T4: case-insensitive
    assert v["decision"] == "pass" and v["stats"]["extension"] == ".xlsx", v
    v = run(book(tmp("macro.xlsm"), [S("A", "")], main_ct="xlsm", vba="ole2"), 77)
    assert v["decision"] == "pass" and v["stats"]["macro_parts"] == ["xl/vbaProject.bin"], v
    v = run(book(tmp("xlm.xlsm"), [S("A", "")], main_ct="xlsm", macrosheet=True), 77)   # Excel 4.0 macro sheet
    assert v["decision"] == "pass", v
    # the producer is not judged
    assert run(book(tmp("lo.xlsx"), [S("A", "")], app="LibreOffice/7.6"), 77)["decision"] == "pass"


def test_77_xlsm_without_macros_fails():
    for name, kw in (("nomacro.xlsm", {}),                                         # T1/T2
                     ("orphan.xlsm", {"vba": "ole2", "vba_related": False}),      # bin in zip, not related
                     ("junkvba.xlsm", {"vba": "junk"})):                          # related, not OLE2
        v = run(book(tmp(name), [S("A", "")], main_ct="xlsm", **kw), 77)
        assert v["decision"] == "fail" and locs(v) == [name], (name, v)
        assert "no VBA project or macro sheet" in v["mistakes"][0]["description"]
        assert v["mistakes"][0]["severity"] == "major"
    v = run(book(tmp("orphan2.xlsm"), [S("A", "")], main_ct="xlsm", vba="ole2", vba_related=False), 77)
    assert v["stats"]["unrelated_vba_files"] == ["xl/vbaProject.bin"] and "not attached" in v["summary"], v


def test_77_mismatches_fail():
    v = run(book(tmp("vba_in.xlsx"), [S("A", "")], vba="ole2"), 77)                # macros need .xlsm
    assert v["decision"] == "fail" and "must be .xlsm" in v["summary"], v
    v = run(book(tmp("xlm_in.xlsx"), [S("A", "")], macrosheet=True), 77)
    assert v["decision"] == "fail", v
    v = run(book(tmp("enabled.xlsx"), [S("A", "")], main_ct="xlsm", vba="ole2"), 77)   # macro-enabled named .xlsx
    assert v["decision"] == "fail" and "declared macro-enabled" in v["summary"], v
    v = run(book(tmp("plain.xlsm"), [S("A", "")], vba="ole2"), 77)                 # plain package named .xlsm
    assert v["decision"] == "fail" and "plain macro-free .xlsx" in v["summary"], v
    v = run(book(tmp("tmpl.xlsx"), [S("A", "")], main_ct="xltx"), 77)              # template content as .xlsx
    assert v["decision"] == "fail" and "template" in v["summary"], v
    v = run(book(tmp("tmpl.xltx"), [S("A", "")], main_ct="xltx"), 77)              # template name
    assert v["decision"] == "fail" and ".xltx" in v["summary"], v


def test_77_non_ooxml_files():
    # .xls (BIFF in OLE2): fails without crashing, under its own name or as .xlsx
    v = run(ole2(tmp("legacy.xls")), 77)
    assert v["decision"] == "fail" and "legacy Excel 97-2003" in v["summary"], v
    v = run(ole2(tmp("legacy_renamed.xlsx")), 77)
    assert v["decision"] == "fail" and "disagree" in v["summary"], v
    # password-encrypted package: cannot be verified under an .xlsx name; .xls name fails on the name
    raises(ole2(tmp("enc.xlsx"), ("EncryptedInfo", "EncryptedPackage")), 77, "password-encrypted")
    assert run(ole2(tmp("enc.xls"), ("EncryptedInfo", "EncryptedPackage")), 77)["decision"] == "fail"
    # CSV text
    assert run(raw(tmp("data.csv"), b"a,b\n1,2\n"), 77)["decision"] == "fail"
    v = run(raw(tmp("data_as.xlsx"), b"a,b\n1,2\n"), 77)
    assert v["decision"] == "fail" and "not a zip or OLE2" in v["summary"], v
    # empty / corrupt files under an allowed name cannot be classified
    raises(raw(tmp("empty.xlsx"), b""), 77, "an empty file")
    raises(raw(tmp("broken.xlsx"), b"PK\x03\x04" + b"\0" * 40), 77, "damaged zip")
    assert run(raw(tmp("empty.xls"), b""), 77)["decision"] == "fail"
    # a zip that is not a workbook (renamed .docx)
    p = tmp("doc.xlsx")
    with zipfile.ZipFile(p, "w") as z:
        z.writestr("[Content_Types].xml", '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
                   '<Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.'
                   'wordprocessingml.document.main+xml"/></Types>')
        z.writestr("_rels/.rels", f'<Relationships xmlns="{PKG_REL}"><Relationship Id="rId1" Type="{REL}/officeDocument" '
                                  f'Target="word/document.xml"/></Relationships>')
        z.writestr("word/document.xml", "<document/>")
    v = run(p, 77)
    assert v["decision"] == "fail" and "not an Excel workbook" in v["summary"], v
    # an unreadable OLE2 directory named .xlsx: could be encrypted .xlsx or .xls -> cannot decide
    raises(raw(tmp("ole_junk.xlsx"), OLE2_MAGIC + b"\xff" * 100), 77, "cannot be read")


def test_77_delivered_filename_from_task_meta():
    staged = book(tmp("ai_attempt.xlsx"), [S("A", "")])
    v = run(staged, 77, {"delivered_filename": "20260910_Model.XLSX"})
    assert v["decision"] == "pass" and v["stats"]["filename_source"] == "task_meta.delivered_filename", v
    assert v["stats"]["delivered_filename"] == "20260910_Model.XLSX" and v["stats"]["staged_file_name"] == "ai_attempt.xlsx"
    v = run(staged, 77, {"delivered_filename": "s3://bucket/x/Model.xlsm"})
    assert v["decision"] == "fail" and locs(v) == ["Model.xlsm"], v
    v = run(staged, 77, {"delivered_filename": "Model.xls"})
    assert v["decision"] == "fail" and locs(v) == ["Model.xls"], v
    other = book(tmp("Model_delivered.xlsx"), [S("A", "")])
    assert run(other, 77, {"delivered_filename": None})["stats"]["filename_source"] == "file_name"
    # a malformed delivered name: the format is judged from the content (Patrick 2026-10-05: every attempt graded)
    for bad in ("", 7, "folder/"):
        v = run(staged, 77, {"delivered_filename": bad})
        assert v["decision"] == "pass" and v["stats"]["delivered_name_unknown"].startswith("malformed"), v
        assert v["stats"]["defaults"]["delivered_name_unknown"]["count"] == 1, v["stats"]


def test_77_staged_name_without_delivered_name_raises():
    # review 2026-10-03 (77-staged-name-fallback): production stages every delivery as ai_attempt.xlsx, whose
    # own name is not the delivered one.  Patrick 2026-10-05 (every attempt graded): without a delivered name the
    # format is judged from the file's content (a zip workbook without a VBA project = .xlsx, with a real one =
    # .xlsm, OLE2 = .xls) and recorded in stats.defaults (it raised until then)
    staged = book(tmp("ai_attempt.xlsx"), [S("A", "")], main_ct="xlsm", vba="ole2")
    for meta in (None, {"delivered_filename": None}, {"delivered_filename_problem": "missing"}):
        v = run(staged, 77, meta)
        assert v["decision"] == "pass" and v["stats"]["judged_from_content_as"] == ".xlsm", v
        assert v["stats"]["delivered_filename"] is None and v["stats"]["delivered_name_unknown"].startswith("missing")
        d = v["stats"]["defaults"]["delivered_name_unknown"]
        assert d["count"] == 1 and "judged from its content as .xlsm" in d["examples"][0], d
    v = run(staged, 77, {"original_filename": "Model.xlsm"})              # the sidecar key is named in the record
    assert v["decision"] == "pass" and "'original_filename'" in v["stats"]["delivered_name_unknown"], v
    os.makedirs(tmp("upper"), exist_ok=True)              # own folder: the file system may ignore case
    upper = book(os.path.join(tmp("upper"), "AI_ATTEMPT.XLSX"), [S("A", "")])
    v = run(upper, 77, {"delivered_filename_problem": "malformed: a JSON list, not an object"})
    assert v["decision"] == "pass" and v["stats"]["judged_from_content_as"] == ".xlsx", v
    assert v["stats"]["delivered_name_unknown"] == "malformed: a JSON list, not an object", v
    # judged from the content: a macro-enabled package without macros, legacy .xls bytes, a damaged zip fail
    os.makedirs(tmp("c1"), exist_ok=True)
    nomacro = book(os.path.join(tmp("c1"), "ai_attempt.xlsx"), [S("A", "")], main_ct="xlsm")
    v = run(nomacro, 77)
    assert v["decision"] == "fail" and v["stats"]["judged_from_content_as"] == ".xlsx", v
    os.makedirs(tmp("c2"), exist_ok=True)
    xls = raw(os.path.join(tmp("c2"), "ai_attempt.xlsx"), OLE2_MAGIC + b"\xff" * 100)
    v = run(xls, 77)
    assert v["decision"] == "fail" and v["stats"]["judged_from_content_as"] == ".xls", v
    os.makedirs(tmp("c3"), exist_ok=True)
    junk = raw(os.path.join(tmp("c3"), "ai_attempt.xlsx"), b"PK\x03\x04 not really a zip")
    v = run(junk, 77)
    assert v["decision"] == "fail" and v["stats"]["judged_from_content_as"] is None, v
    v = run(staged, 77, {"delivered_filename": "Model.xlsm"})
    assert v["decision"] == "pass" and v["stats"]["filename_source"] == "task_meta.delivered_filename", v


def test_77_vba_part_must_be_a_vba_project():
    # review 2026-10-03 (77-vba-signature-only): the OLE2 signature alone is not a VBA project
    for kind, needle in (("magic", "damaged compound-file header"), ("xls", "without a VBA storage"),
                         ("nodir", "no dir stream"), ("noproject", "without a PROJECT stream")):
        v = run(book(tmp(f"vba_{kind}.xlsm"), [S("A", "")], main_ct="xlsm", vba=kind), 77)
        assert v["decision"] == "fail" and needle in v["summary"], (kind, v)
        assert list(v["stats"]["invalid_vba_parts"]) == ["xl/vbaProject.bin"] and not v["stats"]["has_macros"], v
        # ... and the same non-project part does not make an .xlsx a macro workbook
        v = run(book(tmp(f"vba_{kind}.xlsx"), [S("A", "")], vba=kind), 77)
        assert v["decision"] == "pass", (kind, v)
    # a real VBA project structure passes as .xlsm and needs .xlsm
    v = run(book(tmp("vba_real.xlsm"), [S("A", "")], main_ct="xlsm", vba="ole2"), 77)
    assert v["decision"] == "pass" and v["stats"]["macro_parts"] == ["xl/vbaProject.bin"], v
    # a valid header whose directory cannot be walked: whether it holds macros is unknown
    raises(book(tmp("vba_badtree.xlsm"), [S("A", "")], main_ct="xlsm", vba="badtree"), 77, "directory cannot be read")
    raises(book(tmp("vba_badtree.xlsx"), [S("A", "")], vba="badtree"), 77, "directory cannot be read")
    # ... unless a real macro carrier decides anyway (an XLM macro sheet)
    v = run(book(tmp("vba_badtree_xlm.xlsm"), [S("A", "")], main_ct="xlsm", vba="badtree", macrosheet=True), 77)
    assert v["decision"] == "pass" and "xl/vbaProject.bin" in v["stats"]["unreadable_vba_parts"], v


def test_77_unverified_parts_raise():
    # Excel 5.0 dialog sheet as the only macro-like part: whether it needs .xlsm is unverified
    raises(book(tmp("dlg.xlsm"), [S("A", ""), S("Dialog1", "", kind="dialogsheet")], main_ct="xlsm"), 77, "dialog sheet")
    raises(book(tmp("dlg.xlsx"), [S("A", ""), S("Dialog1", "", kind="dialogsheet")]), 77, "dialog sheet")
    # with a VBA project the dialog sheet does not matter
    v = run(book(tmp("dlgvba.xlsm"), [S("A", ""), S("Dialog1", "", kind="dialogsheet")], main_ct="xlsm", vba="ole2"), 77)
    assert v["decision"] == "pass", v
    # a generic main content type under .xlsx / .xlsm: the package type is not declared
    raises(book(tmp("generic.xlsx"), [S("A", "")], main_ct="generic"), 77, "generic content type")


class option:
    """Temporarily set a module-level policy switch."""

    def __init__(self, module, name, value):
        self.module, self.name, self.value = module, name, value

    def __enter__(self):
        self.old = getattr(self.module, self.name)
        setattr(self.module, self.name, self.value)

    def __exit__(self, *exc):
        setattr(self.module, self.name, self.old)


def test_policy_switches():
    # 61: the per-view reading would let zoomScaleNormal decide (not the default)
    p = book(tmp("o_normal_only.xlsx"), [S("A", view(80)), S("B", view(normal=80)), S("C", view(80))])
    with option(c61, "ZOOM_READING", "per_view"):
        v = run(p, 61)
    assert v["decision"] == "pass" and v["stats"]["other_reading"]["decision"] == "fail", v
    p = book(tmp("o_hidden61.xlsx"), [S("A", view(80)), S("H", view(100), state="hidden")])
    with option(c61, "INCLUDE_HIDDEN_SHEETS", False):
        assert run(p, 61)["decision"] == "pass"
    assert run(p, 61)["decision"] == "fail"
    # 62: scroll check (off by default)
    p = book(tmp("o_scroll.xlsx"), [S("A", view(top_left="A200", sels=[(None, "A1", "A1")]))])
    with option(c62, "CHECK_SCROLL", True):
        v = run(p, 62)
    assert locs(v) == ["A!A1"] and "scrolled to A200" in v["mistakes"][0]["description"], v
    p = book(tmp("o_hidden62.xlsx"), [S("A", ""), S("H", view(sels=[(None, "Z99", "Z99")]), state="veryHidden")])
    with option(c62, "INCLUDE_HIDDEN_SHEETS", False):
        assert run(p, 62)["decision"] == "pass"
    # 77: dialog-sheet policy
    xlsm = book(tmp("o_dlg.xlsm"), [S("A", ""), S("D", "", kind="dialogsheet")], main_ct="xlsm")
    xlsx = book(tmp("o_dlg.xlsx"), [S("A", ""), S("D", "", kind="dialogsheet")])
    with option(c77, "DIALOG_SHEET_POLICY", "macro"):
        assert run(xlsm, 77)["decision"] == "pass" and run(xlsx, 77)["decision"] == "fail"
    with option(c77, "DIALOG_SHEET_POLICY", "not_macro"):
        assert run(xlsm, 77)["decision"] == "fail" and run(xlsx, 77)["decision"] == "pass"


TESTS = [test_61_same_level_passes_and_mixed_fails, test_61_missing_zoomscale_is_100,
         test_61_zoomscalenormal_never_decides, test_61_view_modes, test_61_scope_hidden_chart_windows,
         test_61_tie_and_summary, test_61_unreadable_zoom_raises_only_when_it_decides,
         test_61_unknown_zoom_with_certain_fail, test_61_sheet_views_after_sheet_data_is_unknown,
         test_61_both_attributes_differ_reading_switch,
         test_62_a1_and_defaults_pass, test_62_strict_a1_frozen_corner_fails, test_62_active_pane_not_first_selection,
         test_62_unverified_shapes, test_62_undecidable_sheet_with_certain_fail, test_62_absent_sqref_is_schema_default_a1,
         test_62_active_pane_without_selection, test_62_stray_pane_labels_without_pane,
         test_62_openpyxl_unfreeze_keeps_pane_labels, test_62_active_topleft_pane_of_scrolled_view_without_selection,
         test_62_scrolled_pane_without_selection_keeps_corner_candidate,
         test_62_scope_hidden_chart_scroll, test_62_mistake_per_sheet_and_summary,
         test_77_pass_cases, test_77_xlsm_without_macros_fails, test_77_mismatches_fail, test_77_non_ooxml_files,
         test_77_delivered_filename_from_task_meta, test_77_staged_name_without_delivered_name_raises,
         test_77_vba_part_must_be_a_vba_project, test_77_unverified_parts_raise, test_policy_switches]


def main() -> int:
    failed = 0
    try:
        for t in TESTS:
            try:
                t()
                print(f"PASS {t.__name__}")
            except Exception:  # noqa: BLE001
                failed += 1
                print(f"FAIL {t.__name__}")
                traceback.print_exc()
    finally:
        if _TMP and os.path.isdir(_TMP):
            shutil.rmtree(_TMP, ignore_errors=True)
    print(f"\n{len(TESTS) - failed}/{len(TESTS)} tests passed")
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main())
