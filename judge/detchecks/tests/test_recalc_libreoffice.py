"""LibreOffice process tests for core/recalc.py and core/lo_watchdog.py: encoded paths, and a
LibreOffice conversion that can never outlive the grading that started it.

    cd judge && python -m detchecks.tests.test_recalc_libreoffice

ONE test launches the REAL LibreOffice (DETCHECKS_SOFFICE, default ~/.local/bin/soffice) on a
3-cell workbook in a temporary folder whose path holds a space, '%', '%41' and a non-ASCII letter;
it is skipped (and reported SKIP, not passed) where LibreOffice is absent.  On the shared 16 GB
grading machine run this module on its own, like every LibreOffice job.  The other tests use
stand-in 'soffice' scripts (Python, with a tagged child process), run in-process or inside a
stand-in grader process that is then sent SIGTERM / SIGKILL; the memory guard of 2026-10-05
(core/lo_guard.py) is exercised with them too: a soffice that crashes twice and then converts (retry
with the timeout doubled), one that always fails (LibreOfficeUnavailable, retry_later, through grade()
as well), and two grader processes sharing one lock file (their LibreOffice runs never overlap).  The
stand-in tests use a lock file in their temporary folder and no memory wait.  Every process a test
starts carries a unique tag and is swept at the end; temporary folders are removed.
"""
from __future__ import annotations

import json
import os
import re
import shutil
import signal
import subprocess
import sys
import tempfile
import time
import traceback
import zipfile

from detchecks.core import lo_watchdog
from detchecks.core import recalc as R
from detchecks.errors import GradingError, LibreOfficeUnavailable

try:
    import pytest
    _Skip = pytest.skip.Exception

    def skip(msg: str):
        pytest.skip(msg)
except ImportError:  # plain script run without pytest installed
    pytest = None

    class _Skip(Exception):
        pass

    def skip(msg: str):
        raise _Skip(msg)

JUDGE = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
ODD_NAME = "FruitJuice_3-Statement-Model - v2 %41 café %"     # space, '%41', '%', non-ASCII
_TMPS: list = []
_TAGS: list = []


def tmpdir() -> str:
    d = tempfile.mkdtemp(prefix="detchecks_lo_")
    _TMPS.append(d)
    return d


def new_tag() -> str:
    t = f"detchecks_lo_fake_{os.getpid()}_{len(_TAGS)}"
    _TAGS.append(t)
    return t


def procs_with(*markers) -> list:
    table = lo_watchdog.process_table() or []
    me = os.getpid()
    return [(pid, cmd[:160]) for pid, _pp, cmd in table
            if pid != me and not cmd.startswith("ps ") and any(m in cmd for m in markers)]


def wait_until(fn, timeout=15.0, step=0.1):
    end = time.time() + timeout
    while time.time() < end:
        v = fn()
        if v:
            return v
        time.sleep(step)
    return fn()


FAKE_SOFFICE = r'''#!{python}
"""Stand-in soffice: --outdir and the source are file URLs (as libreoffice_recalc passes them)."""
import json, os, shutil, subprocess, sys, time
from urllib.parse import unquote, urlparse
mode, tag = {mode!r}, {tag!r}
args = sys.argv[1:]
def path(u):
    return unquote(urlparse(u).path) if u.startswith("file:") else u
outdir, src = path(args[args.index("--outdir") + 1]), path(args[-1])
def copy():
    shutil.copy(src, os.path.join(outdir, os.path.splitext(os.path.basename(src))[0] + ".xlsx"))
if mode == "copy":
    copy()
    sys.exit(0)
if mode == "fail":
    sys.stderr.write("boom\n")
    sys.exit(3)
if mode == "flaky":                     # fails (no copy) on the first two runs, then converts
    counter = os.path.abspath(sys.argv[0]) + ".count"
    n = int(open(counter).read()) + 1 if os.path.exists(counter) else 1
    open(counter, "w").write(str(n))
    if n < 3:
        sys.stderr.write(f"crash {{n}}\n")
        sys.exit(134)
    copy()
    sys.exit(0)
if mode == "slowcopy":                  # logs when it runs (the lock test), then converts
    t0 = time.time()
    time.sleep(1.0)
    copy()
    with open(os.path.abspath(sys.argv[0]) + ".log", "a") as fh:
        fh.write(json.dumps([t0, time.time(), os.getppid()]) + "\n")
    sys.exit(0)
# hang: a tagged child (no profile in its command line) and a long sleep
subprocess.Popen(["/bin/bash", "-c", "exec -a " + tag + "_child sleep 300"])
time.sleep(299)
'''


def fake_soffice(folder: str, mode: str, tag: str) -> str:
    p = os.path.join(folder, f"soffice_{mode}")
    with open(p, "w") as fh:
        fh.write(FAKE_SOFFICE.format(python=sys.executable, mode=mode, tag=tag))
    os.chmod(p, 0o755)
    return p


def tiny_workbook(path: str):
    """B3 = B1 - B2 with a STALE cached value 999 (LibreOffice must recalculate it to 95)."""
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Model"
    ws["B1"], ws["B2"], ws["B3"] = 100, 5, "=B1-B2"
    tmp = path + ".build.xlsx"
    wb.save(tmp)
    with zipfile.ZipFile(tmp) as zin, zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                s = data.decode()
                s2 = re.sub(r'(<c r="B3"[^>]*>)<f>B1-B2</f>(?:<v\s*/>|<v></v>)?',
                            lambda m: m.group(1) + "<f>B1-B2</f><v>999</v>", s, count=1)
                assert s2 != s
                data = s2.encode()
            zout.writestr(item, data)
    os.remove(tmp)


def b3_value(path: str):
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)
    try:
        return wb["Model"]["B3"].value
    finally:
        wb.close()


# ============================================================================ tests
def test_file_url_encodes_every_path():
    assert R.file_url("/a b/c%41/café") == "file:///a%20b/c%2541/caf%C3%A9"
    from pathlib import Path
    assert R.file_url("x y") == Path(os.path.abspath("x y")).as_uri()          # relative: made absolute first


def test_real_libreoffice_paths_with_space_and_percent():
    """REAL LibreOffice on a 3-cell file in a task folder named like task 28 plus '%41', '%' and
    a non-ASCII letter: the private profile is used (the stale 999 is recalculated to 95), the copy
    lands in the workdir, nothing is written outside the task folder, no process is left, the
    profile is removed (before: exit -6 with a space; with '%41' the profile and the copy went to
    a decoded 'A' folder next to the task folder)."""
    soffice = R.DEFAULT_SOFFICE
    if not os.path.exists(soffice):
        skip(f"LibreOffice not found at {soffice} (set DETCHECKS_SOFFICE)")
    root = tmpdir()
    task = os.path.join(root, ODD_NAME)
    os.makedirs(task)
    src = os.path.join(task, "ai_attempt.xlsx")
    tiny_workbook(src)
    assert b3_value(src) == 999
    workdir = os.path.join(task, "det_checks_recalc")
    pol = R.RecalcPolicy(workdir=workdir, libreoffice_path=soffice, lo_timeout_s=120, excel_allowed=False)
    try:
        plan = R.ensure_values(src, workdir, pol)
        assert plan.source == "libreoffice" and plan.n_gaps == 0 and plan.errors_vetted, plan.summary()
        assert os.path.dirname(plan.value_path).startswith(workdir + os.sep), plan.value_path
        assert b3_value(plan.value_path) == 95, "LibreOffice did not use the private profile (no recalculation)"
        assert sorted(os.listdir(root)) == [ODD_NAME], os.listdir(root)            # nothing outside the task folder
        assert os.listdir(os.path.join(workdir, "_lo_profiles")) == []              # profile removed
        assert procs_with(R.file_url(workdir), workdir) == []
    finally:
        R.reap_profiles(os.path.join(workdir, "_lo_profiles"))


def test_fake_conversion_through_the_watchdog():
    """A normal run goes through the watchdog (exit code passed through) with file URLs, in the
    grader's process group; a failing run's exit code and stderr reach the error message."""
    tag = new_tag()
    root = tmpdir()
    task = os.path.join(root, ODD_NAME)
    os.makedirs(task)
    src = os.path.join(task, "ai_attempt.xlsx")
    tiny_workbook(src)
    out_dir = os.path.join(task, "det_checks_recalc", "h", "libreoffice")
    pol = R.RecalcPolicy(workdir=os.path.join(task, "det_checks_recalc"),
                         libreoffice_path=fake_soffice(root, "copy", tag), lo_timeout_s=60,
                         lo_lock_path=os.path.join(root, "lo.lock"), lo_min_free_pct=0, lo_retries=1)
    out = R.libreoffice_recalc(src, out_dir, pol)
    assert out == os.path.join(out_dir, "ai_attempt.xlsx") and os.path.exists(out)
    assert R._ACTIVE == {} and os.listdir(os.path.join(task, "det_checks_recalc", "_lo_profiles")) == []
    assert R._watchdog_cmd(["soffice"])[:3] == [sys.executable, "-I", "-S"]
    pol.libreoffice_path = fake_soffice(root, "fail", tag)
    try:
        R.libreoffice_recalc(src, out_dir, pol)
    except LibreOfficeUnavailable as e:
        assert "produced no copy" in str(e) and "exit 3" in str(e) and "boom" in str(e), e
        assert "all 2 tries failed" in str(e) and e.retry_later, e
    else:
        raise AssertionError("a failing soffice did not raise")


def test_timeout_kills_the_whole_tree():
    tag = new_tag()
    root = tmpdir()
    src = os.path.join(root, "ai_attempt.xlsx")
    tiny_workbook(src)
    workdir = os.path.join(root, "w d")
    pol = R.RecalcPolicy(workdir=workdir, libreoffice_path=fake_soffice(root, "hang", tag), lo_timeout_s=2,
                         lo_lock_path=os.path.join(root, "lo.lock"), lo_min_free_pct=0, lo_retries=0)
    t0 = time.time()
    try:
        R.libreoffice_recalc(src, os.path.join(workdir, "out"), pol)
    except GradingError as e:
        assert "timed out after 2 s" in str(e), e
    else:
        raise AssertionError("no timeout")
    assert time.time() - t0 < 15
    assert wait_until(lambda: not procs_with(tag, R.file_url(workdir)), 10), procs_with(tag, R.file_url(workdir))


GRADER = r'''
import sys
sys.path.insert(0, {judge!r})
from detchecks.core import recalc as R
pol = R.RecalcPolicy(workdir={workdir!r}, libreoffice_path={soffice!r}, lo_timeout_s=120,
                     lo_lock_path={lock!r}, lo_min_free_pct=0)
R.libreoffice_recalc({src!r}, {out!r}, pol)
'''


def _grader_dies(sig) -> None:
    """Stand-in grader (the reviewer's P20): a process running libreoffice_recalc on a hanging
    soffice that keeps a tagged child; the grader gets `sig`; nothing of LibreOffice may survive."""
    tag = new_tag()
    root = tmpdir()
    src = os.path.join(root, "ai_attempt.xlsx")
    tiny_workbook(src)
    workdir = os.path.join(root, "Task - v2 %41", "det_checks_recalc")
    os.makedirs(workdir)
    code = GRADER.format(judge=JUDGE, workdir=workdir, soffice=fake_soffice(root, "hang", tag), src=src,
                         out=os.path.join(workdir, "h", "libreoffice"), lock=os.path.join(root, "lo.lock"))
    grader = subprocess.Popen([sys.executable, "-c", code], stdout=subprocess.DEVNULL, stderr=subprocess.PIPE)
    try:
        assert wait_until(lambda: procs_with(tag + "_child"), 30), "the stand-in soffice never started"
        assert procs_with(R.file_url(workdir)), "soffice / watchdog not visible with the profile URL"
        grader.send_signal(sig)
        rc = grader.wait(timeout=30)
        assert rc == -sig, (rc, grader.stderr.read().decode("utf-8", "replace")[-800:])
        left = wait_until(lambda: not procs_with(tag, R.file_url(workdir)), 10)
        assert left, f"LibreOffice outlived its grader after {signal.Signals(sig).name}: {procs_with(tag, R.file_url(workdir))}"
    finally:
        if grader.poll() is None:
            grader.kill()
            grader.wait(timeout=10)
        if grader.stderr:
            grader.stderr.close()


def test_sigterm_to_the_grader_kills_libreoffice():
    _grader_dies(signal.SIGTERM)


def test_sigkill_to_the_grader_kills_libreoffice():
    _grader_dies(signal.SIGKILL)


def test_retry_after_two_crashes_then_graded():
    """Patrick 2026-10-05: a LibreOffice run that fails is retried (timeout doubled, memory wait and lock
    again): a stand-in soffice that crashes twice (no copy) and then converts gives the copy on try 3;
    every try ran through the watchdog with its own private profile, all of them removed."""
    tag = new_tag()
    root = tmpdir()
    src = os.path.join(root, "ai_attempt.xlsx")
    tiny_workbook(src)
    workdir = os.path.join(root, "det_checks_recalc")
    soffice = fake_soffice(root, "flaky", tag)
    logs = []
    pol = R.RecalcPolicy(workdir=workdir, libreoffice_path=soffice, lo_timeout_s=30, lo_retries=3,
                         lo_lock_path=os.path.join(root, "lo.lock"), lo_min_free_pct=0, lo_log=logs.append)
    out = R.libreoffice_recalc(src, os.path.join(workdir, "h", "libreoffice"), pol)
    assert os.path.exists(out) and open(soffice + ".count").read() == "3"
    assert sum("try 1 of 4" in m or "try 2 of 4" in m for m in logs) == 2 and "60 s timeout" in logs[0], logs
    assert os.listdir(os.path.join(workdir, "_lo_profiles")) == [] and R._ACTIVE == {}
    # the same through the pipeline: the plan comes from LibreOffice, nothing failed
    os.remove(soffice + ".count")
    plan = R.ensure_values(src, workdir, R.RecalcPolicy(workdir=workdir, libreoffice_path=soffice, lo_timeout_s=30,
                                                        lo_lock_path=os.path.join(root, "lo.lock"), lo_min_free_pct=0,
                                                        use_cache=False))
    assert plan.source == "libreoffice" and open(soffice + ".count").read() == "3"


def test_failing_every_try_fails_loudly_with_retry_later():
    """Every try fails: LibreOfficeUnavailable (retry_later) listing each try; through grade() the value
    checks fail and the GradingError carries retry_later, so the drivers list the attempt to re-run."""
    from detchecks.api import grade
    tag = new_tag()
    root = tmpdir()
    src = os.path.join(root, "ai_attempt.xlsx")
    tiny_workbook(src)
    workdir = os.path.join(root, "det_checks_recalc")
    pol = R.RecalcPolicy(workdir=workdir, libreoffice_path=fake_soffice(root, "fail", tag), lo_timeout_s=30,
                         lo_retries=2, lo_lock_path=os.path.join(root, "lo.lock"), lo_min_free_pct=0)
    try:
        R.libreoffice_recalc(src, os.path.join(workdir, "out"), pol)
    except LibreOfficeUnavailable as e:
        assert e.retry_later and "all 3 tries failed" in str(e) and str(e).count("produced no copy") == 3, e
    else:
        raise AssertionError("no LibreOfficeUnavailable")
    try:
        grade(src, checks=[22, 94, 92], recalc=pol)
    except GradingError as e:
        assert e.retry_later, e
        assert set(e.failures) == {"Error Checks/No formula errors", "Potential Dangers/No white-on-white hiding"}, e
        assert "Potential Dangers/No hidden sheets" in e.verdicts                     # not a value check: graded
        assert all("values unavailable" in m and "all 3 tries failed" in m for m in e.failures.values()), e.failures
    else:
        raise AssertionError("graded without values")


SLOW_GRADER = r'''
import sys
sys.path.insert(0, {judge!r})
from detchecks.core import recalc as R
pol = R.RecalcPolicy(workdir={workdir!r}, libreoffice_path={soffice!r}, lo_timeout_s=60, lo_lock_path={lock!r},
                     lo_min_free_pct=0, lo_retries=0, lo_log=lambda m: print("LOG", m, flush=True))
print(R.libreoffice_recalc({src!r}, {out!r}, pol), flush=True)
'''


def test_lock_serialises_two_graders():
    """Two grading processes recalculating at the same time share the machine-wide lock: the second
    LibreOffice starts only after the first has finished (stand-in soffice logging its run times)."""
    tag = new_tag()
    root = tmpdir()
    soffice = fake_soffice(root, "slowcopy", tag)
    lock = os.path.join(root, "mbabench_libreoffice.lock")
    procs = []
    for k in range(2):
        task = os.path.join(root, f"task {k}")
        os.makedirs(task)
        src = os.path.join(task, "ai_attempt.xlsx")
        tiny_workbook(src)
        wd = os.path.join(task, "det_checks_recalc")
        code = SLOW_GRADER.format(judge=JUDGE, workdir=wd, soffice=soffice, lock=lock, src=src,
                                  out=os.path.join(wd, "h", "libreoffice"))
        procs.append(subprocess.Popen([sys.executable, "-c", code], stdout=subprocess.PIPE, stderr=subprocess.PIPE,
                                      text=True))
    outs = []
    for p in procs:
        so, se = p.communicate(timeout=120)
        assert p.returncode == 0, se[-800:]
        outs.append(so)
    runs = sorted(json.loads(line) for line in open(soffice + ".log"))
    assert len(runs) == 2 and runs[0][1] <= runs[1][0] + 0.05, f"two LibreOffice runs overlapped: {runs}"
    assert any("waiting for the machine-wide LibreOffice lock" in o for o in outs), outs


def test_watchdog_alone_kills_soffice_when_its_parent_dies():
    """The watchdog without the grader's handler: its parent (a shell that execs nothing) is
    SIGKILLed; the watchdog notices within its poll and kills soffice and soffice's child."""
    tag = new_tag()
    root = tmpdir()
    soffice = fake_soffice(root, "hang", tag)
    parent = subprocess.Popen([sys.executable, "-c",
                               "import subprocess, sys, time; "
                               f"subprocess.Popen([sys.executable, '-I', '-S', {R.WATCHDOG_SCRIPT!r}, "
                               f"str(__import__('os').getpid()), '0.2', '--', {soffice!r}, '--outdir', 'x', 'y']); "
                               "time.sleep(300)"])
    try:
        assert wait_until(lambda: procs_with(tag + "_child"), 30)
        parent.kill()
        parent.wait(timeout=10)
        assert wait_until(lambda: not procs_with(tag, soffice), 10), procs_with(tag, soffice)
    finally:
        if parent.poll() is None:
            parent.kill()
            parent.wait(timeout=10)


HANDLER_PROBE = r'''
import signal, sys, threading
sys.path.insert(0, {judge!r})
from detchecks.core import recalc as R
mine = lambda s, f: None
signal.signal(signal.SIGTERM, mine)
signal.signal(signal.SIGHUP, signal.SIG_IGN)
box = {{}}
t = threading.Thread(target=lambda: box.update(r=R.install_termination_reaper()))
t.start(); t.join()
assert box["r"] is False, "installed from a worker thread"
assert R.install_termination_reaper() is True
assert signal.getsignal(signal.SIGTERM) is mine, "an application handler was replaced"
assert signal.getsignal(signal.SIGHUP) == signal.SIG_IGN, "an ignored signal was replaced"
print("ok")
'''


def test_handler_only_where_the_default_action_was():
    out = subprocess.run([sys.executable, "-c", HANDLER_PROBE.format(judge=JUDGE)], capture_output=True, text=True,
                         timeout=60)
    assert out.returncode == 0 and out.stdout.strip() == "ok", out.stderr[-800:]
    fresh = subprocess.run([sys.executable, "-c",
                            f"import sys, signal; sys.path.insert(0, {JUDGE!r}); from detchecks.core import recalc as R; "
                            "R.install_termination_reaper(); "
                            "print(signal.getsignal(signal.SIGTERM) is R._reap_and_die)"],
                           capture_output=True, text=True, timeout=60)
    assert fresh.stdout.strip() == "True", fresh.stderr[-800:]


def test_reap_profiles_matches_url_and_raw_paths_only_below_the_folder():
    root = tmpdir()
    profiles = os.path.join(root, "Model - v2 %41", "det_checks_recalc", "_lo_profiles")
    sleeper = [sys.executable, "-c", "import time; time.sleep(120)"]
    procs = {
        "url": subprocess.Popen(sleeper + [f"-env:UserInstallation={R.file_url(os.path.join(profiles, 'lo_a'))}"]),
        "raw": subprocess.Popen(sleeper + [f"-env:UserInstallation=file://{os.path.join(profiles, 'lo_b')}"]),
        "other": subprocess.Popen(sleeper + [f"-env:UserInstallation={R.file_url(os.path.join(root, 'elsewhere', 'lo_c'))}"]),
        "sibling": subprocess.Popen(sleeper + [f"-env:UserInstallation={R.file_url(profiles + '_x')}/lo_d"]),
    }
    try:
        assert wait_until(lambda: len(procs_with(root)) + len(procs_with(R.file_url(root))) >= 4, 10)
        killed = R.reap_profiles(profiles)
        assert procs["url"].wait(timeout=10) != 0 and procs["raw"].wait(timeout=10) != 0
        assert procs["url"].pid in killed and procs["raw"].pid in killed
        assert procs["other"].poll() is None and procs["sibling"].poll() is None, "an unrelated process was killed"
    finally:
        for p in procs.values():
            if p.poll() is None:
                p.kill()
                p.wait(timeout=10)


TESTS = [test_file_url_encodes_every_path, test_real_libreoffice_paths_with_space_and_percent,
         test_fake_conversion_through_the_watchdog, test_timeout_kills_the_whole_tree,
         test_retry_after_two_crashes_then_graded, test_failing_every_try_fails_loudly_with_retry_later,
         test_lock_serialises_two_graders,
         test_sigterm_to_the_grader_kills_libreoffice, test_sigkill_to_the_grader_kills_libreoffice,
         test_watchdog_alone_kills_soffice_when_its_parent_dies, test_handler_only_where_the_default_action_was,
         test_reap_profiles_matches_url_and_raw_paths_only_below_the_folder]


def _sweep():
    """Kill anything a test started (unique tags, temp folders) and remove the folders."""
    for t in _TAGS:
        for pid, _cmd in procs_with(t):
            lo_watchdog.kill_tree(pid)
    for d in _TMPS:
        for pid, _cmd in procs_with(d, R.file_url(d)):
            lo_watchdog.kill_tree(pid)
        shutil.rmtree(d, ignore_errors=True)


if pytest is not None:
    @pytest.fixture(autouse=True, scope="module")
    def _cleanup_module():
        yield
        _sweep()


def main() -> int:
    failed = skipped = 0
    try:
        for t in TESTS:
            try:
                t()
                print(f"PASS {t.__name__}")
            except _Skip as e:
                skipped += 1
                print(f"SKIP {t.__name__}: {e}")
            except Exception:  # noqa: BLE001
                failed += 1
                print(f"FAIL {t.__name__}")
                traceback.print_exc()
    finally:
        _sweep()
    passed = len(TESTS) - failed - skipped
    print(f"\n{passed}/{len(TESTS)} tests passed" + (f", {skipped} SKIPPED (not passed)" if skipped else ""))
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main())
