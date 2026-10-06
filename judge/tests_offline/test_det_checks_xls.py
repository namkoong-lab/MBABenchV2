"""judge v13 with an attempt delivered in the legacy binary .xls format (the maintainer, 2026-10-05: "just
keep doing whatever v12 did or does"; utils/det_checks.py, "Legacy .xls deliveries").

Run from judge/ - it starts the REAL LibreOffice (paths.libreoffice_path), one run at a
time under the machine-wide guard:
    python tests_offline/test_det_checks_xls.py
(or pytest). Every test is SKIPPED (reported, never counted as passed) where LibreOffice is absent.

The .xls is built once per run: a small openpyxl workbook (a hidden sheet, a merged block, B3 = B1 - B2 = 0)
converted with LibreOffice (--convert-to xls: private profile, encoded file URLs, the watchdog, under the
guard), then B3's cached result in the BIFF FORMULA record is set to a stale 999, so the tests can tell the
value LibreOffice calculates from the value the file carries. Then:
  - the adapter with run_calculation: ONE LibreOffice run converts the delivered .xls to .xlsx in
    det_checks_recalc/xls_converted/; every check except File extension (.xlsx) (77) grades that copy, with
    the values LibreOffice calculated while converting (B3 = 0, so Zeros as dashes (66) fails at B3); 77
    grades the delivered file and name and fails; det_checks.json, the summary and every verdict's stats
    say so; det_checks_recalc/ and every LibreOffice process of the grading are gone;
  - the adapter without run_calculation: fails loudly, names every check but 77, no LibreOffice run;
  - an OLE2 file without a workbook stream is no .xls: never converted, refused as before;
  - grade_from_db.grade_single_attempt end to end with the stub LLM of test_det_checks_single_pass: with
    run_calculation the attempt is graded (the checks convert first, then the judge's own --run-calculation
    re-save, then the LLM); without it the grading fails before the LLM call, with no LibreOffice run.
No DB, S3, Excel or LLM API.
"""
import contextlib
import hashlib
import json
import os
import re
import shutil
import struct
import subprocess
import sys
import tempfile
import time
from pathlib import Path

JUDGE = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(JUDGE))
sys.path.insert(0, str(JUDGE / "tests_offline"))

import openpyxl  # noqa: E402

import test_det_checks_single_pass as SP  # noqa: E402  (v2 config, grade_from_db, the stub LLM, staging helpers)
from detchecks.core import lo_watchdog  # noqa: E402

D, gfd, det_recalc = SP.D, SP.gfd, SP.det_recalc
WEIGHTS_PATH = JUDGE / "prompts" / "rubrics" / "rubric_9_weights.json"
DELIVERED = "Model_final.xls"
STALE = 999.0


def key(no):
    return "/".join(D.DET_CHECK_NAMES[no])


K77 = key(77)


# --------------------------------------------------------------------------- the .xls fixture
def _soffice() -> str:
    path = D.load_settings().libreoffice_path
    if not (path and os.path.exists(path)):
        SP._skip(f"LibreOffice not found at {path!r} (paths.libreoffice_path)")
    return path


def _convert_to_xls(src: Path, out_dir: Path) -> Path:
    """LibreOffice --convert-to xls of `src`, run the way the pipeline runs LibreOffice: private profile,
    encoded file URLs, under the watchdog and the machine-wide guard (utils.det_checks.run_libreoffice);
    the profile is swept and removed."""
    soffice = _soffice()
    out = out_dir / (src.stem + ".xls")

    def once(timeout_s):
        profile = Path(tempfile.mkdtemp(prefix="lo_", dir=out_dir))
        try:
            (profile / "user").mkdir()
            (profile / "user" / "registrymodifications.xcu").write_text(det_recalc._LO_REGISTRY)
            cmd = [soffice, f"-env:UserInstallation={det_recalc.file_url(str(profile))}", "--headless",
                   "--norestore", "--nologo", "--nofirststartwizard", "--calc", "--convert-to", "xls", "--outdir",
                   det_recalc.file_url(str(out_dir)), det_recalc.file_url(str(src))]
            subprocess.run(det_recalc._watchdog_cmd(cmd), stdin=subprocess.DEVNULL, capture_output=True,
                           timeout=timeout_s, env=dict(os.environ, SAL_USE_VCLPLUGIN="svp"))
        finally:
            det_recalc.reap_profiles(str(profile), subtree=False)
            shutil.rmtree(profile, ignore_errors=True)
        if not out.exists():
            raise FileNotFoundError(f"LibreOffice produced no .xls for {src}")
        return out

    return D.run_libreoffice(once, timeout_s=120, what=f"convert {src.name} to .xls (test fixture)")


_XLS = {}


def xls_bytes() -> bytes:
    """The delivered .xls (built once per run; module doc)."""
    if "data" not in _XLS:
        _soffice()
        with tempfile.TemporaryDirectory() as tmp:
            src = Path(tmp) / "model.xlsx"
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Model"
            ws["A1"], ws["B1"] = "Revenue", 100
            ws["A2"], ws["B2"] = "Cost", 100
            ws["A3"], ws["B3"] = "Profit", "=B1-B2"
            ws["A5"] = "Notes"
            ws.merge_cells("A5:C5")
            hidden = wb.create_sheet("Workings")
            hidden["A1"] = "source: case pack"
            hidden.sheet_state = "hidden"
            wb.save(src)
            data = _convert_to_xls(src, Path(tmp)).read_bytes()
        assert data[:8] == bytes.fromhex("d0cf11e0a1b11ae1"), "LibreOffice did not write an OLE2 .xls"
        # B3 (row 2, column 1) FORMULA record (type 0x0006): its cached result LibreOffice computed, 0.0 -> STALE
        hits = list(re.finditer(rb"\x06\x00..\x02\x00\x01\x00.." + re.escape(struct.pack("<d", 0.0)), data, re.S))
        assert len(hits) == 1, f"B3's FORMULA record not found exactly once ({len(hits)})"
        at = hits[0].end() - 8
        _XLS["data"] = data[:at] + struct.pack("<d", STALE) + data[at + 8:]
    return _XLS["data"]


def make_xls_task(root: Path, data: bytes = None) -> Path:
    """A staged task folder: the .xls bytes as ai_attempt.xlsx (production's staging name) + the sidecar."""
    folder = Path(tempfile.mkdtemp(dir=root))
    (folder / D.ATTEMPT_FILENAME).write_bytes(xls_bytes() if data is None else data)
    (folder / "_attempt_origin.json").write_text(json.dumps({"original_filename": DELIVERED, "source": "test"}))
    return folder


def run(folder, **kw):
    kw.setdefault("benchmark", None)
    return D.run_det_checks(folder, rubric_path=SP.RUBRIC_PATH, weights_path=WEIGHTS_PATH, **kw)


@contextlib.contextmanager
def recorded_libreoffice():
    """Record, and let run, every LibreOffice run: the det checks' (detchecks' libreoffice_recalc) and the
    judge's other ones (utils.det_checks.run_libreoffice: --run-calculation, the answer check)."""
    calls = []
    orig_det, orig_run = det_recalc.libreoffice_recalc, D.run_libreoffice

    def _det(src, out_dir, policy):
        calls.append(("det_checks", src, out_dir))
        return orig_det(src, out_dir, policy)

    def _run(run_once, **kw):
        calls.append(("judge", kw.get("what"), None))
        return orig_run(run_once, **kw)

    det_recalc.libreoffice_recalc, D.run_libreoffice = _det, _run
    try:
        yield calls
    finally:
        det_recalc.libreoffice_recalc, D.run_libreoffice = orig_det, orig_run


def libreoffice_left(root: Path) -> list:
    """Processes whose command line names `root` (raw path or file URL): a LibreOffice run of this test that
    is still alive (the conversion's profile, outdir and source all live under root)."""
    markers = (str(root), det_recalc.file_url(str(root)))
    deadline = time.time() + 15
    while True:
        table = lo_watchdog.process_table() or []
        mine = lo_watchdog.ancestors(table)
        left = [(pid, cmd[:200]) for pid, _pp, cmd in table
                if pid not in mine and not cmd.startswith("ps ") and any(m in cmd for m in markers)]
        if not left or time.time() > deadline:
            return left
        time.sleep(0.5)


def strict_json(text):
    def _bad(c):
        raise ValueError(f"non-standard JSON constant {c}")
    return json.loads(text, parse_constant=_bad)


# --------------------------------------------------------------------------- the adapter
def test_run_calculation_grades_libreoffices_conversion():
    with tempfile.TemporaryDirectory() as tmp:
        root = Path(tmp)
        folder = make_xls_task(root)
        attempt = folder / D.ATTEMPT_FILENAME
        assert D.legacy_xls(attempt)
        with recorded_libreoffice() as calls:
            r = run(folder, run_calculation=True)
        # ONE LibreOffice run: the conversion of the delivered file into det_checks_recalc/xls_converted/
        assert [(c[0], Path(c[1]), Path(c[2])) for c in calls] == \
            [("det_checks", attempt, folder / D.RECALC_DIRNAME / D.XLS_COPY_DIRNAME)], calls
        s = r.summary
        settings = D.load_settings()
        graded = set(settings.live) | set(settings.recorded_only)
        assert s["status"] == "ok" and set(s["graded"]) == graded and set(r.harness_verdicts) == {key(n) for n in graded}
        # File extension (.xlsx) (77): the delivered file and its delivered name - fails
        e77 = r.harness_verdicts[K77]
        assert e77["decision"] == "fail" and DELIVERED in e77["summary"], e77["summary"]
        assert e77["stats"]["graded_on"] == D.GRADED_ON_DELIVERED and e77["stats"]["reader_format"] == "xls"
        assert e77["stats"]["delivered_filename"] == DELIVERED and e77["stats"]["container"] == "ole2"
        # every other check: a verdict on LibreOffice's conversion
        for no in graded - {77}:
            e = r.harness_verdicts[key(no)]
            assert e["decision"] in ("pass", "fail") and e["stats"]["graded_on"] == D.GRADED_ON_COPY, (no, e)
        assert r.harness_verdicts[key(92)]["decision"] == "fail", "the hidden sheet survives the conversion"
        assert r.harness_verdicts[key(74)]["decision"] == "fail", "the merged block survives the conversion"
        # values: the ones LibreOffice calculated while converting (B3 = 0), not the stale 999 the .xls carries
        e66 = r.harness_verdicts[key(66)]
        assert e66["decision"] == "fail" and any("B3" in m["location"] for m in e66["mistakes"]), e66
        values = s["values"]
        assert values["source"] == "libreoffice" and values["writer"] == "libreoffice", values
        assert D.XLS_VALUES_NOTE in values["notes"], values["notes"]
        xc = s["xls_conversion"]
        assert {k: v for k, v in xc.items() if k != "conversion_s"} == {
            "delivered_format": "xls", "run_calculation": True, "converted": True, "converted_by": "libreoffice",
            "graded_on_delivered": [77]}, xc
        assert values["timings"]["libreoffice_s"] == xc["conversion_s"] > 0
        # det_checks.json: the delivered file, and the converted copy it was graded on
        art = strict_json((folder / D.ARTEFACT_FILENAME).read_text())
        assert art["status"] == "ok" and art["file"]["path"] == str(attempt)
        assert art["file"]["sha256"] == hashlib.sha256(xls_bytes()).hexdigest()
        block = art["xls_conversion"]
        assert block["converted"] is True and block["run_calculation"] is True
        assert block["copy"] == str(folder / D.RECALC_DIRNAME / D.XLS_COPY_DIRNAME / D.ATTEMPT_FILENAME)
        assert block["copy_bytes"] > 0 and len(block["copy_sha256"]) == 64
        assert set(block["graded_on_copy"]) == graded - {77} and block["graded_on_delivered"] == [77]
        assert art["verdicts"][key(92)]["stats"]["graded_on"] == D.GRADED_ON_COPY
        json.dumps(s, allow_nan=False)
        # nothing left: the copy and the profiles, and every LibreOffice process of this grading
        assert not (folder / D.RECALC_DIRNAME).exists(), list((folder / D.RECALC_DIRNAME).rglob("*"))
        assert libreoffice_left(root) == []


def test_without_run_calculation_the_grading_fails_loudly():
    with tempfile.TemporaryDirectory() as tmp:
        root = Path(tmp)
        folder = make_xls_task(root)
        with recorded_libreoffice() as calls:
            try:
                run(folder, run_calculation=False)
            except D.DetChecksError as e:
                msg = str(e)
                settings = D.load_settings()
                graded = set(settings.live) | set(settings.recorded_only)
                assert set(e.failures) == {key(n) for n in graded - {77}}, sorted(e.failures)
                assert "legacy binary .xls" in msg and "re-run with --run-calculation" in msg, msg
                assert "before the LLM call" in msg and not e.retry_later
                assert e.verdicts[K77]["decision"] == "fail"           # 77 graded the delivered file
            else:
                raise AssertionError("an .xls delivery was graded without --run-calculation")
        assert calls == [], "no LibreOffice run without --run-calculation"
        art = strict_json((folder / D.ARTEFACT_FILENAME).read_text())
        assert art["status"] == "error" and art["stage"] == "grade"
        assert art["xls_conversion"] == {"delivered_format": "xls", "run_calculation": False, "converted": False,
                                         "note": D.XLS_NOT_CONVERTED_NOTE}
        assert not (folder / D.RECALC_DIRNAME).exists()


def test_an_ole2_file_without_a_workbook_stream_is_not_converted():
    """Content, not the name, decides: an OLE2 file whose directory names no Workbook / Book stream (here the
    .xls with that entry renamed) is no .xls workbook, so it is never sent to LibreOffice and is refused as
    before - and nothing else is taken for one."""
    data = xls_bytes()
    name = "Workbook".encode("utf-16-le") + b"\x00\x00"
    assert data.count(name) == 1
    renamed = data.replace(name, "Wxrkbook".encode("utf-16-le") + b"\x00\x00")
    with tempfile.TemporaryDirectory() as tmp:
        root = Path(tmp)
        folder = make_xls_task(root, renamed)
        assert not D.legacy_xls(folder / D.ATTEMPT_FILENAME)
        with recorded_libreoffice() as calls:
            try:
                run(folder, run_calculation=True)
            except D.DetChecksError as e:
                assert "cannot grade a 'xls' file" in str(e) and K77 not in e.failures, str(e)
            else:
                raise AssertionError("an OLE2 file without a workbook stream was graded")
        assert calls == []
        assert "xls_conversion" not in strict_json((folder / D.ARTEFACT_FILENAME).read_text())
        # nothing else is an .xls: a zip workbook, a missing file
        xlsx = root / "book.xlsx"
        openpyxl.Workbook().save(xlsx)
        assert not D.legacy_xls(xlsx) and not D.legacy_xls(root / "missing.xls")
        assert D.legacy_xls(make_xls_task(root) / D.ATTEMPT_FILENAME)


# --------------------------------------------------------------------------- end to end
def _grade(root: Path, attempt: dict, llm, run_calculation: bool, run: str = "run"):
    with SP.env(PATHS_SCRATCH_PATH=root / "judge_scratch"):
        return gfd.grade_single_attempt(
            attempt=attempt, client=llm, rubric_path=str(SP.RUBRIC_PATH), template_path=SP.TEMPLATE,
            agentic_template_path=SP.TEMPLATE, model=SP.MODEL, scratch_run_dir=root / run,
            agentic=True, single_pass=True, max_tool_rounds=4, max_forced_rounds=1,
            no_s3_upload=True, suitability_source_path=root / "annotation.json",
            accuracy_check="harness", det_checks=None, run_calculation=run_calculation,
        )


def _staged(root: Path):
    src = root / "fixture.xls"
    src.write_bytes(xls_bytes())
    return SP._stage(root, src, DELIVERED)


def test_end_to_end_run_calculation_grades_the_xls_attempt():
    """grade_single_attempt with --run-calculation: the checks convert and grade first, then the judge's own
    --run-calculation re-save (v12's code, temp_recalculated/) and the LLM; the grading succeeds."""
    with tempfile.TemporaryDirectory() as tmp:
        root = Path(tmp)
        attempt = _staged(root)
        llm = SP.StubLLM(root / "run")
        with recorded_libreoffice() as calls:
            res = _grade(root, attempt, llm, run_calculation=True)
        assert res["success"], res.get("error")
        assert llm.calls == 2, llm.calls
        task_folder = Path(res["task_folder"])
        # order: the det checks' conversion of the delivered file, then the judge's --run-calculation re-save
        assert [c[0] for c in calls] == ["det_checks", "judge"], calls
        assert Path(calls[0][1]) == task_folder / D.ATTEMPT_FILENAME and "--run-calculation" in calls[1][1], calls
        resaved = task_folder / "temp_recalculated" / D.ATTEMPT_FILENAME
        assert resaved.read_bytes()[:2] == b"PK", "the judge's own re-save is a real .xlsx"
        sr = res["scored_results"]
        assert sr["det_checks"]["xls_conversion"]["converted"] is True
        assert sr["det_checks"]["xls_conversion"]["graded_on_delivered"] == [77]
        checks = sr["accuracy_engine"]["checks"]
        det = {k: v for k, v in checks.items() if v.get("family") == "det_checks"}
        assert det and checks[K77]["decision"] == "fail" and checks[K77]["counted"] is True
        assert checks[K77]["stats"]["graded_on"] == D.GRADED_ON_DELIVERED
        assert all(v["stats"]["graded_on"] == D.GRADED_ON_COPY for k, v in det.items() if k != K77), det.keys()
        assert sr["check_scores"]["Formatting"]["File extension (.xlsx)"]["score"] == 0.0
        assert not res["hard_parse_failures"] and not res["missing_scores"]
        assert not (task_folder / D.RECALC_DIRNAME).exists()
        json.dumps(sr, allow_nan=False)
        gfd.prune_workbook_copies(res)                       # what grade_from_db.main does after a grading
        assert libreoffice_left(root) == []


def test_end_to_end_without_run_calculation_fails_before_the_llm():
    with tempfile.TemporaryDirectory() as tmp:
        root = Path(tmp)
        attempt = _staged(root)
        llm = SP.StubLLM(root / "run")
        with recorded_libreoffice() as calls:
            res = _grade(root, attempt, llm, run_calculation=False)
        assert res["success"] is False and "scores" not in res and llm.calls == 0, res.get("error")
        assert calls == [], f"LibreOffice ran: {calls}"
        err = res["error"]
        assert "deterministic checks could not grade" in err and "re-run with --run-calculation" in err, err
        assert f"delivered as '{DELIVERED}'" in err and "DetChecksError" in res["traceback"]
        task_folder = Path(res["task_folder"])
        assert not (task_folder / "answer_check.json").exists(), "the checks run before the answer check"
        art = strict_json((task_folder / D.ARTEFACT_FILENAME).read_text())
        assert art["status"] == "error" and art["xls_conversion"]["converted"] is False
        assert libreoffice_left(root) == []


if __name__ == "__main__":
    tests = [(n, f) for n, f in sorted(globals().items()) if n.startswith("test_") and callable(f)]
    failed = skipped = 0
    for name, fn in tests:
        try:
            fn()
            print(f"OK   {name}")
        except SP._Skipped as e:
            skipped += 1
            print(f"SKIPPED {name}: {e}")
        except Exception as e:  # noqa: BLE001
            failed += 1
            import traceback
            traceback.print_exc()
            print(f"FAIL {name}: {e}")
    print()
    print(f"{len(tests) - failed - skipped}/{len(tests)} passed" + (f", {skipped} SKIPPED (not passed)" if skipped else ""))
    sys.exit(1 if failed else 0)
