"""Offline: the OVERSIZED RANGE evidence tag (judge/utils/workbook_properties.py).

Rubric_9 check 25 (Reasonable file size) names oversized used ranges as bloat.
The per-sheet line of the properties block tags a sheet whose declared extent
reaches Excel's full width (16,384 columns) or height (1,048,576 rows) while
its content ends far earlier. Deliberately extreme-only: a ratio rule flagged
59 of 471 cached real workbooks, goldens included.

Run from judge/:  python tests_offline/test_oversized_range.py
"""
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import openpyxl  # noqa: E402
from openpyxl.styles import Font  # noqa: E402

from utils.workbook_properties import (  # noqa: E402
    extract_workbook_properties, oversized_range_tag, render_properties_text,
)


def _sheet(max_row, max_column, used_range):
    return {"index": 1, "name": "Solution Model", "kind": "worksheet", "state": "visible",
            "max_row": max_row, "max_column": max_column, "used_range": used_range,
            "n_values": 10, "n_formulas": 5}


def test_full_width_is_flagged():
    tag = oversized_range_tag(_sheet(276, 16384, "A1:EW272"))  # the toy 25 Fail
    assert "OVERSIZED RANGE" in tag and "16,384 columns" in tag and "column 153" in tag and "107x" in tag, tag


def test_full_height_is_flagged():
    tag = oversized_range_tag(_sheet(1048576, 20, "A1:J91"))
    assert "OVERSIZED RANGE" in tag and "1,048,576 rows" in tag and "row 91" in tag, tag


def test_modest_extent_is_not_flagged():
    assert oversized_range_tag(_sheet(276, 302, "A1:EW272")) == ""   # the toy 25 Pass
    assert oversized_range_tag(_sheet(125, 305, "A1:G56")) == ""     # a real golden's Summary sheet


def test_empty_or_unknown_is_not_flagged():
    assert oversized_range_tag(_sheet(0, 0, None)) == ""
    assert oversized_range_tag(_sheet(1048576, 16384, "unknown")) == ""
    assert oversized_range_tag(_sheet(1048576, 16384, "")) == ""


def test_rendered_from_a_real_workbook(tmp_path=Path("/tmp")):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Solution Model"
    ws["A1"] = "x"
    ws["C3"] = 5
    ws.cell(row=1, column=16384).font = Font(bold=True)  # styled, no value -> declared 16,384 wide
    ok = wb.create_sheet("Tidy")
    ok["A1"] = 1
    path = tmp_path / "oversized_range_test.xlsx"
    wb.save(path)
    props = extract_workbook_properties(openpyxl.load_workbook(path), path)
    text = render_properties_text(props)
    lines = {l.strip().split(".")[1].split(":")[0].strip(): l for l in text.splitlines() if ". " in l and "(used " in l}
    assert "OVERSIZED RANGE" in lines["Solution Model"], lines["Solution Model"]
    assert "OVERSIZED RANGE" not in lines["Tidy"], lines["Tidy"]


if __name__ == "__main__":
    for name, fn in list(globals().items()):
        if name.startswith("test_") and callable(fn):
            fn()
            print(f"ok  {name}")
    print("all oversized-range tests passed")
