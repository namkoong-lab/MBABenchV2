"""Offline end-to-end tests of judge v13: the single-pass flow with the deterministic checks.

Run from judge/:  python tests_offline/test_det_checks_single_pass.py
(it opens the attempt and a golden: on the shared 16 GB grading Mac, run it through the
heavy-job guard like any multi-workbook job).

No DB, S3, LLM, LibreOffice or Excel. grade_from_db.grade_single_attempt runs for real —
staging (with the _attempt_origin.json sidecar), the score-neutral answer check, the
deterministic checks, CSV extraction, the formula-cache gate, suitability gating, the
single-pass tool loop, scoring, artefacts — and only the LLM is a stub client: it records
every check in one round (the OPPOSITE of Python's decision on every deterministically
graded check, so each overlay is visible) and stops in the next.

The attempt is a real delivered workbook: corpus attempt 1375 (EasyDCF, 57 KB, saved by
Excel Online, so its own caches are the values and no LibreOffice recalculation is needed),
copied into a temporary folder; never into the repo. Set DETCHECKS_E2E_ATTEMPT to use
another Excel-saved file under 1 MB. The golden is a synthetic openpyxl workbook (the
answer check then stays with the LLM: no question convention).
"""
import contextlib
import copy
import importlib.util
import json
import os
import re
import shutil
import sys
import tempfile
from pathlib import Path
from types import SimpleNamespace

JUDGE = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(JUDGE))

import openpyxl  # noqa: E402

from utils.misc_utils import load_env_var, load_project_configs, project_prefix  # noqa: E402

load_project_configs(benchmark="v2")

import detchecks.core.recalc as det_recalc  # noqa: E402
from openai.types.chat import ChatCompletion  # noqa: E402
from utils import answer_check, rubric_suitability  # noqa: E402
from utils import det_checks as D  # noqa: E402

_spec = importlib.util.spec_from_file_location("grade_from_db", JUDGE / "main_scripts" / "grade_from_db.py")
gfd = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(gfd)

RUBRIC_PATH = JUDGE / "prompts" / "rubrics" / "rubric_9.json"
RUBRIC = json.loads(RUBRIC_PATH.read_text())
FLAT = [(cat, c["name"]) for cat, checks in RUBRIC.items() for c in checks]
TEMPLATE = str(JUDGE / load_env_var("SINGLE_PASS_PROMPT_TEMPLATE"))
MODEL = load_env_var("JUDGE_DEFAULT_GRADER")

ATTEMPT = Path(os.environ.get(
    "DETCHECKS_E2E_ATTEMPT",
    str(Path.home() / "MBABench-deterministic-checks" / "corpus" / "attempts" / "1375"
        / "20260911_080041_1_EasyDCF_Solution_jp_claude_excel_agent_Model.xlsx")))
NOT_APPLICABLE = {51}           # Red font for external links: not applicable to any task today


# --------------------------------------------------------------------------- stubs
class StubLLM:
    """chat.completions.create stand-in. Round 1: record_check for every check number the seed
    lists (deterministic checks get the opposite of Python's decision, read from the
    det_checks.json the adapter wrote BEFORE the call; Final calculation accuracy fails;
    everything else passes) plus one append_mistake per fail. Round 2: stop."""

    def __init__(self, scratch: Path):
        self.scratch = scratch
        self.calls = 0
        self.decisions = {}
        self.chat = SimpleNamespace(completions=SimpleNamespace(create=self.create))

    def _python_decisions(self) -> dict:
        found = list(self.scratch.rglob(D.ARTEFACT_FILENAME))
        assert len(found) == 1, f"det_checks.json must exist before the LLM call: {found}"
        art = json.loads(found[0].read_text())
        return {int(v_no): v["decision"] for v_no, v in
                ((D._NUMBER_BY_KEY[k], v) for k, v in (art.get("verdicts") or {}).items())}

    def create(self, **kwargs):
        self.calls += 1
        if self.calls > 1:
            return _completion(content="All checks recorded.")
        seed = "\n".join(str(m.get("content")) for m in kwargs["messages"] if isinstance(m, dict))
        m = re.search(r"Check numbers you must record \((\d+) checks\): ([\d, ]+)", seed)
        assert m, "check list not found in the seed"
        ids = [int(t) for t in m.group(2).replace(" ", "").split(",") if t]
        assert len(ids) == int(m.group(1))
        python = self._python_decisions()
        calls = []
        for no in ids:
            if no in python:
                decision = "fail" if python[no] == "pass" else "pass"
            else:
                decision = "fail" if no == 1 else "pass"
            self.decisions[no] = decision
            calls.append(("record_check", {"check": str(no), "decision": decision, "summary": "stub"}))
            if decision == "fail":
                calls.append(("append_mistake", {"check": str(no), "location": "'Summary'!A1",
                                                 "description": "stub mistake", "severity": "major"}))
        return _completion(tool_calls=calls)


def _completion(tool_calls=None, content=None):
    msg = {"role": "assistant", "content": content}
    if tool_calls:
        msg["tool_calls"] = [{"id": f"call_{i}", "type": "function",
                              "function": {"name": n, "arguments": json.dumps(a)}}
                             for i, (n, a) in enumerate(tool_calls)]
    return ChatCompletion.model_validate({
        "id": "stub", "object": "chat.completion", "created": 0, "model": "stub",
        "choices": [{"index": 0, "finish_reason": "tool_calls" if tool_calls else "stop", "message": msg}],
        "usage": {"prompt_tokens": 1000, "completion_tokens": 10, "total_tokens": 1010},
    })


@contextlib.contextmanager
def no_libreoffice(det_checks=True):
    """Record (and refuse) any LibreOffice run: neither the deterministic checks' recalculation
    nor the answer check's own may start one in these tests. det_checks=False keeps detchecks'
    real runner (which then fails on a missing binary before launching anything)."""
    calls = []
    orig_det, orig_ac = det_recalc.libreoffice_recalc, answer_check._recalculate_copy

    def _det(src, out_dir, policy):
        calls.append(("det_checks", src))
        raise det_recalc.GradingError("LibreOffice is not allowed in this test")

    def _ac(*a, **k):
        calls.append(("answer_check", a))
        raise RuntimeError("LibreOffice is not allowed in this test")

    if det_checks:
        det_recalc.libreoffice_recalc = _det
    answer_check._recalculate_copy = _ac
    try:
        yield calls
    finally:
        det_recalc.libreoffice_recalc, answer_check._recalculate_copy = orig_det, orig_ac


@contextlib.contextmanager
def env(**overrides):
    prefix = project_prefix()
    saved = {f"{prefix}_{k}": os.environ.get(f"{prefix}_{k}") for k in overrides}
    try:
        for k, v in overrides.items():
            os.environ[f"{prefix}_{k}"] = str(v)
        yield
    finally:
        for name, v in saved.items():
            if v is None:
                os.environ.pop(name, None)
            else:
                os.environ[name] = v


def _golden(path: Path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Model"
    ws["A1"], ws["B1"] = "Enterprise value", 1250.0
    ws["A2"], ws["B2"] = "WACC", 0.085
    wb.save(path)


def _annotation(path: Path):
    rubrics = [{"no": i, "category": cat, "name": name, "conditional": False,
                "verdict": "not_applicable" if i in NOT_APPLICABLE else "applicable"}
               for i, (cat, name) in enumerate(FLAT, 1)]
    path.write_text(json.dumps({"annotator": "julian", "complete": True, "created_at": "2026-10-04",
                                "rubric_version": "test", "rubrics": rubrics,
                                "_staging": {"s3_key": "test/annotation.json"}}))


def _stage(root: Path, attempt_file: Path, delivered_name: str):
    src = root / "src"
    src.mkdir(parents=True, exist_ok=True)
    shutil.copy(attempt_file, src / delivered_name)
    _golden(src / "EasyDCF - Solution.xlsx")
    _annotation(root / "annotation.json")
    return {
        "attempt_id": 1375, "task_id": 20, "task_name": "EasyDCF",
        "agent_model_name": "claude_excel_agent", "agent_model_type": "excel", "agent_failed": False,
        "attempt_files": [{"name": delivered_name, "path": str(src / delivered_name)}],
        "task_solution_files": [{"name": "EasyDCF - Solution.xlsx", "path": str(src / "EasyDCF - Solution.xlsx")}],
        "task_starting_files": None,
    }


def _grade(root: Path, attempt: dict, det_mode, llm: StubLLM):
    return gfd.grade_single_attempt(
        attempt=attempt, client=llm, rubric_path=str(RUBRIC_PATH), template_path=TEMPLATE,
        agentic_template_path=TEMPLATE, model=MODEL, scratch_run_dir=root / "run",
        agentic=True, single_pass=True, max_tool_rounds=4, max_forced_rounds=1,
        no_s3_upload=True, suitability_source_path=root / "annotation.json",
        accuracy_check="harness", det_checks=det_mode,
    )


def _attempt_or_skip():
    if not ATTEMPT.exists():
        print(f"SKIP: corpus attempt not found ({ATTEMPT}); set DETCHECKS_E2E_ATTEMPT")
        return None
    assert ATTEMPT.stat().st_size < 1_000_000, "pick an attempt under 1 MB (maintainer: no large files)"
    return ATTEMPT


# --------------------------------------------------------------------------- tests
def test_end_to_end_live_and_recorded_only():
    attempt_file = _attempt_or_skip()
    if attempt_file is None:
        return
    with tempfile.TemporaryDirectory() as tmp, no_libreoffice() as lo_calls:
        root = Path(tmp)
        attempt = _stage(root, attempt_file, attempt_file.name)
        llm = StubLLM(root / "run")
        # A LibreOffice size limit below the file's size: an Excel-saved file of any size is
        # still graded, because its own caches are the values (no recalculation).
        with env(DET_CHECKS_LIBREOFFICE_MAX_MB=0.01):
            res = _grade(root, attempt, None, llm)          # None = config default (harness)
        assert res["success"], res.get("error")
        assert llm.calls == 2, llm.calls
        assert lo_calls == [], lo_calls
        assert not res["hard_parse_failures"] and not res["missing_scores"] and not res["has_scoring_warnings"]
        assert res["versions"]["JUDGE_VERSION"] == "13" and res["versions"]["PROMPT_VERSION"] == "8"

        out = Path(res["output_dir"])
        task_folder = Path(res["task_folder"])
        sr = res["scored_results"]
        eng = sr["accuracy_engine"]
        art = json.loads((out / D.ARTEFACT_FILENAME).read_text())
        assert art["status"] == "ok" and art["mode"] == "harness"
        assert art["task_meta"] == {"requires_external_links": False, "delivered_filename": attempt_file.name}
        assert art["file"]["path"] == str(task_folder / "ai_attempt.xlsx")
        settings = D.load_settings()
        graded = set(art["graded"])
        assert graded == (set(settings.live) | set(settings.recorded_only)) - NOT_APPLICABLE
        assert set(art["not_applicable"]) == NOT_APPLICABLE
        assert sr["det_checks"]["status"] == "ok" and set(sr["det_checks"]["graded"]) == graded
        assert eng["det_checks_mode"] == "harness" and eng["effective"] == "harness"
        assert sr["total_score"] == eng["total_score_harness"] != eng["total_score_llm"]
        assert res["scores"]["final_score"] == sr["total_score"]
        json.dumps(sr, allow_nan=False)                       # what goes into gradings.scored_results

        live = set(settings.live)
        harness_items = json.loads((out / "ai_judgement_harness.json").read_text())
        pure = json.loads((out / "ai_judgement.json").read_text())
        assert "decided_by" not in json.dumps(pure)
        for key, v in art["verdicts"].items():
            no = D._NUMBER_BY_KEY[key]
            cat, name = D.DET_CHECK_NAMES[no]
            p = eng["checks"][key]
            py, llm_dec = v["decision"], llm.decisions[no]
            assert llm_dec != py
            assert p["decision"] == py and p["llm_decision"] == llm_dec and p["agreed"] is False
            assert p["family"] == "det_checks" and p["check_no"] == no and isinstance(p["stats"], dict)
            score = sr["check_scores"][cat][name]["score"]
            item = next(i for i in harness_items[cat] if i["name"] == name)
            if no in live:
                assert p["engine"] == "harness" and p["live"] is True and p["counted"] is True, key
                assert score == (1.0 if py == "pass" else 0.0), (key, score)
                assert item["decided_by"] == "harness" and item["llm_decision"] == llm_dec
            else:                                             # No formula errors (22), Negatives in parentheses (65)
                assert no in (22, 65)
                assert p["engine"] == "llm" and p["live"] is False and p["counted"] is False, key
                assert p["fallback_reason"] == "recorded only; the LLM verdict stands at scoring"
                assert score == (1.0 if llm_dec == "pass" else 0.0), (key, score)
                assert "decided_by" not in item and item["decision"] == llm_dec
        # gated: Red font for external links (51) is neither graded nor scored
        k51 = "/".join(D.DET_CHECK_NAMES[51])
        assert k51 not in eng["checks"] and "Red font for external links" not in sr["check_scores"]["Formatting"]
        # the answer check still rides along (its family, its own switch)
        assert "Accuracy/Final calculation accuracy" in eng["checks"] and "answer_check" in sr
        # workbook copies go, the det_checks record stays
        gfd.prune_workbook_copies(res)
        assert not (task_folder / "ai_attempt.xlsx").exists()
        assert (task_folder / D.ARTEFACT_FILENAME).exists() and (out / D.ARTEFACT_FILENAME).exists()
        harness_total = sr["total_score"]

        # Shadow mode (--det-checks llm): same verdicts recorded, the LLM's count.
        root2 = root / "shadow"
        attempt2 = _stage(root2, attempt_file, attempt_file.name)
        llm2 = StubLLM(root2 / "run")
        res2 = _grade(root2, attempt2, "llm", llm2)
        assert res2["success"], res2.get("error")
        sr2, eng2 = res2["scored_results"], res2["scored_results"]["accuracy_engine"]
        assert eng2["det_checks_mode"] == "llm" and eng2["effective"] == "llm"
        assert sr2["det_checks"]["mode"] == "llm"
        assert sr2["total_score"] == eng2["total_score_llm"]
        assert abs(eng2["total_score_harness"] - harness_total) < 1e-9, "shadow run carries the v13 total"
        k92 = "/".join(D.DET_CHECK_NAMES[92])
        assert eng2["checks"][k92]["engine"] == "harness" and eng2["checks"][k92]["counted"] is False
        assert sr2["check_scores"]["Potential Dangers"]["No hidden sheets"]["score"] == \
            (1.0 if llm2.decisions[92] == "pass" else 0.0)

        # --det-checks off: nothing runs, the row says so
        root3 = root / "off"
        attempt3 = _stage(root3, attempt_file, attempt_file.name)
        llm3 = StubLLM(root3 / "run")
        res3 = _grade(root3, attempt3, "off", llm3)
        assert res3["success"], res3.get("error")
        sr3 = res3["scored_results"]
        assert sr3["det_checks"] == {"status": "off", "mode": "off"}
        assert sr3["accuracy_engine"]["det_checks_mode"] == "off"
        assert not any(v.get("family") == "det_checks" for v in sr3["accuracy_engine"]["checks"].values())
        assert lo_calls == []


def test_grading_error_fails_before_the_llm_call():
    """A check that cannot grade the file stops the attempt in grade_single_attempt exactly like
    the formula-cache refusal: FAILED, success False, no scores (so no DB row), and the stub
    LLM is never called. Here: an openpyxl-written attempt whose value checks need a
    LibreOffice recalculation on a host where paths.libreoffice_path does not exist."""
    with tempfile.TemporaryDirectory() as tmp:
        root = Path(tmp)
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Model"
        ws["A1"], ws["B1"], ws["B2"] = "Revenue", 100, "=B1*2"    # a formula without a cached value
        built = root / "built.xlsx"
        wb.save(built)
        attempt = _stage(root, built, "Model_final.xlsx")
        llm = StubLLM(root / "run")
        with env(PATHS_LIBREOFFICE_PATH="/nonexistent/LibreOffice/soffice"), \
                no_libreoffice(det_checks=False) as lo_calls:
            res = _grade(root, attempt, None, llm)
        assert llm.calls == 0, "the LLM must not be called"
        assert lo_calls == [], lo_calls
        assert res["success"] is False and "scores" not in res
        err = res["error"]
        task_folder = Path(res["task_folder"])
        assert "deterministic checks could not grade" in err and "before the LLM call" in err
        assert str(task_folder / "ai_attempt.xlsx") in err and "delivered as 'Model_final.xlsx'" in err
        for no in (66, 73, 94):
            assert D.label(no) in err, (no, err)
        assert "LibreOffice not found" in err and "DetChecksError" in res["traceback"]
        art = json.loads((task_folder / D.ARTEFACT_FILENAME).read_text())
        assert art["status"] == "error" and "/".join(D.DET_CHECK_NAMES[66]) in art["failures"]
        assert (task_folder / "answer_check.json").exists(), "the answer check ran first (score-neutral)"
        assert not (task_folder / "judge_results" / "scores.json").exists()
        gfd.prune_workbook_copies(res)                       # what grade_from_db.main does on failure
        assert not (task_folder / "ai_attempt.xlsx").exists() and (task_folder / D.ARTEFACT_FILENAME).exists()


def test_call_sites_are_wired_before_the_llm():
    gfd_src = (JUDGE / "main_scripts" / "grade_from_db.py").read_text()
    gsa = gfd_src.split("def grade_single_attempt")[1].split("\ndef ")[0]
    i_ac, i_det, i_judge = gsa.index("run_answer_check("), gsa.index("run_det_checks("), gsa.index("single_pass_judge_case(")
    assert i_ac < i_det < i_judge
    between = gsa[i_ac:i_det]
    assert "score-neutral by design" in between and "harness_verdicts = ac_result.get(" in between, \
        "run_det_checks must sit after (outside) the answer check's score-neutral try"
    assert gsa.rfind("try:", 0, i_det) > gsa.index("harness_verdicts = ac_result.get("), \
        "run_det_checks must sit inside the try whose except logs FAILED"
    assert "merge_harness_verdicts(" in gsa and "det_checks=det_run.for_judge()" in gsa
    assert "det_checks=args.det_checks" in gfd_src and "add_det_checks_arg(parser)" in gfd_src
    prune = gfd_src.split("def prune_workbook_copies")[1].split("\ndef ")[0]
    assert "RECALC_DIRNAME" in prune
    judge_src = (JUDGE / "main_scripts" / "judge.py").read_text()
    main = judge_src.split("def main(args)")[1]
    j_ac, j_det, j_judge = main.index("run_answer_check("), main.index("run_det_checks("), main.index("single_pass_judge_case(")
    assert j_ac < j_det < j_judge and "det_checks=det_run.for_judge()" in main
    assert "add_det_checks_arg(parser)" in judge_src
    orch = (JUDGE / "main_scripts" / "grade_with_orchestration.py").read_text()
    assert "det_checks=self.det_checks" in orch and "add_det_checks_arg(parser)" in orch
    assert "det_checks=args.det_checks" in orch
    toy = (JUDGE / "main_scripts" / "grade_toy.py").read_text()
    assert "det_checks=args.det_checks" in toy and "add_det_checks_arg(ap)" in toy
    # every DB driver refuses a bad det_checks config before grading anything
    for src in (gfd_src.split("def main(args)")[1], orch.split("def main()")[1], toy.split("def main()")[1]):
        assert "startup_check(rubric_path, args.det_checks)" in src


if __name__ == "__main__":
    tests = [(n, f) for n, f in sorted(globals().items()) if n.startswith("test_") and callable(f)]
    failed = 0
    for name, fn in tests:
        try:
            fn()
            print(f"OK   {name}")
        except Exception as e:  # noqa: BLE001
            failed += 1
            import traceback
            traceback.print_exc()
            print(f"FAIL {name}: {e}")
    print()
    print(f"{len(tests) - failed}/{len(tests)} passed")
    sys.exit(1 if failed else 0)
