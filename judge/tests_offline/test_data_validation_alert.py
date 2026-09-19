"""Offline: data-validation error-alert evidence (judge v11, 2026-09-18).

A validation whose error alert is off never rejects or flags an entry, so
the properties block must say so, and the full rule (operator, both bounds)
must be visible. Run from judge/:  python tests_offline/test_data_validation_alert.py
"""
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import openpyxl  # noqa: E402
from openpyxl.worksheet.datavalidation import DataValidation  # noqa: E402

from utils.misc_utils import load_project_configs  # noqa: E402

load_project_configs(benchmark="v2")

from utils import excel_utils, workbook_properties as wp  # noqa: E402


def _extract(wb):
    d = Path(tempfile.mkdtemp())
    path = d / "t.xlsx"
    wb.save(path)
    excel_utils.process_all_worksheets(str(path), d / "out", quiet=True)
    return wp.load_properties(d / "out")


def _sheet(props, name):
    return next(s for s in props["sheets"] if s["name"] == name)


def _dv(ws, cell, **kw):
    dv = DataValidation(**kw)
    ws.add_data_validation(dv)
    dv.add(cell)
    return dv


def test_extraction_and_render_alert_states():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Assumptions"
    ws["D7"], ws["D43"], ws["E36"], ws["C48"], ws["F2"] = 50, "Base", 2.5, 0.03, 7
    _dv(ws, "D7", type="whole", operator="between", formula1="1", formula2="50", showErrorMessage=False)
    _dv(ws, "D43", type="list", formula1='"Base,Low,High"', showErrorMessage=True)          # errorStyle default = stop
    _dv(ws, "E36", type="decimal", operator="greaterThan", formula1="0", showErrorMessage=True, errorStyle="warning")
    _dv(ws, "C48", type="custom", formula1="C48<0.5", showErrorMessage=False)
    off = wb.create_sheet("AllOff")
    off["A1"], off["A2"] = 1, 2
    _dv(off, "A1", type="whole", operator="between", formula1="1", formula2="1000000", showErrorMessage=False)
    _dv(off, "A2", type="list", formula1='"0,1"', showErrorMessage=False)
    props = _extract(wb)
    s = _sheet(props, "Assumptions")
    by = {d["sqref"]: d for d in s["data_validations"]}
    assert by["D7"]["alert"] is False and by["D7"]["error_style"] is None and by["D7"]["operator"] == "between"
    assert by["D43"]["alert"] is True and by["D43"]["error_style"] == "stop"
    assert by["E36"]["alert"] is True and by["E36"]["error_style"] == "warning"
    text = wp.render_properties_text(props)
    lines = {l.strip() for l in text.splitlines() if "data validation:" in l}
    a = next(l for l in lines if "D7 whole" in l)
    assert a.startswith("data validation: 4 (2 with the error alert ON, 2 OFF): "), a
    assert "D7 whole between 1 and 50 — alert OFF (any entry accepted)" in a, a
    assert 'D43 list "Base,Low,High" — alert ON (stop, entry rejected)' in a, a
    assert "E36 decimal > 0 — alert ON (warning, entry allowed after a prompt)" in a, a
    assert "C48 custom C48<0.5 — alert OFF (any entry accepted)" in a, a
    b = next(l for l in lines if "A1 whole" in l)
    assert b.startswith("data validation: 2 (error alert OFF for ALL — none of them rejects or flags an entry): "), b
    assert "A1 whole between 1 and 1000000 — alert OFF (any entry accepted)" in b, b


def test_older_cache_without_the_flag_renders_unknown_not_off():
    # a v7 cache row: no 'alert' key at all
    assert wp._data_validation_rule({"sqref": "B4", "type": "list", "formula1": "=Lists!A1:A5"}) \
        == "B4 list =Lists!A1:A5 — alert unknown"
    t = wp._data_validation_text([{"sqref": "B4", "type": "list", "formula1": "=Lists!A1:A5"},
                                  {"sqref": "B5", "type": "whole", "operator": "between", "formula1": "1",
                                   "formula2": "9", "alert": True, "error_style": "stop"}])
    assert t.startswith("2 (alert state unknown for 1): ") and "B5 whole between 1 and 9 — alert ON (stop, entry rejected)" in t, t
    assert wp._data_validation_text([]) == "none" and wp._data_validation_text("unknown") == "unknown"
    all_on = wp._data_validation_text([{"sqref": "D9 D55", "type": "decimal", "operator": "greaterThan",
                                        "formula1": "0", "alert": True, "error_style": "stop"}])
    assert all_on.startswith("1 (error alert ON for all): D9 D55 decimal > 0 — alert ON"), all_on


def test_operator_wording():
    r = lambda **k: wp._data_validation_rule({"sqref": "A1", "alert": True, "error_style": "stop", **k})  # noqa: E731
    assert r(type="whole", operator="notBetween", formula1="1", formula2="5").startswith("A1 whole not between 1 and 5")
    assert r(type="decimal", operator="lessThanOrEqual", formula1="1").startswith("A1 decimal <= 1")
    assert r(type="date", operator="equal", formula1="45000").startswith("A1 date = 45000")
    assert r(type="textLength", operator="lessThan", formula1="10").startswith("A1 textLength < 10")
    assert r(type="whole", formula1="1", formula2="50").startswith("A1 whole between 1 and 50")   # Excel's default operator
    assert r(type="whole", formula1="1").startswith("A1 whole 1 — alert")                           # only one bound stored


if __name__ == "__main__":
    for name, fn in list(globals().items()):
        if name.startswith("test_") and callable(fn):
            fn()
            print(f"ok  {name}")
    print("all data-validation alert tests passed")
