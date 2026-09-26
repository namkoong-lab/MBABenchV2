"""Offline: IMPLICIT INTERSECTION evidence (judge v11, 2026-09-18).

Every formula below was written to a workbook by openpyxl (no array marker),
opened in Excel for Mac 16.112 with a full recalculation and read back:
'excel' is the error Excel displayed or 'ok'. Ranges A1:A10 / B1:B10 /
H1:L10 hold numbers; the formula cell is the coordinate given (D20+ lies
outside every range, D5/H5/H20/A20/J20/D6 probe the intersection geometry).
The detector must never flag a formula Excel evaluates fine (precision), and
must flag every #VALUE! except the constructs it deliberately does not model
(KNOWN_MISSES). '#NAME?' rows are functions stored without the _xlfn. prefix
— a different defect, expected NOT to be flagged here.

Run from judge/:  python tests_offline/test_implicit_intersection.py
"""
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import openpyxl  # noqa: E402
from openpyxl.worksheet.formula import ArrayFormula  # noqa: E402

from utils.misc_utils import load_project_configs  # noqa: E402

load_project_configs(benchmark="v2")

from utils import excel_utils, workbook_properties as wp  # noqa: E402
from utils import implicit_intersection as ii  # noqa: E402

# (formula, cell, excel result, written as an array formula)
PROBES = [
    ('=ABS(A1:A10-B1:B10)', 'D5', 'ok', False),
    ('=SUMPRODUCT(ABS(A1:A10-B1:B10))', 'D20', 'ok', False),
    ('=SUMPRODUCT(MAX(ABS(A1:A10-B1:B10)))', 'E20', 'ok', True),
    ('=SUMPRODUCT(MAX(ABS(A1:A10-B1:B10)))', 'D21', '#VALUE!', False),
    ('=MAX(ABS(A1:A10-B1:B10))', 'E21', 'ok', True),
    ('=SUMPRODUCT(MAX(A1:A10-B1:B10))', 'D22', 'ok', False),
    ('=SUM(A1:A10*B1:B10)', 'E22', 'ok', True),
    ('=MAX(ABS(A1:A10-B1:B10))', 'D23', '#VALUE!', False),
    ('=SUM(A1:A10*B1:B10)', 'D24', '#VALUE!', False),
    ('=SUMPRODUCT(A1:A10*B1:B10)', 'D25', 'ok', False),
    ('=MAX(INDEX(ABS(A1:A10-B1:B10),0))', 'D26', 'ok', False),
    ('=SUMPRODUCT(LEN(A1:A10))', 'D27', 'ok', False),
    ('=SUMPRODUCT((A1:A10>3)*B1:B10)', 'D28', 'ok', False),
    ('=SUMPRODUCT(MAX((A1:A10>3)*B1:B10))', 'D29', 'ok', False),
    ('=IF(A1:A10>3,1,0)', 'D30', '#VALUE!', False),
    ('=SUM(IF(A1:A10>3,B1:B10,0))', 'D31', '#VALUE!', False),
    ('=IFERROR(A1:A10/B1:B10,0)', 'D32', 'ok', False),
    ('=MIN(ROUND(A1:A10,2))', 'D33', '#VALUE!', False),
    ('=AVERAGE(ABS(A1:A10))', 'D34', '#VALUE!', False),
    ('=MAX(A1:A10-B1:B10)', 'D35', '#VALUE!', False),
    ('=SUMPRODUCT(--(A1:A10>3))', 'D36', 'ok', False),
    ('=MAX(IF(A1:A10>3,B1:B10))', 'D37', '#VALUE!', False),
    ('=SUM(SQRT(A1:A10))', 'D38', '#VALUE!', False),
    ('=SUMPRODUCT(SQRT(A1:A10))', 'D39', 'ok', False),
    ('=SUMPRODUCT(MAX(SQRT(A1:A10)))', 'D40', '#VALUE!', False),
    ('=SUMPRODUCT(MIN(A1:A10*B1:B10))', 'D41', 'ok', False),
    ('=INDEX(A1:A10*B1:B10,3)', 'D42', 'ok', False),
    ('=MATCH(MAX(A1:A10*B1:B10),A1:A10*B1:B10,0)', 'D43', '#VALUE!', False),
    ('=SUM(ABS(A1:A10))', 'D44', '#VALUE!', False),
    ('=SUMPRODUCT(SUM(ABS(A1:A10)))', 'D45', '#VALUE!', False),
    ('=ABS(A1:A10)', 'D46', '#VALUE!', False),
    ('=SUMPRODUCT(EXP(A1:A10-B1:B10))', 'D47', 'ok', False),
    ('=SUMPRODUCT(MAX(ABS(A1:A10)))', 'D48', '#VALUE!', False),
    ('=SUMPRODUCT(MAX(ABS(A1:A10),ABS(B1:B10)))', 'D49', '#VALUE!', False),
    ('=NPV(0.1,A1:A10*B1:B10)', 'D50', '#VALUE!', False),
    ('=SUMPRODUCT(NPV(0.1,A1:A10*B1:B10))', 'D51', 'ok', False),
    ('=COUNT(A1:A10*B1:B10)', 'D52', 'ok', False),
    ('=SUMPRODUCT(1/(1+A1:A10)^B1:B10)', 'D53', 'ok', False),
    ('=SUM(A1:A10/(1+0.1)^B1:B10)', 'D54', '#VALUE!', False),
    ('=MAX(A1:A10)', 'D55', 'ok', False),
    ('=SUM(A1:A10)', 'D56', 'ok', False),
    ('=ABS(H1:L10)', 'D5', '#VALUE!', False),
    ('=ABS(A1:A10)', 'H5', 'ok', False),
    ('=SUMPRODUCT(MAX(ABS(A1:A10-B1:B10)))', 'D6', 'ok', False),
    ('=ABS(H1:L1)', 'A20', '#VALUE!', False),
    ('=SUMPRODUCT(MAX(ABS(A1:A10)-ABS(B1:B10)))', 'D20', '#VALUE!', False),
    ('=ABS(H1:L10)', 'H20', '#VALUE!', False),
    ('=ABS(H1:L1)', 'J20', 'ok', False),
    ('=SUMPRODUCT(MAX((A1:A10-B1:B10)*1))', 'D21', 'ok', False),
    ('=SUMPRODUCT(MAX(ABS(A1:A10-B1:B10)*1))', 'D22', '#VALUE!', False),
    ('=SUMPRODUCT(MAX(A1:A10-B1:B10,A1:A10*2))', 'D23', 'ok', False),
    ('=SUMPRODUCT(SUM(A1:A10*B1:B10))', 'D24', 'ok', False),
    ('=SUMPRODUCT(MAX(INDEX(ABS(A1:A10-B1:B10),0)))', 'D25', 'ok', False),
    ('=SUMPRODUCT(ABS(MAX(A1:A10-B1:B10)))', 'D26', 'ok', False),
    ('=SUMPRODUCT(MAX(IF(A1:A10>3,B1:B10,0)))', 'D27', '#VALUE!', False),
    ('=SUMPRODUCT(IF(A1:A10>3,B1:B10,0))', 'D28', '#VALUE!', False),
    ('=SUMPRODUCT(ROUND(A1:A10*B1:B10,0))', 'D29', 'ok', False),
    ('=SUMPRODUCT(MAX(ROUND(A1:A10*B1:B10,0)))', 'D30', '#VALUE!', False),
    ('=SUMPRODUCT(LARGE(ABS(A1:A10-B1:B10),1))', 'D31', '#VALUE!', False),
    ('=SUMPRODUCT(AVERAGE(ABS(A1:A10-B1:B10)))', 'D32', '#VALUE!', False),
    ('=SUMPRODUCT(MAX(ABS(A1:A10-B1:B10))*1)', 'D33', '#VALUE!', False),
    ('=MAX(ABS(A1:A10-B1:B10)*1)', 'D34', '#VALUE!', False),
    ('=SUM(INDEX(ABS(A1:A10-B1:B10),0))', 'D35', 'ok', False),
    ('=AVERAGE(A1:A10*B1:B10)', 'D36', '#VALUE!', False),
    ('=SUMPRODUCT(A1:A10*B1:B10)/SUM(A1:A10)', 'D37', 'ok', False),
    ('=IF(SUM(A1:A10)>0,A1:A10*2,0)', 'D38', '#VALUE!', False),
    ('=TEXT(A1:A10,"0")', 'D39', '#VALUE!', False),
    ('=SUMPRODUCT(--(LEN(A1:A10)>0))', 'D40', 'ok', False),
    ('=SUMPRODUCT(MAX(--(A1:A10>3)*B1:B10))', 'D41', 'ok', False),
    ('=SUMPRODUCT(MAX(N(A1:A10)))', 'D42', 'ok', False),
    ('=MIN(A1:A10)+MAX(B1:B10)', 'D43', 'ok', False),
    ('=SUM(A1:A10)*2', 'D44', 'ok', False),
    ('=SUMIF(A1:A10,">3",B1:B10)/COUNTIF(A1:A10,">3")', 'D45', 'ok', False),
    ('=SUMPRODUCT((A1:A10>3)*(B1:B10<8))', 'D46', 'ok', False),
    ('=SUM(--(A1:A10>3))', 'D47', '#VALUE!', False),
    ('=ABS(H1:L10)', 'D48', '#VALUE!', False),
    ('=SUMPRODUCT(MAX(H1:L10-H1:L10))', 'D49', 'ok', False),
    ('=SUMPRODUCT(MAX(ABS(H1:L10)))', 'D50', '#VALUE!', False),
    ('=MAX(H1:L10)', 'D51', 'ok', False),
    ('=SUMPRODUCT(SUMIF(A1:A10,">3",B1:B10)*ABS(A1:A10))', 'D52', 'ok', False),
    ('=SUMPRODUCT(IFERROR(A1:A10/B1:B10,0))', 'D53', 'ok', False),
    ('=SUMPRODUCT(MAX(IFERROR(A1:A10/B1:B10,0)))', 'D54', 'ok', False),
    ('=MATCH(9,A1:A10*B1:B10,0)', 'D55', '#VALUE!', False),
    ('=LOOKUP(2,1/(A1:A10>3),B1:B10)', 'D56', 'ok', False),
    ('=VLOOKUP(3,A1:B10*1,2,0)', 'D57', '#VALUE!', False),
    ('=SUMPRODUCT(MAX(ABS(A1:A10-B1:B10)),1)', 'D58', '#VALUE!', False),
    ('=SUMPRODUCT(MAX(MAX(ABS(A1:A10-B1:B10))))', 'D59', '#VALUE!', False),
    ('=MAX(MAX(A1:A10-B1:B10))', 'D60', '#VALUE!', False),
    ('=SUMPRODUCT(EXP(LN(A1:A10)))', 'D61', 'ok', False),
    ('=SUMPRODUCT(MAX(EXP(A1:A10)))', 'D62', '#VALUE!', False),
    ('=SUMPRODUCT(SUM(EXP(A1:A10)))', 'D63', '#VALUE!', False),
    ('=SUM(EXP(A1:A10))', 'D64', '#VALUE!', False),
    ('=SUMPRODUCT(MAX(A1:A10)*ABS(B1:B10))', 'D65', 'ok', False),
    ('=SUMPRODUCT(MAX(A1:A10*ABS(B1:B10)))', 'D66', '#VALUE!', False),
    ('=SUMPRODUCT(MAX(ABS(A1:A10)*B1:B10))', 'D67', '#VALUE!', False),
    ('=SUMPRODUCT(MAX(ABS(A1:A10-B1:B10)+0))', 'D68', '#VALUE!', False),
    ('=SUMPRODUCT(MAX(0+ABS(A1:A10-B1:B10)))', 'D69', '#VALUE!', False),
    ('=SUMPRODUCT(MAX(SIGN(A1:A10-B1:B10)))', 'D70', '#VALUE!', False),
    ('=SUMPRODUCT(MAX(NOT(A1:A10>3)*B1:B10))', 'D71', '#VALUE!', False),
    ('=SUMPRODUCT(MAX(ISNUMBER(A1:A10)*B1:B10))', 'D72', 'ok', False),
    ('=COUNTIF(A1:A10,">3")', 'D73', 'ok', False),
    ('=SUMPRODUCT(MAX(ABS(A1:A10-B1:B10),0))', 'D74', '#VALUE!', False),
    ('=ABS(Q!A1:A10)', 'H5', 'ok', False),
    ('=AND(ABS(A1:A10)>0)', 'D20', '#VALUE!', False),
    ('=OR(ABS(A1:A10)>5)', 'D21', '#VALUE!', False),
    ('=PRODUCT(ABS(A1:A10))', 'D22', '#VALUE!', False),
    ('=MEDIAN(ABS(A1:A10-B1:B10))', 'D23', '#VALUE!', False),
    ('=STDEV(ABS(A1:A10))', 'D24', '#VALUE!', False),
    ('=SMALL(ABS(A1:A10-B1:B10),1)', 'D25', '#VALUE!', False),
    ('=PERCENTILE(ABS(A1:A10),0.5)', 'D26', '#VALUE!', False),
    ('=SUMSQ(ABS(A1:A10))', 'D27', '#VALUE!', False),
    ('=NPV(0.1,ABS(A1:A10))', 'D28', '#VALUE!', False),
    ('=SUMPRODUCT(NPV(0.1,ABS(A1:A10)))', 'D29', '#VALUE!', False),
    ('=SUMPRODUCT(MATCH(9,A1:A10*B1:B10,0))', 'D30', '#N/A', False),
    ('=SUMPRODUCT(VLOOKUP(3,A1:B10*1,2,0))', 'D31', '#VALUE!', False),
    ('=SUMPRODUCT(MAX(ISERROR(A1:A10)*1))', 'D32', 'ok', False),
    ('=IFS(SUM(A1:A10)>0,A1:A10*2,TRUE,0)', 'D33', '#NAME?', False),
    ('=CHOOSE(1,A1:A10*2,0)', 'D34', '#VALUE!', False),
    ('=SWITCH(1,1,A1:A10*2,0)', 'D35', '#NAME?', False),
    ('=XLOOKUP(9,A1:A10*B1:B10,B1:B10)', 'D36', '#NAME?', False),
    ('=A1:A10', 'D37', '#VALUE!', False),
    ('=SUM(A1:A10)+A1:A10*2', 'D38', '#VALUE!', False),
    ('=SUMPRODUCT(TRANSPOSE(A1:A10)*1)', 'D39', '#VALUE!', False),
    ('=SUM(ROW(A1:A10)*1)', 'D40', 'ok', False),
    ('=SUMPRODUCT(ROW(A1:A10)*1)', 'D41', 'ok', False),
    ('=SUMPRODUCT(MAX(ROW(A1:A10)*1))', 'D42', 'ok', False),
    ('=OFFSET(A1,0,0,10,1)*2', 'D43', '#VALUE!', False),
    ('=SUMPRODUCT(MAX(ABS(OFFSET(A1,0,0,10,1))))', 'D44', '#VALUE!', False),
    ('=SUMPRODUCT(MAX(OFFSET(A1,0,0,10,1)*1))', 'D45', 'ok', False),
    ('=ABS(A:A)', 'D46', 'ok', False),
    ('=ABS(1:1)', 'D47', 'ok', False),
    ('=ABS(Q!A1:A10)', 'D48', '#VALUE!', False),
    ('=SUMPRODUCT(COUNTIF(A1:A10,A1:A10))', 'D49', 'ok', False),
    ('=SUMPRODUCT(MAX(COUNTIF(A1:A10,A1:A10)))', 'D50', 'ok', False),
    ('=SUMPRODUCT(MAX(ABS(A1:A10-B1:B10)),ABS(A1:A10))', 'D51', '#VALUE!', False),
    ('=MAX(SUMPRODUCT(ABS(A1:A10-B1:B10)))', 'D52', 'ok', False),
    ('=MAX(ABS(SUMPRODUCT(A1:A10-B1:B10)))', 'D53', 'ok', False),
    ('=SUMPRODUCT(MAX(A1:A10-B1:B10)-ABS(A1:A10))', 'D54', 'ok', False),
    ('=SUMPRODUCT(MAX(MOD(A1:A10,2)))', 'D55', '#VALUE!', False),
    ('=SUMPRODUCT(MAX(POWER(A1:A10,2)))', 'D56', '#VALUE!', False),
    ('=SUMPRODUCT(MAX(INT(A1:A10)))', 'D57', '#VALUE!', False),
    ('=SUMPRODUCT(MAX(LEN(A1:A10)))', 'D58', '#VALUE!', False),
    ('=SUMPRODUCT(MAX((A1:A10>3)*ROUND(B1:B10,0)))', 'D59', '#VALUE!', False),
    ('=SUMPRODUCT(MAX(TEXT(A1:A10,"0")*1))', 'D60', '#VALUE!', False),
    ('=SUMPRODUCT(SUM(A1:A10)*ABS(B1:B10))', 'D61', 'ok', False),
    ('=MAX(A1:A10*B1:B10,0)', 'D62', '#VALUE!', False),
    ('=AVERAGEIF(A1:A10,">3",B1:B10)+ABS(A1:A10)', 'D63', '#VALUE!', False),
    ('=SUMPRODUCT(MAXA(ABS(A1:A10)))', 'D64', '#VALUE!', False),
    ('=SUMPRODUCT(MINA(ABS(A1:A10)))', 'D65', '#VALUE!', False),
    ('=SUMPRODUCT(COUNT(ABS(A1:A10)))', 'D66', 'ok', False),
    ('=SUMPRODUCT(COUNTA(ABS(A1:A10)))', 'D67', 'ok', False),
    ('=SUMPRODUCT(MAX(IFERROR(ABS(A1:A10),0)))', 'D68', 'ok', False),
    ('=SUMPRODUCT(MAX(EXP(A1:A10)*0+1))', 'D69', '#VALUE!', False),
    ('=SUMPRODUCT(MAX(1/(1+A1:A10)^B1:B10))', 'D70', 'ok', False),
    ('=SUMPRODUCT(MAX(ABS(A1:A10*1)))', 'D71', '#VALUE!', False),
    ('=SUMPRODUCT(MAX((A1:A10)*(B1:B10)))', 'D72', 'ok', False),
    ('=SUMPRODUCT(MIN(ABS(A1:A10-B1:B10)))', 'D73', '#VALUE!', False),
    ('=SUMPRODUCT(PRODUCT(ABS(A1:A10)))', 'D74', '#VALUE!', False),
    ('=SUMPRODUCT(SMALL(A1:A10*B1:B10,1))', 'D75', 'ok', False),
    ('=SUMPRODUCT(SUMSQ(A1:A10-B1:B10))', 'D76', 'ok', False),
    ('=SUMPRODUCT(MEDIAN(A1:A10-B1:B10))', 'D77', 'ok', False),
    ('=SUMPRODUCT(AND(A1:A10>0))', 'D78', 'ok', False),
    ('=SUMPRODUCT(AND(ABS(A1:A10)>0))', 'D79', '#VALUE!', False),
    ('=SUMPRODUCT(MAX(SQRT(ABS(A1:A10-B1:B10))))', 'D80', '#VALUE!', False),
    ('=SUMPRODUCT(SQRT(MAX(ABS(A1:A10-B1:B10))))', 'D81', '#VALUE!', False),
    ('=SUMPRODUCT(MAX(ROUND(A1:A10-B1:B10,0)))', 'D82', '#VALUE!', False),
]

# Excel shows #VALUE! but the walker treats the construct as unmodelled:
# OFFSET results, CHOOSE values, TRANSPOSE. Never flagged by design.
KNOWN_MISSES = {
    "=CHOOSE(1,A1:A10*2,0)",
    "=SUMPRODUCT(TRANSPOSE(A1:A10)*1)",
    "=OFFSET(A1,0,0,10,1)*2",
    "=SUMPRODUCT(MAX(ABS(OFFSET(A1,0,0,10,1))))",
}


def _flagged(formula, cell, is_array):
    return bool(ii.check_formula(formula, cell)) and not is_array


def test_probe_precision_no_false_flags():
    bad = [(f, c, e) for f, c, e, a in PROBES if _flagged(f, c, a) and e != "#VALUE!"]
    assert not bad, bad


def test_probe_recall_every_value_error_is_flagged():
    missed = [(f, c) for f, c, e, a in PROBES
              if e == "#VALUE!" and not a and not _flagged(f, c, a) and f not in KNOWN_MISSES]
    assert not missed, missed
    # the known misses really are misses (so the list does not go stale)
    stale = [f for f in KNOWN_MISSES if not any(p[0] == f and p[2] == "#VALUE!" for p in PROBES)]
    assert not stale, stale
    assert sum(1 for f, c, e, a in PROBES if e == "#VALUE!" and not a) >= 80


def test_grading_1075_d448_and_its_geometry():
    f = "=SUMPRODUCT(MAX(ABS(P413:Y445-P376:Y408)))"
    hits = ii.check_formula(f, "D448")
    assert [h["range"] for h in hits] == ["P413:Y445", "P376:Y408"] and hits[0]["via"] == "-", hits
    assert ii.check_formula(f, "D420")            # row inside, column outside: still no intersection (2-D block)
    # a cell inside the first block (R420) intersects it; only the second block is left without a cell
    assert [h["range"] for h in ii.check_formula(f, "R420")] == ["P376:Y408"]
    assert not ii.check_formula("=SUMPRODUCT(ABS(P413:Y445-P376:Y408))", "D448")   # array-evaluated
    assert ii.check_formula("=MAX(P413:Y445-P376:Y408)", "D448"), "operator inside MAX is scalar at top level"


def test_geometry():
    assert ii.intersection_impossible((1, 1, 1, 10), col=8, row=20)          # one column, row outside
    assert not ii.intersection_impossible((1, 1, 1, 10), col=8, row=5)       # one column, row inside
    assert ii.intersection_impossible((8, 1, 12, 1), col=1, row=20)          # one row, column outside
    assert not ii.intersection_impossible((8, 1, 12, 1), col=10, row=20)     # one row, column inside
    assert ii.intersection_impossible((8, 1, 12, 10), col=4, row=5)          # 2-D: row inside is not enough
    assert ii.intersection_impossible((8, 1, 12, 10), col=8, row=20)         # 2-D: column inside is not enough
    assert not ii.intersection_impossible((8, 1, 12, 10), col=10, row=5)     # 2-D: both inside
    assert not ii.intersection_impossible((1, None, 1, None), col=4, row=20)  # whole column always intersects
    assert not ii.intersection_impossible((None, 1, None, 1), col=4, row=20)  # whole row always intersects
    assert not ii.intersection_impossible((3, 3, 3, 3), col=9, row=9)        # single cell is not a range


def test_unmodelled_constructs_are_never_flagged():
    for f in ["=A1 A2:B3", "=INDEX(A1:A10,0):B5", "=Table1[Col]*2", "=SUM([1]Sheet1!A1:A5)",
              "=Sheet1:Sheet3!A1:A5*2", "=IFERROR(A1:A10/B1:B10,0)", "=_xlfn.XLOOKUP(1,A1:A10*2,B1:B10)",
              "=ABS(A:A)", "=ABS(1:1)", "=SUMPRODUCT(ABS(A1:A10-B1:B10))", "=COUNT(A1:A10*B1:B10)",
              "=ISNUMBER(A1:A10)*1", "=N(A1:A10)", "=CHOOSE(1,A1:A10*2,0)", "=SUM(OFFSET(A1,0,0,10,1)*1)"]:
        assert ii.check_formula(f, "D20") == [], f
    assert ii.check_formula("plain text", "D20") == [] and ii.check_formula(None, "D20") == []


def test_cross_sheet_intersection_is_positional():
    assert ii.check_formula("=ABS(Q!A1:A10)", "D48")        # row 48 outside 1-10
    assert not ii.check_formula("=ABS(Q!A1:A10)", "H5")     # same row on another sheet: Excel intersects
    assert ii.check_formula("=ABS(\'Sheet 1\'!$A$1:$A$10)", "H20")


def _extract(wb):
    d = Path(tempfile.mkdtemp())
    path = d / "t.xlsx"
    wb.save(path)
    excel_utils.process_all_worksheets(str(path), d / "out", quiet=True)
    return wp.load_properties(d / "out")


def test_end_to_end_properties_and_render():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sens_Engine"
    for r in range(376, 446):
        for c in range(16, 26):
            ws.cell(row=r, column=c, value=r * c / 1000)
    ws["D448"] = "=SUMPRODUCT(MAX(ABS(P413:Y445-P376:Y408)))"   # the grading-1075 defect
    ws["D449"] = "=SUMPRODUCT(ABS(P413:Y445-P376:Y408))"        # fine in Excel
    ws["D450"] = ArrayFormula("D450", "=SUMPRODUCT(MAX(ABS(P413:Y445-P376:Y408)))")  # array marker: fine
    ws["D451"] = "=MAX(P413:Y445)"                              # bare range in an aggregator: fine
    clean = wb.create_sheet("Clean")
    clean["A1"] = 1
    clean["A2"] = "=SUM(A1:A1)*2"
    props = _extract(wb)
    assert props["schema"] >= 5
    s = next(x for x in props["sheets"] if x["name"] == "Sens_Engine")
    assert s["implicit_intersection"]["count"] == 1, s["implicit_intersection"]
    ex = s["implicit_intersection"]["examples"][0]
    assert ex["cell"] == "D448" and ex["ranges"] == ["P413:Y445", "P376:Y408"] and ex["via"] == "-"
    c = next(x for x in props["sheets"] if x["name"] == "Clean")
    assert c["implicit_intersection"] == {"count": 0, "examples": []}
    text = wp.render_properties_text(props)
    assert text.count("IMPLICIT INTERSECTION") == 1
    line = next(l for l in text.splitlines() if "IMPLICIT INTERSECTION" in l)
    assert "D448 =SUMPRODUCT(MAX(ABS(P413:Y445-P376:Y408)))" in line and "P413:Y445, P376:Y408" in line
    assert "#VALUE!" in line and "1 plain formula without the array marker uses" in line
    # older caches without the key render nothing and do not crash
    for x in props["sheets"]:
        x.pop("implicit_intersection", None)
    assert "IMPLICIT INTERSECTION" not in wp.render_properties_text(props)
    for x in props["sheets"]:
        x["implicit_intersection"] = "unknown"
    assert "IMPLICIT INTERSECTION" not in wp.render_properties_text(props)


def test_dependents_of_a_flagged_cell_are_listed():
    # judge v12, grading 1092: D448 fed Checks!F19 and the COUNTIF roll-up skipped the erroring row
    wb = openpyxl.Workbook(); eng = wb.active; eng.title = "Sens Engine"
    eng["D448"] = "=SUMPRODUCT(MAX(ABS(P413:Y445-P376:Y408)))"
    eng["D450"] = "=D448*2"                                   # same sheet, unqualified
    eng["D451"] = "=SUM(D440:D449)"                           # a range that contains it: not listed
    eng["D452"] = "=D4480+AD448"                              # look-alikes: not listed
    ck = wb.create_sheet("Checks")
    ck["D19"] = "='Sens Engine'!D448"
    ck["F19"] = "='Sens Engine'!$D$448<0.0001"
    ck["F20"] = "=COUNTIF(F6:F42,FALSE)"
    deps = ii.find_dependents(wb, {"Sens Engine": ["D448"]})
    assert deps == {("Sens Engine", "D448"): ["Sens Engine!D450", "Checks!D19", "Checks!F19"]}, deps
    assert ii.find_dependents(wb, {}) == {} and ii.find_dependents(wb, {"Sens Engine": ["Z1"]}) == {}
    d = Path(tempfile.mkdtemp()); path = d / "t.xlsx"; wb.save(path)
    excel_utils.process_all_worksheets(str(path), d / "out", quiet=True)
    props = wp.load_properties(d / "out")
    ex = next(x for x in props["sheets"] if x["name"] == "Sens Engine")["implicit_intersection"]["examples"]
    assert ex[0]["referenced_by"] == ["Sens Engine!D450", "Checks!D19", "Checks!F19"], ex
    line = next(l for l in wp.render_properties_text(props).splitlines() if "IMPLICIT INTERSECTION" in l)
    assert "referenced by Sens Engine!D450, Checks!D19, Checks!F19: these cells show the error in Excel too" in line
    many = {"count": 1, "examples": [{"cell": "D1", "formula": "=ABS(A1:A9)", "ranges": ["A1:A9"], "via": "ABS",
                                      "referenced_by": [f"S!A{i}" for i in range(11)]}]}
    assert "(+3 more)" in ii.render_lines(many)[0]


def test_render_caps_examples():
    ex = [{"cell": f"D{i}", "formula": "=ABS(A1:A10)", "ranges": ["A1:A10"], "via": "ABS"} for i in range(10)]
    lines = ii.render_lines({"count": 37, "examples": ex})
    assert len(lines) == 1 and "37 plain formulas" in lines[0] and "(+27 more)" in lines[0]
    assert "as a single value in ABS( )" in lines[0]
    assert ii.render_lines({"count": 0, "examples": []}) == [] and ii.render_lines("unknown") == []


if __name__ == "__main__":
    for name, fn in list(globals().items()):
        if name.startswith("test_") and callable(fn):
            fn()
            print(f"ok  {name}")
    print("all implicit-intersection tests passed")
