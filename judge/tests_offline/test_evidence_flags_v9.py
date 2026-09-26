"""Offline: judge v9 evidence flags + retired checks (2026-09-16 walkthrough).

  - WIDE OUTLIER (check 70): render-time from stored column widths
  - PERIOD SERIES scan (checks 108/123/124/125): orientation, OUT OF ORDER,
    unlabeled gaps, VERTICAL PERIOD SERIES, content past an End marker
  - content fit (check 69): NUMERIC exceeds width / TEXT cut off / overflow
  - typed date-like string literal inside a formula (checks 2/10/81)
  - retired checks 28/37/101: forced not_applicable by rule, numbering kept;
    28's own evidence line follows the config list (judge v11, 2026-09-19)

Run from judge/:  python tests_offline/test_evidence_flags_v9.py
"""
import datetime as dt
import json
import os
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import openpyxl  # noqa: E402
from openpyxl.styles import Font  # noqa: E402

from utils.misc_utils import load_project_configs  # noqa: E402

load_project_configs(benchmark="v2")

from utils import excel_utils, rubric_suitability as rs, workbook_properties as wp  # noqa: E402

JUDGE = Path(__file__).resolve().parents[1]
RUBRIC = json.loads((JUDGE / "prompts" / "rubrics" / "rubric_9.json").read_text())
RUBRIC = {k: v for k, v in RUBRIC.items() if k != "CategoryWeights"}


def _extract(wb) -> dict:
    """Full extraction (cells + properties) of an in-memory workbook."""
    d = Path(tempfile.mkdtemp())
    path = d / "t.xlsx"
    wb.save(path)
    excel_utils.process_all_worksheets(str(path), d / "out", quiet=True)
    return wp.load_properties(d / "out")


def _sheet(props, name):
    return next(s for s in props["sheets"] if s["name"] == name)


# ---------------------------------------------------------------------------
# WIDE OUTLIER
# ---------------------------------------------------------------------------

def _widths(runs, used="A1:BZ50", name="Solution Model"):
    return {"name": name, "used_range": used, "default_col_width": 8.43, "hidden_cols": [],
            "column_widths": [{"first": a, "last": b, "value": w} for a, b, w in runs]}


def test_wide_outlier_toy94_shape():
    # A:AZ = 12.9, BA:BD = 34.9, BE:BZ = 12.9 (the toy 94 Fail)
    tags = wp.wide_outlier_tags(_widths([(1, 52, 12.9), (53, 56, 34.9), (57, 78, 12.9)]))
    assert len(tags) == 1 and tags[0].startswith("BA:BD width 34.9 vs neighbours 12.9 (2.7x): WIDE OUTLIER"), tags


def test_wide_outlier_lone_label_column_is_not_compared():
    tags = wp.wide_outlier_tags(_widths([(1, 1, 2.0), (2, 2, 45.0), (3, 30, 12.0)]))
    assert tags == [], tags


def test_wide_outlier_spacer_is_not_a_neighbour():
    # D:E = 30 next to a width-2 spacer, then 12s: the spacer is skipped, 30/12 = 2.5x -> flagged
    tags = wp.wide_outlier_tags(_widths([(1, 2, 12.0), (3, 3, 2.0), (4, 5, 30.0), (6, 6, 2.0), (7, 20, 12.0)]))
    assert len(tags) == 1 and tags[0].startswith("D:E"), tags
    # 24 vs 12 is only 2.0x -> not flagged
    assert wp.wide_outlier_tags(_widths([(1, 2, 12.0), (4, 5, 24.0), (7, 20, 12.0)])) == []


def test_wide_outlier_field_is_not_an_outlier():
    # a model's timeline block (G:DW at 16) beside narrow label columns (A:F at 5.8) is the field, not a group
    assert wp.wide_outlier_tags(_widths([(1, 6, 5.8), (7, 127, 16.0)], used="A1:DW300")) == []


def test_wide_outlier_needs_both_neighbours():
    # a wide pair at the right edge of the used range (long-text note columns; toy 70 Pass)
    assert wp.wide_outlier_tags(_widths([(8, 151, 12.9), (152, 153, 59.7)], used="A1:EW272")) == []
    # a 13-wide pair beside the sheet's margin columns only (Telecom golden)
    assert wp.wide_outlier_tags(_widths([(1, 2, 4.66), (3, 3, 48.9), (4, 4, 13.0), (5, 5, 5.66), (6, 9, 13.0)], used="A1:G122")) == []


def test_wide_outlier_skips_instructions_sheet():
    assert wp.wide_outlier_tags(_widths([(1, 8, 5.2), (9, 10, 23.6), (11, 11, 5.2)], name="Instructions")) == []
    assert wp.wide_outlier_tags(_widths([(1, 8, 5.2), (9, 10, 23.6), (11, 12, 5.2)], name="Model")) != []


# ---------------------------------------------------------------------------
# PERIOD SERIES
# ---------------------------------------------------------------------------

def test_period_keys():
    assert wp._period_key("Q1 2028") == (2028, 3, 0)
    assert wp._period_key("Q3 2028") > wp._period_key("Q2 2028")
    assert wp._period_key("FY2026") == (2026, 12, 31) and wp._period_key("FY26") == (2026, 12, 31)
    # a fiscal-year total column after Q4 is in order; a year row over months is in order
    assert not wp._out_of_order([wp._period_key(x) for x in ("Q3'22", "Q4'22", "FY'22", "Q1'23")])
    assert not wp._out_of_order([wp._period_key(x) for x in ("DEC '25", "DEC '24", "DEC '23")])   # descending = ordered
    assert wp._period_key("Jan-26") == (2026, 1, 0) and wp._period_key("December 2031") == (2031, 12, 0)
    assert wp._period_key(dt.datetime(2027, 1, 31)) == (2027, 1, 31)
    assert wp._period_key(2028) is None and wp._numeric_year(2028) == 2028   # numbers only via runs
    assert wp._period_key("Revenue") is None and wp._period_key("1999 Ford") is None
    assert wp._numeric_year(1234.5) is None and wp._numeric_year(True) is None


def _timeline_wb(out_of_order=False, gap=False, vertical=None, end_marker=None, numeric=False):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Solution Model"
    ws["B2"] = "Year"
    labels = [2026, 2027, 2027, 2028] if numeric else ["Q1 2028", "Q2 2028", "Q3 2028", "Q4 2028"]
    if out_of_order:
        labels[1], labels[2] = labels[2], labels[1]
    for i, lab in enumerate(labels):
        ws.cell(row=2, column=3 + i, value=lab)
    if gap:
        ws["F2"].value = None                    # blank header over F
        ws.cell(row=2, column=7, value="Q1 2029")
    for i in range(5):
        ws.cell(row=3, column=3 + i, value=f"=C1+{i}")
        ws.cell(row=4, column=3 + i, value=100 + i)
    if vertical == "formulas":
        for i, y in enumerate((2028, 2029, 2030, 2031)):
            ws.cell(row=10 + i, column=3, value=y)
            ws.cell(row=10 + i, column=4, value=f"=SUM(C4:G4)*{i + 1}")
    elif vertical == "inputs":
        for i, y in enumerate((2028, 2029, 2030, 2031)):
            ws.cell(row=10 + i, column=3, value=y)
            ws.cell(row=10 + i, column=4, value=0.05 * (i + 1))
    if end_marker:
        ws["B20"] = "End Sheet"
        if end_marker == "content_below":
            ws["C22"] = "Recap"
            ws["C23"] = 5
    return wb


def test_period_series_plain_timeline():
    ps = _sheet(_extract(_timeline_wb()), "Solution Model")["period_series"]
    assert len(ps["horizontal"]) == 1 and ps["horizontal"][0]["range"] == "C2:F2", ps
    assert ps["horizontal"][0]["first"] == "Q1 2028" and ps["horizontal"][0]["last"] == "Q4 2028"
    assert not ps["out_of_order"] and not ps["gaps"] and not ps["vertical"] and ps["end_marker"] is None


def test_period_series_numeric_years_with_repeats():
    ps = _sheet(_extract(_timeline_wb(numeric=True)), "Solution Model")["period_series"]
    assert len(ps["horizontal"]) == 1 and ps["horizontal"][0]["numeric"] is True, ps
    # a row of data values is NOT a timeline
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "S"
    for i, v in enumerate((2010, 2050, 1995, 2030)):
        ws.cell(row=5, column=2 + i, value=v)
    assert _sheet(_extract(wb), "S")["period_series"]["horizontal"] == []


def test_period_series_out_of_order():
    ps = _sheet(_extract(_timeline_wb(out_of_order=True)), "Solution Model")["period_series"]
    assert ps["out_of_order"] and ps["out_of_order"][0]["range"] == "C2:F2", ps
    text = wp.render_properties_text(_extract(_timeline_wb(out_of_order=True)))
    assert "PERIOD SERIES OUT OF ORDER: C2:F2 Q1 2028, Q3 2028, Q2 2028, Q4 2028" in text, text


def test_period_series_unlabeled_gap():
    ps = _sheet(_extract(_timeline_wb(gap=True)), "Solution Model")["period_series"]
    assert ps["gaps"] and ps["gaps"][0]["cells"] == ["F2"], ps
    assert ps["horizontal"][0]["range"] == "C2:G2"


def test_vertical_series_only_with_formulas_beside():
    ps = _sheet(_extract(_timeline_wb(vertical="formulas")), "Solution Model")["period_series"]
    assert len(ps["vertical"]) == 1 and ps["vertical"][0]["range"] == "C10:C13", ps
    ps = _sheet(_extract(_timeline_wb(vertical="inputs")), "Solution Model")["period_series"]
    assert ps["vertical"] == [], ps   # register of typed inputs stays silent
    # no horizontal timeline on the sheet -> never a vertical flag
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Assumptions"
    for i, y in enumerate((2028, 2029, 2030, 2031)):
        ws.cell(row=2 + i, column=1, value=y); ws.cell(row=2 + i, column=2, value=f"=A{2 + i}*2")
    assert _sheet(_extract(wb), "Assumptions")["period_series"]["vertical"] == []


def test_end_marker():
    ps = _sheet(_extract(_timeline_wb(end_marker="clean")), "Solution Model")["period_series"]
    assert ps["end_marker"]["cell"] == "B20" and ps["end_marker"]["rows_below"] == 0
    # no timeline on the sheet -> the marker is not reported at all
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Assumption"
    ws["B5"] = "End Sheet"; ws["B7"] = "data"; ws["B8"] = 1
    assert _sheet(_extract(wb), "Assumption")["period_series"]["end_marker"] is None
    props = _extract(_timeline_wb(end_marker="content_below"))
    ps = _sheet(props, "Solution Model")["period_series"]
    assert ps["end_marker"]["rows_below"] == 2 and ps["end_marker"]["first_row_below"] == 22
    assert 'content continues 2 rows past the "End Sheet" marker at B20 (rows 22-23)' in wp.render_properties_text(props)


# ---------------------------------------------------------------------------
# content fit
# ---------------------------------------------------------------------------

def test_content_fit_classes():
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Model"
    ws.column_dimensions["B"].width = 8.43
    ws.column_dimensions["C"].width = 8.43
    ws.column_dimensions["D"].width = 8.43
    ws["B2"] = 1234567890.0; ws["B2"].number_format = "#,##0"        # 13 chars in 8.43 -> ###
    ws["B3"] = 12.5; ws["B3"].number_format = "#,##0.0"              # fits
    ws["B4"] = 123456789.0                                            # General: never ###
    ws["C6"] = "A rather long text label"; ws["D6"] = 1              # cut off by D6
    ws["C7"] = "A rather long text label"                            # overflows into empty D7
    ws["C8"] = "A rather long text label"; ws["D8"] = 1
    ws["C8"].alignment = openpyxl.styles.Alignment(wrap_text=True)   # wrapped: skipped
    ws.merge_cells("B10:E10"); ws["B10"] = "Merged title that is very long indeed"
    cf = _sheet(_extract(wb), "Model")["column_fit"]
    assert [e["ref"] for e in cf["numeric_overflow"]["examples"]] == ["B2"], cf
    assert [e["ref"] for e in cf["text_cut_off"]["examples"]] == ["C6"], cf
    assert cf["text_overflow_into_empty"] == 1, cf
    text = wp.render_properties_text(_extract(wb))
    # judge v10: the text cut-off class stays in the JSON but is not rendered (toy 69 Pass failed 3/3 on it)
    assert "NUMERIC exceeds width (would render ###): B2" in text and "cut off" not in text and "C6" not in text, text
    ws["B2"] = 1.0
    text = wp.render_properties_text(_extract(wb))
    assert "no numeric value exceeds its column width" in text and "NUMERIC exceeds" not in text, text


def test_display_width_scales_with_font():
    base = wp.display_width("1,234,567")
    assert wp.display_width("1,234,567", 14) > base and wp.display_width("1,234,567", 11, True) > base
    assert wp.display_width("iiii") < wp.display_width("MMMM")


def test_numeric_width_uses_the_cell_font():
    # judge v12, grading 1092: WACC!C31:C39 "3,276,619.94" in Arial 10, column 9.0, shows ### in
    # Excel. Arial 10 digits are as wide as Calibri 11's; scaling by size/11 made them 9% narrow.
    need = wp.numeric_display_width("3,276,619.94", "Arial", 10)
    assert 11.0 < need < 11.3 and need > 9.0 + wp.NUMERIC_FIT_MARGIN_CHARS, need
    assert wp.display_width("3,276,619.94", 10) < 9.0 + wp.NUMERIC_FIT_MARGIN_CHARS   # what missed it
    assert wp.numeric_display_width("73,730.00", "Arial", 10) < 9.0 + wp.NUMERIC_FIT_MARGIN_CHARS   # C30 fits
    # the CashNiagara golden's borderline cells (10 characters, Aptos Narrow 11, column 9.0) stay unflagged
    assert wp.numeric_display_width("456,642.66", "Aptos Narrow", 11) < 9.0 + wp.NUMERIC_FIT_MARGIN_CHARS
    # bold leaves Arial's digits alone and widens Verdana's; an unknown face is measured as Calibri
    assert wp.numeric_display_width("1,234", "Arial", 10, True) == wp.numeric_display_width("1,234", "Arial", 10)
    assert wp.numeric_display_width("1,234", "Verdana", 10, True) > wp.numeric_display_width("1,234", "Verdana", 10)
    assert wp.numeric_display_width("1,234", "No Such Face", 11) == wp.numeric_display_width("1,234", "Calibri", 11)
    assert wp.numeric_display_width("1,234", "Helvetica", 10) == wp.numeric_display_width("1,234", "Arial", 10)
    # the unit is the Normal font's digit in whole pixels: 7 for the three defaults in the corpus
    assert wp.column_unit_px("Calibri", 11) == wp.column_unit_px("Aptos Narrow", 11) == wp.column_unit_px("Arial", 10) == 7.0
    assert wp.column_unit_px("Arial", 11) == 8.0 and wp.column_unit_px(None, None) == 7.0


def test_numeric_fit_end_to_end_arial_10():
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "WACC"
    ws.column_dimensions["C"].width = 9.0
    for r, v in ((30, 73730.0), (31, 5157432.5712), (32, 3253445.5946)):
        ws.cell(r, 3, v).number_format = '#,##0.00;(#,##0.00);"-"'
        ws.cell(r, 3).font = Font(name="Arial", size=10)
    props = _extract(wb)
    cf = _sheet(props, "WACC")["column_fit"]
    assert [e["ref"] for e in cf["numeric_overflow"]["examples"]] == ["C31", "C32"], cf
    assert "NUMERIC exceeds width (would render ###): C31, C32" in wp.render_properties_text(props)
    assert wp.normal_font(wb) == ("Calibri", 11.0) and wp.normal_font(None) == ("Calibri", 11.0)


# ---------------------------------------------------------------------------
# print estimate (check 76, judge v12)
# ---------------------------------------------------------------------------

def test_print_estimate_tags_multi_page_sheets_without_breaks():
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Summary"
    for r in range(1, 134):
        ws.cell(r, 2, f"line {r}")
    ws.print_area = "A1:I133"
    ws.page_setup.orientation = "landscape"
    ws.sheet_properties.pageSetUpPr = openpyxl.worksheet.properties.PageSetupProperties(fitToPage=True)
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0          # one page wide, as many tall as it takes
    short = wb.create_sheet("Checks")
    for r in range(1, 21):
        short.cell(r, 2, r)
    brief = wb.create_sheet("Instructions")
    for r in range(1, 300):
        brief.cell(r, 1, "case text")
    props = _extract(wb)
    pe = _sheet(props, "Summary")["print_estimate"]
    assert pe["basis"] == "print area" and pe["range"] == "A1:I133" and pe["pages_tall"] > 3 and pe["row_breaks_in_range"] == 0, pe
    text = wp.render_properties_text(props)
    line = next(l for l in text.splitlines() if "print estimate (print area A1:I133)" in l)
    assert "MULTI-PAGE, NO MANUAL BREAKS (down)" in line and "fit to 1 wide x any tall" in line, line
    assert "print estimate (used range B1:B20): fits one page" in text        # short sheet: no tag
    assert text.count("print estimate") == 2                                   # the case's Instructions sheet is skipped
    ws.row_breaks.append(openpyxl.worksheet.pagebreak.Break(id=40))
    ws.row_breaks.append(openpyxl.worksheet.pagebreak.Break(id=85))
    text = wp.render_properties_text(_extract(wb))
    line = next(l for l in text.splitlines() if "print estimate (print area A1:I133)" in l)
    assert "manual breaks inside it: 2 row" in line and "MULTI-PAGE" not in line, line
    ws.page_setup.fitToHeight = 1                                              # squeezed onto one page: no tag, the scale shows
    ws.row_breaks = openpyxl.worksheet.pagebreak.RowBreak()
    line = next(l for l in wp.render_properties_text(_extract(wb)).splitlines() if "print estimate (print area A1:I133)" in l)
    assert "fits one page at" in line and "MULTI-PAGE" not in line, line
    for x in props["sheets"]:                                                  # older caches: no key, no line, no crash
        x.pop("print_estimate", None)
    assert "print estimate" not in wp.render_properties_text(props)


# ---------------------------------------------------------------------------
# date literal in formula
# ---------------------------------------------------------------------------

def test_date_literal_formulas():
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Model"
    ws["A1"] = '=IF(B1>"12.12.2028",1,0)'
    ws["A2"] = '=DATEVALUE("2028-12-12")'
    ws["A3"] = '=IF(B1>"12/12/2028",1,0)'
    ws["A4"] = "=B1+1"
    ws["A5"] = dt.datetime(2028, 12, 12)          # a plain date value: not a hit
    ws["A6"] = '="Version 1.2.3"'                 # not a date
    props = _extract(wb)
    assert _sheet(props, "Model")["date_literal_formulas"] == ["A1", "A2", "A3"], props
    assert "formulas with a typed date-like string literal: A1, A2, A3" in wp.render_properties_text(props)


# ---------------------------------------------------------------------------
# retired checks
# ---------------------------------------------------------------------------

def test_retired_checks_config_and_pins():
    retired = rs.retired_check_numbers(RUBRIC)
    assert retired == [28, 37, 101], retired
    flat = rs._flat(RUBRIC)
    assert flat[27] == ("Error Checks", "No unused formatting")
    assert flat[36] == ("Flexibility", "M&A / divestiture flexibility")
    assert flat[100] == ("Purpose & Scope", "Architecture suited to audience")
    assert len(flat) == 132, "the rubric keeps its 132 positions; retirement is a rule, not a deletion"


def test_retired_checks_refuse_on_numbering_drift():
    shifted = {cat: list(checks) for cat, checks in RUBRIC.items()}
    shifted["Accuracy"] = shifted["Accuracy"][1:]     # drop check 1 -> every number shifts
    try:
        rs.retired_check_numbers(shifted)
    except rs.SuitabilityError as e:
        assert "numbering drift" in str(e)
    else:
        raise AssertionError("a renumbered rubric must refuse to retire by position")
    os.environ[f"BIZBENCHJUDGE_{rs.RETIRED_ENV_KEY}"] = "37,99"
    try:
        rs.retired_check_numbers(RUBRIC)
    except rs.SuitabilityError as e:
        assert "no entry in RETIRED_CHECK_NAMES" in str(e)
    else:
        raise AssertionError("an unpinned number must be refused")
    finally:
        os.environ[f"BIZBENCHJUDGE_{rs.RETIRED_ENV_KEY}"] = "28,37,101"


def test_retired_checks_applied_with_and_without_annotation():
    d = Path(tempfile.mkdtemp())
    # (a) no annotation, v1/unknown benchmark -> retired-only gating
    out = rs.load_for_case(d, RUBRIC, benchmark=None)
    assert out is not None and out["provenance"]["retired_checks"] == [28, 37, 101]
    assert out["provenance"]["gated"] is False and out["provenance"]["excluded_count"] == 3
    assert "No unused formatting" in out["excluded"]["Error Checks"]
    assert "M&A / divestiture flexibility" in out["excluded"]["Flexibility"]
    assert "Architecture suited to audience" in out["excluded"]["Purpose & Scope"]
    # (b) skip env -> still retired
    os.environ[f"BIZBENCHJUDGE_{rs.SKIP_ENV}"] = "1"
    os.environ[rs.SKIP_ENV] = "1"
    try:
        out = rs.load_for_case(d, RUBRIC, benchmark="v2")
        assert out["provenance"]["skipped_via_env"] is True and out["provenance"]["retired_checks"] == [28, 37, 101]
    finally:
        os.environ.pop(rs.SKIP_ENV, None)
        os.environ.pop(f"BIZBENCHJUDGE_{rs.SKIP_ENV}", None)
    # (c) a real annotation marking 37 applicable is overridden
    rubrics = [{"no": i, "category": c, "name": n, "verdict": "applicable", "conditional": False}
               for i, (c, n) in enumerate(rs._flat(RUBRIC), 1)]
    rubrics[4]["verdict"] = "not_applicable"      # task's own exclusion of check 5 survives
    (d / rs.STAGED_FILENAME).write_text(json.dumps({"complete": True, "annotator": "t", "rubrics": rubrics}))
    out = rs.load_for_case(d, RUBRIC, benchmark="v2")
    assert out["provenance"]["gated"] is True and out["provenance"]["retired_checks"] == [28, 37, 101]
    assert out["provenance"]["excluded_count"] == 4, out["provenance"]
    eff = rs.build_effective_weights(
        json.loads((JUDGE / "prompts" / "rubrics" / "rubric_9_weights.json").read_text()), out["excluded"])
    assert all(c["name"] != "No unused formatting" for c in eff["Error Checks"])
    assert len(eff["Error Checks"]) == 14, "the other 14 Error Checks items carry the category"
    assert all(c["name"] != "M&A / divestiture flexibility" for c in eff["Flexibility"])
    assert all(c["name"] != "Architecture suited to audience" for c in eff["Purpose & Scope"])


def test_retired_check_evidence_line_follows_the_config():
    """Check 28's own evidence line is left out while 28 is retired and comes
    back when 28 is taken off judge.retired_checks — no code change either way."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"], ws["C3"] = "x", 1
    ws["B2"].font = Font(bold=True)               # styled, empty, inside the used range
    props = _extract(wb)
    line = "styled empty cells in used range: 1 (e.g. B2)"
    assert line in wp.render_properties_text(props), "no retired list given -> everything renders"
    key = f"BIZBENCHJUDGE_{rs.RETIRED_ENV_KEY}"
    try:
        retired = rs.retired_check_numbers(RUBRIC)             # config as shipped: 28 retired
        assert wp.STYLED_EMPTY_CHECK in retired
        assert "styled empty cells" not in wp.render_properties_text(props, retired_checks=retired)
        os.environ[key] = "37,101"                             # 28 taken off the list
        retired = rs.retired_check_numbers(RUBRIC)
        assert retired == [37, 101]
        assert line in wp.render_properties_text(props, retired_checks=retired)
    finally:
        os.environ[key] = "28,37,101"
    # the data is kept either way, so neither direction needs a cache rebuild
    assert props["sheets"][0]["styled_empty_cells"] == {"count": 1, "examples": ["B2"]}


if __name__ == "__main__":
    for name, fn in list(globals().items()):
        if name.startswith("test_") and callable(fn):
            fn()
            print(f"ok  {name}")
    print("all judge v9 evidence-flag / retired-check tests passed")
