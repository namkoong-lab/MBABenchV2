"""Offline: freeze-pane extent and the EXCESSIVE tag (judge v11, check 122).

Run from judge/:  python tests_offline/test_freeze_extent.py
"""
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import openpyxl  # noqa: E402

from utils.misc_utils import load_project_configs  # noqa: E402

load_project_configs(benchmark="v2")

from utils import excel_utils, workbook_properties as wp  # noqa: E402


def _extract(wb):
    d = Path(tempfile.mkdtemp())
    path = d / "t.xlsx"
    wb.save(path)
    excel_utils.process_all_worksheets(str(path), d / "out", quiet=True)
    return wp.load_properties(d / "out")


def _sheet(props, name):
    return next(s for s in props["sheets"] if s["name"] == name)


def test_extent_and_tag():
    wb = openpyxl.Workbook()
    big = wb.active
    big.title = "Checks"
    for r in range(1, 60):
        big.cell(row=r, column=1, value=r)
        big.row_dimensions[r].height = 18
    big.row_dimensions[3].height = 42
    big.row_dimensions[4].hidden = True               # hidden rows take no screen space
    big.freeze_panes = "A35"                          # 34 rows locked: 32 x 18 + 42, one hidden
    ok = wb.create_sheet("Owning")
    for r in range(1, 40):
        ok.cell(row=r, column=1, value=r)
    ok.column_dimensions["A"].width = 30
    ok.column_dimensions["B"].width = 10
    ok.freeze_panes = "C9"                            # 8 default rows, 2 columns = 40 characters
    wide = wb.create_sheet("Wide")
    wide["A1"] = 1
    for col in "ABCDEFGHI":
        wide.column_dimensions[col].width = 25
    wide.freeze_panes = "J2"                          # 225 characters of columns: a full screen
    gold = wb.create_sheet("GoldenLike")
    gold["A1"] = 1
    for col in "ABCDEF":
        gold.column_dimensions[col].width = 26
    gold.freeze_panes = "G6"                          # 156 characters, as three goldens do: no tag
    none = wb.create_sheet("Cover")
    none["A1"] = "x"
    props = _extract(wb)
    e = _sheet(props, "Checks")["freeze_extent"]
    assert e["rows"] == 34 and e["cols"] == 0 and abs(e["rows_pt"] - (32 * 18 + 42)) < 0.01, e
    t = wp.freeze_text(_sheet(props, "Checks"))
    assert t.startswith("A35 (34 rows ≈ 618 pt frozen, 0 columns) EXCESSIVE: 618 pt of rows locked"), t
    o = wp.freeze_text(_sheet(props, "Owning"))
    assert o == "C9 (8 rows ≈ 120 pt frozen, 2 columns ≈ 40 characters wide)", o
    w = wp.freeze_text(_sheet(props, "Wide"))
    assert "EXCESSIVE: 225 characters of columns locked" in w, w
    assert "EXCESSIVE" not in wp.freeze_text(_sheet(props, "GoldenLike"))
    assert wp.freeze_text(_sheet(props, "Cover")) == "none"
    text = wp.render_properties_text(props)
    assert text.count("EXCESSIVE:") == 2 and "freeze panes: none" in text


def test_older_cache_renders_the_cell_alone():
    assert wp.freeze_text({"freeze_panes": "A35"}) == "A35"
    assert wp.freeze_text({"freeze_panes": None}) == "none"
    assert wp.freeze_text({"freeze_panes": "B2", "freeze_extent": {"rows": 1, "cols": 1, "rows_pt": 15.0, "cols_chars": 8.4}}) \
        == "B2 (1 row ≈ 15 pt frozen, 1 column ≈ 8 characters wide)"


if __name__ == "__main__":
    for name, fn in list(globals().items()):
        if name.startswith("test_") and callable(fn):
            fn()
            print(f"ok  {name}")
    print("all freeze-extent tests passed")
