"""Offline: the inflated-extent guard (judge/utils/sheet_extent.py).

A sheet with values in A1:C3 plus one styled-but-empty cell at XFD1 declares
3 x 16,384. Ordinary sheets must come through byte-identical; only a sheet
over the cell budget is narrowed to its last value-bearing column.

Run from judge/:  python tests_offline/test_sheet_extent_guard.py
"""
import io
import os
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import openpyxl  # noqa: E402
from openpyxl.styles import Font  # noqa: E402

from utils.sheet_extent import ENV_MAX_SHEET_CELLS, iteration_bounds  # noqa: E402
from utils.excel_utils import extract_all_cell_data  # noqa: E402


def _inflated_workbook():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "x"
    ws["B2"] = "=A1"
    ws["C3"] = 5
    ws.cell(row=1, column=16384).font = Font(bold=True)  # styled, no value
    return wb


def _both_views(wb):
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    formula = openpyxl.load_workbook(buf)
    buf.seek(0)
    data = openpyxl.load_workbook(buf, data_only=True, read_only=True)
    return formula.active, data[formula.active.title]


def test_declared_extent_is_inflated():
    ws = _inflated_workbook().active
    assert (ws.max_row, ws.max_column) == (3, 16384)


def test_bounds_untouched_within_budget():
    ws = _inflated_workbook().active
    assert iteration_bounds(ws) == (3, 16384, False)
    assert iteration_bounds(ws, max_cells=3 * 16384) == (3, 16384, False)


def test_bounds_narrow_to_value_extent_over_budget():
    ws = _inflated_workbook().active
    assert iteration_bounds(ws, max_cells=1000) == (3, 3, True)


def test_extract_narrows_only_over_budget():
    ws, ws_data = _both_views(_inflated_workbook())
    old = os.environ.pop(ENV_MAX_SHEET_CELLS, None)
    try:
        os.environ[ENV_MAX_SHEET_CELLS] = "1000"
        narrowed = extract_all_cell_data(ws, ws_data)
        assert narrowed["metadata"]["rows"] == 3
        assert narrowed["metadata"]["columns"] == 3
        assert "=A1" in narrowed["full"] and "5" in narrowed["data"]

        os.environ.pop(ENV_MAX_SHEET_CELLS)
        full = extract_all_cell_data(ws, ws_data)
        assert full["metadata"]["rows"] == 3
        assert full["metadata"]["columns"] == 16384
    finally:
        if old is not None:
            os.environ[ENV_MAX_SHEET_CELLS] = old
        else:
            os.environ.pop(ENV_MAX_SHEET_CELLS, None)


if __name__ == "__main__":
    for name, fn in list(globals().items()):
        if name.startswith("test_") and callable(fn):
            fn()
            print(f"ok  {name}")
    print("all sheet-extent guard tests passed")
