"""Offline: rounding statements evidence for check 105 (judge v11, 2026-09-18).

Run from judge/:  python tests_offline/test_rounding_statements.py
"""
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import openpyxl  # noqa: E402

from utils.misc_utils import load_project_configs  # noqa: E402

load_project_configs(benchmark="v2")

from utils import excel_utils, workbook_properties as wp  # noqa: E402


def _extract(wb):
    d = Path(tempfile.mkdtemp())
    path = d / "t.xlsx"
    wb.save(path)
    excel_utils.process_all_worksheets(str(path), d / "out", quiet=True)
    return wp.load_properties(d / "out")


def _sheet(props, name):
    return next(s for s in props["sheets"] if s["name"] == name)


def test_statements_and_round_counts():
    wb = openpyxl.Workbook()
    own = wb.active
    own.title = "Owning"
    own["B3"] = "USD, rounded to $0.01. Costs and cash outflows are negative."
    own["D9"], own["D10"], own["D11"] = "=1000/3", "=D9*1.02", "=SUM(D9:D10)"          # full precision carried
    q = wb.create_sheet("Questions")
    q["A1"] = "Questions (please round your answers to two decimal places)"             # the case's instruction: not a claim
    q["C2"], q["C3"] = "=ROUND(Owning!D9,2)", "=ROUNDUP(Owning!D10,0)"
    hon = wb.create_sheet("Honest")
    hon["B2"] = "All figures are carried unrounded; only the answers are rounded to two decimals"
    hon["C5"] = "=Owning!D9*2"
    ins = wb.create_sheet("Instructions")
    ins["B9"] = "Report the headcount rounded to the next whole number."                 # claim-shaped, but a sheet without formulas
    none = wb.create_sheet("NoStatement")
    none["A1"], none["B1"], none["A2"] = "Revenue", "=ROUND(Owning!D9,0)", "Round moulder / proofing baskets"
    props = _extract(wb)
    o = _sheet(props, "Owning")["rounding_statements"]
    assert o["n_formulas"] == 3 and o["n_round_formulas"] == 0 and o["statements"][0]["cell"] == "B3", o
    qq = _sheet(props, "Questions")["rounding_statements"]
    assert qq["n_round_formulas"] == 2 and qq["statements"] == [], qq
    assert _sheet(props, "NoStatement")["rounding_statements"]["statements"] == []      # a ROUND formula / a product name is no statement
    assert _sheet(props, "Instructions")["rounding_statements"]["statements"][0]["cell"] == "B9"
    text = wp.render_properties_text(props)
    lines = [l.strip() for l in text.splitlines() if l.strip().startswith("rounding statements:")]
    assert len(lines) == 2, lines                                                        # Owning and Honest; Instructions has no formulas
    own_line = next(l for l in lines if "B3" in l)
    assert 'B3 "USD, rounded to $0.01. Costs and cash outflows are negative."' in own_line
    assert own_line.endswith("(ROUND/ROUNDUP/ROUNDDOWN/MROUND): 0 of 3"), own_line
    assert any("carried unrounded" in l and l.endswith(": 0 of 1") for l in lines)


def test_render_tolerates_old_caches_and_caps():
    assert wp._rounding_lines(None) == [] and wp._rounding_lines("unknown") == []
    assert wp._rounding_lines({"statements": [], "n_formulas": 9, "n_round_formulas": 0}) == []
    many = {"statements": [{"cell": f"B{i}", "text": "rounded"} for i in range(1, 10)], "n_formulas": 1200, "n_round_formulas": 35}
    line = wp._rounding_lines(many)[0]
    assert "(+3 more)" in line and line.endswith("35 of 1,200"), line


if __name__ == "__main__":
    for name, fn in list(globals().items()):
        if name.startswith("test_") and callable(fn):
            fn()
            print(f"ok  {name}")
    print("all rounding-statement tests passed")
