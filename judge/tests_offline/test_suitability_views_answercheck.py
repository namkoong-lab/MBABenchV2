"""Offline tests for the 2026-08 judge update (no network, no DB, no S3).

Covers:
  - rubric suitability: selection rule, validation refusals, filter + letter
    remapping, effective-weights renormalization, zero-applicable defense
  - scoring math under filtered weights (renorm is emergent in
    calculate_scores; unscored-check penalty confined to applicable checks)
  - answer check: extractor conventions (B/C columns, split sheets,
    convention_not_found, uncached-formula detection), comparator tolerance
    edges, label alignment fallback
  - context views: *_data.csv stripping (styles gone, number-format
    rendering kept), prompt-file exclusion, category-keyed read_file serving
  - template 5: byte-identical shared prefix across categories
  - prompt propagation: the live GUI/excel prompt files carry exactly the
    rubric_9.json edition

Run:  uv run python judge/tests_offline/test_suitability_views_answercheck.py
"""

import copy
import importlib.util
import json
import sys
import tempfile
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill

JUDGE_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(JUDGE_ROOT))

from utils import rubric_suitability as rs  # noqa: E402
from utils.misc_utils import load_project_configs  # noqa: E402

load_project_configs()

from utils import answer_check as ac  # noqa: E402
from utils.excel_utils import (  # noqa: E402
    prepare_directory_files,
    process_all_worksheets,
)
from utils.prompt_utils import check_letter, compile_prompt  # noqa: E402

_spec = importlib.util.spec_from_file_location(
    "_judge_module", str(JUDGE_ROOT / "main_scripts" / "judge.py")
)
_judge = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(_judge)

RUBRIC = json.loads((JUDGE_ROOT / "prompts/rubrics/rubric_9.json").read_text())
WEIGHTS = json.loads(
    (JUDGE_ROOT / "prompts/rubrics/rubric_9_weights.json").read_text()
)


def make_annotation(rubric, not_applicable=(), conditional=()):
    """Fixture annotation matching `rubric`, with the given `no`s flipped."""
    entries = []
    no = 0
    for cat, checks in rubric.items():
        for c in checks:
            no += 1
            entries.append(
                {
                    "no": no,
                    "category": cat,
                    "name": c["name"],
                    "verdict": (
                        "not_applicable" if no in set(not_applicable) else "applicable"
                    ),
                    "conditional": no in set(conditional),
                    "note": "",
                }
            )
    return {
        "task_id": 999,
        "annotator": "julian",
        "created_at": "20260827T000000Z",
        "rubric_version": "ffe2d7b6c063",
        "complete": True,
        "rubrics": entries,
    }


# ---------------------------------------------------------------- suitability


def test_selection_rule():
    metas = [
        {"key": "a/julian_1.json", "annotator": "julian", "complete": True,
         "created_at": "20260101T000000Z"},
        {"key": "a/julian_2.json", "annotator": "julian", "complete": True,
         "created_at": "20260301T000000Z"},
        {"key": "a/julian_3.json", "annotator": "julian", "complete": False,
         "created_at": "20261231T000000Z"},
        {"key": "a/thomson_1.json", "annotator": "thomson", "complete": True,
         "created_at": "20261231T000000Z"},
    ]
    assert rs.select_annotation_key(metas)["key"] == "a/julian_2.json"
    # created_at tie -> lexicographically larger key, deterministic
    tie = [
        {"key": "a/julian_a.json", "annotator": "julian", "complete": True,
         "created_at": "20260101T000000Z"},
        {"key": "a/julian_b.json", "annotator": "julian", "complete": True,
         "created_at": "20260101T000000Z"},
    ]
    assert rs.select_annotation_key(tie)["key"] == "a/julian_b.json"
    try:
        rs.select_annotation_key([metas[2], metas[3]])
        raise AssertionError("expected SuitabilityError")
    except rs.SuitabilityError:
        pass
    print("OK  selection rule: latest complete julian; deterministic ties; refusals")


def test_validation_refusals():
    good = make_annotation(RUBRIC, not_applicable=[5])
    rs.validate_annotation(good, RUBRIC)  # must not raise

    renamed = copy.deepcopy(good)
    renamed["rubrics"][10]["name"] = "Some Other Check"
    for broken, label in [
        (renamed, "renamed check"),
        ({**good, "rubrics": good["rubrics"][:-1]}, "wrong count"),
        ({**good, "rubrics": []}, "empty"),
    ]:
        try:
            rs.validate_annotation(broken, RUBRIC)
            raise AssertionError(f"expected refusal: {label}")
        except rs.SuitabilityError:
            pass

    bad_verdict = copy.deepcopy(good)
    bad_verdict["rubrics"][0]["verdict"] = "maybe"
    try:
        rs.validate_annotation(bad_verdict, RUBRIC)
        raise AssertionError("expected refusal: unknown verdict")
    except rs.SuitabilityError:
        pass
    print("OK  validation refuses drifted/incomplete/unknown-verdict annotations")


def test_filter_letters_and_weights():
    # Flip two Accuracy checks (nos 1 and 3) and one Structure check.
    struct_first_no = 1 + sum(
        len(RUBRIC[c]) for c in RUBRIC if list(RUBRIC).index(c) < list(RUBRIC).index("Structure")
    )
    ann = make_annotation(RUBRIC, not_applicable=[1, 3, struct_first_no])
    suit = rs.build_suitability(ann, RUBRIC, s3_key="k")
    assert suit["provenance"]["excluded_count"] == 3
    assert len(suit["applicable"]["Accuracy"]) == 1
    # Letters over the FILTERED list stay contiguous from A.
    filtered = [
        c for c in RUBRIC["Accuracy"] if c["name"] in set(suit["applicable"]["Accuracy"])
    ]
    letters = [check_letter(i) for i in range(len(filtered))]
    assert letters == ["A"]
    assert filtered[0]["name"] == RUBRIC["Accuracy"][1]["name"]

    eff = rs.build_effective_weights(WEIGHTS, suit["excluded"])
    assert len(eff["Accuracy"]) == 1
    assert eff["CategoryWeights"] == WEIGHTS["CategoryWeights"]  # untouched
    # Scoring: the surviving Accuracy check passes -> category 100 despite
    # the exclusions; a missing applicable check is penalized, an excluded
    # one is not (no unscored warning for it).
    responses = {"Accuracy": [{"name": filtered[0]["name"], "mistakes": []}]}
    scores = _judge.calculate_scores(responses, eff, max_mistakes=1)
    assert scores["criteria_scores"]["Accuracy"]["normalized_score"] == 100.0
    assert "Accuracy" not in scores["scoring_warnings"]["unscored_checks"]
    print("OK  filter + contiguous letters + emergent renorm + confined penalty")


def test_zero_applicable_defense():
    nos = list(range(1, len(RUBRIC["Accuracy"]) + 1))  # all of Accuracy
    ann = make_annotation(RUBRIC, not_applicable=nos)
    suit = rs.build_suitability(ann, RUBRIC)
    eff = rs.build_effective_weights(WEIGHTS, suit["excluded"])
    cw = eff["CategoryWeights"][0]
    assert "Accuracy" not in cw and "Accuracy" not in eff
    assert abs(sum(cw.values()) - 1.0) < 1e-9
    print("OK  zero-applicable category dropped; CategoryWeights renormalized to 1")


# --------------------------------------------------------------- answer check


def _wb(rows, sheet="Questions", answer_col=2, header_row=1, extra_sheets=()):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet
    ws.cell(row=header_row, column=1, value="Question")
    ws.cell(row=header_row, column=answer_col, value="Answers")
    for i, (label, value) in enumerate(rows, start=header_row + 1):
        ws.cell(row=i, column=1, value=label)
        ws.cell(row=i, column=answer_col, value=value)
    for name, rows2 in extra_sheets:
        ws2 = wb.create_sheet(name)
        ws2["A1"] = "Question"
        ws2["B1"] = "Answers"
        for i, (label, value) in enumerate(rows2, start=2):
            ws2.cell(row=i, column=1, value=label)
            ws2.cell(row=i, column=2, value=value)
    return wb


def test_extractor_conventions():
    with tempfile.TemporaryDirectory() as td:
        td = Path(td)
        # Column C variant (column B holds a label, like tasks 19/22/32/38/45)
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Questions"
        ws["A1"], ws["B1"], ws["C1"] = "Question", "Part", "Answers"
        ws["A2"], ws["B2"], ws["C2"] = "Q1", "Model", 7
        p = td / "c_col.xlsx"
        wb.save(p)
        got = ac.extract_answers(p, allow_recalc=False)
        assert got["status"] == "ok" and got["rows"][0]["value"] == 7

        # Split sheets (task 54 style) concatenate in name order
        wb = _wb([("Q1", 1)], sheet="Questions Task 1",
                 extra_sheets=[("Questions Task 2", [("Q2", 2)])])
        p = td / "split.xlsx"
        wb.save(p)
        got = ac.extract_answers(p, allow_recalc=False)
        assert [r["value"] for r in got["rows"]] == [1, 2]

        # No Questions sheet -> convention_not_found
        wb = openpyxl.Workbook()
        wb.active.title = "Model"
        p = td / "none.xlsx"
        wb.save(p)
        assert ac.extract_answers(p, allow_recalc=False)["status"] == "convention_not_found"

        # Header outside the A-F scan window -> convention_not_found
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Questions"
        ws["H1"] = "Answers"
        p = td / "far.xlsx"
        wb.save(p)
        assert ac.extract_answers(p, allow_recalc=False)["status"] == "convention_not_found"

        # Uncached-formula detection (no recalc attempted offline)
        wb = _wb([("Q1", "=1+1")])
        p = td / "formula.xlsx"
        wb.save(p)
        n = ac._uncached_answer_cells(p, [("Questions", 2, 2)])
        assert n == 1
    print("OK  extractor: C column, split sheets, not-found paths, uncached detection")


def test_comparator_edges():
    m = ac.compare_values
    assert m(0.0, 0.0)["verdict"] == "match"                      # zero
    assert m(0.0, 5e-10)["verdict"] == "match"                    # abs floor
    assert m(0.0, 2e-9)["verdict"] == "mismatch"                  # beyond floor
    assert m(1e12, 1e12 * (1 + 5e-7))["verdict"] == "match"       # rel, large
    assert m(1e12, 1e12 * (1 + 5e-6))["verdict"] == "mismatch"
    assert m(0.15, 15.0)["verdict"] == "mismatch"                 # percent vs decimal
    assert m("Yes", " yes ")["verdict"] == "match"                # text normalize
    assert m("Yes", "No")["verdict"] == "mismatch"
    assert m("42", 42.0)["verdict"] == "match"                    # numeric string
    r = m("apples", 42.0)
    assert r["verdict"] == "mismatch" and r["mismatch_type"] == "type_mismatch"
    assert m(1.0, None)["verdict"] == "missing"

    # Label alignment with positional fallback
    sol = {"rows": [
        {"sheet": "Q", "row": 2, "label": "Alpha", "value": 1},
        {"sheet": "Q", "row": 3, "label": "Beta", "value": 2},
    ]}
    att = {"rows": [
        {"sheet": "Q", "row": 2, "label": "Beta", "value": 2},
        {"sheet": "Q", "row": 3, "label": "Gamma", "value": 9},
    ]}
    cmp_result = ac.compare_answer_sets(sol, att)
    by_label = {q["label"]: q for q in cmp_result["questions"]}
    assert by_label["Beta"]["verdict"] == "match"          # label-matched
    assert by_label["Alpha"].get("label_mismatch") is True  # positional fallback
    assert cmp_result["n_label_mismatch"] == 1
    print("OK  comparator: tolerance edges, text/type rules, label fallback")


# ------------------------------------------------------------- context views


def test_data_views_and_serving():
    with tempfile.TemporaryDirectory() as td:
        td = Path(td)
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Model"
        c = ws["A1"]
        c.value = 1234.5678
        c.number_format = "#,##0.00"
        c.font = Font(bold=True, color="FFFF0000")
        c.fill = PatternFill("solid", start_color="FFFFFF00")
        ws["B1"] = "=A1*2"
        ws.merge_cells("A3:B3")
        xlsx = td / "model.xlsx"
        wb.save(xlsx)
        out = td / "csvs"
        process_all_worksheets(str(xlsx), out, quiet=True)

        full = (out / "Model_full.csv").read_text()
        data = (out / "Model_data.csv").read_text()
        assert "bold" in full and "bgcolor" in full
        assert "bold" not in data and "bgcolor" not in data and "textcolor" not in data
        assert "1,234.57" in data                      # number format rendered
        assert "FORMULA:=A1*2" in data
        assert not any(
            n.endswith("_data.csv") for n in prepare_directory_files(str(out))
        )

        class TC:
            class function:
                name = "read_file"
                arguments = json.dumps(dict(
                    source="attempt", filename="Model_full.csv",
                    start_row=1, end_row=1, start_col="A", end_col="B",
                ))

        d = str(out)
        assert "bold" in _judge._execute_read_file(TC, d, None)  # legacy v4
        assert "bold" not in _judge._execute_read_file(
            TC, d, None, category="Accuracy", format_notes=set()
        )
        notes = set()
        first = _judge._execute_read_file(TC, d, None, category="Formatting",
                                          format_notes=notes)
        again = _judge._execute_read_file(TC, d, None, category="Formatting",
                                          format_notes=notes)
        assert "bold" in first and "MERGED CELLS" in first
        assert "MERGED CELLS" not in again                # once per sheet
        struct = _judge._execute_read_file(TC, d, None, category="Structure",
                                           format_notes=set())
        assert "bold" not in struct and "MERGED CELLS" in struct
    print("OK  data views: stripped/kept content, exclusion, serving matrix")


def test_template5_shared_prefix():
    tpl = str(JUDGE_ROOT / "prompts/agentic_judge_template_5.yaml")
    base = dict(
        rubric_checks_text="Check A: X\nPass: y\nFail: z",
        check_letters_text="A",
        attempt_files_text="    a_full.csv\n      Dimensions: 5 rows x 4 columns",
        solution_files_text="    s_full.csv",
        prior_findings=None,
    )
    s1 = compile_prompt(tpl, category="Accuracy", **base)[0]
    s2 = compile_prompt(
        tpl, category="Formatting",
        **{**base, "prior_findings": "Findings from prior categories (for reference):\n{}\n\n"},
    )[0]
    assert s1[0]["content"] == s2[0]["content"], "system message must be static"
    u1, u2 = s1[1]["content"], s2[1]["content"]
    import os as _os

    prefix = _os.path.commonprefix([u1, u2])
    assert "-" * 80 in prefix, "shared prefix must cover all static user blocks"
    assert u1[len(prefix):].startswith("Accuracy") or u2[len(prefix):].startswith(
        "Formatting"
    )
    print(f"OK  template_5: shared prefix covers static blocks ({len(prefix)} chars)")


def test_prompt_files_carry_rubric9():
    spec = importlib.util.spec_from_file_location(
        "_builder",
        str(JUDGE_ROOT / "operation_scripts" / "build_rubric_9_from_xlsx.py"),
    )
    builder = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(builder)
    expected_block = builder.render_rubric_block(RUBRIC)
    for p in builder.LIVE_PROMPT_FILES + list(builder.EXCEL_COPIES.values()):
        text = p.read_text()
        got = text[text.index(builder.RUBRIC_MARKER):]
        assert got == expected_block, f"{p} rubric block != rubric_9.json rendering"
    print("OK  live prompt files carry exactly the rubric_9.json edition (4 files)")


if __name__ == "__main__":
    test_selection_rule()
    test_validation_refusals()
    test_filter_letters_and_weights()
    test_zero_applicable_defense()
    test_extractor_conventions()
    test_comparator_edges()
    test_data_views_and_serving()
    test_template5_shared_prefix()
    test_prompt_files_carry_rubric9()
    print("ALL 2026-08 JUDGE UPDATE OFFLINE CHECKS PASSED")
