"""Offline tests for the deterministic-checks adapter (utils/det_checks.py, judge v13) and
the scoring layer it feeds (judge._apply_harness_verdicts / _finalize_case).

Run from judge/:  python tests_offline/test_det_checks.py
or  uv run pytest tests_offline/test_det_checks.py -q
(it builds several small workbooks: on the shared 16 GB grading Mac, run it through the
heavy-job guard like any multi-workbook job).

No DB, S3, LLM, LibreOffice or Excel: workbooks are built with openpyxl in temporary
folders, and the one LibreOffice call the recalculation pipeline would make for an
openpyxl-written file is replaced by a fake that copies the file (the toys hold constants
only, so the copy carries the same values) and records the policy it was given.
"""
import atexit
import contextlib
import copy
import json
import os
import shutil
import signal
import subprocess
import sys
import tempfile
import time
from pathlib import Path

JUDGE = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(JUDGE))

import openpyxl  # noqa: E402

from utils.misc_utils import load_project_configs, project_prefix  # noqa: E402

load_project_configs()

import importlib.util  # noqa: E402

import detchecks.api as det_api  # noqa: E402
import detchecks.core.recalc as det_recalc  # noqa: E402
from detchecks.checks import REGISTRY  # noqa: E402
from detchecks.core.sheet import ExcelError  # noqa: E402
from utils import det_checks as D  # noqa: E402
from utils import rubric_suitability, workbook_properties  # noqa: E402

_spec = importlib.util.spec_from_file_location("_judge_module", str(JUDGE / "main_scripts" / "judge.py"))
judge = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(judge)

RUBRIC_PATH = JUDGE / "prompts" / "rubrics" / "rubric_9.json"
WEIGHTS_PATH = JUDGE / "prompts" / "rubrics" / "rubric_9_weights.json"
RUBRIC = json.loads(RUBRIC_PATH.read_text())
WEIGHTS = json.loads(WEIGHTS_PATH.read_text())
FLAT = [(cat, c["name"]) for cat, checks in RUBRIC.items() for c in checks]

# The maintainer's decisions (2026-10-04) — the config must say exactly this.
LIVE = {29, 47, 49, 50, 51, 61, 62, 66, 69, 70, 73, 74, 77, 80, 87, 92, 93, 94, 95}
RECORDED_ONLY = {22, 65}


def key(no):
    return "/".join(D.DET_CHECK_NAMES[no])


# --------------------------------------------------------------------------- fixtures
@contextlib.contextmanager
def env(**overrides):
    """Temporarily set BIZBENCHJUDGE_<KEY> config values (None removes one)."""
    prefix = project_prefix()
    saved = {}
    try:
        for k, v in overrides.items():
            name = f"{prefix}_{k}"
            saved[name] = os.environ.get(name)
            if v is None:
                os.environ.pop(name, None)
            else:
                os.environ[name] = str(v)
        yield
    finally:
        for name, v in saved.items():
            if v is None:
                os.environ.pop(name, None)
            else:
                os.environ[name] = v


@contextlib.contextmanager
def fake_libreoffice():
    """Replace detchecks' LibreOffice runner: copy the file, record (src, out_dir, policy)."""
    calls = []
    original = det_recalc.libreoffice_recalc

    def _fake(src, out_dir, policy):
        calls.append({"src": src, "out_dir": out_dir, "policy": policy})
        os.makedirs(out_dir, exist_ok=True)
        out = os.path.join(out_dir, os.path.splitext(os.path.basename(src))[0] + ".xlsx")
        shutil.copy(src, out)
        return out

    det_recalc.libreoffice_recalc = _fake
    try:
        yield calls
    finally:
        det_recalc.libreoffice_recalc = original


@contextlib.contextmanager
def patched_grade(fn):
    original = det_api.grade
    det_api.grade = fn
    try:
        yield
    finally:
        det_api.grade = original


def make_task(root: Path, *, hidden=False, negative=False, origin="Model_v3.xlsx", corrupt=False) -> Path:
    """A staged task folder: ai_attempt.xlsx (+ _attempt_origin.json unless origin is None)."""
    folder = Path(tempfile.mkdtemp(dir=root))
    if corrupt:
        (folder / "ai_attempt.xlsx").write_bytes(b"PK\x03\x04 this is not a workbook")
    else:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Model"
        ws["A1"], ws["B1"] = "Revenue", 100
        ws["A2"], ws["B2"] = "Cost", (-5 if negative else 5)
        notes = wb.create_sheet("Notes")
        notes["A1"] = "source: case pack"
        if hidden:
            notes.sheet_state = "hidden"
        wb.save(folder / "ai_attempt.xlsx")
    if origin is not None:
        (folder / workbook_properties.ORIGIN_FILENAME).write_text(
            json.dumps({"original_filename": origin, "source": "test", "attempt_id": 0}))
    return folder


def stage_annotation(folder: Path, not_applicable=()) -> Path:
    """A complete 'julian' rubric-suitability annotation with the given numbers not_applicable."""
    rubrics = [{"no": i, "category": cat, "name": name,
                "verdict": "not_applicable" if i in not_applicable else "applicable", "conditional": False}
               for i, (cat, name) in enumerate(FLAT, 1)]
    doc = {"annotator": "julian", "complete": True, "created_at": "2026-10-04", "rubric_version": "test",
           "rubrics": rubrics, "_staging": {"s3_key": "test/annotation.json"}}
    p = folder / rubric_suitability.STAGED_FILENAME
    p.write_text(json.dumps(doc))
    return p


def run(folder, **kw):
    kw.setdefault("benchmark", None)
    return D.run_det_checks(folder, rubric_path=RUBRIC_PATH, weights_path=WEIGHTS_PATH, **kw)


def strict_json(text):
    def _bad(c):
        raise ValueError(f"non-standard JSON constant {c}")
    return json.loads(text, parse_constant=_bad)


# --------------------------------------------------------------------------- config / pins
def test_pins_registry_rubric_weights_and_config():
    for no, cls in REGISTRY.items():
        assert no in D.DET_CHECK_NAMES, f"registered check {no} has no pin"
        assert cls.key == key(no), (no, cls.key, D.DET_CHECK_NAMES[no])
    for no, pin in D.DET_CHECK_NAMES.items():
        assert FLAT[no - 1] == pin, f"pin {no} {pin} != rubric_9 {FLAT[no - 1]}"
        cat, name = pin
        assert any(e["name"] == name for e in WEIGHTS[cat]), f"{pin} not weighted"
    s = D.load_settings()
    assert s.enabled is True, "det_checks.enabled must be true for judge v13"
    assert set(s.live) == LIVE, sorted(set(s.live) ^ LIVE)
    assert set(s.recorded_only) == RECORDED_ONLY
    assert s.excel_recalc is False, "Excel recalculation must stay OFF (maintainer 2026-10-04)"
    assert s.libreoffice_path == os.environ[f"{project_prefix()}_PATHS_LIBREOFFICE_PATH"]
    assert s.libreoffice_timeout_s == 600
    assert s.libreoffice_max_mb == 10, "LibreOffice size limit is 10 MB (maintainer 2026-10-04)"
    plan = D.configured_checks(s, RUBRIC)
    assert {n for n, live in plan.items() if live} == LIVE
    assert {n for n, live in plan.items() if not live} == RECORDED_ONLY
    assert D.label(92) == "No hidden sheets (92)" and D.label_for_key(key(77)) == "File extension (.xlsx) (77)"


def test_config_errors_refuse_to_grade():
    s = D.load_settings()
    cases = {
        "overlap": s.__class__(**{**s.__dict__, "live": (92,), "recorded_only": (92,)}),
        "unpinned": s.__class__(**{**s.__dict__, "live": (23,), "recorded_only": ()}),
        "not live in detchecks": s.__class__(**{**s.__dict__, "live": (65,), "recorded_only": ()}),
    }
    for what, bad in cases.items():
        try:
            D.configured_checks(bad, RUBRIC)
        except D.DetChecksConfigError as e:
            assert "(" in str(e), e
        else:
            raise AssertionError(f"{what}: no DetChecksConfigError")
    drifted = copy.deepcopy(RUBRIC)
    pd = drifted["Potential Dangers"]
    i = next(i for i, c in enumerate(pd) if c["name"] == "No hidden sheets")
    pd[i], pd[i + 1] = pd[i + 1], pd[i]                  # 92 and 93 swap places
    try:
        D.configured_checks(s, drifted, "/x/task/rubric.json")
    except D.DetChecksConfigError as e:
        assert "numbering drift" in str(e) and str(e).startswith("No hidden sheets (92) is pinned to (")
        assert "the loaded rubric (/x/task/rubric.json) has ('Potential Dangers', 'No hidden rows/columns')" in str(e)
    else:
        raise AssertionError("rubric drift not refused")
    small = {"Accuracy": RUBRIC["Accuracy"][:5]}         # a v1-sized rubric: none of the checks exists
    assert D.configured_checks(s, small) == {}
    for bad_env in ({"DET_CHECKS_LIVE": "92,x"}, {"DET_CHECKS_LIVE": "92,92"},
                    {"DET_CHECKS_LIBREOFFICE_MAX_MB": "0"}):
        with env(**bad_env):
            try:
                D.load_settings()
            except D.DetChecksConfigError:
                pass
            else:
                raise AssertionError(f"bad config accepted: {bad_env}")


def test_startup_check():
    assert D.startup_check(RUBRIC_PATH) == "harness"
    assert D.startup_check(RUBRIC_PATH, "llm") == "llm"
    assert D.startup_check("/nonexistent/rubric.json", "off") == "off"       # off reads nothing
    with env(DET_CHECKS_LIVE="29,65"):
        try:
            D.startup_check(RUBRIC_PATH)
        except D.DetChecksConfigError as e:
            assert "Negatives in parentheses (65)" in str(e)
        else:
            raise AssertionError("a bad config passed the startup check")


def test_resolve_mode():
    s = D.load_settings()
    assert D.resolve_mode(None, s) == "harness"
    assert D.resolve_mode("llm", s) == "llm" and D.resolve_mode("off", s) == "off"
    assert D.resolve_mode(None, s.__class__(**{**s.__dict__, "enabled": False})) == "off"
    try:
        D.resolve_mode("shadow", s)
    except ValueError:
        pass
    else:
        raise AssertionError("unknown mode accepted")


# --------------------------------------------------------------------------- verdicts
def test_live_fail_and_live_pass():
    with tempfile.TemporaryDirectory() as tmp, fake_libreoffice():
        root = Path(tmp)
        bad = run(make_task(root, hidden=True))
        good = run(make_task(root, hidden=False))
    e = bad.harness_verdicts[key(92)]
    assert e["engine"] == "harness" and e["live"] is True and e["family"] == "det_checks"
    assert e["decision"] == "fail" and e["check_no"] == 92 and e["fallback_reason"] is None
    assert e["n_mistakes"] == 1 and len(e["mistakes"]) == 1 and "Notes" in e["mistakes"][0]["location"]
    assert isinstance(e["stats"], dict) and e["summary"]
    g = good.harness_verdicts[key(92)]
    assert g["engine"] == "harness" and g["decision"] == "pass" and g["mistakes"] == [] and g["n_mistakes"] == 0
    # every configured check was graded (no annotation staged, benchmark None: retired rule only)
    assert set(good.summary["graded"]) == LIVE | RECORDED_ONLY and good.summary["not_applicable"] == []
    assert good.summary["checks"][key(92)] == {"check_no": 92, "engine": "harness", "decision": "pass",
                                              "live": True, "n_mistakes": 0,
                                              "summary": good.harness_verdicts[key(92)]["summary"]}
    assert bad.mode == "harness" and bad.for_judge()["mode"] == "harness"


def test_recorded_only_check_never_overlays():
    with tempfile.TemporaryDirectory() as tmp, fake_libreoffice():
        root = Path(tmp)
        r = run(make_task(root, negative=True))
        e = r.harness_verdicts[key(65)]
        assert e["engine"] == "llm" and e["live"] is False and e["decision"] == "fail"
        assert e["fallback_reason"] == det_api.LIVE_NOTE and e["n_mistakes"] == 1
        assert r.harness_verdicts[key(22)]["engine"] == "llm"           # the other v3-bucket check
        # the config can also hold a live-capable check back: 92 recorded only
        with env(DET_CHECKS_LIVE="29,47", DET_CHECKS_RECORDED_ONLY="22,65,92"):
            r2 = run(make_task(root, hidden=True))
        e2 = r2.harness_verdicts[key(92)]
        assert e2["engine"] == "llm" and e2["live"] is False and e2["decision"] == "fail"
        assert set(r2.summary["graded"]) == {22, 29, 47, 65, 92}


def test_suitability_gate_matches_the_judge():
    with tempfile.TemporaryDirectory() as tmp, fake_libreoffice():
        root = Path(tmp)
        folder = make_task(root, hidden=True)
        stage_annotation(folder, not_applicable={51, 92, 65})
        r = run(folder, benchmark="v2")
        assert key(92) not in r.harness_verdicts and key(51) not in r.harness_verdicts
        assert set(r.summary["not_applicable"]) == {51, 65, 92}
        # the judge's own gate on the same folder (single_pass_judge_case: load_for_case +
        # build_effective_weights) scores exactly the pinned checks the adapter graded
        suit = rubric_suitability.load_for_case(folder, RUBRIC, "v2")
        eff = rubric_suitability.build_effective_weights(WEIGHTS, suit["excluded"])
        judged = {(cat, e["name"]) for cat, es in eff.items() if cat != "CategoryWeights" for e in es}
        assert {n for n, pin in D.DET_CHECK_NAMES.items() if pin in judged} == set(r.summary["graded"])
        # every configured check gated out: nothing runs, so even an unreadable file passes through
        dead = make_task(root, corrupt=True)
        stage_annotation(dead, not_applicable=LIVE | RECORDED_ONLY)
        r0 = run(dead, benchmark="v2")
        assert r0.harness_verdicts == {} and r0.summary["status"] == "no_applicable_checks"
        assert strict_json((dead / D.ARTEFACT_FILENAME).read_text())["status"] == "no_applicable_checks"
        # v2 without an annotation refuses exactly like the judge does
        try:
            run(make_task(root), benchmark="v2")
        except rubric_suitability.SuitabilityError:
            pass
        else:
            raise AssertionError("v2 grading without an annotation was not refused")


def test_task_meta_file_extension():
    with tempfile.TemporaryDirectory() as tmp, fake_libreoffice():
        root = Path(tmp)
        folder = make_task(root, origin="Deliverable.xls")
        r = run(folder)
        e = r.harness_verdicts[key(77)]
        assert e["engine"] == "harness" and e["decision"] == "fail"
        assert e["stats"]["delivered_filename"] == "Deliverable.xls"
        assert e["stats"]["filename_source"] == "task_meta.delivered_filename"
        art = strict_json((folder / D.ARTEFACT_FILENAME).read_text())
        assert art["task_meta"] == {"requires_external_links": False, "delivered_filename": "Deliverable.xls"}
        assert r.summary["delivered_filename"] == "Deliverable.xls"
        ok = run(make_task(root, origin="Model.XLSX"))
        assert ok.harness_verdicts[key(77)]["decision"] == "pass"
        # no sidecar: File extension (.xlsx) (77) raises, loudly, naming the check and the file
        lost = make_task(root, origin=None)
        try:
            run(lost)
        except D.DetChecksError as err:
            msg = str(err)
            assert "File extension (.xlsx) (77)" in msg and str(lost / "ai_attempt.xlsx") in msg
            assert workbook_properties.ORIGIN_FILENAME in msg and "before the LLM call" in msg
            assert set(err.failures) == {key(77)} and isinstance(err, det_api.GradingError)
        else:
            raise AssertionError("missing delivered filename did not raise")
        art = strict_json((lost / D.ARTEFACT_FILENAME).read_text())
        assert art["status"] == "error" and key(77) in art["failures"]
        assert art["task_meta"] == {"requires_external_links": False} and art["file"]["origin_problem"] == "missing"


def test_malformed_origin_sidecar_fails_loudly():
    """A _attempt_origin.json that is not JSON, not an object, or has no usable 'original_filename'
    fails File extension (.xlsx) (77) as MALFORMED (not missing), naming the sidecar file, with
    det_checks.json written (before: a JSON list crashed with AttributeError, no artefact)."""
    bad = {"not JSON": "{not json", "a JSON list": '["Model.xlsx"]', "a number name": '{"original_filename": 123}',
           "an empty name": '{"original_filename": ""}', "no name": '{"source": "s3://x"}'}
    with tempfile.TemporaryDirectory() as tmp, fake_libreoffice():
        root = Path(tmp)
        for what, raw in bad.items():
            folder = make_task(root, origin=None)
            sidecar = folder / workbook_properties.ORIGIN_FILENAME
            sidecar.write_text(raw)
            try:
                run(folder)
            except D.DetChecksError as err:
                msg = str(err)
                assert set(err.failures) == {key(77)}, (what, err.failures)
                assert "File extension (.xlsx) (77)" in msg and f"{sidecar} is malformed" in msg, (what, msg)
                assert "missing" not in msg.split(str(sidecar))[1].split(")")[0], (what, msg)
            else:
                raise AssertionError(f"{what}: no DetChecksError")
            art = strict_json((folder / D.ARTEFACT_FILENAME).read_text())
            assert art["status"] == "error" and art["stage"] == "grade" and key(77) in art["failures"], what
            assert art["file"]["origin_problem"].startswith("malformed") and art["file"]["origin_sidecar"] is False
        # File extension (.xlsx) (77) gated out for the task: nothing needs the delivered name (recorded)
        folder = make_task(root, origin=None)
        (folder / workbook_properties.ORIGIN_FILENAME).write_text('["Model.xlsx"]')
        stage_annotation(folder, not_applicable={77})
        r = run(folder, benchmark="v2")
        assert r.summary["status"] == "ok" and key(77) not in r.harness_verdicts
        art = strict_json((folder / D.ARTEFACT_FILENAME).read_text())
        assert art["status"] == "ok" and art["file"]["origin_problem"] == "malformed: a JSON list, not an object"
    _, problem = D.read_origin(Path("/nonexistent/folder"))
    assert problem == "missing"


def test_artefact_written_in_every_case():
    """det_checks.json is written whatever stops the run: a config error, a rubric drift in the task
    folder's rubric.json, a suitability refusal (each re-raised as is, with the stage recorded)."""
    with tempfile.TemporaryDirectory() as tmp, fake_libreoffice():
        root = Path(tmp)
        cfg = make_task(root)
        with env(DET_CHECKS_LIVE="29,65"):
            try:
                run(cfg)
            except D.DetChecksConfigError as e:
                assert "Negatives in parentheses (65)" in str(e)
            else:
                raise AssertionError("bad config accepted")
        art = strict_json((cfg / D.ARTEFACT_FILENAME).read_text())
        assert art["status"] == "error" and art["stage"] == "config" and "DetChecksConfigError" in art["error"]
        drift = make_task(root)
        rb = copy.deepcopy(RUBRIC)
        pd = rb["Potential Dangers"]
        i = next(i for i, c in enumerate(pd) if c["name"] == "No hidden sheets")
        pd[i], pd[i + 1] = pd[i + 1], pd[i]
        (drift / "rubric.json").write_text(json.dumps(rb))
        try:
            run(drift)
        except D.DetChecksConfigError as e:
            assert str(e).startswith("No hidden sheets (92) is pinned to (") and str(drift / "rubric.json") in str(e), e
        else:
            raise AssertionError("drift in the task folder's rubric.json accepted")
        art = strict_json((drift / D.ARTEFACT_FILENAME).read_text())
        assert art["status"] == "error" and art["stage"] == "config" and art["rubric"] == str(drift / "rubric.json")
        unsuit = make_task(root)
        try:
            run(unsuit, benchmark="v2")                      # v2 without an annotation
        except rubric_suitability.SuitabilityError:
            pass
        else:
            raise AssertionError("no SuitabilityError")
        art = strict_json((unsuit / D.ARTEFACT_FILENAME).read_text())
        assert art["status"] == "error" and art["stage"] == "gate" and "SuitabilityError" in art["error"]


def test_recalc_reads_the_delivered_file_with_the_configured_policy():
    with tempfile.TemporaryDirectory() as tmp, fake_libreoffice() as calls:
        folder = make_task(Path(tmp), negative=True)
        (folder / "temp_recalculated").mkdir()
        shutil.copy(folder / "ai_attempt.xlsx", folder / "temp_recalculated" / "ai_attempt.xlsx")
        run(folder)
        assert len(calls) == 1, calls
        c = calls[0]
        assert Path(c["src"]) == folder / "ai_attempt.xlsx"            # never temp_recalculated/
        p = c["policy"]
        assert p.libreoffice_path == os.environ[f"{project_prefix()}_PATHS_LIBREOFFICE_PATH"]
        assert p.excel_allowed is False and p.lo_timeout_s == 600
        assert Path(p.workdir) == folder / D.RECALC_DIRNAME
        assert Path(c["out_dir"]).is_relative_to(folder / D.RECALC_DIRNAME)


def test_recalc_dir_is_removed_when_the_checks_finish():
    """det_checks_recalc/ (LibreOffice copy + profiles) goes as soon as run_det_checks returns or
    raises, whatever the driver (grade_with_orchestration never pruned it); det_checks.json stays."""
    with tempfile.TemporaryDirectory() as tmp, fake_libreoffice() as calls:
        ok = make_task(Path(tmp), negative=True)
        run(ok)
        assert len(calls) == 1 and Path(calls[0]["out_dir"]).is_relative_to(ok / D.RECALC_DIRNAME)
        assert not (ok / D.RECALC_DIRNAME).exists(), list((ok / D.RECALC_DIRNAME).rglob("*"))
        assert strict_json((ok / D.ARTEFACT_FILENAME).read_text())["status"] == "ok"
    original = det_recalc.libreoffice_recalc

    def _copy_then_fail(src, out_dir, policy):                 # a copy on disk, then a failure
        original_out = os.path.join(out_dir, "ai_attempt.xlsx")
        os.makedirs(out_dir, exist_ok=True)
        shutil.copy(src, original_out)
        raise det_recalc.GradingError("LibreOffice produced no usable copy (test)")

    def _boom(path, *, checks, task_meta, recalc):            # anything else, after the workdir exists
        os.makedirs(os.path.join(recalc.workdir, "_lo_profiles", "lo_x"), exist_ok=True)
        raise RuntimeError("engine bug (test)")

    with tempfile.TemporaryDirectory() as tmp:
        bad = make_task(Path(tmp), negative=True)
        det_recalc.libreoffice_recalc = _copy_then_fail
        try:
            try:
                run(bad)
            except D.DetChecksError as e:
                assert "no usable copy" in str(e)
            else:
                raise AssertionError("no DetChecksError")
        finally:
            det_recalc.libreoffice_recalc = original
        assert not (bad / D.RECALC_DIRNAME).exists()
        assert strict_json((bad / D.ARTEFACT_FILENAME).read_text())["status"] == "error"
        crash = make_task(Path(tmp))
        with patched_grade(_boom):
            try:
                run(crash)
            except D.DetChecksError as e:
                assert "engine bug" in str(e)
            else:
                raise AssertionError("no DetChecksError")
        assert not (crash / D.RECALC_DIRNAME).exists()
        assert (crash / D.ARTEFACT_FILENAME).exists() and (crash / "ai_attempt.xlsx").exists()


def test_missing_libreoffice_fails_loudly_without_launching_anything():
    with tempfile.TemporaryDirectory() as tmp:
        folder = make_task(Path(tmp), negative=True)
        with env(PATHS_LIBREOFFICE_PATH="/nonexistent/LibreOffice/soffice"):
            try:
                run(folder)
            except D.DetChecksError as err:
                msg = str(err)
                for no in (65, 66, 73, 94):                            # value checks
                    assert D.label(no) in msg, (no, msg)
                assert "LibreOffice not found" in msg and str(folder / "ai_attempt.xlsx") in msg
                assert "paths.libreoffice_path" in msg
                assert key(92) not in err.failures                     # structure checks were fine
            else:
                raise AssertionError("missing LibreOffice did not raise")
        assert strict_json((folder / D.ARTEFACT_FILENAME).read_text())["status"] == "error"


def test_libreoffice_size_limit_fails_loudly_without_running_it():
    with tempfile.TemporaryDirectory() as tmp, fake_libreoffice() as calls:
        folder = make_task(Path(tmp), negative=True)
        size = (folder / "ai_attempt.xlsx").stat().st_size
        with env(DET_CHECKS_LIBREOFFICE_MAX_MB=str(size / 2 / 1_000_000)):
            try:
                run(folder)
            except D.DetChecksError as err:
                msg = str(err)
                assert "not graded: too large" in msg and "det_checks.libreoffice_max_mb" in msg
                assert D.label(65) in msg and D.label(66) in msg and str(folder / "ai_attempt.xlsx") in msg
                assert key(92) not in err.failures and key(77) not in err.failures
            else:
                raise AssertionError("a file over the LibreOffice size limit was graded")
        assert calls == [], "LibreOffice must not run on a file over the limit"
        # under the limit the same file goes to LibreOffice as usual
        with env(DET_CHECKS_LIBREOFFICE_MAX_MB=str(size * 2 / 1_000_000)):
            run(folder)
        assert len(calls) == 1


def test_mode_off_runs_nothing():
    def _boom(*a, **k):
        raise AssertionError("detchecks must not run in mode off")
    with tempfile.TemporaryDirectory() as tmp, patched_grade(_boom):
        folder = make_task(Path(tmp), corrupt=True)
        r = run(folder, mode="off")
        assert r.mode == "off" and r.harness_verdicts == {} and r.summary == {"status": "off", "mode": "off"}
        assert strict_json((folder / D.ARTEFACT_FILENAME).read_text())["status"] == "off"
        with env(DET_CHECKS_ENABLED="false"):
            assert run(folder).mode == "off"


def test_mode_llm_records_the_same_entries():
    with tempfile.TemporaryDirectory() as tmp, fake_libreoffice():
        folder = make_task(Path(tmp), hidden=True)
        h = run(folder, mode="harness")
        s = run(folder, mode="llm")
    assert s.mode == "llm" and s.for_judge()["mode"] == "llm"
    strip = lambda hv: {k: {kk: vv for kk, vv in v.items() if kk != "stats"} for k, v in hv.items()}  # noqa: E731
    assert strip(s.harness_verdicts) == strip(h.harness_verdicts)   # the switch acts at scoring


def test_json_safety():
    nasty = {
        "nan": float("nan"), "inf": float("inf"), ("Sheet", 3): {1, 2}, 7: frozenset({"b", "a"}),
        "err": ExcelError("#N/A"), "path": Path("/x/y.xlsx"), "tuple": (1, 2.5, None),
        "bytes": b"\xffabc", "mixed_set": {1, "a"}, "nested": [{"v": -float("inf")}],
    }
    text = json.dumps(D.json_safe(nasty), allow_nan=False)
    back = strict_json(text)
    assert back["nan"] == "nan" and back["('Sheet', 3)"] == [1, 2] and back["7"] == ["a", "b"]
    assert back["err"] == "#N/A" and type(D.json_safe(ExcelError("#N/A"))) is str

    def _crafted(path, *, checks, task_meta, recalc):
        return {key(92): {"engine": "harness", "decision": "fail", "summary": "s",
                          "mistakes": [{"location": "'Notes'", "description": "d", "severity": "major"}],
                          "stats": {"n_mistakes": 1, "ratio": float("nan"), ("a", 1): {3, 1},
                                    "err": ExcelError("#REF!"), "where": Path(path)},
                          "live": True}}
    with tempfile.TemporaryDirectory() as tmp, patched_grade(_crafted), env(DET_CHECKS_LIVE="92", DET_CHECKS_RECORDED_ONLY=""):
        folder = make_task(Path(tmp))
        r = run(folder)
        json.dumps(r.harness_verdicts, allow_nan=False)
        json.dumps(r.summary, allow_nan=False)
        json.dumps(r.for_judge(), allow_nan=False)
        st = r.harness_verdicts[key(92)]["stats"]
        assert st["ratio"] == "nan" and st["('a', 1)"] == [1, 3] and st["err"] == "#REF!"
        strict_json((folder / D.ARTEFACT_FILENAME).read_text())


def test_code_sha_hashes_only_the_grading_code():
    """code_sha fingerprints detchecks/__init__.py, api.py, errors.py, core/ and checks/ only: a
    probe dropped into the git-ignored scratch/ (or tests/, tools/, docs/, out/) must not move it,
    so the same commit always records the same code_sha; any change to the grading code must."""
    import detchecks

    root = Path(detchecks.__file__).resolve().parent
    files = [p.relative_to(root).as_posix() for p in D.code_sha_files(root)]
    assert {"__init__.py", "api.py", "errors.py", "core/recalc.py", "core/lo_watchdog.py", "checks/c92.py"} <= set(files)
    assert all(f in D.CODE_SHA_FILES or f.startswith(("core/", "checks/")) for f in files), files
    assert D.code_sha() == D.code_sha_of(root)
    with tempfile.TemporaryDirectory() as tmp:
        copy = Path(tmp) / "detchecks"
        for f in files:
            (copy / f).parent.mkdir(parents=True, exist_ok=True)
            shutil.copy(root / f, copy / f)
        base = D.code_sha_of(copy)
        assert base == D.code_sha(), "the same grading code elsewhere gives the same code_sha"
        for extra in ("scratch/review_v13/ops/probe.py", "tests/test_new.py", "tools/new_tool.py", "docs/x.py",
                      "out/y.py", "core/__pycache__/recalc.cpython-312.py"):
            (copy / extra).parent.mkdir(parents=True, exist_ok=True)
            (copy / extra).write_text("x = 1\n")
        assert D.code_sha_of(copy) == base, "a file outside the grading code moved code_sha"
        (copy / "checks" / "c92.py").write_text((copy / "checks" / "c92.py").read_text() + "\n# changed\n")
        assert D.code_sha_of(copy) != base
    # and in the real package folder: a new file under scratch/ leaves it unchanged
    probe = root / "scratch" / f"_code_sha_probe_{os.getpid()}.py"
    probe.parent.mkdir(exist_ok=True)
    try:
        probe.write_text("print('probe')\n")
        assert D.code_sha_of(root) == D.code_sha()
    finally:
        probe.unlink(missing_ok=True)


def test_db_payload_stores_values_once_and_caps_lists():
    """scored_results carries the recalculation block once (det_checks.values; each value check's
    stats.values is a reference to it), lists cut to DB_LIST_CAP with their length recorded, the
    Python summary per check (capped); det_checks.json keeps every list and every copy in full."""
    gaps = [{"sheet": "S", "ref": f"A{i}", "value": "#VALUE!", "formula": f"SUM(UNIQUE_GAP_{i:02d}:B9)",
             "functions": ["range"], "why": "w"} for i in range(25)]
    shared = {"writer": "openpyxl", "source": "libreoffice", "value_path": "/tmp/x/det_checks_recalc/h/lo/a.xlsx",
              "value_writer": "libreoffice", "n_gaps": 25, "gap_functions": ["range"], "gaps": gaps,
              "n_lo_errors": 25, "timings": {"libreoffice_s": 2.0}, "lo_version": "LibreOffice/25.8", "notes": ["n"]}

    def _verdict(no, decision="pass", **stats):
        return {"decision": decision, "summary": f"py {no} " + "x" * (400 if no == 66 else 0),
                "mistakes": [], "stats": {"n_mistakes": 0, **stats}, "live": no not in (22, 65)}

    def _crafted(path, *, checks, task_meta, recalc):
        return {key(22): _verdict(22, values=shared), key(65): _verdict(65, values=shared),
                key(66): _verdict(66, values=shared,
                                  examples=[f"E{i}" for i in range(40)],
                                  per_sheet=[{"sheet": "S", "rows": list(range(15))}]),
                key(92): _verdict(92, hidden_sheets=["a", "b"])}

    with tempfile.TemporaryDirectory() as tmp, patched_grade(_crafted), \
            env(DET_CHECKS_LIVE="66,92", DET_CHECKS_RECORDED_ONLY="22,65"):
        folder = make_task(Path(tmp))
        r = run(folder)
        art = strict_json((folder / D.ARTEFACT_FILENAME).read_text())
    vals = r.summary["values"]
    assert vals["n_gaps"] == 25 and len(vals["gaps"]) == D.DB_LIST_CAP and vals["db_capped"] == {"gaps": 25}
    assert "value_path" not in vals and vals["source"] == "libreoffice" and vals["timings"] == {"libreoffice_s": 2.0}
    for no in (22, 65, 66):
        assert r.harness_verdicts[key(no)]["stats"]["values"] == {"ref": D.VALUES_REF}, no
    st66 = r.harness_verdicts[key(66)]["stats"]
    assert st66["examples"] == [f"E{i}" for i in range(D.DB_LIST_CAP)]
    assert st66["per_sheet"][0]["rows"] == list(range(D.DB_LIST_CAP))
    assert st66["db_capped"] == {"examples": 40, "per_sheet[0].rows": 15}
    assert r.harness_verdicts[key(92)]["stats"] == {"n_mistakes": 0, "hidden_sheets": ["a", "b"]}
    assert r.summary["checks"][key(92)]["summary"] == "py 92 "
    assert len(r.summary["checks"][key(66)]["summary"]) == D.DB_SUMMARY_CAP
    assert r.summary["db_list_cap"] == D.DB_LIST_CAP
    # the bundle keeps everything
    assert len(art["values"]["gaps"]) == 25 and art["values"]["value_path"] == shared["value_path"]
    assert all(len(art["verdicts"][key(no)]["stats"]["values"]["gaps"]) == 25 for no in (22, 65, 66))
    assert len(art["verdicts"][key(66)]["stats"]["examples"]) == 40
    # through the scoring layer: the block lands in scored_results exactly once
    res, _out = _finalize(_llm_judgement(), r.harness_verdicts, "harness", r.for_judge())
    text = json.dumps(res["score_results"], allow_nan=False)
    assert text.count("UNIQUE_GAP_00") == 1 and "UNIQUE_GAP_24" not in text
    p66 = res["score_results"]["accuracy_engine"]["checks"][key(66)]
    assert p66["stats"]["values"] == {"ref": D.VALUES_REF} and res["score_results"]["det_checks"]["values"] == vals


def test_merge_harness_verdicts():
    a = {"Accuracy/Final calculation accuracy": {"engine": "harness"}}
    b = {key(92): {"engine": "harness"}}
    merged = D.merge_harness_verdicts(a, b)
    assert set(merged) == set(a) | set(b) and D.merge_harness_verdicts(None, None) == {}
    try:
        D.merge_harness_verdicts(a, {"Accuracy/Final calculation accuracy": {}})
    except ValueError:
        pass
    else:
        raise AssertionError("overlapping keys accepted")


def test_reaper_kills_only_this_gradings_libreoffice():
    # a task folder named like task 28 plus '%41': soffice gets its profile as an encoded file URL
    with tempfile.TemporaryDirectory() as tmp:
        workdir = Path(tmp) / "FruitJuice_3-Statement-Model - v2 %41__task_28" / D.RECALC_DIRNAME
        profile = workdir / "_lo_profiles" / "lo_test"
        profile.mkdir(parents=True)
        sleeper = [sys.executable, "-c", "import time; time.sleep(60)"]
        ours_url = subprocess.Popen(sleeper + [f"-env:UserInstallation={profile.as_uri()}"])
        ours_raw = subprocess.Popen(sleeper + [f"-env:UserInstallation=file://{profile}"], start_new_session=True)
        other = subprocess.Popen(sleeper + [f"-env:UserInstallation={(Path(tmp) / 'elsewhere' / 'lo_x').as_uri()}"],
                                 start_new_session=True)
        try:
            time.sleep(0.3)
            killed = D._reap_libreoffice(workdir)
            assert ours_url.pid in killed and ours_raw.pid in killed and other.pid not in killed, killed
            assert ours_url.wait(timeout=10) != 0 and ours_raw.wait(timeout=10) != 0
            assert other.poll() is None, "an unrelated process was killed"
        finally:
            for p in (ours_url, ours_raw, other):
                if p.poll() is None:
                    p.kill()
                    p.wait(timeout=10)


_STAND_IN_GRADER = r'''
import os, sys
sys.path.insert(0, {judge!r})
os.chdir({judge!r})
from utils.misc_utils import load_project_configs, project_prefix
load_project_configs()
os.environ[project_prefix() + "_PATHS_LIBREOFFICE_PATH"] = {soffice!r}
from utils import det_checks as D
D.run_det_checks({folder!r}, rubric_path={rubric!r}, weights_path={weights!r}, benchmark=None)
'''


def _libreoffice_dies_with_its_grader(sig):
    """The reviewer's P20: a grading process whose LibreOffice step hangs (a stand-in soffice that
    keeps a tagged child) gets `sig`; neither soffice nor its child may survive it (before: SIGTERM
    left soffice re-parented to pid 1 with no timeout)."""
    from detchecks.tests import test_recalc_libreoffice as LT

    tag = LT.new_tag()
    with tempfile.TemporaryDirectory() as tmp:
        root = Path(tmp) / "FruitJuice - v2 %41"
        root.mkdir()
        folder = make_task(root, negative=True)                     # openpyxl file: value checks need LibreOffice
        soffice = LT.fake_soffice(tmp, "hang", tag)
        code = _STAND_IN_GRADER.format(judge=str(JUDGE), soffice=soffice, folder=str(folder),
                                       rubric=str(RUBRIC_PATH), weights=str(WEIGHTS_PATH))
        grader = subprocess.Popen([sys.executable, "-c", code], stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        marker = (folder / D.RECALC_DIRNAME / "_lo_profiles").as_uri()
        try:
            assert LT.wait_until(lambda: LT.procs_with(tag + "_child"), 60), "the stand-in soffice never started"
            assert LT.procs_with(marker), "soffice not visible with its encoded profile URL"
            grader.send_signal(sig)
            assert grader.wait(timeout=30) == -sig
            gone = LT.wait_until(lambda: not LT.procs_with(tag, marker), 10)
            assert gone, f"LibreOffice outlived its grading ({sig!r}): {LT.procs_with(tag, marker)}"
        finally:
            if grader.poll() is None:
                grader.kill()
                grader.wait(timeout=10)
            for pid, _cmd in LT.procs_with(tag, marker):
                LT.lo_watchdog.kill_tree(pid)


def test_libreoffice_dies_with_a_sigterm_to_the_grading():
    _libreoffice_dies_with_its_grader(signal.SIGTERM)


def test_libreoffice_dies_with_a_sigkill_to_the_grading():
    _libreoffice_dies_with_its_grader(signal.SIGKILL)


# --------------------------------------------------------------------------- scoring layer
def _llm_judgement(fail=(), skip=()):
    """An LLM judgement over all of rubric_9: pass everywhere except `fail` (by number);
    numbers in `skip` are not recorded at all."""
    out = {}
    for no, (cat, name) in enumerate(FLAT, 1):
        if no in skip:
            continue
        failing = no in fail
        out.setdefault(cat, []).append({
            "check": str(no), "name": name, "decision": "fail" if failing else "pass", "summary": "llm",
            "mistakes": [{"location": "S!A1", "description": "llm", "severity": "major"}] if failing else [],
        })
    return out


def _det_entry(no, decision, live=True):
    mistakes = [{"location": "'Notes'", "description": "py", "severity": "major"}] if decision == "fail" else []
    return {"engine": "harness" if live else "llm", "decision": decision, "summary": f"py {no}",
            "mistakes": mistakes, "fallback_reason": None if live else det_api.LIVE_NOTE,
            "family": "det_checks", "live": live, "check_no": no, "n_mistakes": len(mistakes),
            "stats": {"n_mistakes": len(mistakes), "values": {"source": "cache"}}}


FA = "Accuracy/Final calculation accuracy"
HARNESS = {
    FA: {"engine": "harness", "decision": "pass", "summary": "75/75", "mistakes": [],
         "fallback_reason": None, "n_questions": 75, "n_match": 75},
    key(92): _det_entry(92, "pass"),                  # LLM fails it below
    key(61): _det_entry(61, "fail"),                  # LLM passes it
    key(74): _det_entry(74, "fail"),                  # LLM never recorded it
    key(65): _det_entry(65, "fail", live=False),      # recorded only; LLM passes
    key(22): _det_entry(22, "pass", live=False),      # recorded only; LLM fails
}
LLM_FAIL = {1, 92, 22}


def test_apply_harness_verdicts_with_deterministic_entries():
    llm = _llm_judgement(fail=LLM_FAIL, skip={74})
    overlaid, prov = judge._apply_harness_verdicts(llm, HARNESS, WEIGHTS)
    p92 = prov[key(92)]
    assert p92["engine"] == "harness" and p92["agreed"] is False and p92["llm_decision"] == "fail"
    assert p92["family"] == "det_checks" and p92["live"] is True and p92["check_no"] == 92
    assert p92["n_mistakes"] == 0 and p92["stats"]["values"]["source"] == "cache"
    item92 = next(i for i in overlaid["Potential Dangers"] if i["name"] == "No hidden sheets")
    assert item92["decided_by"] == "harness" and item92["decision"] == "pass" and item92["llm_decision"] == "fail"
    # recorded only: Python decision + agreement recorded, the LLM item untouched
    p65 = prov[key(65)]
    assert p65["engine"] == "llm" and p65["decision"] == "fail" and p65["llm_decision"] == "pass"
    assert p65["agreed"] is False and p65["live"] is False and p65["fallback_reason"] == det_api.LIVE_NOTE
    item65 = next(i for i in overlaid["Formatting"] if i["name"] == "Negatives in parentheses")
    assert item65["decision"] == "pass" and "decided_by" not in item65
    assert prov[key(22)]["agreed"] is False and prov[key(61)]["agreed"] is False
    # a check the LLM never recorded is inserted under its rubric number
    item74 = next(i for i in overlaid["Formatting"] if i["name"] == "No merged cells")
    assert item74["check"] == "74" and item74["decided_by"] == "harness" and item74["llm_decision"] is None
    assert prov[key(74)]["agreed"] is None
    # the answer-check entry keeps its v12 provenance
    assert prov[FA]["engine"] == "harness" and prov[FA]["n_match"] == 75 and "family" not in prov[FA]


_FINALIZE_ROOT = tempfile.mkdtemp(prefix="detchecks_finalize_")
atexit.register(shutil.rmtree, _FINALIZE_ROOT, True)


def _finalize(llm, hv, accuracy_engine, det_checks):
    out = Path(tempfile.mkdtemp(dir=_FINALIZE_ROOT))
    log = out / "cache.log"
    log.write_text("")
    tt = {"evaluations": {}, "total_message_size": 0, "total_message_size_with_images": 0, "total_tokens": 0,
          "total_prompt_tokens": 0, "total_completion_tokens": 0, "total_cost": 0.0}
    res = judge._finalize_case(
        all_responses=copy.deepcopy(llm), output_dir=out, weights_data=WEIGHTS, token_tracking=tt,
        model="test/model", attempt_model="agent", task_folder_name="t", golden_solution_files={},
        ai_attempt_files={}, context_file_path=None, start_time=time.time(), cache_log_path=str(log),
        versions={"JUDGE_VERSION": "13", "PROMPT_VERSION": "8", "RUBRIC_VERSION": "9", "RUBRIC_WEIGHT_VERSION": "9"},
        agentic=True, harness_verdicts=hv, accuracy_engine=accuracy_engine, det_checks=det_checks,
    )
    return res, out


def test_finalize_case_switches():
    llm = _llm_judgement(fail=LLM_FAIL, skip={74})
    with tempfile.TemporaryDirectory() as tmp:
        artefact = Path(tmp) / "det_checks.json"
        artefact.write_text(json.dumps({"status": "ok", "verdicts": {}}))
        summary = {"status": "ok", "mode": None, "graded": [22, 61, 65, 74, 92], "not_applicable": [],
                   "checks": {}, "code_sha": "x", "values": {"source": "cache"}}
        results = {}
        for acc in ("harness", "llm"):
            for det in ("harness", "llm"):
                dc = {"mode": det, "summary": {**summary, "mode": det}, "artefact": str(artefact)}
                results[(acc, det)] = _finalize(llm, HARNESS, acc, dc)
        expected = {("harness", "harness"): "harness", ("harness", "llm"): "mixed",
                    ("llm", "harness"): "mixed", ("llm", "llm"): "llm"}
        totals_llm, totals_h = set(), set()
        for (acc, det), (res, out) in results.items():
            sr = res["score_results"]
            eng = sr["accuracy_engine"]
            assert eng["effective"] == expected[(acc, det)], (acc, det, eng["effective"])
            assert eng["mode"] == acc and eng["det_checks_mode"] == det
            totals_llm.add(eng["total_score_llm"])
            totals_h.add(eng["total_score_harness"])
            counted = {k for k, v in eng["checks"].items() if v["counted"]}
            want = ({FA} if acc == "harness" else set()) | ({key(92), key(61), key(74)} if det == "harness" else set())
            assert counted == want, (acc, det, counted)
            assert not eng["checks"][key(65)]["counted"] and not eng["checks"][key(22)]["counted"]
            # recorded-only checks always score the LLM's verdict
            assert sr["check_scores"]["Formatting"]["Negatives in parentheses"]["score"] == 1.0
            assert sr["check_scores"]["Error Checks"]["No formula errors"]["score"] == 0.0
            # live checks follow Python only when det_checks counts
            hid = sr["check_scores"]["Potential Dangers"]["No hidden sheets"]["score"]
            assert hid == (1.0 if det == "harness" else 0.0), (acc, det, hid)
            assert sr["det_checks"]["mode"] == det and (out / "det_checks.json").exists()
            meta = json.loads((out / "_metadata.json").read_text())
            assert meta["det_checks"]["mode"] == det and meta["det_checks"]["graded"] == summary["graded"]
            strict_json((out / "scores.json").read_text())
            json.dumps(sr, allow_nan=False)
            assert "decided_by" not in json.dumps(json.loads((out / "ai_judgement.json").read_text()))
            if expected[(acc, det)] == "harness":
                assert sr["total_score"] == eng["total_score_harness"]
            if expected[(acc, det)] == "llm":
                assert sr["total_score"] == eng["total_score_llm"]
        assert len(totals_llm) == 1 and len(totals_h) == 1, "both totals are the same under every switch"
        # mixed: exactly the counted family overlaid
        res, _ = results[("llm", "harness")]
        only_det = {k: v for k, v in HARNESS.items() if k != FA}
        overlaid, _ = judge._apply_harness_verdicts(llm, only_det, WEIGHTS)
        want_total = judge.calculate_scores(overlaid, WEIGHTS, max_mistakes=judge.RUBRIC_MAX_MISTAKES)["total_score"]
        assert abs(res["score_results"]["total_score"] - want_total) < 1e-9
        assert res["final_score"] == res["score_results"]["total_score"]


def test_finalize_case_without_det_checks_keeps_the_v12_shape():
    llm = _llm_judgement(fail={1})
    res, out = _finalize(llm, {FA: HARNESS[FA]}, "harness", None)
    sr = res["score_results"]
    assert "det_checks" not in sr and "det_checks_mode" not in sr["accuracy_engine"]
    assert sr["accuracy_engine"]["effective"] == "harness" and not (out / "det_checks.json").exists()
    assert "det_checks" not in json.loads((out / "_metadata.json").read_text())


if __name__ == "__main__":
    tests = [(n, f) for n, f in sorted(globals().items()) if n.startswith("test_") and callable(f)]
    failed = 0
    for name, fn in tests:
        try:
            fn()
            print(f"OK   {name}")
        except Exception as e:  # noqa: BLE001
            failed += 1
            import traceback
            traceback.print_exc()
            print(f"FAIL {name}: {e}")
    print()
    print(f"{len(tests) - failed}/{len(tests)} passed")
    sys.exit(1 if failed else 0)
